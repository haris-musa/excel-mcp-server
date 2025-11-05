import os
import sys
from pathlib import Path
from typing import Any, List, Optional

from dotenv import load_dotenv
from agno.agent import Agent
from agno.models.openai import OpenAILike

# Resolve project paths
SCRIPT_DIR = Path(__file__).parent
PROJECT_ROOT = SCRIPT_DIR.parents[1]

# Load environment variables from .env relative to this script
DOTENV_PATH = SCRIPT_DIR / "config" / ".env"
load_dotenv(dotenv_path=DOTENV_PATH)

# Ensure `src` is on sys.path so we can import project modules
SRC_PATH = PROJECT_ROOT / "src"
if str(SRC_PATH) not in sys.path:
    sys.path.append(str(SRC_PATH))

from excel_mcp.data import get_sheet_as_markdown_table, read_excel_range, write_data  # type: ignore
from openpyxl import load_workbook

# --- Tool Definitions ---


def list_sheets(file_path: str) -> List[str]:
    """列出 Excel 文件中的所有工作表名称。"""
    try:
        wb = load_workbook(file_path, read_only=True)
        return wb.sheetnames
    except Exception as e:
        return [f"Error listing sheets: {e}"]


def read_range(file_path: str, sheet_name: str, cell_range: str) -> List[List[Any]]:
    """读取指定工作表和范围的数据。范围格式如 'A1:D1'。"""
    try:
        start_cell, end_cell = cell_range.split(":")
        return read_excel_range(file_path, sheet_name, start_cell, end_cell)
    except Exception as e:
        return [[f"Error reading range: {e}"]]


def write_range(file_path: str, sheet_name: str, start_cell: str, data: List[List[Any]]) -> str:
    """从指定单元格开始，将数据写入工作表。"""
    try:
        write_data(file_path, sheet_name, data, start_cell)
        return f"成功将数据写入到工作表 '{sheet_name}' 的 '{start_cell}' 位置。"
    except Exception as e:
        return f"Error writing range: {e}"


def read_excel_as_markdown(file_path: str, sheet_name: Optional[str] = None) -> str:
    """读取整个工作表并以 Markdown 表格形式返回，用于概览。"""
    try:
        if not sheet_name:
            sheet_name = list_sheets(file_path)[0]
        return get_sheet_as_markdown_table(file_path, sheet_name)
    except Exception as e:
        return f"Error reading as markdown: {e}"


# --- Agent Build ---


def build_agent() -> Agent:
    """构建具备 Excel 读写能力的 Agno Agent。"""
    zhipu_api_key = os.getenv("ZHIPU_API_KEY")
    if not zhipu_api_key or zhipu_api_key == "YOUR_ZHIPU_API_KEY":
        raise RuntimeError("ZHIPU_API_KEY 未配置，请在 .env 文件中设置")

    base_url = "https://open.bigmodel.cn/api/coding/paas/v4"
    # base_url = "http://127.0.0.1:3456/"
    model_id = os.getenv("MODEL_NAME") or "glm-4.6"

    model = OpenAILike(id=model_id, api_key=zhipu_api_key, base_url=base_url)

    tools = [list_sheets, read_range, write_range, read_excel_as_markdown]

    instructions = (
        "你是一个强大的 Excel 助手，能够理解和编辑 Excel 文件。"
        "1. **分析任务**: 首先理解用户的修改请求。"
        "2. **探索文件**: 使用 `list_sheets` 查看有哪些工作表，使用 `read_excel_as_markdown` 概览表格结构。"
        "3. **定位和读取**: 使用 `read_range` 精确读取需要修改的单元格范围（例如标题行）。"
        "4. **思考和修改**: 在内部完成文本的翻译或修改。"
        "5. **执行写入**: 使用 `write_range` 将修改后的新数据写回原位。"
        "6. **确认结果**: 回复用户，告知操作已完成。"
    )

    return Agent(
        model=model,
        tools=tools,
        instructions=instructions,
        markdown=True,
        debug_mode=True,
    )


if __name__ == "__main__":
    agent = build_agent()
    excel_path = r"F:\WS_Python\excel-mcp-server\HUA_VVP_RA_INT_05-2025_072339.xlsx"
    prompt = f"请帮我修改这个 Excel 文件：'{excel_path}'。" "具体操作：请将第一个工作表的标题行翻译成中文。"

    response = agent.run(prompt)
    print(response.content)
