"""Dynamic array formulas: stored the way Excel stores them, and kept across later edits."""

import re
from pathlib import Path
from typing import NamedTuple
from xml.etree import ElementTree

import pytest
from mcp import Client
from openpyxl import Workbook, load_workbook

from excel_mcp.config import Limits, Settings
from excel_mcp.server import create_server
from tests.conftest import ToolCall, error_text
from tests.package_builders import value_metadata_workbook
from tests.package_support import (
    FIXTURES,
    assert_package_is_consistent,
    content_type,
    copy_fixture,
    read_parts,
    relationships,
    sheet_part,
    text,
)

pytestmark = pytest.mark.anyio

BOOK = {"path": "book.xlsx"}
MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
# The formulas of tests/fixtures/excel_dynamic_arrays.xlsx, which Excel saved.
FORMULAS = [
    "=FILTER(Data!A2:B6,Data!B2:B6>4)",
    None,
    "=SORT(Data!B2:B6)",
    "=SEQUENCE(3)",
    "=SUM(Data!B2:B6*2)",
    "=UNIQUE(Data!A2:A6)",
]


class Anchor(NamedTuple):
    """The formula cell of a spill as the file stores it."""

    kind: str
    cm: str | None
    attributes: dict[str, str]
    formula: str | None
    value: str | None


def _anchors(parts: dict[str, bytes], sheet: str) -> dict[str, Anchor]:
    root = ElementTree.fromstring(parts[sheet_part(parts, "Out")])
    found = {}
    for cell in root.iter(f"{MAIN}c"):
        formula = cell.find(f"{MAIN}f")
        if formula is not None:
            value = cell.find(f"{MAIN}v")
            found[cell.get("r", "")] = Anchor(
                cell.get("t", "n"),
                cell.get("cm"),
                dict(formula.attrib),
                formula.text,
                None if value is None else value.text,
            )
    return found


@pytest.fixture
def numbers(files: Path) -> Path:
    """The Data sheet of the Excel fixture, with an empty Out sheet."""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in [["Item", "Qty"], ["a", 5], ["b", 3], ["c", 9], ["d", 1], ["e", 7]]:
        sheet.append(row)
    workbook.create_sheet("Out")
    path = files / "book.xlsx"
    workbook.save(path)
    return path


async def test_arrays_are_stored_exactly_as_excel_stores_them(
    call: ToolCall, numbers: Path
) -> None:
    await call("write_range", **BOOK, sheet="Out", start_cell="A1", rows=[FORMULAS])
    ours = read_parts(numbers)
    excel = read_parts(FIXTURES / "excel_dynamic_arrays.xlsx")

    assert _anchors(ours, "Out") == _anchors(excel, "Out")
    assert len(_anchors(ours, "Out")) == 5
    # Excel ends the XML declaration with CR LF; the text of the part is the same otherwise.
    ours_xml = text(ours, "xl/metadata.xml").split("\n", 1)[1]
    assert ours_xml == text(excel, "xl/metadata.xml").split("\n", 1)[1]
    assert content_type(ours, "xl/metadata.xml") == content_type(excel, "xl/metadata.xml")
    kinds = [kind for kind, _ in relationships(ours, "xl/workbook.xml").values()]
    assert kinds.count("sheetMetadata") == 1
    assert_package_is_consistent(ours)


async def test_spilled_cells_hold_the_results_and_read_back(call: ToolCall, numbers: Path) -> None:
    await call("write_range", **BOOK, sheet="Out", start_cell="A1", rows=[FORMULAS])
    sheet = load_workbook(numbers)["Out"]
    cells = text(read_parts(numbers), sheet_part(read_parts(numbers), "Out"))

    assert [sheet.cell(row, 3).value for row in range(2, 6)] == [3, 5, 7, 9]
    assert [sheet.cell(row, 6).value for row in range(2, 6)] == ["b", "c", "d", "e"]
    assert '<c r="F2" t="str"><v>b</v></c>' in cells  # how Excel stores spilled text
    assert sheet["A1"].value.ref == "A1:B3"  # type: ignore[union-attr]
    data = await call("read_range", **BOOK, sheet="Out")
    assert data["values"][0] == ["a", 5, 1, 1, 50, "a"]
    assert data["values"][2] == ["e", 7, 5, 3, None, "c"]
    formulas = await call("read_range", **BOOK, sheet="Out", range="A1:A1", mode="formulas")
    assert formulas["values"] == [["=_xlfn._xlws.FILTER(Data!A2:B6,Data!B2:B6>4)"]]


async def test_ordinary_formulas_stay_ordinary(call: ToolCall, numbers: Path) -> None:
    rows = [["=SUM(Data!B2:B6)", "=INDEX(Data!B2:B6,MATCH(9,Data!B2:B6,0))", "=Data!B2*2"]]
    await call("write_range", **BOOK, sheet="Out", start_cell="A1", rows=rows)
    parts = read_parts(numbers)

    assert "xl/metadata.xml" not in parts
    assert 't="array"' not in text(parts, sheet_part(parts, "Out"))
    assert " cm=" not in text(parts, sheet_part(parts, "Out"))


async def test_a_blocked_spill_is_reported_and_leaves_the_cells_alone(
    call: ToolCall, numbers: Path
) -> None:
    await call("write_range", **BOOK, sheet="Out", start_cell="A3", rows=[["in the way"]])
    result = await call(
        "write_range", **BOOK, sheet="Out", start_cell="A1", rows=[["=SORT(Data!B2:B6)"]]
    )

    assert result["blocked"] == ["A1"]
    assert load_workbook(numbers)["Out"]["A3"].value == "in the way"
    anchor = _anchors(read_parts(numbers), "Out")["A1"]
    assert anchor.attributes["ref"] == "A1" and anchor.cm == "1" and anchor.value is None
    data = await call("read_range", **BOOK, sheet="Out")
    assert data["values"][0] == ["#SPILL!"]
    await call("write_range", **BOOK, sheet="Out", start_cell="A3", rows=[[None]])
    data = await call("read_range", **BOOK, sheet="Out", range="A1:A5")
    assert [row[0] for row in data["values"]] == [1, 3, 5, 7, 9]


async def test_a_result_the_calculator_cannot_size_is_left_to_excel(
    call: ToolCall, numbers: Path
) -> None:
    await call(
        "write_range", **BOOK, sheet="Out", start_cell="A1", rows=[['=TEXTSPLIT("a,b",",")']]
    )
    anchor = _anchors(read_parts(numbers), "Out")["A1"]

    assert anchor.cm == "1"  # still a dynamic array formula
    assert anchor.attributes["ref"] == "A1"
    assert anchor.value is None  # no result is made up


async def test_replacing_a_formula_replaces_what_it_spilled(call: ToolCall, numbers: Path) -> None:
    await call("write_range", **BOOK, sheet="Out", start_cell="A1", rows=[["=SEQUENCE(5)"]])
    await call("write_range", **BOOK, sheet="Out", start_cell="A1", rows=[["=SEQUENCE(2)"]])
    sheet = load_workbook(numbers)["Out"]

    assert [sheet.cell(row, 1).value for row in range(2, 6)] == [2, None, None, None]
    assert sheet["A1"].value.ref == "A1:A2"  # type: ignore[union-attr]


async def test_clearing_the_formula_clears_what_it_spilled(call: ToolCall, numbers: Path) -> None:
    await call("write_range", **BOOK, sheet="Out", start_cell="A1", rows=[["=SEQUENCE(4)"]])
    await call("write_range", **BOOK, sheet="Out", start_cell="A1", rows=[["plain"]])
    sheet = load_workbook(numbers)["Out"]

    assert sheet["A1"].value == "plain"
    assert [sheet.cell(row, 1).value for row in range(2, 5)] == [None, None, None]
    assert "xl/metadata.xml" not in read_parts(numbers) or " cm=" not in text(
        read_parts(numbers), sheet_part(read_parts(numbers), "Out")
    )


async def test_arrays_made_by_excel_survive_edits_and_moves(call: ToolCall, files: Path) -> None:
    copy_fixture(files, "excel_dynamic_arrays.xlsx")
    original = _anchors(read_parts(files / "book.xlsx"), "Out")
    await call("write_range", **BOOK, sheet="Data", start_cell="D1", rows=[[1]])
    await call("insert_rows_or_columns", **BOOK, sheet="Out", axis="rows", at=1, count=2)
    moved = _anchors(read_parts(files / "book.xlsx"), "Out")

    assert sorted(moved) == ["A3", "C3", "D3", "E3", "F3"]
    refs = [moved[f"{c}3"].attributes["ref"] for c in "ACDEF"]
    assert refs == ["A3:B5", "C3:C7", "D3:D5", "E3", "F3:F7"]
    for old, new in zip(original, moved, strict=True):
        assert original[old].cm == moved[new].cm == "1"
        assert original[old].value == moved[new].value
    data = await call("read_range", **BOOK, sheet="Out")
    assert data["values"][0][:2] == ["a", 5]

    await call("clear_range", **BOOK, sheet="Out", range="C3")
    sheet = load_workbook(files / "book.xlsx")["Out"]
    assert [sheet.cell(row, 3).value for row in range(3, 8)] == [None] * 5
    assert sheet["D3"].value.ref == "D3:D5"  # type: ignore[union-attr]


async def test_dynamic_array_metadata_joins_what_the_file_has(call: ToolCall, files: Path) -> None:
    value_metadata_workbook(files / "book.xlsx")
    await call("write_range", **BOOK, sheet="Sheet", start_cell="D1", rows=[["=SEQUENCE(2)"]])
    parts = read_parts(files / "book.xlsx")
    metadata = text(parts, "xl/metadata.xml")

    assert re.findall(r'<metadataType name="(\w+)"', metadata) == ["XLRICHVALUE", "XLDAPR"]
    assert '<metadataTypes count="2">' in metadata and '<valueMetadata count="1">' in metadata
    assert re.search(r'<cellMetadata count="1"><bk><rc t="2" v="0"/></bk></cellMetadata>', metadata)
    sheet = text(parts, sheet_part(parts, "Sheet"))
    assert re.search(r'<c r="D1" cm="1">', sheet) and re.search(r'<c r="A1"[^>]* vm="1"', sheet)


async def test_a_spill_larger_than_one_call_may_write_is_refused(files: Path) -> None:
    workbook = Workbook()
    workbook.worksheets[0].title = "Out"
    workbook.save(files / "book.xlsx")
    server = create_server(Settings(allowed_dirs=[files], limits=Limits(max_cells=50)))

    async with Client(server) as client:
        result = await client.call_tool(
            "write_range",
            {**BOOK, "sheet": "Out", "start_cell": "A1", "rows": [["=SEQUENCE(100)"]]},
        )

    assert result.is_error
    assert "returns 100 values" in error_text(result)


async def test_a_copied_sheet_copies_its_dynamic_arrays(call: ToolCall, files: Path) -> None:
    copy_fixture(files, "excel_dynamic_arrays.xlsx")
    await call("copy_sheet", **BOOK, sheet="Out", new_name="Copy")
    parts = read_parts(files / "book.xlsx")
    root = ElementTree.fromstring(parts[sheet_part(parts, "Copy")])
    anchors = {}
    for cell in root.iter(f"{MAIN}c"):
        formula = cell.find(f"{MAIN}f")
        if formula is not None:
            anchors[cell.get("r")] = (cell.get("cm"), formula.get("ref"))

    assert anchors == {
        "A1": ("1", "A1:B3"),
        "C1": ("1", "C1:C5"),
        "D1": ("1", "D1:D3"),
        "E1": ("1", "E1"),
        "F1": ("1", "F1:F5"),
    }
