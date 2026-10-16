"""PivotTables over sources that mix text, numbers, dates, booleans and errors, and formulas.

tests/test_pivot_golden.py checks the figures against Excel; these check the cache Excel
writes for such fields, the tool contract, the errors, and slicers and timelines.
"""

import datetime as dt
import shutil
import zipfile
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.pivot.cache import CacheDefinition

from excel_mcp.calc.values import ExcelError
from excel_mcp.config import Limits
from excel_mcp.operations import pivot
from excel_mcp.operations.pivot import PivotRequest
from excel_mcp.operations.pivot_fields import distinct, item_text
from excel_mcp.operations.pivot_index import sheet_pivots
from excel_mcp.operations.pivot_options import PivotValue
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

FIXTURES = Path(__file__).parent / "fixtures"
NA = "#N/A"


def book(files: Path, rows: list[list[Any]], name: str = "m.xlsx") -> str:
    workbook = Workbook()
    workbook.worksheets[0].title = "Data"
    for row in rows:
        workbook["Data"].append(row)
    workbook.create_sheet("Out")
    workbook.save(files / name)
    return name


def source_of(rows: list[list[Any]]) -> str:
    return f"Data!A1:{chr(64 + len(rows[0]))}{len(rows)}"


async def make(call: ToolCall, path: str, rows: list[list[Any]], **arguments: Any) -> Any:
    return await call(
        "create_pivot_table",
        path=path,
        source=source_of(rows),
        sheet="Out",
        at="A1",
        **arguments,
    )


def grid(files: Path, path: str, ref: str) -> list[list[Any]]:
    return [[cell.value for cell in row] for row in load_workbook(files / path)["Out"][ref]]


def part(files: Path, path: str, name: str) -> str:
    with zipfile.ZipFile(files / path) as archive:
        return archive.read(name).decode()


# What Excel writes in the cache for each kind of column: the sharedItems flags that differ
# from the defaults, as recorded from Excel itself (scripts/pivot_golden.py has the PivotTables).
DATE = dt.datetime(2023, 3, 15)
DATES = {"containsSemiMixedTypes": "0", "containsNonDate": "0", "containsDate": "1"}
NUMBERS = {"containsSemiMixedTypes": "0", "containsString": "0", "containsNumber": "1"}
WITH_NUMBERS = {"containsMixedTypes": "1", "containsNumber": "1", "containsInteger": "1"}
FLAGS = {
    "integers": ([1, 2, 3], {**NUMBERS, "containsInteger": "1"}),
    "decimals": ([1.5, 2, 3], NUMBERS),
    "integers and a blank": (
        [1, None, 3],
        {
            "containsString": "0",
            "containsBlank": "1",
            "containsNumber": "1",
            "containsInteger": "1",
        },
    ),
    "text": (["a", "b"], {}),
    "booleans": ([True, False], {}),
    "errors": ([NA, "#DIV/0!"], {}),
    "text and booleans": (["a", True], {}),
    "text and a blank": (["a", None], {"containsBlank": "1"}),
    "numbers and text": ([1, "a"], WITH_NUMBERS),
    "numbers and booleans": ([1, True], WITH_NUMBERS),
    "numbers and errors": ([1, NA], WITH_NUMBERS),
    "dates": ([DATE, DATE + dt.timedelta(days=1)], {**DATES, "containsString": "0"}),
    "dates and a blank": (
        [DATE, None],
        {"containsNonDate": "0", "containsDate": "1", "containsString": "0", "containsBlank": "1"},
    ),
    "dates and text": ([DATE, "a"], {"containsDate": "1", "containsMixedTypes": "1"}),
    "dates and booleans": ([DATE, True], {"containsDate": "1", "containsMixedTypes": "1"}),
    "dates and numbers": (
        [DATE, 5],
        {
            "containsSemiMixedTypes": "0",
            "containsString": "0",
            "containsDate": "1",
            "containsMixedTypes": "1",
        },
    ),
    "times": ([dt.time(6), dt.time(12)], {**DATES, "containsString": "0"}),
    "durations": (
        [dt.timedelta(hours=36), dt.timedelta(hours=1.5)],
        {**DATES, "containsString": "0"},
    ),
    "only blanks": (
        [None, None],
        {"containsNonDate": "0", "containsString": "0", "containsBlank": "1"},
    ),
}


def cache_of(values: list[Any], formats: list[str] | None = None) -> CacheDefinition:
    """The cache of a PivotTable that counts the one column holding ``values``."""
    workbook = Workbook()
    source = workbook.worksheets[0]
    source.append(["X"])
    for value, number_format in zip(values, formats or ["General"] * len(values), strict=True):
        source.append([value])
        source.cell(source.max_row, 1).number_format = number_format
    request = PivotRequest(
        ["X"], [], [PivotValue(field="X", function="count")], [], [], [], "tabular", True,
        "columns", None,
    )  # fmt: skip
    pivot.create_pivot(
        workbook,
        source,
        f"A1:A{len(values) + 1}",
        workbook.create_sheet("Out"),
        "A1",
        request,
        Limits(),
        lambda sheet, area: {},
    )
    return sheet_pivots(workbook["Out"])[0].cache


@pytest.mark.parametrize("kind", FLAGS)
def test_cache_flags_are_those_excel_writes(kind: str) -> None:
    values, expected = FLAGS[kind]
    shared = cache_of(values).cacheFields[0].sharedItems
    assert shared is not None
    flags = {k: v for k, v in shared.to_tree().attrib.items() if k.startswith("contains")}
    assert flags == expected


def test_a_field_with_one_number_format_keeps_it_else_general() -> None:
    assert cache_of([1.5, 2], ["0.00", "0.00"]).cacheFields[0].numFmtId == 2
    assert cache_of([1.5, "a"], ["0.00", "0.00"]).cacheFields[0].numFmtId == 2
    assert cache_of([1.5, "a"], ["0.00", "General"]).cacheFields[0].numFmtId == 0
    assert cache_of([1.5, None], ["0.00", "General"]).cacheFields[0].numFmtId == 0


def test_items_sort_as_excel_sorts_them() -> None:
    values = [
        None,
        DATE,
        ExcelError("#DIV/0!"),
        ExcelError(NA),
        True,
        False,
        "b",
        "B",
        "a",
        10,
        1.5,
        3,
        "3",
    ]
    shared = distinct(values)
    shown = [shared.values[item] for item in shared.order]
    assert shown == [
        1.5,
        3,
        10,
        "3",
        "a",
        "b",
        False,
        True,
        ExcelError(NA),
        ExcelError("#DIV/0!"),
        DATE,
        None,
    ]
    assert item_text(True) == "TRUE" and item_text(dt.datetime(1899, 12, 30, 6)) == "06:00:00"


# Figures: what Excel shows for a source of every type (the first case of the golden fixture)

MIXED = [
    ["K", "V"],
    ["b", 1],
    [3, 2],
    [dt.datetime(2023, 3, 15), 3.5],
    [True, 4],
    [NA, 5],
    [None, 6],
    ["a", "x"],
    [1.5, None],
    ["B", True],
    [10, NA],
    [False, 8],
]
FUNCTIONS = [{"field": "V", "function": f} for f in ("sum", "count", "average", "min", "max")]


async def test_row_labels_of_every_type_and_figures_over_text_booleans_and_errors(
    call: ToolCall, files: Path
) -> None:
    path = book(files, MIXED)
    await make(call, path, MIXED, row_fields=["K"], value_fields=FUNCTIONS)
    assert grid(files, path, "A2:F13") == [
        ["K", "Sum of V", "Count of V", "Average of V", "Min of V", "Max of V"],
        [1.5, None, None, None, None, None],
        [3, 2, 1, 2, 2, 2],
        [10, NA, 1, NA, NA, NA],
        ["a", 0, 1, "#DIV/0!", 0, 0],
        ["b", 1, 2, 1, 1, 1],
        ["FALSE", 8, 1, 8, 8, 8],
        ["TRUE", 4, 1, 4, 4, 4],
        [NA, 5, 1, 5, 5, 5],
        [dt.datetime(2023, 3, 15), 3.5, 1, 3.5, 3.5, 3.5],
        ["(blank)", 6, 1, 6, 6, 6],
        ["Grand Total", NA, NA, NA, NA, NA],
    ]


async def test_records_are_written_with_the_type_of_each_cell(call: ToolCall, files: Path) -> None:
    path = book(files, MIXED)
    await make(call, path, MIXED, row_fields=["K"], value_fields=FUNCTIONS[:1])
    records = part(files, path, "xl/pivotCache/pivotCacheRecords1.xml")
    for item in (
        '<x v="0" />',
        '<n v="3.5" />',
        '<s v="x" />',
        '<b v="1" />',
        '<e v="#N/A" />',
        "<m />",
    ):
        assert item in records
    cache = part(files, path, "xl/pivotCache/pivotCacheDefinition1.xml")
    assert '<b v="1" />' in cache and '<e v="#N/A" />' in cache and "<d v=" in cache
    assert 'containsMixedTypes="1"' in cache


async def test_errors_make_totals_errors_even_for_counts(call: ToolCall, files: Path) -> None:
    rows = [["G", "V"], ["a", 5], ["a", NA], ["b", 3], ["b", 4]]
    path = book(files, rows)
    await make(call, path, rows, row_fields=["G"], value_fields=FUNCTIONS[:2])
    assert grid(files, path, "A3:C5") == [["a", NA, 2], ["b", 7, 2], ["Grand Total", NA, NA]]


async def test_the_first_error_of_a_total_is_the_first_in_table_order(
    call: ToolCall, files: Path
) -> None:
    rows = [["G", "V"], ["z", "#NUM!"], ["a", NA], ["b", 1]]
    path = book(files, rows)
    await make(call, path, rows, row_fields=["G"], value_fields=FUNCTIONS[:1])
    assert grid(files, path, "A2:B5") == [["a", NA], ["b", 1], ["z", "#NUM!"], ["Grand Total", NA]]


async def test_a_calculated_field_over_text_uses_zero_and_over_errors_the_last_error(
    call: ToolCall, files: Path
) -> None:
    rows = [["G", "V", "W"], ["a", "x", 2], ["b", "#DIV/0!", 4], ["b", NA, 5]]
    path = book(files, rows)
    await make(
        call,
        path,
        rows,
        row_fields=["G"],
        value_fields=[{"field": "C"}],
        calculated_fields=[{"name": "C", "formula": "V*2+W"}],
    )
    assert grid(files, path, "A2:B3") == [["a", 2], ["b", NA]]


async def test_dates_are_summed_as_serial_numbers_and_times_are_dates(
    call: ToolCall, files: Path
) -> None:
    rows = [
        ["T", "D"],
        [dt.time(6), dt.datetime(2024, 1, 15)],
        [dt.time(12), dt.datetime(2024, 1, 15)],
    ]
    path = book(files, rows)
    await make(call, path, rows, row_fields=["T"], value_fields=[{"field": "D"}])
    sheet = load_workbook(files / path)["Out"]
    assert sheet["B2"].value == sheet["B3"].value == dt.datetime(2024, 1, 15)
    assert sheet["A2"].value == dt.time(6) and sheet["A2"].number_format == "h:mm:ss"


async def test_dates_in_a_field_of_other_types_show_in_the_host_date_format(
    call: ToolCall, files: Path
) -> None:
    rows = [["K", "V"], ["a", 1], [dt.datetime(2024, 1, 15), 2], [dt.datetime(2024, 1, 15, 9), 3]]
    path = book(files, rows)
    await make(call, path, rows, row_fields=["K"], value_fields=[{"field": "V"}])
    sheet = load_workbook(files / path)["Out"]
    assert sheet["A2"].value == "a"
    assert (sheet["A3"].number_format, sheet["A4"].number_format) == ("mm-dd-yy", "m/d/yy h:mm")


# Formulas


async def test_formulas_use_the_results_excel_stored(call: ToolCall, files: Path) -> None:
    shutil.copy(FIXTURES / "excel_pivot_source.xlsx", files / "f.xlsx")
    await call(
        "create_pivot_table",
        path="f.xlsx",
        source="Data!A1:C6",
        sheet="Report",
        at="A1",
        row_fields=["Key"],
        value_fields=[{"field": "Val"}, {"field": "Val", "function": "count"}],
    )
    sheet = load_workbook(files / "f.xlsx")["Report"]
    assert [[c.value for c in row] for row in sheet["A2:C6"]] == [
        ["Key", "Sum of Val", "Count of Val"],
        ["k1", NA, 2],
        ["k2", 14, 2],
        [dt.datetime(2024, 1, 15), 9, 1],
        ["Grand Total", NA, NA],
    ]


async def test_formulas_without_stored_results_are_calculated(call: ToolCall, files: Path) -> None:
    rows = [["K", "V"], ['="a"&"1"', "=1+1"], ["a1", "=2*3"], ["b", "=1/0"], ["=TRUE", 5]]
    path = book(files, rows)
    await make(call, path, rows, row_fields=["K"], value_fields=[{"field": "V"}])
    assert grid(files, path, "A2:B5") == [
        ["a1", 8],
        ["b", "#DIV/0!"],
        ["TRUE", 5],
        ["Grand Total", "#DIV/0!"],
    ]


async def test_a_formula_that_cannot_be_calculated_is_reported(
    call_error: ToolCall, files: Path
) -> None:
    rows = [["K", "V"], ["a", "=RAND()"]]
    path = book(files, rows)
    message = await call_error(
        "create_pivot_table",
        path=path,
        source="Data!A1:B2",
        sheet="Out",
        at="A1",
        row_fields=["K"],
        value_fields=[{"field": "V", "function": "count"}],
    )
    assert "B2" in message and "Excel recalculates" in message


# Grouping, slicers and timelines


@pytest.mark.parametrize(
    ("setting", "rows"),
    [
        ({"group_numbers": {"by": 5}}, [["X"], [1], ["a"], [3]]),
        ({"group_numbers": {"by": 5}}, [["X"], [1], [True], [3]]),
        ({"group_dates": ["years"]}, [["X"], [dt.datetime(2024, 1, 1)], ["a"]]),
        ({"group_dates": ["years"]}, [["X"], [dt.datetime(2024, 1, 1)], [5]]),
        ({"group_dates": ["years"]}, [["X"], [dt.time(6)], [dt.time(7)]]),
    ],
)
async def test_mixed_fields_cannot_be_grouped_as_in_excel(
    call_error: ToolCall, files: Path, setting: dict[str, Any], rows: list[list[Any]]
) -> None:
    path = book(files, [[*row, 1] if row[0] != "X" else [*row, "V"] for row in rows])
    message = await call_error(
        "create_pivot_table",
        path=path,
        source=f"Data!A1:B{len(rows)}",
        sheet="Out",
        at="A1",
        row_fields=["X"],
        value_fields=[{"field": "V"}],
        field_settings=[{"field": "X", **setting}],
    )
    assert "cannot be grouped" in message


async def test_a_slicer_on_a_mixed_field_lists_and_selects_its_items(
    call: ToolCall, files: Path
) -> None:
    rows = [
        ["K", "G", "V"],
        ["a", "x", 1],
        [True, "x", 2],
        [NA, "y", 3],
        [dt.datetime(2024, 1, 15), "y", 4],
        [None, "x", 5],
        [3, "y", 6],
    ]
    path = book(files, rows)
    await make(call, path, rows, row_fields=["G"], value_fields=[{"field": "V"}], name="P")
    await call(
        "add_slicer",
        path=path,
        sheet="Out",
        target={"sheet": "Out", "name": "P"},
        field="K",
        at="E3",
        selected_items=["a", "TRUE", "(blank)", "2024-01-15"],
    )
    assert grid(files, path, "A2:B4") == [["x", 8], ["y", 4], ["Grand Total", 12]]
    cache = part(files, path, "xl/pivotCache/pivotCacheDefinition1.xml")
    assert '<b v="1" />' in cache and '<e v="#N/A" />' in cache
    definition = part(files, path, "xl/pivotTables/pivotTable1.xml")
    assert definition.count(' h="1"') == 2
    listed = await call("describe_sheet", path=path, sheet="Out")
    assert "K" in str(listed)


async def test_slicers_over_formula_sources_recalculate_the_pivot_table(
    call: ToolCall, files: Path
) -> None:
    rows = [
        ["K", "F", "V"],
        ['="a"&"1"', "=1>2", "=1+1"],
        ["a1", "=1>2", "=2*3"],
        ["b", "=TRUE", 4],
    ]
    path = book(files, rows)
    await make(call, path, rows, row_fields=["F"], value_fields=[{"field": "V"}], name="P")
    await call(
        "add_slicer",
        path=path,
        sheet="Out",
        target={"sheet": "Out", "name": "P"},
        field="K",
        at="E3",
        selected_items=["a1"],
    )
    assert grid(files, path, "A2:B3") == [["FALSE", 8], ["Grand Total", 8]]


async def test_a_timeline_needs_a_field_of_dates_only(
    call_error: ToolCall, call: ToolCall, files: Path
) -> None:
    rows = [
        ["D", "M", "V"],
        [dt.datetime(2024, 1, 15), dt.datetime(2024, 1, 1), 1],
        [dt.datetime(2024, 2, 20), "a", 2],
    ]
    path = book(files, rows)
    await make(call, path, rows, row_fields=["D"], value_fields=[{"field": "V"}], name="P")
    target = {"sheet": "Out", "name": "P"}
    message = await call_error(
        "add_slicer",
        path=path,
        sheet="Out",
        target=target,
        field="M",
        at="E3",
        timeline={"level": "months"},
    )
    assert "not a date field" in message
    await call(
        "add_slicer",
        path=path,
        sheet="Out",
        target=target,
        field="D",
        at="E3",
        timeline={"level": "months", "start": "2024-01-01", "end": "2024-01-31"},
    )
    assert grid(files, path, "A2:B3") == [[dt.datetime(2024, 1, 15), 1], ["Grand Total", 1]]
