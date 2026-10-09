"""transform_range (remove duplicates, text to columns) and replace_cells."""

import datetime as dt
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

BOOK = {"path": "t.xlsx", "sheet": "Data"}
RED = PatternFill("solid", fgColor="FFFF0000")


@pytest.fixture
def book(files: Path) -> Path:
    workbook = Workbook()
    workbook.worksheets[0].title = "Data"
    workbook.create_sheet("Other")
    path = files / "t.xlsx"
    workbook.save(path)
    return path


def data(path: Path, name: str = "Data") -> Worksheet:
    return load_workbook(path)[name]


async def transform(call: ToolCall, range: str, **spec: Any) -> dict[str, str]:
    return await call("transform_range", **BOOK, range=range, transform=spec)


# -- remove duplicates -----------------------------------------------------------------------


async def test_remove_duplicates_matches_excel(call: ToolCall, book: Path) -> None:
    # The same data went through Excel's Remove Duplicates on column A with a header.
    rows = [["h", "n"], ["A", 1], ["a", 2], ["A ", 3], [1, 4], ["1", 5], ["A", 6]]
    rows += [[None, 7], [None, 8], [1.0, 9]]
    await call("write_range", **BOOK, at="A1", rows=rows)
    workbook = load_workbook(book)
    for row in range(2, 11):
        workbook["Data"].cell(row, 1).fill = RED
    workbook["Data"]["B3"].comment = Comment("moves with its row", "x")
    workbook.save(book)

    message = await transform(
        call, "A1:B10", operation="remove_duplicates", columns=["h"], has_header=True
    )
    assert message["note"] == "Removed 4 duplicate rows."
    sheet = data(book)
    assert [[sheet.cell(r, c).value for c in (1, 2)] for r in range(1, 11)] == [
        ["h", "n"],
        ["A", 1],
        ["A ", 3],
        [1, 4],
        ["1", 5],
        [None, 7],
        [None, None],
        [None, None],
        [None, None],
        [None, None],
    ]
    assert sheet["A2"].fill.start_color.rgb == "FFFF0000"
    assert sheet["A6"].fill.start_color.rgb == "FFFF0000"
    assert sheet["A7"].fill.fill_type is None
    assert sheet["B7"].comment is None


async def test_remove_duplicates_compares_only_the_chosen_columns(
    call: ToolCall, book: Path
) -> None:
    rows = [
        ["Region", "Product", "Units"],
        ["N", "x", 1],
        ["N", "y", 1],
        ["S", "x", 1],
        ["N", "x", 2],
    ]
    await call("write_range", **BOOK, at="A1", rows=rows)
    message = await transform(call, "A1:C5", operation="remove_duplicates", columns=["Region", "B"])
    assert message["note"] == "Removed 1 duplicate rows."
    sheet = data(book)
    assert [[c.value for c in row] for row in sheet["A1:C5"]] == [
        ["Region", "Product", "Units"],
        ["N", "x", 1],
        ["N", "y", 1],
        ["S", "x", 1],
        [None, None, None],
    ]


async def test_remove_duplicates_defaults_to_every_column(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, at="A1", rows=[["a", 1], ["a", 1], ["a", 2]])
    message = await transform(call, "A1:B3", operation="remove_duplicates", has_header=False)
    assert message["note"] == "Removed 1 duplicate rows."
    assert [[c.value for c in row] for row in data(book)["A1:B3"]] == [
        ["a", 1],
        ["a", 2],
        [None, None],
    ]


async def test_remove_duplicates_without_a_header_and_with_formulas(
    call: ToolCall, book: Path
) -> None:
    await call(
        "write_range",
        **BOOK,
        at="A1",
        rows=[[2, "=A1*2"], [1, "=A2*2"], [4, "=A3*2"], [1, "=A4*2"], [2, "=A5*2"]],
    )
    message = await transform(
        call, "A1:B5", operation="remove_duplicates", columns=["B"], has_header=False
    )
    assert message["note"] == "Removed 2 duplicate rows."
    sheet = data(book)
    # The results are 4, 2, 8, 2, 4: the last two repeat earlier ones.
    assert [sheet.cell(r, 2).value for r in range(1, 6)] == ["=A1*2", "=A2*2", "=A3*2", None, None]
    assert [sheet.cell(r, 1).value for r in range(1, 6)] == [2, 1, 4, None, None]


async def test_remove_duplicates_errors(call: ToolCall, call_error: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, at="A1", rows=[["a"], ["a"]])
    await call("merge_cells", **BOOK, range="C1:D1")
    spec = {"operation": "remove_duplicates"}
    assert "merged" in await call_error("transform_range", **BOOK, range="A1:D2", transform=spec)
    assert "Column 'zz' is not in" in await call_error(
        "transform_range", **BOOK, range="A1:A2", transform={**spec, "columns": ["zz"]}
    )
    assert "does not use delimiters" in await call_error(
        "transform_range", **BOOK, range="A1:A2", transform={**spec, "delimiters": [","]}
    )


# -- text to columns -------------------------------------------------------------------------


async def test_text_to_columns_types_pieces_like_excel(call: ToolCall, book: Path) -> None:
    text = "a|007|5%|TRUE|=1+1|2026-01-31| 12 |1.5|1,234|x y|-4"
    await call("write_range", **BOOK, at="A1", rows=[[text], ["no delimiter"], [42]])
    message = await transform(call, "A1:A3", operation="text_to_columns", delimiters=["|"])
    assert message["note"] == "Split 2 cells."
    sheet = data(book)
    values = [sheet.cell(1, col).value for col in range(1, 12)]
    assert values == [
        "a",
        7,
        0.05,
        True,
        "=1+1",
        dt.datetime(2026, 1, 31),
        12,
        1.5,
        1234,
        "x y",
        -4,
    ]
    assert sheet["E1"].data_type == "f"
    assert sheet["C1"].number_format == "0%"
    assert sheet["I1"].number_format == "#,##0"
    assert sheet["F1"].number_format == "yyyy-mm-dd"
    assert sheet["A2"].value == "no delimiter" and sheet["B2"].value is None
    assert sheet["A3"].value == 42


@pytest.mark.parametrize(
    ("text", "options", "expected"),
    [
        ('a,"b,c",d', {"delimiters": ["comma"]}, ["a", "b,c", "d"]),
        ('a,"b,c",d', {"delimiters": ["comma"], "text_qualifier": ""}, ["a", '"b', 'c"', "d"]),
        ('a,"say ""hi""",d', {"delimiters": ["comma"]}, ["a", 'say "hi"', "d"]),
        ("a,'b,c'", {"delimiters": ["comma"], "text_qualifier": "'"}, ["a", "b,c"]),
        ("a,,b", {"delimiters": ["comma"]}, ["a", None, "b"]),
        ("a,,b", {"delimiters": ["comma"], "merge_delimiters": True}, ["a", "b"]),
        ("a b;c\td", {"delimiters": ["space", "semicolon", "tab"]}, ["a", "b", "c", "d"]),
        ("a-b", {"delimiters": ["-"]}, ["a", "b"]),
        ("AB12XYZ", {"fixed_widths_chars": [2, 2]}, ["AB", 12, "XYZ"]),
        ("AB", {"fixed_widths_chars": [5]}, ["AB"]),
    ],
)
async def test_text_to_columns_splitting(
    call: ToolCall, book: Path, text: str, options: dict[str, Any], expected: list[object]
) -> None:
    await call("write_range", **BOOK, at="A1", rows=[[text]])
    await transform(call, "A1", operation="text_to_columns", **options)
    sheet = data(book)
    assert [sheet.cell(1, col).value for col in range(1, len(expected) + 1)] == expected


async def test_text_to_columns_keeps_text_cells_as_text(call: ToolCall, book: Path) -> None:
    workbook = load_workbook(book)
    workbook["Data"]["A1"] = "007,1"
    workbook["Data"]["B1"].number_format = "@"
    workbook.save(book)
    await transform(call, "A1", operation="text_to_columns", delimiters=[","])
    sheet = data(book)
    assert (sheet["A1"].value, sheet["B1"].value) == (7, "1")


@pytest.mark.parametrize(
    ("range", "spec", "message"),
    [
        ("A1:B2", {"delimiters": [","]}, "single column"),
        ("A1", {}, "either delimiters or fixed_widths_chars"),
        (
            "A1",
            {"delimiters": [","], "fixed_widths_chars": [1]},
            "either delimiters or fixed_widths_chars",
        ),
        ("A1", {"delimiters": [",,"]}, "single character"),
        ("A1:A2", {"delimiters": [","]}, "B1 is not empty"),
        ("A2", {"delimiters": [","]}, "WEBSERVICE is not allowed"),
        ("A1", {"delimiters": [","], "direction": "down"}, "does not use direction"),
    ],
)
async def test_invalid_text_to_columns(
    call: ToolCall, call_error: ToolCall, book: Path, range: str, spec: dict[str, Any], message: str
) -> None:
    await call(
        "write_range",
        **BOOK,
        at="A1",
        rows=[["a,b"], ['x,=WEBSERVICE("https://example.test")']],
    )
    await call("write_range", **BOOK, at="B1", rows=[["in the way"]])
    before = book.read_bytes()
    result = await call_error(
        "transform_range",
        **BOOK,
        range=range,
        transform={"operation": "text_to_columns", **spec},
    )
    assert message in result
    assert book.read_bytes() == before


# -- replace ---------------------------------------------------------------------------------


async def replace(call: ToolCall, query: str, replacement: str, **options: Any) -> dict[str, int]:
    result = await call(
        "replace_cells", path="t.xlsx", query=query, replacement=replacement, **options
    )
    return result["replaced"]


@pytest.fixture
def texts(book: Path) -> Path:
    workbook = load_workbook(book)
    sheet = workbook["Data"]
    for row, value in enumerate(["a", 100, "abc", "Hello", "hello world", "HELLO"], start=1):
        sheet.cell(row, 1, value)
    sheet["B1"] = "=SUM(A2:A3)+SUM(A5:A6)"
    sheet["B2"] = dt.date(2026, 1, 31)
    sheet["B3"] = True
    sheet["C1"] = 1.5
    sheet["C1"].number_format = "0.00"
    workbook["Other"]["A1"] = "hello"
    workbook.save(book)
    return book


async def test_replace_retypes_results_like_excel(call: ToolCall, texts: Path) -> None:
    # Checked against Excel's Range.Replace: "a" -> "1" gives a number, 100 -> 155.
    assert await replace(call, "a", "1", sheet="Data", exact=True) == {"Data": 1}
    assert await replace(call, "0", "5", sheet="Data") == {"Data": 1}
    sheet = data(texts)
    assert (sheet["A1"].value, sheet["A2"].value) == (1, 155)
    assert await replace(call, "b", "=1+1", sheet="Data") == {"Data": 1}
    assert data(texts)["A3"].value == "a=1+1c"


async def test_replace_matching_options(call: ToolCall, texts: Path) -> None:
    assert await replace(call, "hello", "bye", sheet="Data", exact=True) == {"Data": 2}
    assert [data(texts).cell(r, 1).value for r in (4, 5, 6)] == ["bye", "hello world", "bye"]
    assert await replace(call, "HELLO", "x", sheet="Data", case_sensitive=True) == {}
    assert await replace(call, "hello", "x", sheet="Data", case_sensitive=True) == {"Data": 1}
    assert data(texts)["A5"].value == "x world"


async def test_replace_covers_every_sheet_by_default(call: ToolCall, texts: Path) -> None:
    assert await replace(call, "hello", "bye") == {"Data": 3, "Other": 1}
    assert data(texts, "Other")["A1"].value == "bye"


async def test_replace_changes_formula_text_but_not_dates_booleans_or_formats(
    call: ToolCall, texts: Path
) -> None:
    assert await replace(call, "SUM", "AVERAGE", sheet="Data") == {"Data": 1}
    assert await replace(call, "A5:A6", "A7:A8", sheet="Data") == {"Data": 1}
    sheet = data(texts)
    assert sheet["B1"].value == "=AVERAGE(A2:A3)+AVERAGE(A7:A8)"
    assert sheet["B1"].data_type == "f"
    assert await replace(call, "1", "9", sheet="Data") == {"Data": 2}
    sheet = data(texts)
    assert sheet["B2"].value == dt.datetime(2026, 1, 31)
    assert sheet["B3"].value is True
    assert (sheet["A2"].value, sheet["C1"].value) == (900, 9.5)
    assert sheet["C1"].number_format == "0.00"


async def test_replace_in_formulas_can_be_switched_off(call: ToolCall, texts: Path) -> None:
    assert await replace(call, "SUM", "AVERAGE", sheet="Data", in_formulas=False) == {}
    assert data(texts)["B1"].value == "=SUM(A2:A3)+SUM(A5:A6)"


async def test_replace_keeps_function_prefixes(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, at="A1", rows=[["=IFS(C1>2,1,TRUE,0)"]])
    assert await replace(call, "C1", "D1", sheet="Data") == {"Data": 1}
    assert data(book)["A1"].value == "=_xlfn.IFS(D1>2,1,TRUE,0)"
    assert data(book)["A1"].data_type == "f"


async def test_replaced_formulas_go_through_the_formula_gate(
    call_error: ToolCall, texts: Path
) -> None:
    before = texts.read_bytes()
    message = await call_error(
        "replace_cells",
        path="t.xlsx",
        query="SUM",
        replacement='WEBSERVICE("https://example.test/"&A1)',
        sheet="Data",
    )
    assert "WEBSERVICE is not allowed" in message
    message = await call_error(
        "replace_cells", path="t.xlsx", query="abc", replacement='=INDIRECT("A1")', sheet="Data"
    )
    assert "INDIRECT is not allowed" in message
    assert texts.read_bytes() == before


async def test_replace_into_a_text_cell_stays_text(call: ToolCall, book: Path) -> None:
    workbook = load_workbook(book)
    workbook["Data"]["A1"] = "x1"
    workbook["Data"]["A1"].number_format = "@"
    workbook.save(book)
    await replace(call, "x", "00", sheet="Data")
    assert data(book)["A1"].value == "001"


async def test_replace_with_nothing_found_leaves_the_file_alone(
    call: ToolCall, texts: Path
) -> None:
    assert await replace(call, "zzz", "y") == {}


async def test_replace_rejects_unknown_sheets_and_empty_queries(
    call_error: ToolCall, texts: Path
) -> None:
    assert "not found" in await call_error(
        "replace_cells", path="t.xlsx", query="a", replacement="b", sheet="Nope"
    )
    assert await call_error("replace_cells", path="t.xlsx", query="", replacement="b")


@pytest.mark.parametrize(
    ("tool", "arguments"),
    [
        ("replace_cells", {"query": "a", "replacement": "b"}),
        (
            "transform_range",
            {"sheet": "Data", "range": "A1:A2", "transform": {"operation": "remove_duplicates"}},
        ),
    ],
)
async def test_new_tools_stay_inside_the_allowed_folder(
    call_error: ToolCall, tool: str, arguments: dict[str, Any]
) -> None:
    assert "outside" in await call_error(tool, path="../secret.xlsx", **arguments)
