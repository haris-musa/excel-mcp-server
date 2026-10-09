"""The calculator stays within time and memory limits on hostile formulas."""

import re
import time
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table

from excel_mcp.calc.parser import MAX_NESTING, parse
from excel_mcp.calc.values import UncalculableError
from excel_mcp.calc.wildcard import Wildcard
from tests.conftest import ToolCall
from tests.test_calc import make_workbook, read_values

pytestmark = pytest.mark.anyio


def _regex(pattern: str) -> re.Pattern[str]:
    """The reference meaning of a wildcard pattern, as a regular expression."""
    out, index = [], 0
    while index < len(pattern):
        character = pattern[index]
        if character == "~" and index + 1 < len(pattern):
            index += 1
            out.append(re.escape(pattern[index]))
        else:
            out.append({"*": ".*", "?": "."}.get(character) or re.escape(character))
        index += 1
    return re.compile("".join(out), re.IGNORECASE | re.DOTALL)


@pytest.mark.parametrize("pattern", ["a*b", "*a?b*", "~*a", "a~", "?", "*", "", "A*B*C", "a*a*a"])
@pytest.mark.parametrize(
    "text", ["", "a", "ab", "AXB", "*a", "a~", "aab", "abcabc", "xaxbx", "aaa"]
)
def test_wildcards_agree_with_the_regular_expression_meaning(pattern: str, text: str) -> None:
    expected = _regex(pattern)
    assert Wildcard(pattern).fullmatch(text) == bool(expected.fullmatch(text))
    found = Wildcard(pattern).search(text, 0)
    match = expected.search(text)
    assert found == (match.start() if match else None)


def test_wildcards_match_in_linear_time() -> None:
    started = time.perf_counter()
    assert not Wildcard("a*" * 40 + "b").fullmatch("a" * 20_000)
    assert Wildcard("*" * 40 + "b").search("a" * 20_000, 0) is None
    assert time.perf_counter() - started < 2


async def test_wildcard_criteria_cannot_stall_a_read(call: ToolCall, files: Path) -> None:
    pattern = "a*" * 40 + "b"
    make_workbook(files / "calc.xlsx", {"A1": "a" * 5000, "B1": f'=COUNTIF(A1,"{pattern}")'})
    started = time.perf_counter()

    data = await read_values(call, "B1")

    assert data["values"] == [[0]]
    assert time.perf_counter() - started < 5


def test_deeply_nested_formulas_are_refused_when_parsed() -> None:
    assert parse("=" + "(" * MAX_NESTING + "1" + ")" * MAX_NESTING)
    for formula in ("=" + "(" * 200 + "1" + ")" * 200, "=" + "ABS(" * 100 + "1" + ")" * 100):
        with pytest.raises(UncalculableError, match="nested"):
            parse(formula)
    with pytest.raises(UncalculableError, match="nested"):
        parse("=" + "-" * 5000 + "1")
    with pytest.raises(UncalculableError, match="too long"):
        parse("=" + "1+" * 5000 + "1")


@pytest.mark.parametrize(
    "formula",
    [
        "=" + "(" * 3000 + "1" + ")" * 3000,
        "=" + "1+" * 2000 + "1",
        '=REPT("x",1000000000)',
        '=LEN(REPT("abcdefgh",5000))',
        '=LEN(SUBSTITUTE(REPT("a",30000),"a",REPT("b",30000)))',
        "=SEQUENCE(100000,100000)",
        "=SUM(SEQUENCE(100000)*SEQUENCE(1,100000))",
        "=SUM(VSTACK(SEQUENCE(100000),SEQUENCE(100000)))",
        "=COMBIN(1000000000,500000000)",
        "=PERMUT(1000000000,1000000000)",
        "=CUMIPMT(0.05,1000000000,100,1,999999999,0)",
        "=DDB(1000,1,1000000000,500000000)",
        "=UNICHAR(5000000)",
        "=DATE(99999999,1,1)",
        "=EDATE(1,1000000000)",
        "=EOMONTH(1,-1000000000)",
        "=WORKDAY(1,100000000)",
        "=2^1024",
        "=10^400*10^400",
        "=FACT(1E300)",
        "=ROUND(1E300,-1E9)",
        '=TEXT(1,REPT("0",30000))',
        '=REPT("x",30000)&REPT("y",30000)',
    ],
)
async def test_hostile_formulas_end_as_values_or_uncalculated(
    call: ToolCall, files: Path, formula: str
) -> None:
    make_workbook(files / "calc.xlsx", {"A1": formula})
    started = time.perf_counter()

    data = await read_values(call, "A1")

    assert time.perf_counter() - started < 10
    assert data["values"] in ([[]], [], [[None]]) or isinstance(
        data["values"][0][0], int | float | str
    )


async def test_text_longer_than_a_cell_is_a_value_error(call: ToolCall, files: Path) -> None:
    make_workbook(
        files / "calc.xlsx",
        {
            "A1": '=REPT("x",32768)',
            "A2": '=LEN(REPT("x",32767))',
            "A3": '=REPT("x",30000)&"yyyyyyyy"',
        },
    )

    data = await read_values(call, "A1:A3")

    assert data["values"] == [["#VALUE!"], [32767], ["x" * 30000 + "yyyyyyyy"]]


async def test_tables_spills_and_indirect_stay_within_limits(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Units", "Price"])
    for number in range(1, 50):
        sheet.append([number, number * 2])
    sheet.add_table(Table(displayName="Sales", ref="A1:B50"))
    sheet["D1"] = ArrayFormula("D1", "=_xlfn.SEQUENCE(1000000)")
    sheet["D2"] = ArrayFormula("D2", "=_xlfn.ANCHORARRAY(D2)")
    formulas = [
        '=SUM(INDIRECT("A1:XFD1048576"))',
        '=SUM(INDIRECT("R1C1:R1048576C16384",FALSE))',
        '=SUM(INDIRECT("1:1048576"))',
        '=SUM(INDIRECT("Sales[Units]"))',
        "=SUM(_xlfn.ANCHORARRAY(D1))",
        "=SUM(_xlfn.ANCHORARRAY(D2))",
        "=SUM(Sales[Units]*Sales[Price]*Sales[Units])",
    ]
    for number, formula in enumerate(formulas, start=1):
        sheet.cell(row=number, column=6, value=formula)
    workbook.save(files / "calc.xlsx")
    started = time.perf_counter()

    data = await read_values(call, "F1:F7")

    assert time.perf_counter() - started < 10
    assert data["values"][3] == [1225]
    assert set(data["uncalculated"]) >= {"F1", "F2", "F3", "F5", "F6"}
