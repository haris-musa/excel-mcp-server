"""Table options: totals rows, calculated columns, header, banding, filter buttons, resizing.

The expected XML is what Microsoft Excel saved when the same steps were done in it."""

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from excel_mcp.structured import parse_structured, qualify_formula
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

SHEET = {"path": "sales.xlsx", "sheet": "Data"}
REVENUE = {"name": "Revenue", "formula": "=[@Units]*[@Price]"}


def _table(path: Path, name: str = "Sales") -> Any:
    return load_workbook(path)["Data"].tables[name]


def _cells(path: Path, *cells: str) -> list[Any]:
    sheet = load_workbook(path)["Data"]
    return [sheet[cell].value for cell in cells]


async def _create(call: ToolCall, **options: Any) -> None:
    await call("create_table", **SHEET, range="A1:D5", name="Sales", options=options)


async def test_totals_row_is_what_excel_writes(call: ToolCall, sample: Path) -> None:
    await _create(
        call,
        totals_row=True,
        columns=[
            {"name": "Units", "total": "sum"},
            {"name": "Price", "total": "average"},
            {"name": "Product", "total": "=COUNTA([Product])+1"},
        ],
    )
    table = _table(sample)
    assert (table.ref, table.totalsRowCount, table.autoFilter.ref) == ("A1:D6", 1, "A1:D5")
    columns = {c.name: c for c in table.tableColumns}
    assert columns["Region"].totalsRowLabel == "Total"
    assert columns["Units"].totalsRowFunction == "sum"
    assert columns["Price"].totalsRowFunction == "average"
    assert columns["Product"].totalsRowFunction == "custom"
    assert columns["Product"].totalsRowFormula.attr_text == "COUNTA(Sales[Product])+1"
    assert _cells(sample, "A6", "B6", "C6", "D6") == [
        "Total",
        "=COUNTA(Sales[Product])+1",
        "=SUBTOTAL(109,Sales[Units])",
        "=SUBTOTAL(101,Sales[Price])",
    ]


@pytest.mark.parametrize(
    ("function", "stored", "number"),
    [
        ("sum", "sum", 109),
        ("average", "average", 101),
        ("count", "count", 103),
        ("count_numbers", "countNums", 102),
        ("max", "max", 104),
        ("min", "min", 105),
        ("std_dev", "stdDev", 107),
        ("var", "var", 110),
    ],
)
async def test_total_functions(
    call: ToolCall, sample: Path, function: str, stored: str, number: int
) -> None:
    await _create(call, totals_row=True, columns=[{"name": "Units", "total": function}])
    column = next(c for c in _table(sample).tableColumns if c.name == "Units")
    assert column.totalsRowFunction == stored
    assert _cells(sample, "C6") == [f"=SUBTOTAL({number},Sales[Units])"]


async def test_totals_can_be_changed_labelled_removed_and_the_row_turned_off(
    call: ToolCall, sample: Path
) -> None:
    await _create(call, totals_row=True, columns=[{"name": "Units", "total": "sum"}])
    await call(
        "edit_table",
        **SHEET,
        table="Sales",
        options={
            "columns": [
                {"name": "Units", "total": "none"},
                {"name": "Price", "total_label": "avg"},
                {"name": "Region", "total": "count"},
            ]
        },
    )
    columns = {c.name: c for c in _table(sample).tableColumns}
    assert columns["Units"].totalsRowFunction is None
    assert columns["Price"].totalsRowLabel == "avg"
    assert columns["Region"].totalsRowLabel is None
    assert _cells(sample, "A6", "C6", "D6") == ["=SUBTOTAL(103,Sales[Region])", None, "avg"]
    await call("edit_table", **SHEET, table="Sales", options={"totals_row": False})
    table = _table(sample)
    assert (table.ref, table.totalsRowCount) == ("A1:D5", None)
    assert all(c.totalsRowFunction is None and c.totalsRowLabel is None for c in table.tableColumns)
    assert _cells(sample, "A6", "D6") == [None, None]


async def test_the_totals_row_needs_empty_cells_and_totals_need_the_row(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await call("write_range", **SHEET, start_cell="B6", rows=[["x"]])
    await _create(call)
    assert "B6 is not empty" in await call_error(
        "edit_table", **SHEET, table="Sales", options={"totals_row": True}
    )
    assert "totals_row" in await call_error(
        "edit_table",
        **SHEET,
        table="Sales",
        options={"columns": [{"name": "Units", "total": "sum"}]},
    )
    assert "Columns: Region" in await call_error(
        "edit_table",
        **SHEET,
        table="Sales",
        options={"columns": [{"name": "Nope", "formula": "=1"}]},
    )


async def test_a_calculated_column_is_what_excel_writes(call: ToolCall, sample: Path) -> None:
    await call("write_range", **SHEET, start_cell="E1", rows=[["Revenue"]])
    await call("create_table", **SHEET, range="A1:E5", name="Sales", options={"columns": [REVENUE]})
    column = _table(sample).tableColumns[4]
    assert column.calculatedColumnFormula.attr_text == (
        "Sales[[#This Row],[Units]]*Sales[[#This Row],[Price]]"
    )
    formula = "=Sales[[#This Row],[Units]]*Sales[[#This Row],[Price]]"
    assert _cells(sample, "E2", "E5") == [formula, formula]
    read = await call("read_range", **SHEET, range="E2:E5")
    assert [row[0] for row in read["values"]] == [15.0, 7.5, 14.0, 6.0]


async def test_calculated_columns_fill_new_rows(call: ToolCall, sample: Path) -> None:
    await call("write_range", **SHEET, start_cell="E1", rows=[["Revenue"]])
    await call("create_table", **SHEET, range="A1:E5", name="Sales", options={"columns": [REVENUE]})
    await call("insert_rows_or_columns", **SHEET, axis="rows", at=3, count=2)
    formula = "=Sales[[#This Row],[Units]]*Sales[[#This Row],[Price]]"
    assert _cells(sample, "E3", "E4", "E7") == [formula] * 3
    await call("edit_table", **SHEET, table="Sales", options={}, range="A1:E9")
    assert _cells(sample, "E8", "E9") == [formula] * 2
    assert _table(sample).ref == "A1:E9"


async def test_relative_references_in_a_calculated_column_fill_down(
    call: ToolCall, sample: Path
) -> None:
    await call("write_range", **SHEET, start_cell="E1", rows=[["Revenue"]])
    await call(
        "create_table",
        **SHEET,
        range="A1:E5",
        name="Sales",
        options={"columns": [{"name": "Revenue", "formula": "=C2*D2"}]},
    )
    assert _table(sample).tableColumns[4].calculatedColumnFormula.attr_text == "C2*D2"
    assert _cells(sample, "E2", "E5") == ["=C2*D2", "=C5*D5"]
    await call("edit_table", **SHEET, table="Sales", options={}, range="A1:E7")
    assert _cells(sample, "E6", "E7") == ["=C6*D6", "=C7*D7"]
    await call("insert_rows_or_columns", **SHEET, axis="rows", at=3, count=1)
    assert _cells(sample, "E3") == ["=C3*D3"]


async def test_a_new_totals_row_is_what_excel_makes(call: ToolCall, sample: Path) -> None:
    await _create(call, totals_row=True)
    table = _table(sample)
    columns = table.tableColumns
    assert (columns[0].totalsRowLabel, columns[3].totalsRowFunction) == ("Total", "sum")
    assert _cells(sample, "A6", "B6", "D6") == ["Total", None, "=SUBTOTAL(109,Sales[Price])"]
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="A1",
        rows=[["a", "b"], [1, "x"], [2, "y"]],
    )
    await call(
        "create_table",
        path="sales.xlsx",
        sheet="Report",
        range="A1:B3",
        name="Small",
        options={"totals_row": True},
    )
    report = load_workbook(sample)["Report"]
    assert [report["A4"].value, report["B4"].value] == ["Total", "=SUBTOTAL(103,Small[b])"]
    await call(
        "write_range", path="sales.xlsx", sheet="Report", start_cell="D1", rows=[["n"], [1], [2]]
    )
    await call(
        "create_table",
        path="sales.xlsx",
        sheet="Report",
        range="D1:D3",
        name="One",
        options={"totals_row": True},
    )
    report = load_workbook(sample)["Report"]
    assert report["D4"].value == "=SUBTOTAL(109,One[n])"


async def test_resizing_moves_the_totals_row_like_excel(call: ToolCall, sample: Path) -> None:
    await _create(call, totals_row=True)
    await call("edit_table", **SHEET, table="Sales", options={}, range="A1:D8")
    table = _table(sample)
    assert (table.ref, table.autoFilter.ref) == ("A1:D8", "A1:D7")
    assert _cells(sample, "A6", "D6", "A8", "D8") == [
        None,
        None,
        "Total",
        "=SUBTOTAL(109,Sales[Price])",
    ]
    await call("edit_table", **SHEET, table="Sales", options={}, range="A1:D4")
    assert _table(sample).ref == "A1:D5"  # the range given is the data's
    assert _cells(sample, "A5", "D5") == ["Total", "=SUBTOTAL(109,Sales[Price])"]


async def test_cut_off_rows_go_below_the_totals_row(call: ToolCall, sample: Path) -> None:
    await _create(call, totals_row=True)
    await call("edit_table", **SHEET, table="Sales", options={}, range="A1:D4")
    assert _table(sample).ref == "A1:D5"
    assert _cells(sample, "A4", "A5", "A6", "C6") == ["North", "Total", "South", 3]


async def test_formulas_pass_the_safety_check(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_table",
        **SHEET,
        range="A1:D5",
        options={"columns": [{"name": "Price", "formula": '=WEBSERVICE("http://x")'}]},
    )
    assert "WEBSERVICE" in message


async def test_options_are_what_excel_writes(call: ToolCall, sample: Path) -> None:
    await _create(
        call,
        style="TableStyleMedium2",
        striped_rows=False,
        striped_columns=True,
        first_column=True,
        last_column=True,
        filter_button=False,
    )
    table = _table(sample)
    style = table.tableStyleInfo
    assert (style.name, style.showRowStripes, style.showColumnStripes) == (
        "TableStyleMedium2",
        False,
        True,
    )
    assert (style.showFirstColumn, style.showLastColumn) == (True, True)
    assert table.autoFilter is None
    await call("edit_table", **SHEET, table="sales", options={"filter_button": True})
    assert _table(sample).autoFilter.ref == "A1:D5"


async def test_header_row_off_and_on(call: ToolCall, sample: Path) -> None:
    await _create(call)
    await call("edit_table", **SHEET, table="Sales", options={"header_row": False})
    table = _table(sample)
    assert (table.ref, table.headerRowCount, table.autoFilter) == ("A2:D5", 0, None)
    assert _cells(sample, "A1", "D1", "A2") == [None, None, "North"]
    await call("edit_table", **SHEET, table="Sales", options={"header_row": True})
    table = _table(sample)
    assert (table.ref, table.headerRowCount, table.autoFilter.ref) == ("A1:D5", 1, "A1:D5")
    assert _cells(sample, "A1", "D1") == ["Region", "Price"]


async def test_resize(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    await _create(call)
    await call("write_range", **SHEET, start_cell="E1", rows=[["Extra"], [1]])
    await call("edit_table", **SHEET, table="Sales", options={}, range="A1:E7")
    table = _table(sample)
    assert (table.ref, [c.name for c in table.tableColumns][-1]) == ("A1:E7", "Extra")
    assert table.autoFilter.ref == "A1:E7"
    await call("edit_table", **SHEET, table="Sales", options={}, range="A1:B3")
    table = _table(sample)
    assert (table.ref, len(table.tableColumns), table.autoFilter.ref) == ("A1:B3", 2, "A1:B3")
    assert _cells(sample, "C1", "D5") == ["Units", 2]  # Excel leaves the cells as they are
    await call("edit_table", **SHEET, table="Sales", options={}, range="A1:D3")
    assert [c.name for c in _table(sample).tableColumns] == ["Region", "Product", "Units", "Price"]
    assert "top-left" in await call_error(
        "edit_table", **SHEET, table="Sales", options={}, range="B1:D3"
    )


async def test_a_new_column_is_named_like_excel(call: ToolCall, sample: Path) -> None:
    report = {"path": "sales.xlsx", "sheet": "Report"}
    await call("write_range", **report, start_cell="A1", rows=[["x", "y"], [1, 2]])
    await call("create_table", **report, range="A1:B2", name="Small", options={})
    await call("edit_table", **report, table="Small", options={}, range="A1:C2")
    sheet = load_workbook(sample)["Report"]
    assert sheet["C1"].value == "Column1"
    assert [c.name for c in sheet.tables["Small"].tableColumns] == ["x", "y", "Column1"]


async def test_resize_and_edit_errors(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    await _create(call, totals_row=True)
    assert "no table 'Nope'. Tables: Sales" in await call_error(
        "edit_table", **SHEET, table="Nope", options={}
    )
    assert "filter_button needs a header row" in await call_error(
        "edit_table",
        **SHEET,
        table="Sales",
        options={"header_row": False, "filter_button": True},
    )


async def test_the_calculator_reads_structured_references(call: ToolCall, sample: Path) -> None:
    await call("write_range", **SHEET, start_cell="E1", rows=[["Revenue"]])
    await call(
        "create_table",
        **SHEET,
        range="A1:E5",
        name="Sales",
        options={
            "totals_row": True,
            "columns": [REVENUE, {"name": "Revenue", "total": "sum"}],
        },
    )
    await call(
        "write_range",
        **SHEET,
        start_cell="G1",
        rows=[
            ["=Sales[[#Totals],[Revenue]]"],
            ["=SUM(Sales[Units])"],
            ["=Sales[[#Headers],[Price]]"],
        ],
    )
    read = await call("read_range", **SHEET, range="G1:G3")
    assert [row[0] for row in read["values"]] == [42.5, 25.0, "Price"]


async def test_edit_table_stays_inside_the_workbook_folder(
    call_error: ToolCall, sample: Path
) -> None:
    assert "outside" in await call_error(
        "edit_table", path="../outside.xlsx", sheet="Data", table="Sales", options={}
    )


@pytest.mark.parametrize(
    ("text", "parsed"),
    [
        ("Sales[Price]", ("Sales", (), "Price", None)),
        ("[@Price]", (None, ("#This Row",), "Price", None)),
        ("[@[Unit Price]]", (None, ("#This Row",), "Unit Price", None)),
        ("Sales[[#Totals],[A B]:[C]]", ("Sales", ("#Totals",), "A B", "C")),
        ("Sales[[#Headers],[#Data]]", ("Sales", ("#Headers", "#Data"), None, None)),
        ("Sales[#all]", ("Sales", ("#All",), None, None)),
        ("Sales[Col'[1']]", ("Sales", (), "Col[1]", None)),
    ],
)
def test_structured_references(text: str, parsed: tuple[Any, ...]) -> None:
    found = parse_structured(text)
    assert found is not None
    assert (found.table, found.items, found.first, found.last) == parsed


def test_references_to_a_table_are_stored_in_full() -> None:
    assert qualify_formula("=[@Price]*SUM([Qty])+A1+Other[@x]", "Sales") == (
        "=Sales[[#This Row],[Price]]*SUM(Sales[Qty])+A1+Other[[#This Row],[x]]"
    )
    assert parse_structured("Sales[Price") is None
    assert parse_structured("A1") is None
