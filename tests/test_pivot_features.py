"""PivotTable options: formats, show values as, sorting, grouping, calculated fields, layouts.

tests/test_pivot_golden.py checks the figures against Excel; these check the tool contract,
the definitions that are written, and the errors.
"""

import datetime as dt
import shutil
import zipfile
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from excel_mcp.operations.pivot_index import sheet_pivots
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

FIXTURES = Path(__file__).parent / "fixtures"
PIVOT = {
    "path": "sales.xlsx",
    "source": "Data!A1:D5",
    "sheet": "Report",
    "at": "A1",
}
UNITS = [{"field": "Units"}]


def grid(path: Path, ref: str, sheet: str = "Report") -> list[list[Any]]:
    return [[cell.value for cell in row] for row in load_workbook(path)[sheet][ref]]


def pivot_xml(path: Path, index: int = 1) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(f"xl/pivotTables/pivotTable{index}.xml").decode()


def cache_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("xl/pivotCache/pivotCacheDefinition1.xml").decode()


# Number formats and show values as


async def test_number_format_is_stored_and_shown(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "Price", "function": "average", "number_format": '0.0 "kg"'}],
    )
    sheet = load_workbook(sample)["Report"]
    assert sheet["B2"].number_format == '0.0 "kg"'
    field = sheet_pivots(sheet)[0].dataFields[0]
    assert field.numFmtId is not None and field.numFmtId >= 164


async def test_percent_of_total(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "Units", "show_as": "percent_of_total"}],
    )
    assert grid(sample, "A1:B4") == [
        ["Region", "Sum of Units"],
        ["North", 17 / 25],
        ["South", 8 / 25],
        ["Grand Total", 1],
    ]
    sheet = load_workbook(sample)["Report"]
    assert sheet["B2"].number_format == "0.00%"
    assert sheet_pivots(sheet)[0].dataFields[0].showDataAs == "percentOfTotal"


async def test_same_field_twice_gets_numbered_captions(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "Units"}, {"field": "Units", "show_as": "percent_of_total"}],
    )
    names = [field.name for field in sheet_pivots(load_workbook(sample)["Report"])[0].dataFields]
    assert names == ["Sum of Units", "Sum of Units2"]


async def test_difference_and_running_total_along_a_field(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        column_fields=["Product"],
        value_fields=[{"field": "Units", "show_as": "difference_from", "base_field": "Product"}],
    )
    assert grid(sample, "A3:D4") == [["North", None, -3, None], ["South", None, -2, None]]
    field = sheet_pivots(load_workbook(sample)["Report"])[0].dataFields[0]
    assert (field.showDataAs, field.baseField, field.baseItem) == ("difference", 1, 1048828)


async def test_rank_is_written_as_an_extension_and_survives_edits(
    call: ToolCall, sample: Path
) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "Units", "show_as": "rank_descending", "base_field": "Region"}],
    )
    assert grid(sample, "A2:B3") == [["North", 1], ["South", 2]]
    assert 'pivotShowAs="rankDescending"' in pivot_xml(sample)
    await call("write_range", path="sales.xlsx", sheet="Report", at="F1", rows=[[1]])
    assert 'pivotShowAs="rankDescending"' in pivot_xml(sample)


async def test_excel_authored_extensions_survive_edits(call: ToolCall, files: Path) -> None:
    shutil.copy(FIXTURES / "excel_pivot_ext.xlsx", files / "ext.xlsx")
    before = pivot_xml(files / "ext.xlsx")
    assert "hideValuesRow" in before and 'pivotShowAs="rankAscending"' in before
    await call("write_range", path="ext.xlsx", sheet="Report", at="F1", rows=[["x"]])
    after = pivot_xml(files / "ext.xlsx")
    assert "hideValuesRow" in after and 'pivotShowAs="rankAscending"' in after


@pytest.mark.parametrize(
    ("value", "message"),
    [
        ({"show_as": "running_total"}, "needs base_field"),
        ({"show_as": "percent_of_total", "base_field": "Region"}, "takes no base_field"),
        ({"show_as": "running_total", "base_field": "Price"}, "rows or columns"),
        ({"base_field": "Region"}, "need show_as"),
        (
            {"show_as": "running_total", "base_field": "Region", "base_item": "North"},
            "takes no base_item",
        ),
        (
            {"show_as": "difference_from", "base_field": "Region", "base_item": "Mars"},
            "not an item",
        ),
    ],
)
async def test_invalid_show_as(
    call_error: ToolCall, sample: Path, value: dict[str, str], message: str
) -> None:
    result = await call_error(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "Units", **value}],
    )
    assert message in result


async def test_rank_is_not_offered_with_values_in_rows(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        values_in="rows",
        value_fields=[
            {"field": "Units", "show_as": "rank_ascending", "base_field": "Region"},
            {"field": "Price"},
        ],
    )
    assert "values_in" in message


# Sorting


async def test_sort_by_label_descending(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=UNITS,
        field_settings=[{"field": "Region", "sort": "descending"}],
    )
    assert [row[0] for row in grid(sample, "A2:A4")] == ["South", "North", "Grand Total"]
    assert 'sortType="descending"' in pivot_xml(sample)


async def test_sort_by_value(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Product"],
        value_fields=[{"field": "Units"}, {"field": "Price", "function": "max"}],
        field_settings=[{"field": "Product", "sort": "descending", "sort_by": "Sum of Units"}],
    )
    assert [row[0] for row in grid(sample, "A3:A4")] == ["Apples", "Pears"]
    xml = pivot_xml(sample)
    assert "autoSortScope" in xml and 'sortType="descending"' in xml


async def test_sort_by_needs_a_values_field(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=UNITS,
        field_settings=[{"field": "Region", "sort_by": "Total"}],
    )
    assert "Sum of Units" in message


# Date and number groups


@pytest.fixture
def dated(files: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in [
        ["Day", "Qty", "Amount"],
        [dt.date(2024, 1, 15), 3, 10.5],
        [dt.date(2024, 2, 20), 8, 2.0],
        [dt.date(2025, 1, 9), 12, 4.0],
        [dt.date(2025, 3, 30), 5, 1.0],
    ]:
        sheet.append(row)
    workbook.create_sheet("Out")
    path = files / "dated.xlsx"
    workbook.save(path)
    return path


DATED = {
    "path": "dated.xlsx",
    "source": "Data!A1:C5",
    "sheet": "Out",
    "at": "A1",
}


async def test_date_groups_become_row_levels(call: ToolCall, dated: Path) -> None:
    await call(
        "create_pivot_table",
        **DATED,
        row_fields=["Day"],
        value_fields=[{"field": "Qty"}],
        field_settings=[{"field": "Day", "group_dates": ["months", "years"]}],
    )
    assert grid(dated, "A1:C9", "Out") == [
        ["Years (Day)", "Months (Day)", "Sum of Qty"],
        ["2024", "Jan", 3],
        [None, "Feb", 8],
        ["2024 Total", None, 11],
        ["2025", "Jan", 12],
        [None, "Mar", 5],
        ["2025 Total", None, 17],
        ["Grand Total", None, 28],
        [None, None, None],
    ]
    cache = cache_xml(dated)
    assert 'groupBy="years"' in cache and 'groupBy="months"' in cache
    assert 'name="Years (Day)"' in cache and "databaseField" in cache


async def test_number_groups(call: ToolCall, dated: Path) -> None:
    await call(
        "create_pivot_table",
        **DATED,
        row_fields=["Qty"],
        value_fields=[{"field": "Amount"}],
        field_settings=[{"field": "Qty", "group_numbers": {"by": 5, "start": 0, "end": 15}}],
    )
    assert grid(dated, "A1:B5", "Out") == [
        ["Qty", "Sum of Amount"],
        ["0-4", 10.5],
        ["5-9", 3.0],
        ["10-15", 4.0],
        ["Grand Total", 17.5],
    ]
    assert 'groupInterval="5' in cache_xml(dated)


@pytest.mark.parametrize(
    ("fields", "message"),
    [
        ([{"field": "Qty", "group_dates": ["years"]}], "grouped by date"),
        ([{"field": "Day", "group_numbers": {"by": 5}}], "into ranges"),
        ([{"field": "Day", "group_dates": ["years"], "show_items": ["2024"]}], "date groups"),
        ([{"field": "Amount", "sort": "ascending"}], "not used"),
    ],
)
async def test_invalid_groups(
    call_error: ToolCall, dated: Path, fields: list[dict[str, Any]], message: str
) -> None:
    result = await call_error(
        "create_pivot_table",
        **DATED,
        row_fields=["Day", "Qty"][: 1 if fields[0]["field"] == "Day" else 2],
        value_fields=[{"field": "Amount"}],
        field_settings=fields,
    )
    assert message in result


# Calculated fields


async def test_calculated_field(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "Revenue"}],
        calculated_fields=[{"name": "Revenue", "formula": "=Units*Price"}],
    )
    assert grid(sample, "A1:B4") == [
        ["Region", "Sum of Revenue"],
        ["North", 17 * 3.5],
        ["South", 8 * 3.5],
        ["Grand Total", 25 * 7],
    ]
    assert 'formula="Units*Price"' in cache_xml(sample)


async def test_calculated_fields_can_use_earlier_ones_and_quoted_names(
    call: ToolCall, sample: Path
) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "Both"}],
        calculated_fields=[
            {"name": "Double", "formula": "'Units' * 2"},
            {"name": "Both", "formula": "Double + Price"},
        ],
    )
    assert grid(sample, "A2:B3") == [["North", 34 + 3.5], ["South", 16 + 3.5]]


@pytest.mark.parametrize(
    ("formula", "message"),
    [
        ("WEBSERVICE(Units)", "not allowed"),
        ("Units+Mars", "not a field"),
        ("Region", "not numbers"),
        ("Data!A1", "not a field"),
        ("Units +", "not valid"),
    ],
)
async def test_calculated_field_formulas_are_checked(
    call_error: ToolCall, sample: Path, formula: str, message: str
) -> None:
    result = await call_error(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "X"}],
        calculated_fields=[{"name": "X", "formula": formula}],
    )
    assert message in result


async def test_calculated_fields_are_summed_only(call_error: ToolCall, sample: Path) -> None:
    result = await call_error(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "X", "function": "average"}],
        calculated_fields=[{"name": "X", "formula": "Units"}],
    )
    assert "only be summed" in result


# Layouts, subtotals, values in rows


async def test_compact_layout(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region", "Product"],
        value_fields=UNITS,
        layout="compact",
    )
    assert grid(sample, "A1:B8") == [
        ["Row Labels", "Sum of Units"],
        ["North", 17],
        ["Apples", 10],
        ["Pears", 7],
        ["South", 8],
        ["Apples", 5],
        ["Pears", 3],
        ["Grand Total", 25],
    ]
    sheet = load_workbook(sample)["Report"]
    assert sheet["A3"].alignment.indent == 1
    assert 'compact="0"' not in pivot_xml(sample)


async def test_outline_and_tabular_without_subtotals(call: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook.create_sheet("Outline")
    workbook.create_sheet("Tab")
    workbook.save(sample)
    await call(
        "create_pivot_table",
        **{**PIVOT, "sheet": "Outline"},
        row_fields=["Region", "Product"],
        value_fields=UNITS,
        layout="outline",
    )
    assert grid(sample, "A1:C7", "Outline") == [
        ["Region", "Product", "Sum of Units"],
        ["North", None, 17],
        [None, "Apples", 10],
        [None, "Pears", 7],
        ["South", None, 8],
        [None, "Apples", 5],
        [None, "Pears", 3],
    ]
    await call(
        "create_pivot_table",
        **{**PIVOT, "sheet": "Tab"},
        row_fields=["Region", "Product"],
        value_fields=UNITS,
        subtotals=False,
    )
    assert grid(sample, "A1:C6", "Tab") == [
        ["Region", "Product", "Sum of Units"],
        ["North", "Apples", 10],
        [None, "Pears", 7],
        ["South", "Apples", 5],
        [None, "Pears", 3],
        ["Grand Total", None, 25],
    ]
    assert 'defaultSubtotal="0"' in pivot_xml(sample, 2)


async def test_values_in_rows(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        value_fields=[{"field": "Units"}, {"field": "Price", "function": "average"}],
        values_in="rows",
    )
    assert grid(sample, "A1:C8") == [
        ["Region", "Values", None],
        ["North", "Sum of Units", 17],
        [None, "Average of Price", 1.75],
        ["South", "Sum of Units", 8],
        [None, "Average of Price", 1.75],
        ["Total Sum of Units", None, 25],
        ["Total Average of Price", None, 1.75],
        [None, None, None],
    ]
    assert 'dataOnRows="1"' in pivot_xml(sample)


# Filters with chosen items


async def test_filter_with_one_item(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        filter_fields=["Product"],
        value_fields=UNITS,
        field_settings=[{"field": "Product", "show_items": ["pears"]}],
    )
    assert grid(sample, "A1:B6") == [
        ["Product", "Pears"],
        [None, None],
        ["Region", "Sum of Units"],
        ["North", 7],
        ["South", 3],
        ["Grand Total", 10],
    ]
    xml = pivot_xml(sample)
    assert 'item="1"' in xml and 'h="1"' in xml


async def test_hidden_row_items_leave_out_their_data(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region", "Product"],
        value_fields=UNITS,
        field_settings=[{"field": "Region", "show_items": ["South"]}],
    )
    assert grid(sample, "A2:C5") == [
        ["South", "Apples", 5],
        [None, "Pears", 3],
        ["South Total", None, 8],
        ["Grand Total", None, 8],
    ]
    assert 'h="1"' in pivot_xml(sample)


async def test_filter_with_several_of_three_items(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in [["Kind", "Group", "N"], ["a", "x", 1], ["b", "x", 2], ["c", "y", 4]]:
        sheet.append(row)
    workbook.create_sheet("Out")
    workbook.save(files / "kinds.xlsx")
    await call(
        "create_pivot_table",
        path="kinds.xlsx",
        source="Data!A1:C4",
        sheet="Out",
        at="A1",
        row_fields=["Group"],
        filter_fields=["Kind"],
        value_fields=[{"field": "N"}],
        field_settings=[{"field": "Kind", "show_items": ["a", "b"]}],
    )
    assert grid(files / "kinds.xlsx", "A1:B5", "Out") == [
        ["Kind", "(Multiple Items)"],
        [None, None],
        ["Group", "Sum of N"],
        ["x", 3],
        ["Grand Total", 3],
    ]
    xml = pivot_xml(files / "kinds.xlsx")
    assert 'multipleItemSelectionAllowed="1"' in xml and 'item="' not in xml.split("pageFields")[1]


async def test_invalid_field_settings(call_error: ToolCall, sample: Path) -> None:
    base = {**PIVOT, "row_fields": ["Region"], "value_fields": UNITS}
    message = await call_error(
        "create_pivot_table", **base, field_settings=[{"field": "Region", "show_items": ["Mars"]}]
    )
    assert "no item 'Mars'" in message and "North" in message
    message = await call_error(
        "create_pivot_table", **base, field_settings=[{"field": "Product", "sort": "descending"}]
    )
    assert "not used" in message
    message = await call_error(
        "create_pivot_table",
        **base,
        field_settings=[
            {"field": "Region", "show_items": ["North"]},
            {"field": "region", "sort": "ascending"},
        ],
    )
    assert "twice" in message
