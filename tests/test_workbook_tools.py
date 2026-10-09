import base64
from pathlib import Path

import pytest
from mcp import Client
from mcp.types import BlobResourceContents, EmbeddedResource
from openpyxl import load_workbook
from openpyxl.chart import BarChart
from openpyxl.workbook.defined_name import DefinedName

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


async def test_create_workbook(call: ToolCall, files: Path) -> None:
    created = await call("create_workbook", path="new/book.xlsx", sheets=["Input", "Output"])
    assert created == {"path": "new/book.xlsx"}
    assert load_workbook(files / "new" / "book.xlsx").sheetnames == ["Input", "Output"]


async def test_create_workbook_refuses_to_overwrite(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    assert "already exists" in await call_error("create_workbook", path="sales.xlsx")
    await call("create_workbook", path="sales.xlsx", overwrite=True)
    assert load_workbook(sample).sheetnames == ["Sheet1"]


async def test_create_workbook_rejects_macro_formats(call_error: ToolCall) -> None:
    assert "cannot contain macros" in await call_error("create_workbook", path="m.xlsm")


async def test_create_workbook_rejects_duplicate_sheet_names(call_error: ToolCall) -> None:
    assert "already exists" in await call_error("create_workbook", path="b.xlsx", sheets=["A", "a"])


async def test_describe_workbook(call: ToolCall, sample: Path) -> None:
    info = await call("describe_workbook", path="sales.xlsx")
    assert info == {
        "sheets": [
            {"name": "Data", "used_range": "A1:D5", "active": True},
            {"name": "Report", "used_range": "A1"},
        ],
        "doc_properties": {"author": "openpyxl"},
        "calculation": {"full_calc_on_load": True},
    }


async def test_used_range_ignores_formatted_empty_cells(call: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook["Data"]["H40"].number_format = "0.00"
    workbook.save(sample)
    info = await call("describe_workbook", path="sales.xlsx")
    assert info["sheets"][0]["used_range"] == "A1:D5"


async def test_missing_and_invalid_files(call_error: ToolCall, files: Path) -> None:
    assert "does not exist" in await call_error("describe_workbook", path="missing.xlsx")
    (files / "fake.xlsx").write_text("Region,Units\nNorth,10\n")
    assert "not an Excel workbook" in await call_error("describe_workbook", path="fake.xlsx")


async def test_paths_outside_the_directory_are_rejected(call_error: ToolCall) -> None:
    assert "outside" in await call_error("describe_workbook", path="../secret.xlsx")
    assert "Excel file" in await call_error("create_workbook", path="evil.bat")


async def test_list_workbooks(call: ToolCall, sample: Path, files: Path) -> None:
    (files / "sub").mkdir()
    (files / "sub" / "b.xlsx").write_bytes(sample.read_bytes())
    (files / "notes.txt").write_text("x")
    flat = await call("list_workbooks")
    assert flat == {"sales.xlsx": sample.stat().st_size}
    nested = await call("list_workbooks", recursive=True)
    assert list(nested) == ["sales.xlsx", "sub/b.xlsx"]


async def test_export_and_import_round_trip(
    client: Client, call: ToolCall, call_error: ToolCall, sample: Path, files: Path
) -> None:
    result = await client.call_tool("export_workbook", {"path": "sales.xlsx"})
    assert not result.is_error
    embedded = result.content[1]
    assert isinstance(embedded, EmbeddedResource)
    assert isinstance(embedded.resource, BlobResourceContents)
    blob = embedded.resource.blob
    assert base64.b64decode(blob) == sample.read_bytes()

    await call("import_workbook", path="copy.xlsx", content_base64=blob)
    assert (files / "copy.xlsx").read_bytes() == sample.read_bytes()
    assert "already exists" in await call_error(
        "import_workbook", path="copy.xlsx", content_base64=blob
    )


async def test_import_rejects_non_workbooks(call_error: ToolCall) -> None:
    content = base64.b64encode(b"not a workbook").decode()
    assert "not an .xlsx" in await call_error(
        "import_workbook", path="x.xlsx", content_base64=content
    )
    assert "base64" in await call_error("import_workbook", path="x.xlsx", content_base64="%%")


async def test_describe_workbook_flags_hidden_sheets_and_names(
    call: ToolCall, sample: Path
) -> None:
    workbook = load_workbook(sample)
    workbook["Report"].sheet_state = "hidden"
    workbook.defined_names["Totals"] = DefinedName("Totals", attr_text="Data!$C$2:$C$5")
    workbook.create_chartsheet("Chart").add_chart(BarChart())
    workbook.save(sample)
    info = await call("describe_workbook", path="sales.xlsx")
    assert info["sheets"] == [
        {"name": "Data", "used_range": "A1:D5", "active": True},
        {"name": "Report", "used_range": "A1", "visibility": "hidden"},
    ]
    assert info["defined_names"] == [{"name": "Totals", "refers_to": "Data!$C$2:$C$5"}]
