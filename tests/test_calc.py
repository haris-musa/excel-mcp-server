"""The formula calculator behind read_range, and the function prefixes written to files."""

import datetime as dt
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from excel_mcp.xlfn import add_prefixes
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


def make_workbook(path: Path, cells: dict[str, int | str]) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for ref, value in cells.items():
        sheet[ref] = value
    workbook.save(path)


async def read_values(call: ToolCall, ref: str) -> dict[str, Any]:
    return await call("read_range", path="calc.xlsx", sheet="Data", range=ref)


async def test_formulas_without_a_stored_result_are_calculated(call: ToolCall, files: Path) -> None:
    make_workbook(
        files / "calc.xlsx",
        {"A1": 10, "A2": 20, "B1": "=SUM(A1:A2)", "B2": "=B1*2", "B3": '=IF(B2>50,"big","small")'},
    )

    data = await read_values(call, "A1:B3")

    assert data["values"] == [[10, 30], [20, 60], [None, "big"]]
    assert "uncalculated" not in data


async def test_stored_results_are_kept(call: ToolCall, files: Path) -> None:
    source = files / "calc.xlsx"
    make_workbook(source, {"A1": 1, "B1": "=A1+1"})
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"]
    assert b"<v></v>" in sheet or b"<v />" in sheet
    parts["xl/worksheets/sheet1.xml"] = sheet.replace(b"<v></v>", b"<v>99</v>").replace(
        b"<v />", b"<v>99</v>"
    )
    with ZipFile(source, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    data = await read_values(call, "B1")

    assert data["values"] == [[99]]


async def test_unsupported_functions_are_listed_not_guessed(call: ToolCall, files: Path) -> None:
    make_workbook(files / "calc.xlsx", {"A1": 1, "B1": "=FOO(A1)", "B2": "=RAND()", "B3": "=A1+1"})

    data = await read_values(call, "B1:B3")

    assert data["values"] == [[], [], [2]]
    assert data["uncalculated"] == {"B1": "FOO", "B2": "RAND"}


async def test_blocked_functions_are_never_calculated(call: ToolCall, files: Path) -> None:
    make_workbook(
        files / "calc.xlsx", {"B1": '=WEBSERVICE("http://example.com")', "B2": '=INDIRECT("A1")'}
    )

    data = await read_values(call, "B1:B2")

    assert data["values"] == []
    assert set(data["uncalculated"]) == {"B1", "B2"}


async def test_circular_references_are_uncalculated(call: ToolCall, files: Path) -> None:
    make_workbook(files / "calc.xlsx", {"A1": "=B1", "B1": "=A1"})

    data = await read_values(call, "A1:B1")

    assert data["values"] == []
    assert data["uncalculated"] == {"A1": "circular reference", "B1": "circular reference"}


async def test_long_dependency_chains(call: ToolCall, files: Path) -> None:
    cells: dict[str, int | str] = {"A1": 1}
    cells.update({f"A{row}": f"=A{row - 1}+1" for row in range(2, 301)})
    make_workbook(files / "calc.xlsx", cells)

    data = await read_values(call, "A300")

    assert data["values"] == [[300]]


async def test_work_is_bounded(call: ToolCall, files: Path) -> None:
    make_workbook(files / "calc.xlsx", {"A1": "=SUMPRODUCT(A2:A2000000,A2:A2000000)"})

    data = await read_values(call, "A1")

    assert data["values"] == []
    assert data["uncalculated"] == {"A1": "too much to calculate"}


async def test_dates_come_back_as_dates(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet["A1"] = "=DATE(2024,2,29)+1"
    sheet["A1"].number_format = "yyyy-mm-dd"
    workbook.save(files / "calc.xlsx")

    data = await read_values(call, "A1")

    assert data["values"] == [["2024-03-01"]]


async def test_errors_are_values(call: ToolCall, files: Path) -> None:
    make_workbook(files / "calc.xlsx", {"A1": "=1/0", "A2": '=IFERROR(A1,"x")'})

    data = await read_values(call, "A1:A2")

    assert data["values"] == [["#DIV/0!"], ["x"]]


async def test_other_sheets_and_names(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    other = workbook.create_sheet("My Sheet")
    other["A1"] = 5
    other["A2"] = 6
    sheet["A1"] = "=SUM('My Sheet'!A1:A2)*Rate"
    workbook.defined_names["Rate"] = DefinedName("Rate", attr_text="'My Sheet'!$A$1")
    workbook.save(files / "calc.xlsx")

    data = await read_values(call, "A1")

    assert data["values"] == [[55]]


async def test_written_formulas_get_storage_prefixes(call: ToolCall, sample: Path) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        at="A1",
        rows=[['=IFS(Data!C2>5,"a",TRUE,"b")', "=SORT(Data!C2:C5)"], ["=SUM(Data!C2:C5)"]],
    )

    sheet = load_workbook(sample)["Report"]
    assert sheet["A1"].value == '=_xlfn.IFS(Data!C2>5,"a",TRUE,"b")'
    assert sheet["B1"].value.text == "=_xlfn._xlws.SORT(Data!C2:C5)"
    assert sheet["A2"].value == "=SUM(Data!C2:C5)"
    data = await call("read_range", path="sales.xlsx", sheet="Report")
    assert data["values"] == [["a", 3], [25, 5], [None, 7], [None, 10]]


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ("=IFS(A1,1)", "=_xlfn.IFS(A1,1)"),
        ("=_xlfn.IFS(A1,1)", "=_xlfn.IFS(A1,1)"),
        ('=TEXTJOIN(",",TRUE,"IFS(")', '=_xlfn.TEXTJOIN(",",TRUE,"IFS(")'),
        ("=STDEV.S(A1:A3)+STDEV(A1:A3)", "=_xlfn.STDEV.S(A1:A3)+STDEV(A1:A3)"),
        ("=SUM(A1)", "=SUM(A1)"),
        # As Excel itself stores the formulas (Formula2 saved to a file).
        (
            "=MAP(A1:A5,LAMBDA(x,x*2))",
            "=_xlfn.MAP(A1:A5,_xlfn.LAMBDA(_xlpm.x,_xlpm.x*2))",
        ),
        (
            "=BYROW(A1:B5,LAMBDA(r,SUM(r)))",
            "=_xlfn.BYROW(A1:B5,_xlfn.LAMBDA(_xlpm.r,SUM(_xlpm.r)))",
        ),
        (
            "=REDUCE(0,A1:A5,LAMBDA(a,b,a+b))",
            "=_xlfn.REDUCE(0,A1:A5,_xlfn.LAMBDA(_xlpm.a,_xlpm.b,_xlpm.a+_xlpm.b))",
        ),
        ("=LAMBDA(x,ISOMITTED(x))(1)", "=_xlfn.LAMBDA(_xlpm.x,_xlfn.ISOMITTED(_xlpm.x))(1)"),
        (
            "=BYCOL(A1:B5,LAMBDA(c,SUM(c)))",
            "=_xlfn.BYCOL(A1:B5,_xlfn.LAMBDA(_xlpm.c,SUM(_xlpm.c)))",
        ),
        ("=VALUETOTEXT(A1)", "=_xlfn.VALUETOTEXT(A1)"),
        ("=ARRAYTOTEXT(A1:B2)", "=_xlfn.ARRAYTOTEXT(A1:B2)"),
    ],
)
def test_add_prefixes(formula: str, expected: str) -> None:
    assert add_prefixes(formula) == expected


async def test_today_and_now(call: ToolCall, files: Path) -> None:
    make_workbook(files / "calc.xlsx", {"A1": "=TODAY()", "A2": "=INT(NOW())"})

    data = await read_values(call, "A1:A2")

    today = (dt.date.today() - dt.date(1899, 12, 30)).days
    assert data["values"] == [[today], [today]]


async def test_running_balance_of_2000_rows(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in range(1, 2001):
        sheet.cell(row, 1, row % 7 + 1)
        sheet.cell(row, 2, "=A1" if row == 1 else f"=B{row - 1}+A{row}")
    workbook.save(files / "calc.xlsx")

    data = await read_values(call, "B2000")

    assert data["values"] == [[sum(row % 7 + 1 for row in range(1, 2001))]]


async def test_amortization_schedule_read_from_its_last_row(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet["A1"], sheet["A2"], sheet["A3"] = 250000, 0.065, 360
    for month in range(1, 361):
        row = month + 4
        sheet[f"A{row}"] = month
        sheet[f"B{row}"] = "=$A$1" if month == 1 else f"=F{row - 1}"
        sheet[f"D{row}"] = f"=-IPMT($A$2/12,A{row},$A$3,$A$1)"
        sheet[f"E{row}"] = f"=-PPMT($A$2/12,A{row},$A$3,$A$1)"
        sheet[f"F{row}"] = f"=B{row}-E{row}"
    workbook.save(files / "calc.xlsx")

    data = await read_values(call, "F364")

    assert abs(data["values"][0][0]) < 1e-6
