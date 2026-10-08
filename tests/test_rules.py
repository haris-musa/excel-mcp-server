"""Data validation and conditional formatting rules, checked in the saved file."""

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

BOOK = {"path": "sales.xlsx", "sheet": "Data"}


def _validations(path: Path) -> list[DataValidation]:
    return load_workbook(path)["Data"].data_validations.dataValidation


def _rules(path: Path, sheet: str = "Data") -> list[Any]:
    worksheet: Worksheet = load_workbook(path)[sheet]
    return [rule for entry in worksheet.conditional_formatting for rule in entry.rules]


async def _validate(call: ToolCall, rule: dict[str, Any], range: str = "E2:E5") -> None:
    await call("add_data_validation", **BOOK, range=range, rule=rule)


@pytest.mark.parametrize(
    ("source", "stored"),
    [
        ("=$A$2:$A$5", "$A$2:$A$5"),
        ("=Report!$A:$A", "Report!$A:$A"),
        ("=Regions", "Regions"),
        ("='Report'!$A$1:$E$1", "'Report'!$A$1:$E$1"),
    ],
)
async def test_list_source_from_cells_or_a_name(
    call: ToolCall, sample: Path, source: str, stored: str
) -> None:
    await _validate(call, {"type": "list", "source": source})
    (validation,) = _validations(sample)
    assert validation.type == "list"
    assert validation.formula1 == stored


async def test_list_needs_exactly_one_source_of_values(call_error: ToolCall, sample: Path) -> None:
    assert "either options or source" in await call_error(
        "add_data_validation", **BOOK, range="E2", rule={"type": "list"}
    )
    both = {"type": "list", "options": ["a"], "source": "=$A$2:$A$5"}
    assert "either options or source" in await call_error(
        "add_data_validation", **BOOK, range="E2", rule=both
    )
    assert "only for list" in await call_error(
        "add_data_validation",
        **BOOK,
        range="E2",
        rule={"type": "whole", "operator": "equal", "minimum": "1", "source": "=$A$2:$A$5"},
    )


async def test_list_source_must_be_one_row_or_column(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "add_data_validation", **BOOK, range="E2", rule={"type": "list", "source": "=A1:B5"}
    )
    assert "one row or column" in message


@pytest.mark.parametrize(
    "source",
    ['=INDIRECT("A1")', "=[1]Sheet1!$A$1:$A$5", "=Elsewhere!$A$1:$A$5", '=WEBSERVICE("x")'],
)
async def test_list_source_goes_through_the_formula_gate(
    call_error: ToolCall, sample: Path, source: str
) -> None:
    before = sample.read_bytes()
    await call_error(
        "add_data_validation", **BOOK, range="E2", rule={"type": "list", "source": source}
    )
    assert sample.read_bytes() == before


async def test_time_and_date_bounds_are_stored_as_serial_numbers(
    call: ToolCall, sample: Path
) -> None:
    await _validate(
        call,
        {"type": "time", "operator": "between", "minimum": "09:30", "maximum": "17:00:00"},
    )
    await _validate(
        call,
        {"type": "date", "operator": "greaterThan", "minimum": "2026-01-31"},
        range="F2:F5",
    )
    await _validate(
        call,
        {"type": "date", "operator": "lessThan", "minimum": "DATE(2030,1,1)"},
        range="G2:G5",
    )
    time, date, formula = _validations(sample)
    assert (time.type, time.operator) == ("time", "between")
    assert float(str(time.formula1)) == pytest.approx(9.5 / 24)
    assert float(str(time.formula2)) == pytest.approx(17 / 24)
    assert date.formula1 == "46053"
    assert formula.formula1 == "DATE(2030,1,1)"


@pytest.mark.parametrize(
    ("kind", "minimum"),
    [("date", "2026-02-30"), ("time", "25:00")],
)
async def test_invalid_date_or_time_is_rejected(
    call_error: ToolCall, sample: Path, kind: str, minimum: str
) -> None:
    message = await call_error(
        "add_data_validation",
        **BOOK,
        range="E2",
        rule={"type": kind, "operator": "equal", "minimum": minimum},
    )
    assert f"not a valid {kind}" in message


async def test_messages_and_alert_style(call: ToolCall, sample: Path) -> None:
    await _validate(
        call,
        {
            "type": "whole",
            "operator": "between",
            "minimum": "1",
            "maximum": "10",
            "prompt_title": "Quantity",
            "prompt": "Enter 1 to 10",
            "error_style": "warning",
            "error_title": "Check it",
            "error_message": "Outside 1-10",
        },
    )
    await _validate(call, {"type": "custom", "formula": "=E2>0"}, range="F2:F5")
    warning, default = _validations(sample)
    assert warning.errorStyle == "warning"
    assert (warning.promptTitle, warning.prompt) == ("Quantity", "Enter 1 to 10")
    assert (warning.errorTitle, warning.error) == ("Check it", "Outside 1-10")
    assert warning.showInputMessage and warning.showErrorMessage
    assert default.errorStyle in (None, "stop")
    assert not default.showInputMessage


async def test_message_lengths_follow_excels_limits(call_error: ToolCall, sample: Path) -> None:
    rule = {"type": "custom", "formula": "=E2>0", "error_title": "x" * 33}
    assert "at most 32" in await call_error("add_data_validation", **BOOK, range="E2", rule=rule)
    rule = {"type": "custom", "formula": "=E2>0", "prompt": "x" * 256}
    assert "at most 255" in await call_error("add_data_validation", **BOOK, range="E2", rule=rule)


# -- conditional formats ---------------------------------------------------------------------

RED = {"fill_color": "#FFC7CE"}


@pytest.mark.parametrize(
    ("rule", "attributes"),
    [
        (
            {"type": "top", "count": 3},
            {"type": "top10", "rank": 3, "percent": None, "bottom": None},
        ),
        (
            {"type": "bottom", "count": 10, "percent": True},
            {"type": "top10", "rank": 10, "percent": True, "bottom": True},
        ),
        ({"type": "above_average"}, {"type": "aboveAverage", "aboveAverage": None}),
        (
            {"type": "below_average", "include_equal": True},
            {"type": "aboveAverage", "aboveAverage": False, "equalAverage": True},
        ),
        (
            {"type": "above_average", "std_dev": 2},
            {"type": "aboveAverage", "stdDev": 2},
        ),
        ({"type": "duplicate"}, {"type": "duplicateValues"}),
        ({"type": "unique"}, {"type": "uniqueValues"}),
    ],
)
async def test_ranking_and_uniqueness_rules(
    call: ToolCall, sample: Path, rule: dict[str, Any], attributes: dict[str, Any]
) -> None:
    await call("add_conditional_format", **BOOK, range="C2:C5", rule={**rule, **RED})
    (saved,) = _rules(sample)
    for name, value in attributes.items():
        assert getattr(saved, name) == value, name
    assert saved.dxf.fill.bgColor.rgb == "FFFFC7CE"


@pytest.mark.parametrize(
    ("rule", "kind", "operator", "formula"),
    [
        (
            {"type": "contains_text", "text": 'ab"c'},
            "containsText",
            "containsText",
            'NOT(ISERROR(SEARCH("ab""c",B2)))',
        ),
        (
            {"type": "not_contains_text", "text": "x"},
            "notContainsText",
            "notContains",
            'ISERROR(SEARCH("x",B2))',
        ),
        (
            {"type": "begins_with", "text": "Ap"},
            "beginsWith",
            "beginsWith",
            'LEFT(B2,LEN("Ap"))="Ap"',
        ),
        (
            {"type": "ends_with", "text": "rs"},
            "endsWith",
            "endsWith",
            'RIGHT(B2,LEN("rs"))="rs"',
        ),
    ],
)
async def test_text_rules_write_what_excel_writes(
    call: ToolCall, sample: Path, rule: dict[str, Any], kind: str, operator: str, formula: str
) -> None:
    await call("add_conditional_format", **BOOK, range="B2:B5", rule={**rule, **RED})
    (saved,) = _rules(sample)
    assert (saved.type, saved.operator, saved.formula) == (kind, operator, [formula])
    assert saved.text == rule["text"]


@pytest.mark.parametrize(
    ("period", "formula"),
    [
        ("yesterday", "FLOOR(D2,1)=TODAY()-1"),
        ("today", "FLOOR(D2,1)=TODAY()"),
        ("last7Days", "AND(TODAY()-FLOOR(D2,1)<=6,FLOOR(D2,1)<=TODAY())"),
        (
            "thisWeek",
            "AND(TODAY()-ROUNDDOWN(D2,0)<=WEEKDAY(TODAY())-1,"
            "ROUNDDOWN(D2,0)-TODAY()<=7-WEEKDAY(TODAY()))",
        ),
        ("nextMonth", "AND(MONTH(D2)=MONTH(EDATE(TODAY(),0+1)),YEAR(D2)=YEAR(EDATE(TODAY(),0+1)))"),
    ],
)
async def test_date_rules(call: ToolCall, sample: Path, period: str, formula: str) -> None:
    await call(
        "add_conditional_format",
        **BOOK,
        range="D2:D5",
        rule={"type": "date", "period": period, **RED},
    )
    (saved,) = _rules(sample)
    assert (saved.type, saved.timePeriod, saved.formula) == ("timePeriod", period, [formula])


@pytest.mark.parametrize(
    ("kind", "stored", "formula"),
    [
        ("blanks", "containsBlanks", "LEN(TRIM(B3))=0"),
        ("no_blanks", "notContainsBlanks", "LEN(TRIM(B3))>0"),
        ("errors", "containsErrors", "ISERROR(B3)"),
        ("no_errors", "notContainsErrors", "NOT(ISERROR(B3))"),
    ],
)
async def test_blank_and_error_rules_use_the_top_left_cell(
    call: ToolCall, sample: Path, kind: str, stored: str, formula: str
) -> None:
    await call("add_conditional_format", **BOOK, range="B3:C5", rule={"type": kind, **RED})
    (saved,) = _rules(sample)
    assert (saved.type, saved.formula) == (stored, [formula])


@pytest.mark.parametrize(
    ("rule", "icons", "thresholds", "kind"),
    [
        ({"icon_set": "3Arrows"}, "3Arrows", [0, 33, 67], "percent"),
        ({"icon_set": "4Rating"}, "4Rating", [0, 25, 50, 75], "percent"),
        ({"icon_set": "5Quarters"}, "5Quarters", [0, 20, 40, 60, 80], "percent"),
        (
            {"icon_set": "3Flags", "thresholds": [5, 7], "threshold_type": "number"},
            "3Flags",
            [0, 5, 7],
            "num",
        ),
    ],
)
async def test_icon_sets(
    call: ToolCall,
    sample: Path,
    rule: dict[str, Any],
    icons: str,
    thresholds: list[float],
    kind: str,
) -> None:
    await call("add_conditional_format", **BOOK, range="C2:C5", rule={"type": "icon_set", **rule})
    (saved,) = _rules(sample)
    assert saved.iconSet.iconSet == icons
    assert [point.val for point in saved.iconSet.cfvo] == [float(value) for value in thresholds]
    assert [point.type for point in saved.iconSet.cfvo] == ["percent"] + [kind] * (
        len(thresholds) - 1
    )
    assert not saved.iconSet.reverse and saved.iconSet.showValue is None


async def test_icon_set_reverse_and_icon_only(call: ToolCall, sample: Path) -> None:
    rule = {"type": "icon_set", "icon_set": "3Symbols", "reverse": True, "icon_only": True}
    await call("add_conditional_format", **BOOK, range="C2:C5", rule=rule)
    (saved,) = _rules(sample)
    assert saved.iconSet.reverse is True
    assert saved.iconSet.showValue is False


@pytest.mark.parametrize(
    ("rule", "message"),
    [
        ({"type": "icon_set", "icon_set": "3Arrows", "thresholds": [30]}, "needs 2 thresholds"),
        ({"type": "icon_set", "icon_set": "3Arrows", "thresholds": [70, 30]}, "ascending"),
        ({"type": "icon_set"}, "need an icon_set"),
        ({"type": "top", **RED}, "need a count"),
        ({"type": "top", "count": 101, "percent": True, **RED}, "at most 100"),
        ({"type": "contains_text", **RED}, "needs text"),
        ({"type": "date", **RED}, "need a period"),
        ({"type": "duplicate"}, "fill_color or font_color"),
        (
            {"type": "color_scale", "colors": ["#FF0000", "#00FF00"], "text": "x"},
            "does not use text",
        ),
        ({"type": "unique", "icon_only": True, **RED}, "does not use icon_only"),
    ],
)
async def test_invalid_rules_are_rejected(
    call_error: ToolCall, sample: Path, rule: dict[str, Any], message: str
) -> None:
    before = sample.read_bytes()
    assert message in await call_error("add_conditional_format", **BOOK, range="C2:C5", rule=rule)
    assert sample.read_bytes() == before


async def test_stop_if_true_and_priority(call: ToolCall, sample: Path) -> None:
    for text in ("a", "b", "c"):
        await call(
            "add_conditional_format",
            **BOOK,
            range="B2:B5",
            rule={"type": "contains_text", "text": text, **RED},
        )
    await call(
        "add_conditional_format",
        **BOOK,
        range="B2:B5",
        rule={"type": "begins_with", "text": "d", "stop_if_true": True, "priority": 2, **RED},
    )
    ordered = sorted(_rules(sample), key=lambda rule: rule.priority)
    assert [rule.text for rule in ordered] == ["a", "d", "b", "c"]
    assert [rule.priority for rule in ordered] == [1, 2, 3, 4]
    assert [bool(rule.stopIfTrue) for rule in ordered] == [False, True, False, False]


async def test_priorities_continue_after_existing_rules(call: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    from openpyxl.formatting.rule import CellIsRule

    rule = CellIsRule(operator="equal", formula=["1"])
    rule.priority = 7
    workbook["Data"].conditional_formatting.add("A1", rule)
    workbook.save(sample)
    await call("add_conditional_format", **BOOK, range="C2:C5", rule={"type": "unique", **RED})
    assert sorted(rule.priority for rule in _rules(sample)) == [7, 8]


async def test_conditional_format_formulas_go_through_the_gate(
    call_error: ToolCall, sample: Path
) -> None:
    rule = {"type": "cell_value", "operator": "equal", "values": ['INDIRECT("A1")'], **RED}
    assert "not allowed" in await call_error(
        "add_conditional_format", **BOOK, range="C2:C5", rule=rule
    )
