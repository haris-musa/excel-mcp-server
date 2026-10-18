"""trace_formula (precedents, dependents, Evaluate Formula) and goal_seek."""

from pathlib import Path
from typing import Any

import pytest
from mcp import Client
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


@pytest.fixture
def model(files: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Model"
    inputs = workbook.create_sheet("Inputs")
    inputs["B1"], inputs["B2"] = 0.1, 1000
    workbook.defined_names["Rate"] = DefinedName("Rate", attr_text="Inputs!$B$1")
    for ref, value in {
        "A1": 5,
        "A2": 7,
        "A3": "=A1+A2",
        "A4": "=A3*Rate",
        "A5": "=SUM(A1:A4)",
        "A6": "=Inputs!B2+A3",
        "A7": '=INDIRECT("A"&1)',
        "A8": "=OFFSET(A1,1,0)",
        "A10": "=INDIRECT(B10)",
        "B10": '=CUBEMEMBER("a","b")',
    }.items():
        sheet[ref] = value
    path = files / "model.xlsx"
    workbook.save(path)
    return path


async def trace(call: ToolCall, range: str, direction: str, **more: Any) -> dict[str, Any]:
    return await call(
        "trace_formula",
        path="model.xlsx",
        sheet="Model",
        range=range,
        direction=direction,
        **more,
    )


async def test_precedents_resolve_names_and_other_sheets(call: ToolCall, model: Path) -> None:
    assert await trace(call, "A5", "precedents") == {"cells": {"Model": ["A1:A4"]}}
    assert await trace(call, "A4", "precedents") == {"cells": {"Model": ["A3"], "Inputs": ["B1"]}}
    assert await trace(call, "A6", "precedents") == {"cells": {"Model": ["A3"], "Inputs": ["B2"]}}


async def test_precedents_can_follow_every_level(call: ToolCall, model: Path) -> None:
    result = await trace(call, "A5", "precedents", all_levels=True)
    assert result == {"cells": {"Model": ["A1:A4"], "Inputs": ["B1"]}}
    result = await trace(call, "A6", "precedents", all_levels=True)
    assert result == {"cells": {"Model": ["A1:A3"], "Inputs": ["B2"]}}


async def test_precedents_of_a_range_are_the_union(call: ToolCall, model: Path) -> None:
    result = await trace(call, "A3:A4", "precedents")
    assert result == {"cells": {"Model": ["A1:A3"], "Inputs": ["B1"]}}


async def test_indirect_and_offset_point_at_what_they_compute(call: ToolCall, model: Path) -> None:
    assert await trace(call, "A7", "precedents") == {"cells": {"Model": ["A1"]}}
    assert await trace(call, "A8", "precedents") == {"cells": {"Model": ["A2"]}}


async def test_an_indirect_that_cannot_be_computed_is_unknown(call: ToolCall, model: Path) -> None:
    result = await trace(call, "A10", "precedents")
    assert result == {"cells": {"Model": ["B10"]}, "unknown": ["Model!A10"]}
    dependents = await trace(call, "A1", "dependents")
    assert dependents["unknown"] == ["Model!A10"]


async def test_dependents_across_the_workbook(call: ToolCall, model: Path) -> None:
    direct = await trace(call, "A1", "dependents")
    assert direct["cells"] == {"Model": ["A3", "A5", "A7"]}
    every = await trace(call, "A1", "dependents", all_levels=True)
    assert every["cells"] == {"Model": ["A3:A7"]}
    rate = await call(
        "trace_formula",
        path="model.xlsx",
        sheet="Inputs",
        range="B1",
        direction="dependents",
        all_levels=True,
    )
    assert rate["cells"] == {"Model": ["A4:A5"]}
    block = await trace(call, "A1:A2", "dependents")
    assert block["cells"] == {"Model": ["A3", "A5", "A7:A8"]}


async def test_tables_and_spill_references(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in [["Item", "Amount"], ["a", 1], ["b", 2], ["c", 3]]:
        sheet.append(row)
    sheet.add_table(Table(displayName="Sales", ref="A1:B4"))
    sheet["D1"] = "=SUM(Sales[Amount])"
    sheet["F1"] = ArrayFormula("F1:F3", "=_xlfn.SEQUENCE(3)")
    sheet["G1"] = "=SUM(_xlfn.ANCHORARRAY(F1))"
    workbook.save(files / "t.xlsx")

    def trace_cell(cell: str, direction: str = "precedents") -> Any:
        return call("trace_formula", path="t.xlsx", sheet="Data", range=cell, direction=direction)

    assert await trace_cell("D1") == {"cells": {"Data": ["B2:B4"]}}
    assert await trace_cell("G1") == {"cells": {"Data": ["F1:F3"]}}
    assert await trace_cell("B2", "dependents") == {"cells": {"Data": ["D1"]}}
    assert await trace_cell("F1", "dependents") == {"cells": {"Data": ["G1"]}}
    assert await trace_cell("B2:B4", "dependents") == {"cells": {"Data": ["D1"]}}


async def test_evaluate_shows_each_step_in_excels_order(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet["A1"], sheet["A2"] = 7, "x"
    sheet["B1"] = '=IF(A1>5,A1*2,1/0)&"!"'
    sheet["B2"] = "=IFERROR(1/0,A2)"
    sheet["B3"] = "=SUM(A1,2)*(A1-3)"
    workbook.save(files / "e.xlsx")

    async def steps(cell: str) -> list[list[str]]:
        result = await call(
            "trace_formula", path="e.xlsx", sheet="Data", range=cell, direction="evaluate"
        )
        return [[step["expr"], step["value"]] for step in result["steps"]]

    assert await steps("B1") == [
        ["A1", "7"],
        ["A1>5", "TRUE"],
        ["A1", "7"],
        ["A1*2", "14"],
        ["IF(A1>5,A1*2,1/0)", "14"],
        ['IF(A1>5,A1*2,1/0)&"!"', '"14!"'],
    ]
    assert await steps("B2") == [
        ["1/0", "#DIV/0!"],
        ["A2", '"x"'],
        ["IFERROR(1/0,A2)", '"x"'],
    ]
    assert await steps("B3") == [
        ["A1", "7"],
        ["SUM(A1,2)", "9"],
        ["A1", "7"],
        ["A1-3", "4"],
        ["SUM(A1,2)*(A1-3)", "36"],
    ]


async def test_results_are_capped_with_a_count(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in range(1, 601, 2):
        sheet[f"A{row}"] = row
        sheet[f"C{row}"] = f"=A{row}"
    workbook.save(files / "many.xlsx")

    result = await call(
        "trace_formula", path="many.xlsx", sheet="Data", range="A1:A600", direction="dependents"
    )

    assert len(result["cells"]["Data"]) == 200
    assert result["more"] == 100


async def test_invalid_requests_are_errors(call_error: ToolCall, model: Path) -> None:
    async def fails(range: str, direction: str, **more: Any) -> str:
        return await call_error(
            "trace_formula",
            path="model.xlsx",
            sheet="Model",
            range=range,
            direction=direction,
            **more,
        )

    assert "holds no formula" in await fails("A1", "precedents")
    assert "holds no formula" in await fails("A1", "evaluate")
    assert "one cell" in await fails("A3:A4", "evaluate")
    assert "one cell" in await fails("A3", "evaluate", all_levels=True)
    assert "Invalid range" in await fails("not a range", "precedents")
    assert "Nowhere" in await call_error(
        "trace_formula", path="model.xlsx", sheet="Nowhere", range="A1", direction="dependents"
    )


async def test_an_uncalculable_formula_is_reported(
    call: ToolCall, call_error: ToolCall, files: Path
) -> None:
    workbook = Workbook()
    workbook.worksheets[0]["A1"] = '=CUBEMEMBER("x","y")'
    workbook.save(files / "cube.xlsx")

    message = await call_error(
        "trace_formula", path="cube.xlsx", sheet="Sheet", range="A1", direction="evaluate"
    )

    assert "cannot be evaluated here" in message


async def test_trace_refuses_paths_outside_the_sandbox(call_error: ToolCall) -> None:
    message = await call_error(
        "trace_formula", path="../x.xlsx", sheet="S", range="A1", direction="precedents"
    )
    assert "outside" in message or "not allowed" in message


async def test_trace_is_read_only_and_goal_seek_is_destructive(client: Client) -> None:
    tools = {tool.name: tool for tool in (await client.list_tools()).tools}
    trace_tool, seek_tool = tools["trace_formula"], tools["goal_seek"]
    assert trace_tool.annotations and trace_tool.annotations.read_only_hint
    assert seek_tool.annotations and seek_tool.annotations.destructive_hint
    assert not seek_tool.annotations.read_only_hint


async def test_trace_does_not_change_the_file(call: ToolCall, model: Path) -> None:
    before = model.read_bytes()
    await trace(call, "A5", "precedents", all_levels=True)
    await trace(call, "A1", "dependents", all_levels=True)
    await trace(call, "A5", "evaluate")
    assert model.read_bytes() == before


# -- Goal Seek ----------------------------------------------------------------------------


def goal_workbook(files: Path, cells: dict[str, Any], name: str = "goal.xlsx") -> Path:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for ref, value in cells.items():
        sheet[ref] = value
    path = files / name
    workbook.save(path)
    return path


async def seek(call: ToolCall, cell: str, target: float, changing: str) -> dict[str, Any]:
    return await call(
        "goal_seek",
        path="goal.xlsx",
        sheet="Data",
        cell=cell,
        target=target,
        changing_cell=changing,
    )


async def test_goal_seek_finds_the_rate_of_a_loan(call: ToolCall, files: Path) -> None:
    path = goal_workbook(files, {"A1": 0.05, "B1": "=PMT(A1/12,360,-200000)", "C1": "keep"})

    result = await seek(call, "B1", 1200, "A1")

    assert result["converged"] is True
    assert result["value"] == pytest.approx(0.0600699, abs=1e-6)
    assert result["result"] == pytest.approx(1200, abs=0.001)
    assert 1 <= result["iterations"] <= 10
    sheet = load_workbook(path)["Data"]
    assert sheet["A1"].value == result["value"]
    assert sheet["B1"].value == "=PMT(A1/12,360,-200000)"
    assert sheet["C1"].value == "keep"


async def test_goal_seek_follows_chains_of_formulas_and_names(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    workbook.defined_names["Price"] = DefinedName("Price", attr_text="Data!$C$1")
    for ref, value in {
        "A1": 0,
        "C1": 10,
        "C2": 6,
        "C3": 5000,
        "B1": "=A1*Price",
        "B2": "=A1*C2",
        "B3": "=B1-B2-C3",
    }.items():
        sheet[ref] = value
    path = files / "goal.xlsx"
    workbook.save(path)

    result = await seek(call, "B3", 0, "A1")

    assert result["converged"] is True
    assert result["value"] == pytest.approx(1250)
    assert load_workbook(path)["Data"]["A1"].value == pytest.approx(1250)


async def test_goal_seek_treats_a_blank_input_as_zero(call: ToolCall, files: Path) -> None:
    goal_workbook(files, {"B1": "=A1*3+1"})

    result = await seek(call, "B1", 10, "A1")

    assert result["value"] == pytest.approx(3)


async def test_goal_seek_that_does_not_converge_changes_nothing(
    call: ToolCall, files: Path
) -> None:
    path = goal_workbook(files, {"A1": 1, "B1": "=A1^2"})
    before = path.read_bytes()

    result = await seek(call, "B1", -1, "A1")

    assert result["converged"] is False
    assert "may not have found a solution" in result["note"]
    assert result["iterations"] == 100
    assert path.read_bytes() == before


async def test_goal_seek_can_find_a_root_from_either_side(call: ToolCall, files: Path) -> None:
    goal_workbook(files, {"A1": -1, "B1": "=A1^2-4"})

    result = await seek(call, "B1", 0, "A1")

    assert result["value"] == pytest.approx(-2, abs=1e-3)


async def test_goal_seek_survives_values_the_function_cannot_take(
    call: ToolCall, files: Path
) -> None:
    goal_workbook(files, {"A1": 0, "B1": "=EXP(A1)"})

    result = await seek(call, "B1", 1000, "A1")

    assert result["converged"] is True
    assert result["value"] == pytest.approx(6.9077553, abs=1e-6)


async def test_goal_seek_refusals(call_error: ToolCall, files: Path) -> None:
    goal_workbook(
        files,
        {
            "A1": 1,
            "A2": "text",
            "A3": "=A1+1",
            "B1": "=A1*2",
            "B2": "=5",
            "B3": '=A1&"x"',
            "B4": '=CUBEMEMBER("a","b")+A1',
            "B5": "=A1/0",
        },
    )

    async def refused(cell: str, changing: str) -> str:
        return await call_error(
            "goal_seek", path="goal.xlsx", sheet="Data", cell=cell, target=3, changing_cell=changing
        )

    assert "holds no formula" in await refused("A1", "A1")
    assert "holds a formula" in await refused("B1", "A3")
    assert "not a number" in await refused("B1", "A2")
    assert "does not depend on A1" in await refused("B2", "A1")
    assert "must give a number" in await refused("B3", "A1")
    assert "must give a number" in await refused("B5", "A1")
    assert "cannot be calculated here" in await refused("B4", "A1")


async def test_goal_seek_refuses_merged_cells_and_other_sheets(
    call_error: ToolCall, files: Path
) -> None:
    path = goal_workbook(files, {"A1": 1, "B1": "=A1*2"})
    workbook = load_workbook(path)
    workbook["Data"].merge_cells("D1:E1")
    workbook.save(path)

    merged = await call_error(
        "goal_seek", path="goal.xlsx", sheet="Data", cell="B1", target=3, changing_cell="E1"
    )
    missing = await call_error(
        "goal_seek", path="goal.xlsx", sheet="Nowhere", cell="B1", target=3, changing_cell="A1"
    )

    assert "merged" in merged
    assert "Nowhere" in missing


async def test_goal_seek_refuses_paths_outside_the_sandbox(call_error: ToolCall) -> None:
    message = await call_error(
        "goal_seek", path="../x.xlsx", sheet="S", cell="B1", target=1, changing_cell="A1"
    )
    assert "outside" in message or "not allowed" in message


async def test_goal_seek_refuses_a_target_that_reads_a_spill_body(
    call_error: ToolCall, files: Path
) -> None:
    path = goal_workbook(files, {"A1": 3, "B1": "=C2*2"})
    workbook = load_workbook(path)
    workbook["Data"]["C1"] = ArrayFormula("C1:C3", "=_xlfn.SEQUENCE(A1)")
    workbook.save(path)

    message = await call_error(
        "goal_seek", path="goal.xlsx", sheet="Data", cell="B1", target=8, changing_cell="A1"
    )

    assert "spilling formula" in message


async def test_goal_seek_stores_a_file_that_recalculates_to_the_result(
    call: ToolCall, files: Path
) -> None:
    path = goal_workbook(files, {"A1": 0.01, "B1": "=A1*A1*100+A1"})

    result = await seek(call, "B1", 5, "A1")
    stored = await call("read_range", path="goal.xlsx", sheet="Data", range="A1:B1")

    assert stored["values"][0][1] == pytest.approx(result["result"])
    assert abs(stored["values"][0][1] - 5) < 0.001
    assert path.exists()
