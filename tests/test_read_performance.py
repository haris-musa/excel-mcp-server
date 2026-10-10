"""Reads of workbooks with stored results stream; only an uncalculated formula loads it all."""

import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from excel_mcp.workspace import Workspace
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


def make_stored(path: Path) -> None:
    """A workbook as Excel saves it: results are stored, and "" is a text result without content."""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in range(1, 6):
        sheet.cell(row, 1, row)
        sheet.cell(row, 2, f"=A{row}+1")
        sheet.cell(row, 3, f'=IF(A{row}>9,"x","")')
    workbook.save(path)
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    xml = parts["xl/worksheets/sheet1.xml"].decode()
    xml = re.sub(r'(<c r="B(\d)"[^>]*>)<f>(.*?)</f><v ?/?>(</v>)?', r"\1<f>\3</f><v>\2</v>", xml)
    xml = re.sub(
        r'<c r="(C\d)"([^>]*)>(<f>.*?</f>)<v ?/?>(</v>)?', r'<c r="\1" t="str"\2>\3<v/>', xml
    )
    assert xml.count("<v>") >= 10 and xml.count('t="str"') == 5
    parts["xl/worksheets/sheet1.xml"] = xml.encode()
    with ZipFile(path, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)


@pytest.fixture
def full_loads(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    loaded: list[str] = []
    original = Workspace._load  # pyright: ignore[reportPrivateUsage]

    def record(self: Workspace, path: Path, *, data_only: bool, stream: bool = False):
        if not stream:
            loaded.append(path.name)
        return original(self, path, data_only=data_only, stream=stream)

    monkeypatch.setattr(Workspace, "_load", record)
    return loaded


async def test_reads_with_stored_results_never_load_the_whole_workbook(
    call: ToolCall, files: Path, full_loads: list[str]
) -> None:
    make_stored(files / "stored.xlsx")

    data = await call("read_range", path="stored.xlsx", sheet="Data")
    await call("read_range", path="stored.xlsx", sheet="Data", mode="formulas")
    await call("find_cells", path="stored.xlsx", query="3")
    await call("describe_workbook", path="stored.xlsx")

    assert data["values"][0] == [1, 1, ""]
    assert full_loads == []


async def test_an_uncalculated_formula_loads_the_workbook_to_calculate_it(
    call: ToolCall, files: Path, full_loads: list[str]
) -> None:
    workbook = Workbook()
    workbook.worksheets[0].title = "Data"
    workbook.worksheets[0]["A1"] = 1
    workbook.worksheets[0]["B1"] = "=A1+1"
    workbook.save(files / "fresh.xlsx")

    data = await call("read_range", path="fresh.xlsx", sheet="Data")

    assert data["values"] == [[1, 2]]
    assert full_loads != []
