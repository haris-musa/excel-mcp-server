from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


def stored_links(path: Path, sheet: str = "Data") -> dict[str, tuple[str | None, str | None]]:
    worksheet = load_workbook(path)[sheet]
    return {
        cell.coordinate: (cell.hyperlink.target, cell.hyperlink.location)
        for row in worksheet.iter_rows()
        for cell in row
        if cell.hyperlink
    }


async def write_link(call: ToolCall, target: str, **link: object) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Data",
        at="F1",
        rows=[["Docs"]],
        links=[{"cell": "F1", "target": target, **link}],
    )


async def test_external_link_with_display_text_and_tooltip(call: ToolCall, sample: Path) -> None:
    await write_link(call, "https://example.com/a?b=1", tooltip="Open the docs")
    cell = load_workbook(sample)["Data"]["F1"]
    assert cell.value == "Docs"
    assert cell.hyperlink.target == "https://example.com/a?b=1"
    assert cell.hyperlink.tooltip == "Open the docs"
    assert cell.style == "Hyperlink"
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert details["hyperlinks"] == [
        {"cell": "F1", "target": "https://example.com/a?b=1", "tooltip": "Open the docs"}
    ]


async def test_mailto_link(call: ToolCall, sample: Path) -> None:
    await write_link(call, "mailto:team@example.com?subject=Hello there")
    assert stored_links(sample) == {"F1": ("mailto:team@example.com?subject=Hello there", None)}


async def test_internal_links_to_cells_and_names(call: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook.create_sheet("Sheet 2")
    workbook.defined_names["Totals"] = DefinedName("Totals", attr_text="Data!$C$2:$C$5")
    workbook.save(sample)
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Data",
        at="F1",
        rows=[["Go", "Totals"]],
        links=[{"cell": "F1", "target": "#'Sheet 2'!b2"}, {"cell": "G1", "target": "#Totals"}],
    )
    assert stored_links(sample) == {"F1": (None, "'Sheet 2'!B2"), "G1": (None, "Totals")}
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert details["hyperlinks"][0] == {"cell": "F1", "target": "#'Sheet 2'!B2"}


@pytest.mark.parametrize(
    "target",
    [
        "file:///C:/secret.xlsx",
        "\\\\server\\share\\a.xlsx",
        "C:\\Users\\me\\a.xlsx",
        "javascript:alert(1)",
        "ftp://example.com/a",
        "https:///nohost",
        "mailto:",
        "example.com",
        "https://exam\nple.com/",
        "#Nowhere!A1",
        "#NoSuchName",
        "#Data!A1:",
    ],
)
async def test_unsafe_or_invalid_targets_are_rejected(
    call_error: ToolCall, sample: Path, target: str
) -> None:
    await call_error(
        "write_range",
        path="sales.xlsx",
        sheet="Data",
        at="F1",
        rows=[["x"]],
        links=[{"cell": "F1", "target": target}],
    )
    assert stored_links(sample) == {}


async def test_link_outside_the_written_block_is_rejected(
    call_error: ToolCall, sample: Path
) -> None:
    message = await call_error(
        "write_range",
        path="sales.xlsx",
        sheet="Data",
        at="F1",
        rows=[["x"]],
        links=[{"cell": "F2", "target": "https://example.com"}],
    )
    assert "outside the written" in message


async def test_existing_links_survive_edits_that_move_cells(call: ToolCall, sample: Path) -> None:
    await write_link(call, "https://example.com/")
    await call("insert_rows_or_columns", path="sales.xlsx", sheet="Data", axis="rows", start=1)
    assert list(stored_links(sample)) == ["F2"]
    await call("insert_rows_or_columns", path="sales.xlsx", sheet="Data", axis="columns", start=1)
    assert list(stored_links(sample)) == ["G2"]
    await call("delete_rows_or_columns", path="sales.xlsx", sheet="Data", axis="rows", start=1)
    assert list(stored_links(sample)) == ["G1"]
    await call("copy_sheet", path="sales.xlsx", sheet="Data", new_name="Copy")
    assert list(stored_links(sample, "Copy")) == ["G1"]


async def test_links_are_kept_when_copied_and_sorted(call: ToolCall, sample: Path) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Data",
        at="F1",
        rows=[["b"], ["a"]],
        links=[
            {"cell": "F1", "target": "https://b.example/"},
            {"cell": "F2", "target": "https://a.example/"},
        ],
    )
    await call("copy_range", path="sales.xlsx", sheet="Data", range="F1:F2", at="H1")
    await call(
        "sort_range",
        path="sales.xlsx",
        sheet="Data",
        range="F1:F2",
        sort_by=[{"column": "F"}],
        has_header=False,
    )
    links = stored_links(sample)
    assert links["F1"] == ("https://a.example/", None)
    assert links["H1"] == ("https://b.example/", None)


async def test_clear_all_removes_the_link_but_clear_contents_keeps_it(
    call: ToolCall, sample: Path
) -> None:
    await write_link(call, "https://example.com/")
    await call("clear_range", path="sales.xlsx", sheet="Data", range="F1", clear="contents")
    assert list(stored_links(sample)) == ["F1"]
    await call("clear_range", path="sales.xlsx", sheet="Data", range="F1", clear="all")
    assert stored_links(sample) == {}


async def test_links_stay_inside_the_sandbox(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "write_range",
        path="../sales.xlsx",
        sheet="Data",
        at="F1",
        rows=[["x"]],
        links=[{"cell": "F1", "target": "https://example.com"}],
    )
    assert "outside" in message
