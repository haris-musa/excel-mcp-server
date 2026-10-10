"""Edits that would grow a workbook past what the server accepts are refused."""

import zipfile
from pathlib import Path

import pytest
from mcp import Client
from openpyxl import load_workbook
from PIL import Image

from excel_mcp.config import Limits, Settings
from excel_mcp.server import create_server
from tests.conftest import error_text

pytestmark = pytest.mark.anyio


async def _error(files: Path, limits: Limits, tool: str, **arguments: object) -> str:
    async with Client(create_server(Settings(allowed_dirs=[files], limits=limits))) as client:
        result = await client.call_tool(tool, arguments)
    assert result.is_error
    return error_text(result)


async def test_copy_sheet_refuses_sheets_with_too_many_cells(sample: Path, files: Path) -> None:
    message = await _error(
        files,
        Limits(max_copy_cells=10),
        "copy_sheet",
        path="sales.xlsx",
        sheet="Data",
        new_name="C",
    )
    assert "20 cells; at most 10" in message
    assert load_workbook(sample).sheetnames == ["Data", "Report"]


async def test_edits_that_make_the_file_too_large_leave_it_unchanged(
    sample: Path, files: Path
) -> None:
    Image.effect_noise((300, 300), 90).convert("RGB").save(files / "noise.png")
    before = sample.read_bytes()
    message = await _error(
        files,
        Limits(max_file_bytes=len(before) + 1000),
        "insert_image",
        path="sales.xlsx",
        sheet="Report",
        image_path="noise.png",
        at="A1",
    )
    assert "limit is" in message
    assert sample.read_bytes() == before


def _zip_bomb(path: Path, megabytes: int) -> None:
    """A workbook whose sheet is mostly spaces: tiny on disk, huge when unpacked."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        with archive.open("xl/worksheets/sheet1.xml", "w", force_zip64=True) as sheet:
            for _ in range(megabytes):
                sheet.write(b" " * (1 << 20))


@pytest.mark.parametrize(
    ("tool", "arguments"),
    [
        ("describe_workbook", {}),
        ("read_range", {"sheet": "S"}),
        ("write_range", {"sheet": "S", "at": "A1", "rows": [[1]]}),
    ],
)
async def test_a_zip_bomb_already_on_disk_is_refused_before_it_is_read(
    files: Path, tool: str, arguments: dict[str, object]
) -> None:
    _zip_bomb(files / "bomb.xlsx", 50)
    message = await _error(files, Limits(), tool, path="bomb.xlsx", **arguments)
    assert "safety limit" in message
    assert "compressed too densely" in message


async def test_an_oversized_unpacked_workbook_names_the_flag_that_raises_the_limit(
    files: Path,
) -> None:
    _zip_bomb(files / "big.xlsx", 3)
    limits = Limits(max_file_bytes=1 << 20, max_unpack_factor=1, max_compression_ratio=10**6)
    message = await _error(files, limits, "describe_workbook", path="big.xlsx")
    assert "--max-file-mb" in message
