"""The used range read from a sheet's XML text is the one a pass over its cells finds."""

import datetime as dt
import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excel_mcp.operations import cells
from excel_mcp.operations.used_range import used_range_in
from excel_mcp.refs import CellRange
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


def sheet_xml(path: Path) -> bytes:
    with ZipFile(path) as archive:
        return archive.read("xl/worksheets/sheet1.xml")


def scanned(path: Path) -> CellRange:
    """The used range found by parsing every cell."""
    workbook = load_workbook(path, read_only=True)
    sheet = workbook.worksheets[0]
    found = [
        (row, col)
        for row, line in enumerate(sheet.iter_rows(values_only=True), start=1)
        for col, value in enumerate(line, start=1)
        if value is not None
    ]
    workbook.close()
    if not found:
        return CellRange(1, 1, 1, 1)
    rows, columns = [row for row, _ in found], [col for _, col in found]
    return CellRange(min(rows), min(columns), max(rows), max(columns))


def assert_same(path: Path) -> None:
    assert used_range_in(sheet_xml(path)) == scanned(path)


def test_values_formulas_dates_and_text_of_any_kind(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet["C3"] = 0
    sheet["E4"] = "text"
    sheet["B6"] = "=C3+1"
    sheet["D8"] = dt.date(2026, 1, 31)
    sheet["F9"] = True
    sheet["B10"] = ""
    sheet["G12"] = "#N/A"
    workbook.save(tmp_path / "kinds.xlsx")

    assert used_range_in(sheet_xml(tmp_path / "kinds.xlsx")) == CellRange(3, 2, 12, 7)
    assert_same(tmp_path / "kinds.xlsx")


def test_cells_that_only_have_a_style_do_not_count(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet["B2"] = 1
    sheet["C3"] = 2
    sheet["A1"].font = Font(bold=True)
    sheet["Z50"].number_format = "0.00"
    sheet.row_dimensions[80].height = 30
    workbook.save(tmp_path / "styled.xlsx")

    assert used_range_in(sheet_xml(tmp_path / "styled.xlsx")) == CellRange(2, 2, 3, 3)
    assert_same(tmp_path / "styled.xlsx")


def test_empty_sheets(tmp_path: Path) -> None:
    Workbook().save(tmp_path / "empty.xlsx")
    styled = Workbook()
    styled.worksheets[0]["C3"].number_format = "0%"
    styled.save(tmp_path / "formatted.xlsx")

    assert used_range_in(sheet_xml(tmp_path / "empty.xlsx")) == CellRange(1, 1, 1, 1)
    assert used_range_in(sheet_xml(tmp_path / "formatted.xlsx")) == CellRange(1, 1, 1, 1)


def test_stored_results_and_shared_strings_as_excel_writes_them() -> None:
    xml = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        b'<row r="2" spans="1:3"><c r="B2" t="s"><v>0</v></c><c r="C2"><f>1+1</f><v>2</v></c></row>'
        b'<row r="3"><c r="A3" s="1"/><c r="D3" t="str"><f>"a"</f><v></v></c></row>'
        b'<row r="9"><c r="A9" s="2"><v></v></c></row></sheetData></worksheet>'
    )

    assert used_range_in(xml) == CellRange(2, 2, 3, 4)


def test_inline_strings_with_and_without_content() -> None:
    xml = (
        b'<worksheet xmlns="x"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t></t></is></c>'
        b'<c r="B1" t="inlineStr"/></row><row r="2"><c r="C2" t="inlineStr"><is/></c></row>'
        b"</sheetData></worksheet>"
    )

    assert used_range_in(xml) == CellRange(1, 1, 2, 3)


@pytest.mark.parametrize(
    "xml",
    [
        b'<worksheet><sheetData><row r="2"><c r="A2"><v>1</v></c></row>'
        b'<row r="1"><c r="A1"><v>1</v></c></row></sheetData></worksheet>',
        b'<worksheet><sheetData><row><c r="A1"><v>1</v></c></row></sheetData></worksheet>',
        b'<worksheet><sheetData><row r="1"><c s="1" r="A1"><v>1</v></c></row></sheetData></worksheet>',
        b'<worksheet><sheetData><row r="1"><c><v>1</v></c></row></sheetData></worksheet>',
        b"<worksheet></worksheet>",
    ],
)
def test_layouts_it_does_not_read_are_left_to_the_cell_pass(xml: bytes) -> None:
    assert used_range_in(xml) is None


async def test_describe_and_read_agree_with_the_cell_pass(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in range(5, 30):
        sheet.cell(row, 3, row)
        sheet.cell(row, 4 + row % 3, f"=C{row}")
    sheet["A1"].number_format = "0.0"
    workbook.save(files / "book.xlsx")

    described = await call("describe_workbook", path="book.xlsx")
    column = await call("read_range", path="book.xlsx", sheet="Data", range="C:C")

    assert described["sheets"][0]["used_range"] == str(scanned(files / "book.xlsx"))
    assert column["range"] == "C5:C29"


async def test_a_wrong_dimension_does_not_matter(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    workbook.worksheets[0]["B2"] = 1
    workbook.worksheets[0]["D7"] = 2
    workbook.save(files / "book.xlsx")
    with ZipFile(files / "book.xlsx") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    xml = parts["xl/worksheets/sheet1.xml"].decode()
    parts["xl/worksheets/sheet1.xml"] = re.sub(r'<dimension ref="[^"]*"', '<dimension ref="A1"', xml).encode()
    with ZipFile(files / "book.xlsx", "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    described = await call("describe_workbook", path="book.xlsx")

    assert described["sheets"][0]["used_range"] == "B2:D7"


def test_the_cell_pass_is_used_for_workbooks_of_stored_results_only(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = Workbook()
    workbook.worksheets[0]["B2"] = 1
    workbook.save(files / "book.xlsx")
    monkeypatch.setattr(cells, "used_range_in", lambda _: pytest.fail("read the XML text"))
    stored = load_workbook(files / "book.xlsx", read_only=True, data_only=True)

    assert cells.streamed_used_range(stored.worksheets[0]) == CellRange(2, 2, 2, 2)
