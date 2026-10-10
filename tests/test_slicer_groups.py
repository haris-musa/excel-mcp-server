"""Slicers on fields of Excel-made PivotTables, and on grouped dates and numbers."""

import datetime as dt
import itertools
import re
from collections.abc import Callable, Coroutine
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from tests.package_support import copy_fixture, relationships
from tests.slicer_support import ToolCall, add, named, parts_of

pytestmark = pytest.mark.anyio

BY_DATE = {"sheet": "Pivot", "name": "ByDate"}
BY_REGION = {"sheet": "Pivot", "name": "ByRegion"}
CallError = Callable[..., Coroutine[Any, Any, str]]
REGIONS = ["North", "South", "East", "West"]
SALES = [
    [REGIONS[i % 4], "ABC"[i % 3], dt.date(2025, 1, 15) + dt.timedelta(days=i * 50), 10 * (i + 1)]
    for i in range(14)
]


@pytest.fixture
def excel_book(files: Path) -> str:
    """Made by Excel: ByRegion and ByDate (dates grouped by month, quarter and year) have
    caches of their own; ByDate's holds Product and Sales without shared items."""
    return str(copy_fixture(files, "excel_pivot_groups.xlsx"))


def _cache(parts: dict[str, bytes], field: str) -> str:
    return next(c for c in named(parts, "xl/slicerCaches/") if f'sourceName="{field}"' in c)


def _shared_items(definition: str) -> list[str]:
    return re.findall(r"<sharedItems\b[^>]*?(?:/>|>.*?</sharedItems>)", definition, re.S)


def _check_caches(path: str) -> None:
    """Every cache keeps its own records, and its shared items are values, never indexes."""
    parts = parts_of(path)
    targets = []
    for name in sorted(n for n in parts if re.fullmatch(r"xl/pivotCache/pivotCacheDef\w+\.xml", n)):
        definition = parts[name].decode()
        [(kind, records)] = relationships(parts, name).values()
        assert kind == "pivotCacheRecords"
        targets.append(records)
        assert not any("<x " in items or "<x/>" in items for items in _shared_items(definition))
        stored = len(re.findall(r'<cacheField\b(?![^>]*databaseField="0")', definition))
        first_record = re.search(r"<r>(.*?)</r>", parts[records].decode(), re.S)
        assert first_record and len(re.findall(r"<\w+ ", first_record[1])) == stored
    assert len(set(targets)) == len(targets)


@pytest.mark.parametrize("write_first", [False, True])
@pytest.mark.parametrize("order", list(itertools.permutations(["Region", "Product", "Sales"])))
async def test_slicers_in_any_order_on_the_fields_of_an_excel_cache(
    call: ToolCall, excel_book: str, order: tuple[str, ...], write_first: bool
) -> None:
    if write_first:
        await call("write_range", path=excel_book, sheet="Data", at="G1", rows=[[1]])
    for number, field in enumerate(order):
        await add(call, excel_book, target=BY_DATE, field=field, at=f"P{3 + 18 * number}")
        _check_caches(excel_book)
    await add(call, excel_book, target=BY_REGION, field="Product", at="X3", selected_items=["A"])
    _check_caches(excel_book)
    parts = parts_of(excel_book)
    definition = next(
        parts[n].decode() for n in parts if "Months (Date)" in parts[n].decode()[:3000]
    )
    assert '<sharedItems count="3"><s v="A" /><s v="B" /><s v="C" /></sharedItems>' in definition


async def test_every_pivot_cache_is_saved_with_its_own_records(call: ToolCall, files: Path) -> None:
    path = str(files / "caches.xlsx")
    await call("create_workbook", path=path, sheets=["Data", "Pivot"])
    await call("write_range", path=path, sheet="Data", at="A1", rows=[["A", "B", "C"], [1, 2, 3]])
    for number, source in enumerate(["Data!A1:C2", "Data!A1:B2"]):
        await call(
            "create_pivot_table",
            path=path,
            source=source,
            row_fields=["A"],
            value_fields=[{"field": "B"}],
            sheet="Pivot",
            at=f"A{3 + 10 * number}",
        )
    _check_caches(path)
    parts = parts_of(path)
    assert b'<n v="2" /></r>' in parts["xl/pivotCache/pivotCacheRecords2.xml"]


async def test_slicers_and_a_timeline_on_a_grouped_date_are_written_as_excel_writes_them(
    call: ToolCall, excel_book: str
) -> None:
    for number, field in enumerate(["Date", "Years (Date)", "Quarters (Date)", "Months (Date)"]):
        await add(call, excel_book, target=BY_DATE, field=field, at=f"P{3 + 18 * number}")
    await add(call, excel_book, target=BY_DATE, field="Date", at="X3", timeline={"level": "years"})
    parts = parts_of(excel_book)
    on = [f'<i x="{x}" s="1"/>' for x in range(12)]
    date = _cache(parts, "Date")
    assert 'name="Slicer_Date" sourceName="Date"' in date
    assert f'<items count="12">{"".join(on)}</items>' in date
    years = _cache(parts, "Years (Date)")
    assert 'name="Slicer_Years__Date" sourceName="Years (Date)"' in years
    assert (
        '<items count="4"><i x="1" s="1"/><i x="2" s="1"/><i x="0" s="1" nd="1"/>'
        '<i x="3" s="1" nd="1"/></items>' in years
    )
    quarters = _cache(parts, "Quarters (Date)")
    assert 'name="Slicer_Quarters__Date"' in quarters
    assert '<i x="4" s="1"/><i x="0" s="1" nd="1"/><i x="5" s="1" nd="1"/>' in quarters
    months = _cache(parts, "Months (Date)")
    assert 'name="Slicer_Months__Date"' in months
    assert '<i x="12" s="1"/><i x="0" s="1" nd="1"/><i x="13" s="1" nd="1"/>' in months
    timeline = named(parts, "xl/timelineCaches/")[0]
    assert 'sourceName="Date"' in timeline and 'filterType="unknown"' in timeline
    assert "refreshOnLoad" not in parts["xl/pivotCache/pivotCacheDefinition2.xml"].decode()


async def test_choosing_group_items_hides_the_others_and_asks_excel_to_refresh(
    call: ToolCall, excel_book: str
) -> None:
    result = await add(
        call, excel_book, target=BY_DATE, field="Years (Date)", selected_items=["2025"]
    )
    assert "recalculates them when the file is opened" in result["note"]
    await add(
        call,
        excel_book,
        target=BY_DATE,
        field="Quarters (Date)",
        selected_items=["Qtr1", "Qtr2"],
        sort="descending",
        at="P20",
    )
    parts = parts_of(excel_book)
    first, second = _cache(parts, "Years (Date)"), _cache(parts, "Quarters (Date)")
    assert (
        '<items count="4"><i x="1" s="1"/><i x="2"/><i x="0" nd="1"/><i x="3" nd="1"/></items>'
        in first
    )
    assert '<tabular pivotCacheId="' in second and 'sortOrder="descending"' in second
    assert (
        '<items count="6"><i x="4"/><i x="3"/><i x="2" s="1"/><i x="1" s="1"/>'
        '<i x="5" nd="1"/><i x="0" nd="1"/></items>' in second
    )
    pivot = next(n for n in parts if "pivotTables/pivotTable" in n and b"ByDate" in parts[n])
    fields = re.findall(r"<pivotField\b.*?</pivotField>", parts[pivot].decode(), re.S)
    hidden = [re.findall(r'<item[^>]*\bh="1"[^>]*\bx="(\d+)"', f) for f in fields[-2:]]
    assert hidden == [["0", "3", "4", "5"], ["0", "2", "3"]]
    assert 'refreshOnLoad="1"' in parts["xl/pivotCache/pivotCacheDefinition2.xml"].decode()
    info = await call("describe_sheet", path=excel_book, sheet="Pivot")
    assert info["pivot_tables"][1]["date_groups"] == [
        "Months (Date)",
        "Quarters (Date)",
        "Years (Date)",
    ]
    shown = {s["name"]: s["selected_items"] for s in info["slicers"]}
    assert shown == {"Years (Date)": ["2025"], "Quarters (Date)": ["Qtr1", "Qtr2"]}


async def test_a_timeline_on_a_grouped_date_limits_the_period(
    call: ToolCall, excel_book: str
) -> None:
    await add(
        call,
        excel_book,
        target=BY_DATE,
        field="Date",
        timeline={"start": "2025-03-01", "end": "2025-09-30"},
    )
    parts = parts_of(excel_book)
    assert (
        '<selection startDate="2025-03-01T00:00:00" endDate="2025-09-30T00:00:00"/>'
        in (named(parts, "xl/timelineCaches/")[0])
    )
    pivot = next(
        parts[n].decode() for n in parts if "pivotTables/pivotTable" in n and b"ByDate" in parts[n]
    )
    assert '<filter fld="2" type="dateBetween"' in pivot
    assert 'refreshOnLoad="1"' in parts["xl/pivotCache/pivotCacheDefinition2.xml"].decode()


@pytest.mark.parametrize(
    ("arguments", "message"),
    [
        (
            {"field": "Product", "timeline": {}},
            "'Product' is not a date field, so it cannot have a timeline. Date fields: 'Date'",
        ),
        (
            {"field": "Months (Date)", "timeline": {}},
            "'Months (Date)' is not a date field, so it cannot have a timeline. "
            "Date fields: 'Date'",
        ),
        (
            {"field": "Month"},
            "Field 'Month' not found. Fields: 'Region', 'Product', 'Date', 'Sales', "
            "'Months (Date)', 'Quarters (Date)', 'Years (Date)'.",
        ),
        (
            {"field": "Years (Date)", "selected_items": ["2031"]},
            "Field 'Years (Date)' has no item '2031'. Items: ['<01-01-25', '2025', '2026'",
        ),
    ],
)
async def test_errors_on_grouped_fields_say_what_to_use_instead(
    call_error: CallError, excel_book: str, arguments: dict[str, Any], message: str
) -> None:
    error = await call_error(
        "add_slicer", path=excel_book, sheet="Pivot", target=BY_DATE, at="P3", **arguments
    )
    assert message in error


# -- PivotTables this server made: the figures follow the items ---------------------------------


async def _made(call: ToolCall, files: Path, **settings: Any) -> str:
    path = str(files / "made.xlsx")
    await call("create_workbook", path=path, sheets=["Data", "Pivot"])
    header = [["Region", "Product", "Date", "Amount"]]
    rows = [[r, p, d.isoformat(), a] for r, p, d, a in SALES]
    await call("write_range", path=path, sheet="Data", at="A1", rows=header + rows)
    await call(
        "create_pivot_table",
        path=path,
        source="Data!A1:D15",
        row_fields=["Date"] if "group_numbers" not in settings else ["Amount"],
        value_fields=[{"field": "Amount"}],
        field_settings=[
            {"field": "Date", "group_dates": ["years", "quarters", "months"]}
            if "group_numbers" not in settings
            else {"field": "Amount", "group_numbers": settings["group_numbers"]}
        ],
        sheet="Pivot",
        at="A3",
        name="ByDate",
    )
    return path


def _grand_total(path: str) -> Any:
    cells = load_workbook(path)["Pivot"]
    [row] = [r for r in cells.iter_rows(min_row=3) if r[0].value == "Grand Total"]
    return next(c.value for c in reversed(row) if c.value is not None)


def _total(keep: Callable[[dt.date, str], bool]) -> int:
    return sum(amount for region, _, day, amount in SALES if keep(day, region))


async def test_group_slicers_and_a_timeline_filter_a_pivot_table_made_here(
    call: ToolCall, files: Path
) -> None:
    path = await _made(call, files)
    assert _grand_total(path) == _total(lambda *_: True)
    await add(call, path, target=BY_DATE, field="Years (Date)", selected_items=["2025"], at="P3")
    assert _grand_total(path) == _total(lambda day, _: day.year == 2025)
    await add(
        call, path, target=BY_DATE, field="Months (Date)", selected_items=["Mar", "Jun", "Nov"],
        at="P20",
    )  # fmt: skip
    in_months = lambda day, _: day.year == 2025 and day.month in (3, 6, 11)  # noqa: E731
    assert _grand_total(path) == _total(in_months)
    result = await add(
        call,
        path,
        target=BY_DATE,
        field="Date",
        timeline={"start": "2025-04-01", "end": "2025-12-31"},
        at="X3",
    )
    assert "note" not in result
    assert _grand_total(path) == _total(lambda day, r: in_months(day, r) and day.month > 3)
    await add(call, path, target=BY_DATE, field="Region", selected_items=["East", "West"], at="X20")
    assert _grand_total(path) == _total(
        lambda day, r: in_months(day, r) and day.month > 3 and r in ("East", "West")
    )


async def test_dates_of_a_grouped_field_can_be_chosen_one_by_one(
    call: ToolCall, files: Path
) -> None:
    path = await _made(call, files)
    wanted = [SALES[1][2].isoformat(), SALES[4][2].isoformat()]
    await add(call, path, target=BY_DATE, field="Date", selected_items=wanted, at="P3")
    assert _grand_total(path) == SALES[1][3] + SALES[4][3]
    await add(call, path, target=BY_DATE, field="Years (Date)", selected_items=["2025"], at="P20")
    assert _grand_total(path) == SALES[1][3] + SALES[4][3]
    await call("delete_slicer", path=path, sheet="Pivot", name="Date")
    assert _grand_total(path) == _total(lambda day, _: day.year == 2025)


async def test_a_slicer_works_on_a_field_grouped_into_ranges(call: ToolCall, files: Path) -> None:
    path = await _made(call, files, group_numbers={"by": 50, "start": 0, "end": 150})
    await add(call, path, target=BY_DATE, field="Amount", selected_items=["0-49", "50-99"], at="P3")
    assert _grand_total(path) == sum(a for *_, a in SALES if a < 100)
    cache = named(parts_of(path), "xl/slicerCaches/")[0]
    # as in Excel, ranges are listed sorted as text: 0-49, 100-149, 50-99
    assert '<i x="1" s="1"/><i x="3"/><i x="2" s="1"/><i x="0" nd="1"/><i x="4" nd="1"/>' in cache
    assert 'sourceName="Amount"' in cache
