"""Shared by the slicer tests: a workbook with a table and a PivotTable, and readers."""

from collections.abc import Callable, Coroutine
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from tests.package_support import read_parts, text

ToolCall = Callable[..., Coroutine[Any, Any, Any]]
CallError = Callable[..., Coroutine[Any, Any, str]]
SALES = [
    ["Region", "Product", "Date", "Amount"],
    ["North", "Apples", "2025-01-10", 10],
    ["South", "Apples", "2025-02-10", 20],
    ["North", "Pears", "2025-03-10", 30],
    ["South", "Pears", "2025-04-10", 40],
    ["East", "Kiwis", "2025-05-10", 50],
    ["West", "Kiwis", "2025-06-10", 60],
    ["East", "Apples", "2025-07-10", 70],
    ["West", "Pears", "2025-08-10", 80],
]
PIVOT = {"sheet": "Pivot", "name": "PivotSales"}
TABLE = {"sheet": "Data", "name": "Sales"}


async def build_book(call: ToolCall, files: Path) -> str:
    """Sales in a table on Data, summarized by product in a PivotTable on Pivot."""
    path = str(files / "book.xlsx")
    await call("create_workbook", path=path, sheets=["Data", "Pivot"])
    await call("write_range", path=path, sheet="Data", start_cell="A1", rows=SALES)
    await call("create_table", path=path, sheet="Data", range="A1:D9", name="Sales")
    await call(
        "create_pivot_table",
        path=path,
        source_sheet="Data",
        source_range="A1:D9",
        rows=["Product"],
        values=[{"field": "Amount"}],
        target_sheet="Pivot",
        target_cell="A3",
        name="PivotSales",
    )
    return path


def pivot_cells(path: str, sheet: str = "Pivot") -> list[list[Any]]:
    cells = load_workbook(path)[sheet]
    return [[c.value for c in row] for row in cells.iter_rows(min_row=3, max_row=8, max_col=2)]


def parts_of(path: str) -> dict[str, bytes]:
    return read_parts(Path(path))


def named(parts: dict[str, bytes], folder: str) -> list[str]:
    return sorted(text(parts, n) for n in parts if n.startswith(folder) and n.endswith(".xml"))


async def add(call: ToolCall, path: str, **arguments: Any) -> str:
    arguments.setdefault("sheet", "Pivot")
    arguments.setdefault("source", PIVOT)
    arguments.setdefault("cell", "E3")
    return await call("add_slicer", path=path, **arguments)
