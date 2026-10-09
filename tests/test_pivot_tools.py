import datetime as dt
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_mcp.operations.pivot_index import sheet_pivots
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

PIVOT = {
    "path": "sales.xlsx",
    "source": "Data!A1:D5",
    "sheet": "Report",
    "at": "A3",
}


def _grid(path: Path, ref: str) -> list[list[object]]:
    sheet = load_workbook(path)["Report"]
    return [[cell.value for cell in row] for row in sheet[ref]]


async def test_rows_and_values(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table", **PIVOT, row_fields=["region"], value_fields=[{"field": "Units"}]
    )
    assert _grid(sample, "A3:B6") == [
        ["Region", "Sum of Units"],
        ["North", 17],
        ["South", 8],
        ["Grand Total", 25],
    ]
    pivot = sheet_pivots(load_workbook(sample)["Report"])[0]
    assert pivot.location.ref == "A3:B6"
    assert [item.t for item in pivot.rowItems] == ["data", "data", "grand"]
    assert pivot.cache.cacheSource.worksheetSource.ref == "A1:D5"
    assert pivot.cache.records.count == 4


async def test_subtotals_and_several_values(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region", "Product"],
        value_fields=[{"field": "Units"}, {"field": "Price", "function": "average"}],
    )
    assert _grid(sample, "A3:D11") == [
        [None, None, "Values", None],
        ["Region", "Product", "Sum of Units", "Average of Price"],
        ["North", "Apples", 10, 1.5],
        [None, "Pears", 7, 2],
        ["North Total", None, 17, 1.75],
        ["South", "Apples", 5, 1.5],
        [None, "Pears", 3, 2],
        ["South Total", None, 8, 1.75],
        ["Grand Total", None, 25, 1.75],
    ]


async def test_column_field_with_totals(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **PIVOT,
        row_fields=["Region"],
        column_fields=["Product"],
        value_fields=[{"field": "Units", "function": "max"}],
    )
    assert _grid(sample, "A3:E7") == [
        ["Max of Units", "Product", None, None, None],
        ["Region", "Apples", "Pears", "Grand Total", None],
        ["North", 10, 7, 10, None],
        ["South", 5, 3, 5, None],
        ["Grand Total", 10, 7, 10, None],
    ]


async def test_filters_sit_above_the_table(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table",
        **{**PIVOT, "at": "A1"},
        row_fields=["Region"],
        filter_fields=["Product"],
        value_fields=[{"field": "Units"}],
    )
    assert _grid(sample, "A1:B2") == [["Product", "(All)"], [None, None]]
    assert sheet_pivots(load_workbook(sample)["Report"])[0].location.ref == "A3:B6"


async def test_blanks_dates_and_text_case(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in [
        ["Who", "Day", "N"],
        ["ann", dt.date(2026, 1, 2), 1],
        ["Ann", dt.date(2026, 1, 1), None],
        [None, dt.date(2026, 1, 1), 4],
    ]:
        sheet.append(row)
    workbook.create_sheet("Out")
    workbook.save(files / "d.xlsx")
    await call(
        "create_pivot_table",
        path="d.xlsx",
        source="Data!A1:C4",
        row_fields=["Who"],
        column_fields=["Day"],
        value_fields=[{"field": "N", "function": "count"}],
        sheet="Out",
        at="A1",
    )
    out = load_workbook(files / "d.xlsx")["Out"]
    assert [cell.value for cell in out[2]] == [
        "Who",
        dt.datetime(2026, 1, 1),
        dt.datetime(2026, 1, 2),
        "Grand Total",
    ]
    assert [cell.value for cell in out[3]] == ["ann", None, 1, 1]
    assert [cell.value for cell in out[4]] == ["(blank)", 1, None, 1]


async def test_pivot_survives_other_edits(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table", **PIVOT, row_fields=["Region"], value_fields=[{"field": "Units"}]
    )
    await call("write_range", path="sales.xlsx", sheet="Report", at="F1", rows=[[1]])
    await call("create_sheet", path="sales.xlsx", new_name="Other")
    sheet = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert sheet["pivot_tables"] == [
        {"name": "PivotTable1", "range": "A3:B6", "source": "Data!A1:D5"}
    ]
    with zipfile.ZipFile(sample) as archive:
        assert "xl/pivotTables/pivotTable1.xml" in archive.namelist()
        assert "xl/pivotCache/pivotCacheRecords1.xml" in archive.namelist()


async def test_two_pivots_get_separate_caches(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table", **PIVOT, row_fields=["Region"], value_fields=[{"field": "Units"}]
    )
    await call(
        "create_pivot_table",
        **{**PIVOT, "at": "A12"},
        row_fields=["Product"],
        value_fields=[{"field": "Units"}],
        name="Second",
    )
    pivots = sheet_pivots(load_workbook(sample)["Report"])
    assert [pivot.name for pivot in pivots] == ["PivotTable1", "Second"]
    assert len({pivot.cacheId for pivot in pivots}) == 2


async def test_delete_pivot_table(call: ToolCall, sample: Path) -> None:
    await call(
        "create_pivot_table", **PIVOT, row_fields=["Region"], value_fields=[{"field": "Units"}]
    )
    await call("delete_pivot_table", path="sales.xlsx", sheet="Report", name="pivottable1")
    assert sheet_pivots(load_workbook(sample)["Report"]) == []
    assert _grid(sample, "A3:B4") == [[None, None], [None, None]]
    with zipfile.ZipFile(sample) as archive:
        assert not [name for name in archive.namelist() if "pivot" in name]


async def test_rejects_bad_arguments(call_error: ToolCall, sample: Path) -> None:
    base = {**PIVOT, "value_fields": [{"field": "Units"}]}
    message = await call_error("create_pivot_table", **base, row_fields=["Country"])
    assert "Available fields" in message
    message = await call_error(
        "create_pivot_table", **base, row_fields=["Region"], column_fields=["Region"]
    )
    assert "once" in message
    message = await call_error(
        "create_pivot_table",
        **{**PIVOT, "value_fields": [{"field": "Region"}]},
        row_fields=["Product"],
    )
    assert "count" in message
    message = await call_error(
        "create_pivot_table", **{**base, "source": "Data!A2:D5"}, row_fields=["Region"]
    )
    assert "Header cell" in message
    message = await call_error(
        "create_pivot_table",
        **{**base, "sheet": "Data", "at": "B3"},
        row_fields=["Region"],
    )
    assert "own source" in message or "already holds" in message


async def test_rejects_overlap_and_unknown_pivot(
    call_error: ToolCall, call: ToolCall, sample: Path
) -> None:
    await call(
        "create_pivot_table", **PIVOT, row_fields=["Region"], value_fields=[{"field": "Units"}]
    )
    message = await call_error(
        "create_pivot_table",
        **{**PIVOT, "at": "B4"},
        row_fields=["Product"],
        value_fields=[{"field": "Units"}],
    )
    assert "overlaps the PivotTable" in message
    message = await call_error("delete_pivot_table", path="sales.xlsx", sheet="Report", name="Nope")
    assert "PivotTable1" in message


async def test_rejects_formulas_and_mixed_columns(call_error: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook["Data"]["C2"] = "=1+1"
    workbook.save(sample)
    message = await call_error(
        "create_pivot_table", **PIVOT, row_fields=["Region"], value_fields=[{"field": "Units"}]
    )
    assert "formula" in message
    workbook["Data"]["C2"] = "ten"
    workbook.save(sample)
    message = await call_error(
        "create_pivot_table", **PIVOT, row_fields=["Region"], value_fields=[{"field": "Units"}]
    )
    assert "mixes" in message


async def test_pivot_paths_are_confined(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_pivot_table",
        **{**PIVOT, "path": "../sales.xlsx"},
        row_fields=["Region"],
        value_fields=[{"field": "Units"}],
    )
    assert "outside" in message.lower() or "not allowed" in message.lower()
