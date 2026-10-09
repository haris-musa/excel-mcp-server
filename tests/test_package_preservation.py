"""What an edit through the server keeps of content openpyxl cannot model.

The fixtures were saved by Microsoft Excel; see tests/fixtures/README.md.
"""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.conftest import ToolCall
from tests.package_builders import threaded_workbook, value_metadata_workbook
from tests.package_support import (
    assert_package_is_consistent,
    copy_fixture,
    read_parts,
    relationships,
    sheet_ids,
    sheet_part,
    text,
)

pytestmark = pytest.mark.anyio

BOOK = {"path": "book.xlsx"}
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"


async def _edit_and_read(call: ToolCall, files: Path, fixture: str, sheet: str) -> dict[str, bytes]:
    copy_fixture(files, fixture)
    await call("write_range", **BOOK, sheet=sheet, at="Z1", rows=[[1]])
    parts = read_parts(files / "book.xlsx")
    assert_package_is_consistent(parts)
    load_workbook(files / "book.xlsx").close()
    return parts


async def test_sparklines_and_extended_rules_survive_edits(call: ToolCall, files: Path) -> None:
    copy_fixture(files, "excel_sparklines.xlsx")
    original = text(read_parts(files / "book.xlsx"), "xl/worksheets/sheet2.xml")

    await call("write_range", **BOOK, sheet="Data", at="Z1", rows=[[1]])
    await call("insert_rows_or_columns", **BOOK, sheet="Data", axis="rows", start=3)
    parts = read_parts(files / "book.xlsx")
    sheet = text(parts, sheet_part(parts, "Data"))

    assert sheet.count("<x14:sparklineGroup ") == 3
    assert sheet.count("<x14:sparkline>") == 21  # each group grows over the inserted row
    assert sheet.count("<x14:cfRule ") == 2
    assert "<xm:f>Lists!$A$1:$A$3</xm:f>" in sheet
    assert sheet.count("<conditionalFormatting ") == 3
    assert len(re.findall(r"<ext [^>]*uri=", sheet)) == len(re.findall(r"<ext [^>]*uri=", original))
    assert_package_is_consistent(parts)


async def test_data_bar_stays_linked_to_its_extended_rule(call: ToolCall, files: Path) -> None:
    parts = await _edit_and_read(call, files, "excel_sparklines.xlsx", "Data")
    sheet = text(parts, sheet_part(parts, "Data"))

    linked = re.findall(r"<x14:id>(\{[^}]*\})</x14:id>", sheet)
    assert len(linked) == 1
    assert f'<x14:cfRule type="dataBar" id="{linked[0]}"' in sheet


async def test_new_chart_types_and_shapes_survive(call: ToolCall, files: Path) -> None:
    parts = await _edit_and_read(call, files, "excel_chartex.xlsx", "Data")
    drawing = text(parts, "xl/drawings/drawing1.xml")
    sheet_rels = relationships(parts, sheet_part(parts, "Data"))
    drawing_rels = relationships(parts, "xl/drawings/drawing1.xml")

    assert any(kind == "drawing" for kind, _ in sheet_rels.values())
    assert (
        drawing.count('graphicData uri="http://schemas.microsoft.com/office/drawing/2014/chartex"')
        == 4
    )
    assert sum(1 for kind, _ in drawing_rels.values() if kind == "chartEx") == 4
    for _, target in drawing_rels.values():
        assert target in parts
    for number in range(1, 5):
        chart = relationships(parts, f"xl/charts/chartEx{number}.xml")
        assert {kind for kind, _ in chart.values()} == {"chartStyle", "chartColorStyle"}
    assert "A note" in drawing and 'name="Rect"' in drawing
    assert drawing.count("<c:chart ") == 1  # the column chart openpyxl rewrote
    ids = re.findall(r"<(?:xdr:)?cNvPr id=\"(\d+)\"", drawing)
    assert len(ids) == len(set(ids) - {"0"}) + ids.count("0")
    assert text(parts, "xl/workbook.xml").count('name="_xlchart.v1.') >= 6


async def test_slicers_and_timelines_follow_the_sheet_numbers(call: ToolCall, files: Path) -> None:
    copy_fixture(files, "excel_slicers.xlsx")
    await call("create_sheet", **BOOK, new_name="First", position=1)
    parts = read_parts(files / "book.xlsx")
    ids = sheet_ids(parts)

    assert ids == {"First": 1, "Pivot": 2, "Data": 3}
    caches = {n: text(parts, n) for n in parts if "slicerCaches/" in n or "timelineCaches/" in n}
    assert len(caches) == 4
    pivot_users = [xml for xml in caches.values() if "<pivotTables>" in xml]
    assert len(pivot_users) == 3
    for xml in pivot_users:
        assert re.search(rf'<pivotTable tabId="{ids["Pivot"]}" name="PivotSales"/>', xml)
    table_cache = next(xml for xml in caches.values() if "tableSlicerCache" in xml)
    assert 'tableId="1"' in table_cache
    assert_package_is_consistent(parts)


async def test_slicer_parts_stay_reachable_from_their_owners(call: ToolCall, files: Path) -> None:
    parts = await _edit_and_read(call, files, "excel_slicers.xlsx", "Data")
    workbook = text(parts, "xl/workbook.xml")

    kinds = {"slicerCache", "timelineCache"}
    reachable = {t for k, t in relationships(parts, "xl/workbook.xml").values() if k in kinds}
    assert len(reachable) == 4
    for rel_id in re.findall(r'<x1[45]:(?:slicerCache|timelineCacheRef) r:id="(rId\d+)"', workbook):
        assert relationships(parts, "xl/workbook.xml")[rel_id][1] in reachable
    for sheet in ("Pivot", "Data"):
        part = sheet_part(parts, sheet)
        kinds = {kind for kind, _ in relationships(parts, part).values()}
        assert "slicer" in kinds
        assert text(parts, part).count("slicerList") >= 2
    for name in ("Slicer_Product", "Slicer_Region", "Slicer_Region1", "NativeTimeline_Date"):
        assert f'definedName name="{name}"' in workbook
    pivot = text(parts, "xl/pivotCache/pivotCacheDefinition1.xml")
    assert 'pivotCacheDefinition pivotCacheId="' in pivot  # what slicer caches find the cache by
    assert "<x14:pivotTableDefinition" in text(parts, "xl/pivotTables/pivotTable1.xml")


async def test_form_controls_and_notes_share_the_vml_drawing(call: ToolCall, files: Path) -> None:
    parts = await _edit_and_read(call, files, "excel_controls.xlsx", "Data")
    part = sheet_part(parts, "Data")
    sheet = text(parts, part)
    vml_names = [t for k, t in relationships(parts, part).values() if k == "vmlDrawing"]

    assert len(vml_names) == 1
    vml = text(parts, vml_names[0])
    assert 'ObjectType="Checkbox"' in vml and 'ObjectType="Drop"' in vml
    assert 'ObjectType="Note"' in vml
    shape_ids = re.findall(r'(?:id|o:spid)="(_x0000_s\d+)"', vml)
    assert sorted(shape_ids) == ["_x0000_s1025", "_x0000_s1026", "_x0000_s1027"]
    controls = re.findall(r'<control shapeId="(\d+)"', sheet)
    assert controls == ["1025", "1026"]
    assert all(f"_x0000_s{number}" in vml for number in controls)
    assert sum(1 for k, _ in relationships(parts, part).values() if k == "ctrlProp") == 2
    assert "<legacyDrawing " in sheet and "<controls>" in sheet


async def test_threaded_comments_follow_their_notes(call: ToolCall, files: Path) -> None:
    threaded_workbook(files / "book.xlsx")
    await call("insert_rows_or_columns", **BOOK, sheet="Data", axis="rows", start=1, count=2)
    parts = read_parts(files / "book.xlsx")
    threads = text(parts, "xl/threadedComments/threadedComment1.xml")

    assert re.findall(r'<threadedComment ref="([A-Z]+\d+)"', threads) == ["B4", "B4", "C5"]
    assert "xl/persons/person.xml" in parts
    kinds = {k for k, _ in relationships(parts, sheet_part(parts, "Data")).values()}
    assert "threadedComment" in kinds
    assert_package_is_consistent(parts)

    await call("delete_note", **BOOK, sheet="Data", cell="B4")
    threads = text(read_parts(files / "book.xlsx"), "xl/threadedComments/threadedComment1.xml")
    assert re.findall(r'<threadedComment ref="([A-Z]+\d+)"', threads) == ["C5"]

    await call("delete_note", **BOOK, sheet="Data", cell="C5")
    parts = read_parts(files / "book.xlsx")
    assert not [n for n in parts if "threadedComments/" in n]
    assert "xl/persons/person.xml" in parts


async def test_cell_metadata_follows_the_values_it_describes(call: ToolCall, files: Path) -> None:
    value_metadata_workbook(files / "book.xlsx")
    await call("write_range", **BOOK, sheet="Sheet", at="B1", rows=[["changed"]])
    parts = read_parts(files / "book.xlsx")
    sheet = text(parts, sheet_part(parts, "Sheet"))

    assert re.search(r'<c r="A1"[^>]* vm="1"', sheet)
    assert not re.search(r'<c r="B1"[^>]* vm=', sheet)
    assert "<valueMetadata" in text(parts, "xl/metadata.xml")
    assert_package_is_consistent(parts)


@pytest.mark.parametrize(
    "fixture",
    ["excel_sparklines.xlsx", "excel_chartex.xlsx", "excel_slicers.xlsx", "excel_controls.xlsx"],
)
async def test_editing_again_changes_nothing_that_was_preserved(
    call: ToolCall, files: Path, fixture: str
) -> None:
    copy_fixture(files, fixture)
    sheet = "Pivot" if "slicers" in fixture else "Data"
    await call("write_range", **BOOK, sheet=sheet, at="Z1", rows=[[1]])
    first = read_parts(files / "book.xlsx")
    await call("write_range", **BOOK, sheet=sheet, at="Z2", rows=[[2]])
    await call("write_range", **BOOK, sheet=sheet, at="Z3", rows=[[3]])
    third = read_parts(files / "book.xlsx")

    assert sorted(first) == sorted(third)
    for name, data in first.items():
        if "worksheets/sheet" in name and not name.endswith(".rels"):
            tail = lambda xml: xml.decode()[xml.decode().rindex("</sheetData>") :]  # noqa: E731
            assert tail(data) == tail(third[name]), name
        elif not name.startswith("xl/worksheets/") and name not in (
            "docProps/core.xml",
            "xl/charts/chart1.xml",
        ):
            assert data == third[name], name


async def test_a_table_slicer_follows_its_table_when_tables_are_renumbered(
    call: ToolCall, files: Path
) -> None:
    copy_fixture(files, "excel_slicers.xlsx")
    await call("write_range", **BOOK, sheet="Pivot", at="H1", rows=[["a", "b"], [1, 2]])
    await call("create_table", **BOOK, sheet="Pivot", range="H1:I2", name="Extra")
    parts = read_parts(files / "book.xlsx")

    ids = {}
    for name in parts:
        if name.startswith("xl/tables/"):
            table = text(parts, name)
            ids[re.search(r'\bname="(\w+)"', table)[1]] = re.search(r'\bid="(\d+)"', table)[1]  # type: ignore[index]
    caches = [text(parts, n) for n in parts if "slicerCaches/" in n]
    table_slicer = next(xml for xml in caches if "tableSlicerCache" in xml)

    assert ids == {"Extra": "1", "Sales": "2"}
    assert f'tableId="{ids["Sales"]}"' in table_slicer
