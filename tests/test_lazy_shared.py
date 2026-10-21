# pyright: reportAttributeAccessIssue=false, reportOptionalMemberAccess=false
# The tests look at openpyxl internals and at what the lazy cells hold.
"""Sheets with shared formulas (what Excel writes when a formula is filled down) are edited
without parsing the rows an edit leaves alone, and give what loading the whole sheet gives.

Blocks are four rows here, so every group spans several blocks.
"""

import re
import shutil
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from excel_mcp import lazy_edit
from excel_mcp.cell_budget import CellBudget
from tests.conftest import ToolCall, lazily
from tests.package_support import read_parts, sheet_part
from tests.test_lazy_edit import snapshot

pytestmark = pytest.mark.anyio

ROWS = 40
FILLED = 24  # the amounts are filled down to here only
_CELL = re.compile(r'<c r="([A-Z]+)(\d+)"[^>]*>(<f[^>]*/>|<f[^>]*>[^<]*</f>)')
_ATTRIBUTES = re.compile(r'(\w+)="([^"]*)"')


def make_shared_book(path: Path) -> None:
    """Amounts (D) and labels (E) filled down, and a grid (F:G) filled down and right."""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Id", "Qty", "Price", "Amount", "Label", "Gx", "Gy"])
    for row in range(2, ROWS + 1):
        sheet.append(
            [
                row - 1,
                row % 7,
                row * 1.5,
                f"=B{row}*C{row}",
                f'=IF(D{row}>10,"big & bold","small")',
                f"=$A{row}*B$1",
                f"=$A{row}*C$1",
            ]
        )
    workbook.save(path)
    share(path, 0, "D", 2, FILLED)
    share(path, 1, "E", 2, ROWS)
    share(path, 2, "FG", 2, ROWS)


def share(path: Path, group: int, columns: str, first: int, last: int, sheet: str = "Data") -> None:
    """Rewrite the cells of a range the way Excel stores a formula filled over it."""
    parts = read_parts(path)
    name = sheet_part(parts, sheet)
    ref = f"{columns[0]}{first}:{columns[-1]}{last}"

    def convert(found: re.Match[str]) -> str:
        column, row = found[1], int(found[2])
        if column not in columns or not first <= row <= last:
            return found[0]
        text = re.search(r">([^<]*)</f>", found[3])
        if (column, row) == (columns[0], first):
            formula = f'<f t="shared" ref="{ref}" si="{group}">{text[1] if text else ""}</f>'
        else:
            formula = f'<f t="shared" si="{group}"/>'
        return found[0].replace(found[3], formula)

    xml = parts[name].decode()
    xml = _CELL.sub(convert, xml)
    xml = re.sub(r"(</f>|<f[^>]*/>)<v ?/>", r"\1<v>1</v>", xml)
    parts[name] = xml.encode()
    with ZipFile(path, "w") as archive:
        for part, content in parts.items():
            archive.writestr(part, content)


def groups(path: Path, sheet: str = "Data") -> dict[str, dict[str, Any]]:
    """The shared formulas a saved sheet has: each one's master, its range, its children."""
    xml = read_parts(path)[sheet_part(read_parts(path), sheet)].decode()
    found: dict[str, dict[str, Any]] = {}
    for cell in _CELL.finditer(xml):
        attributes = dict(_ATTRIBUTES.findall(cell[3].split(">")[0]))
        if attributes.get("t") != "shared":
            continue
        group = found.setdefault(attributes["si"], {"master": None, "children": []})
        if "ref" in attributes:
            group["master"] = (cell[1] + cell[2], attributes["ref"])
        else:
            group["children"].append(cell[1] + cell[2])
    return found


def assert_groups_are_whole(path: Path) -> None:
    """Every child has a master, and it is in the range of that master."""
    for si, group in groups(path).items():
        if group["children"]:
            assert group["master"], f"children of {si} have no master"
            first, last = group["master"][1].split(":")
            top, bottom = int(first[1:]), int(last[1:])
            assert all(top <= int(child[1:]) <= bottom for child in group["children"])


EDITS: list[tuple[str, dict[str, Any]]] = [
    ("write_range", {"sheet": "Data", "at": "D22", "rows": [[5]]}),
    ("write_range", {"sheet": "Data", "at": "E22", "rows": [["=A22+1"]]}),
    ("write_range", {"sheet": "Data", "at": "D2", "rows": [[7]]}),
    ("write_range", {"sheet": "Data", "at": "E2", "rows": [["=A2&B2"]]}),
    ("write_range", {"sheet": "Data", "at": "F2", "rows": [["=1+1"]]}),
    ("write_range", {"sheet": "Data", "at": "A2", "rows": [[100, 8]]}),
    ("write_range", {"sheet": "Data", "at": "A41", "rows": [[1, 2, 3]]}),
    (
        "format_range",
        {"sheet": "Data", "range": "D2:E40", "style": {"bold": True, "fill_color": "FFFF00"}},
    ),
    ("clear_range", {"sheet": "Data", "range": "D2:D3", "clear": "all"}),
    ("clear_range", {"sheet": "Data", "range": "E30:E31", "clear": "contents"}),
    ("copy_range", {"sheet": "Data", "range": "D2:E3", "at": "H44"}),
    ("copy_range", {"sheet": "Data", "range": "D30:E31", "at": "D10"}),
    ("insert_rows_or_columns", {"sheet": "Data", "axis": "rows", "start": 10, "count": 3}),
    ("delete_rows_or_columns", {"sheet": "Data", "axis": "rows", "start": 2, "count": 1}),
    ("insert_rows_or_columns", {"sheet": "Data", "axis": "columns", "start": 2, "count": 1}),
    (
        "create_chart",
        {"sheet": "Data", "chart_type": "column", "source": "Data!A1:B12", "at": "J2"},
    ),
]


@pytest.mark.parametrize("edit", range(len(EDITS)))
async def test_every_edit_gives_what_loading_the_whole_sheet_gives(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch, edit: int
) -> None:
    make_shared_book(files / "lazy.xlsx")
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")
    tool, arguments = EDITS[edit]

    with lazily(monkeypatch):
        await call(tool, path="lazy.xlsx", **arguments)
    await call(tool, path="full.xlsx", **arguments)

    assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx")
    assert_groups_are_whole(files / "lazy.xlsx")


def test_a_block_with_children_and_not_their_master_has_their_formulas(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_shared_book(files / "book.xlsx")

    with lazily(monkeypatch):
        book, edit = lazy_edit.load(files / "book.xlsx", CellBudget(3_000_000))
        assert edit is not None
        cells = book["Data"]._cells
        assert cells[30, 4].value == "=B30*C30"
        assert cells[30, 5].value == '=IF(D30>10,"big & bold","small")'
        assert (cells[30, 6].value, cells[30, 7].value) == ("=$A30*B$1", "=$A30*C$1")
        assert cells[2, 4].value == "=B2*C2"
        assert cells.loaded == {30 // 4, 2 // 4}


async def test_the_values_read_back_are_the_same(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_shared_book(files / "lazy.xlsx")
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")

    with lazily(monkeypatch):
        lazy = await call("read_range", path="lazy.xlsx", sheet="Data", range="D20:G26")
    full = await call("read_range", path="full.xlsx", sheet="Data", range="D20:G26")

    assert lazy["values"] == full["values"]


async def test_an_edit_in_the_middle_keeps_the_masters_and_the_rows_it_did_not_parse(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_shared_book(files / "book.xlsx")

    with lazily(monkeypatch):
        await call("write_range", path="book.xlsx", sheet="Data", at="D22", rows=[[5]])

    found = groups(files / "book.xlsx")
    assert found["0"]["master"] == ("D2", f"D2:D{FILLED}")
    # Rows 20 to 23 are the block that was parsed; they are written as ordinary formulas.
    assert not {"D20", "D21", "D22", "D23"} & set(found["0"]["children"])
    assert "D19" in found["0"]["children"]
    assert "D3" in found["0"]["children"] and f"D{FILLED}" in found["0"]["children"]
    assert found["2"]["master"] == ("F2", "F2:G40")
    saved = load_workbook(files / "book.xlsx")["Data"]
    assert (saved["D22"].value, saved["D23"].value) == (5, "=B23*C23")


async def test_an_edit_at_the_master_row_writes_the_master_again(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_shared_book(files / "book.xlsx")

    with lazily(monkeypatch):
        await call("write_range", path="book.xlsx", sheet="Data", at="A2", rows=[[100]])

    found = groups(files / "book.xlsx")
    assert found["0"]["master"] == ("D2", f"D2:D{FILLED}")
    assert found["0"]["children"] == [f"D{row}" for row in range(4, FILLED + 1)]
    assert_groups_are_whole(files / "book.xlsx")


async def test_overwriting_a_master_parses_what_its_range_covers(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_shared_book(files / "book.xlsx")

    with lazily(monkeypatch):
        await call("write_range", path="book.xlsx", sheet="Data", at="D2", rows=[[7]])

    found = groups(files / "book.xlsx")
    assert "0" not in found
    assert "1" in found and "2" in found
    saved = load_workbook(files / "book.xlsx")["Data"]
    assert (saved["D2"].value, saved["D24"].value) == (7, "=B24*C24")


async def test_when_every_row_is_rewritten_no_shared_formula_is_left(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_shared_book(files / "book.xlsx")

    with lazily(monkeypatch):
        await call("insert_rows_or_columns", path="book.xlsx", sheet="Data", axis="rows", start=5)

    assert groups(files / "book.xlsx") == {}
