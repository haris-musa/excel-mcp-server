"""Details of the files the server writes: macros on import, document properties, date columns."""

import base64
import zipfile
from pathlib import Path

import pytest
from mcp import Client
from openpyxl import load_workbook

from excel_mcp.config import Settings
from excel_mcp.server import create_server
from tests.conftest import ToolCall
from tests.package_support import read_parts

pytestmark = pytest.mark.anyio


def _core(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("docProps/core.xml").decode()


async def test_importing_a_macro_workbook_as_xlsx_drops_the_macros(files: Path) -> None:
    async with Client(create_server(Settings(allowed_dirs=[files], allow_vba_write=True))) as c:
        await c.call_tool("create_workbook", {"path": "macro.xlsm"})
        encoded = base64.b64encode((files / "macro.xlsm").read_bytes()).decode()
        for name in ("plain.xlsx", "kept.xlsm"):
            await c.call_tool("import_workbook", {"path": name, "content_base64": encoded})

    plain = read_parts(files / "plain.xlsx")
    assert not [name for name in plain if "vbaProject" in name]
    assert b"vbaProject" not in plain["[Content_Types].xml"]
    assert b"macroEnabled" not in plain["[Content_Types].xml"]
    assert b"vbaProject" not in plain["xl/_rels/workbook.xml.rels"]
    assert load_workbook(files / "plain.xlsx").sheetnames
    assert "xl/vbaProject.bin" in read_parts(files / "kept.xlsm")


async def test_new_workbooks_name_no_creator_and_edits_keep_the_one_there(
    call: ToolCall, files: Path
) -> None:
    await call("create_workbook", path="new.xlsx")
    assert "creator" not in _core(files / "new.xlsx")

    await call(
        "set_workbook_settings", path="new.xlsx", settings={"doc_properties": {"author": "Ada"}}
    )
    await call("write_range", path="new.xlsx", sheet="Sheet1", at="A1", rows=[[1]])
    assert "<dc:creator>Ada</dc:creator>" in _core(files / "new.xlsx")


async def test_a_file_without_a_creator_gets_none_added(call: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook.properties.creator = None
    workbook.save(sample)
    await call("write_range", path="sales.xlsx", sheet="Report", at="A1", rows=[[1]])
    assert "creator" not in _core(sample)


async def test_dates_widen_default_width_columns_only(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Report",
        layout={"column_widths_chars": {"C": 30}},
    )
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        at="A1",
        rows=[["2024-01-15", "2024-01-15T09:30:00", "2024-01-15", "text"]],
    )
    widths = load_workbook(sample)["Report"].column_dimensions
    assert widths["A"].width > 10 and widths["B"].width > 18
    assert widths["C"].width == 30  # a width that was set stays
    assert "D" not in widths or (widths["D"].width == 13 and not widths["D"].customWidth)

    await call("write_range", path="sales.xlsx", sheet="Report", at="A2", rows=[["2024-02-01"]])
    assert load_workbook(sample)["Report"].column_dimensions["A"].width > 10
