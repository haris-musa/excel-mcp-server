"""Sparklines and the conditional formats Excel keeps in its 2010 extension.

Expected XML is what Microsoft Excel saved for the same settings (set through its object
model), without the ids it generates.
"""

import re
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from excel_mcp import package
from excel_mcp.package import extensions as ext
from excel_mcp.package import state_of
from excel_mcp.package.references import rewrite_lines
from tests.conftest import ToolCall
from tests.package_support import assert_package_is_consistent, read_parts, sheet_part, text
from tests.test_package_references import _edit, _formula

pytestmark = pytest.mark.anyio

BOOK = {"path": "sales.xlsx"}
COLORS = (
    '<x14:colorSeries rgb="FF376092"/><x14:colorNegative rgb="FFD00000"/>'
    '<x14:colorAxis rgb="FF000000"/><x14:colorMarkers rgb="FFD00000"/>'
    '<x14:colorFirst rgb="FFD00000"/><x14:colorLast rgb="FFD00000"/>'
    '<x14:colorHigh rgb="FFD00000"/><x14:colorLow rgb="FFD00000"/>'
)


def _sheet_xml(path: Path, sheet: str = "Data") -> str:
    parts = read_parts(path)
    return text(parts, sheet_part(parts, sheet))


def _groups(path: Path, sheet: str = "Data") -> list[str]:
    return re.findall(
        r"<x14:sparklineGroup\b.*?</x14:sparklineGroup>", _sheet_xml(path, sheet), re.S
    )


def _x14_rules(path: Path, sheet: str = "Data") -> list[str]:
    return re.findall(
        r"<x14:conditionalFormatting\b.*?</x14:conditionalFormatting>",
        _sheet_xml(path, sheet),
        re.S,
    )


def _line(items: dict[str, str]) -> str:
    return (
        "<x14:sparklines>"
        + "".join(
            f"<x14:sparkline><xm:f>{f}</xm:f><xm:sqref>{c}</xm:sqref></x14:sparkline>"
            for c, f in items.items()
        )
        + "</x14:sparklines>"
    )


async def test_a_line_group_is_what_excel_writes(call: ToolCall, sample: Path) -> None:
    result = await call("add_sparklines", **BOOK, sheet="Data", location="F2:F5", data="C2:D5")
    assert "Added 4 line sparklines to Data!F2:F5" in result
    (group,) = _groups(sample)
    cells = {f"F{row}": f"Data!C{row}:D{row}" for row in range(2, 6)}
    expected = (
        f'<x14:sparklineGroup displayEmptyCellsAs="gap">{COLORS}{_line(cells)}</x14:sparklineGroup>'
    )
    assert group == expected
    assert_package_is_consistent(read_parts(sample))


async def test_every_option_is_what_excel_writes(call: ToolCall, sample: Path) -> None:
    style: dict[str, Any] = {
        "type": "column",
        "colors": {"series": "FF0000", "negative": "00FF00", "high": "0000FF", "axis": "808080"},
        "show": ["high", "low", "first", "last", "negative"],
        "show_axis": True,
        "axis_min": -5,
        "axis_max": 10,
        "empty_cells": "zero",
        "hidden": True,
    }
    await call("add_sparklines", **BOOK, sheet="Data", location="F2:F3", data="C2:D3", style=style)
    (group,) = _groups(sample)
    assert group.startswith(
        '<x14:sparklineGroup manualMax="10" manualMin="-5" type="column" high="1" low="1" '
        'first="1" last="1" negative="1" displayXAxis="1" displayHidden="1" '
        'minAxisType="custom" maxAxisType="custom">'
        '<x14:colorSeries rgb="FFFF0000"/><x14:colorNegative rgb="FF00FF00"/>'
        '<x14:colorAxis rgb="FF808080"/>'
    )
    assert '<x14:colorHigh rgb="FF0000FF"/>' in group


async def test_win_loss_dates_and_line_options(call: ToolCall, sample: Path) -> None:
    await call(
        "add_sparklines",
        **BOOK,
        sheet="Data",
        location="F2:F3",
        data="C2:D3",
        style={
            "type": "win_loss",
            "show": ["negative"],
            "axis_min": "same",
            "empty_cells": "connect",
            "right_to_left": True,
        },
    )
    await call(
        "add_sparklines",
        **BOOK,
        sheet="Data",
        location="G2",
        data="C2:D2",
        style={"show": ["markers"], "line_weight": 2.25, "dates": "C1:D1"},
    )
    line, win_loss = _groups(sample)  # Excel puts the newest group first
    assert line.startswith(
        '<x14:sparklineGroup lineWeight="2.25" dateAxis="1" displayEmptyCellsAs="gap" markers="1">'
    )
    assert "<xm:f>Data!C1:D1</xm:f><x14:sparklines>" in line
    assert win_loss.startswith(
        '<x14:sparklineGroup type="stacked" displayEmptyCellsAs="span" negative="1" '
        'minAxisType="group" rightToLeft="1">'
    )


async def test_sparklines_are_listed_replaced_and_deleted(call: ToolCall, sample: Path) -> None:
    await call("add_sparklines", **BOOK, sheet="Data", location="F2:F5", data="C2:D5")
    await call(
        "add_sparklines",
        **BOOK,
        sheet="Data",
        location="G2:G3",
        data="C2:D3",
        style={"type": "column", "show": ["high"], "colors": {"high": "FF0000"}},
    )
    listed = (await call("describe_sheet", **BOOK, sheet="Data"))["sparklines"]
    assert listed[0] == {
        "type": "column",
        "colors": {"high": "FF0000"},
        "show": ["high"],
        "sparklines": {"G2": "Data!C2:D2", "G3": "Data!C3:D3"},
    }
    assert listed[1]["sparklines"]["F5"] == "Data!C5:D5"

    replaced = await call("add_sparklines", **BOOK, sheet="Data", location="F4:F5", data="A4:B5")
    assert "Replaced 2 existing" in replaced
    await call("delete_sparklines", **BOOK, sheet="Data", range="G2:G3")
    listed = (await call("describe_sheet", **BOOK, sheet="Data"))["sparklines"]
    assert [g["sparklines"] for g in listed] == [
        {"F4": "Data!A4:B4", "F5": "Data!A5:B5"},
        {"F2": "Data!C2:D2", "F3": "Data!C3:D3"},
    ]
    await call("delete_sparklines", **BOOK, sheet="Data", range="F1:F9")
    assert "sparklines" not in await call("describe_sheet", **BOOK, sheet="Data")
    assert ext.SPARKLINES not in _sheet_xml(sample)


async def test_columns_of_data_when_the_cell_count_matches_them(
    call: ToolCall, sample: Path
) -> None:
    await call("add_sparklines", **BOOK, sheet="Data", location="F1:G1", data="C2:D5")
    listed = (await call("describe_sheet", **BOOK, sheet="Data"))["sparklines"]
    assert listed[0]["sparklines"] == {"F1": "Data!C2:C5", "G1": "Data!D2:D5"}


async def test_data_on_another_sheet(call: ToolCall, sample: Path) -> None:
    await call("add_sparklines", **BOOK, sheet="Report", location="A1:A4", data="Data!C2:D5")
    listed = (await call("describe_sheet", **BOOK, sheet="Report"))["sparklines"]
    assert listed[0]["sparklines"]["A1"] == "Data!C2:D2"


@pytest.mark.parametrize(
    ("arguments", "message"),
    [
        ({"location": "F2:G3", "data": "C2:D3"}, "one row or one column"),
        ({"location": "F2:F4", "data": "C2:D5"}, "3 cells need 3 rows or columns"),
        ({"location": "F2", "data": "Missing!C2:D2"}, "Sheet 'Missing' not found"),
        (
            {"location": "F2", "data": "C2:D2", "style": {"type": "column", "show": ["markers"]}},
            "line",
        ),
        ({"location": "F2", "data": "C2:D2", "style": {"dates": "C1:E1"}}, "dates must be"),
        ({"location": "F2", "data": "C2:D2", "style": {"axis_min": "other"}}, "axis_min"),
        ({"location": "F2", "data": "C2:D2", "style": {"colors": {"series": "red"}}}, "color"),
    ],
)
async def test_invalid_sparklines_are_rejected(
    call_error: ToolCall, sample: Path, arguments: dict[str, Any], message: str
) -> None:
    before = sample.read_bytes()
    assert message in await call_error("add_sparklines", **BOOK, sheet="Data", **arguments)
    assert sample.read_bytes() == before


async def test_deleting_what_is_not_there_is_an_error(call_error: ToolCall, sample: Path) -> None:
    message = await call_error("delete_sparklines", **BOOK, sheet="Data", range="F2:F3")
    assert "no sparklines" in message


async def test_sparklines_stay_inside_the_allowed_folder(call_error: ToolCall) -> None:
    message = await call_error(
        "add_sparklines", path="../x.xlsx", sheet="Data", location="F2", data="C2:D2"
    )
    assert "outside" in message or "not allowed" in message


async def test_sparklines_survive_edits_and_sheet_copies(call: ToolCall, sample: Path) -> None:
    await call("add_sparklines", **BOOK, sheet="Data", location="F2:F3", data="C2:D3")
    await call("write_range", **BOOK, sheet="Data", start_cell="H1", rows=[[1]])
    await call("copy_sheet", **BOOK, sheet="Data", new_name="My Copy")
    (original,) = _groups(sample)
    (copy,) = _groups(sample, "My Copy")
    assert "<xm:f>Data!C2:D2</xm:f>" in original
    assert "<xm:f>'My Copy'!C2:D2</xm:f>" in copy
    assert_package_is_consistent(read_parts(sample))


async def test_sparklines_follow_inserted_rows(call: ToolCall, sample: Path) -> None:
    await call("add_sparklines", **BOOK, sheet="Data", location="F3:F4", data="C3:D4")
    workbook = load_workbook(sample)
    package.capture(sample, workbook, 50_000_000)
    edit = _edit("Data", "rows", 2)
    rewrite_lines(workbook, edit, _formula(edit))
    xml = state_of(workbook).sheet(workbook["Data"]).extensions[ext.SPARKLINES]
    assert re.findall(r"<xm:f>(.*?)</xm:f><xm:sqref>(.*?)</xm:sqref>", xml) == [
        ("Data!C4:D4", "F4"),
        ("Data!C5:D5", "F5"),
    ]


BAR_NEGATIVES = {
    "colors": ["#00FF00"],
    "bar_fill": "solid",
    "border_color": "FF0000",
    "negative_color": "0000FF",
    "negative_border_color": "FFFF00",
    "axis": "middle",
    "axis_color": "808080",
    "bar_direction": "right_to_left",
    "min_type": "number",
    "min_value": -5,
    "max_type": "percent",
    "max_value": 90,
    "hide_values": True,
}


async def test_a_default_data_bar_is_what_excel_writes(call: ToolCall, sample: Path) -> None:
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="C2:C5",
        rule={"type": "data_bar", "colors": ["#638EC6"]},
    )
    xml = _sheet_xml(sample)
    guid = re.search(r"<x14:id>(\{[^}]*\})</x14:id>", xml)[1]  # type: ignore[index]
    assert (
        '<cfRule type="dataBar" priority="1"><dataBar><cfvo type="min" /><cfvo type="max" />'
        '<color rgb="FF638EC6" /></dataBar>'
    ) in xml
    assert _x14_rules(sample) == [
        '<x14:conditionalFormatting xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
        f'<x14:cfRule type="dataBar" id="{guid}"><x14:dataBar minLength="0" maxLength="100" '
        'negativeBarColorSameAsPositive="1" axisPosition="none"><x14:cfvo type="min"/>'
        '<x14:cfvo type="max"/></x14:dataBar></x14:cfRule><xm:sqref>C2:C5</xm:sqref>'
        "</x14:conditionalFormatting>"
    ]
    assert_package_is_consistent(read_parts(sample))


async def test_a_data_bar_with_every_option_is_what_excel_writes(
    call: ToolCall, sample: Path
) -> None:
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="C2:C5",
        rule={"type": "data_bar", **BAR_NEGATIVES},
    )
    xml = _sheet_xml(sample)
    assert '<dataBar showValue="0"><cfvo type="num" val="-5" />' in xml
    assert '<cfvo type="percent" val="90" /><color rgb="FF00FF00" />' in xml
    (rule,) = _x14_rules(sample)
    assert (
        '<x14:dataBar minLength="0" maxLength="100" border="1" gradient="0" '
        'direction="rightToLeft" negativeBarBorderColorSameAsPositive="0" axisPosition="middle">'
        '<x14:cfvo type="num"><xm:f>-5</xm:f></x14:cfvo>'
        '<x14:cfvo type="percent"><xm:f>90</xm:f></x14:cfvo>'
        '<x14:borderColor rgb="FFFF0000"/><x14:negativeFillColor rgb="FF0000FF"/>'
        '<x14:negativeBorderColor rgb="FFFFFF00"/><x14:axisColor rgb="FF808080"/></x14:dataBar>'
    ) in rule


async def test_automatic_and_formula_scale_points(call: ToolCall, sample: Path) -> None:
    rule = {
        "type": "data_bar",
        "colors": ["#638EC6"],
        "min_type": "automatic",
        "max_type": "formula",
        "max_value": "=$F$1",
        "axis": "automatic",
        "border_color": "000000",
    }
    await call("add_conditional_format", **BOOK, sheet="Data", range="C2:C5", rule=rule)
    assert '<cfvo type="min" /><cfvo type="formula" val="$F$1" />' in _sheet_xml(sample)
    (extended,) = _x14_rules(sample)
    assert '<x14:cfvo type="autoMin"/><x14:cfvo type="formula"><xm:f>$F$1</xm:f>' in extended
    assert '<x14:axisColor rgb="FF000000"/>' in extended and "axisPosition" not in extended


async def test_icon_sets_that_only_exist_in_the_extension(call: ToolCall, sample: Path) -> None:
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="C2:C5",
        rule={"type": "icon_set", "icon_set": "3Stars"},
    )
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="D2:D5",
        rule={
            "type": "icon_set",
            "icon_set": "3Arrows",
            "icons": ["none", "3Flags:3", "4TrafficLights:1"],
            "hide_values": True,
            "thresholds": [3, 80],
            "threshold_type": "percentile",
        },
    )
    stars, custom = _x14_rules(sample)
    assert (
        '<x14:cfRule type="iconSet" priority="1" id="'
    ) in stars and '<x14:iconSet iconSet="3Stars"><x14:cfvo type="percent"><xm:f>0</xm:f>' in stars
    assert (
        '<x14:iconSet iconSet="3Arrows" showValue="0" custom="1">'
        '<x14:cfvo type="percent"><xm:f>0</xm:f></x14:cfvo>'
        '<x14:cfvo type="percentile"><xm:f>3</xm:f></x14:cfvo>'
        '<x14:cfvo type="percentile"><xm:f>80</xm:f></x14:cfvo>'
        '<x14:cfIcon iconSet="NoIcons" iconId="0"/><x14:cfIcon iconSet="3Flags" iconId="2"/>'
        '<x14:cfIcon iconSet="4TrafficLights" iconId="0"/></x14:iconSet>'
    ) in custom
    assert 'priority="2"' in custom
    assert "<cfRule" not in _sheet_xml(sample).split("<extLst>")[0]
    listed = (await call("describe_sheet", **BOOK, sheet="Data"))["conditional_formats"]
    assert listed == [{"range": "C2:C5", "type": "iconSet"}, {"range": "D2:D5", "type": "iconSet"}]


@pytest.mark.parametrize(
    ("rule", "message"),
    [
        ({"type": "data_bar", "colors": ["#FF0000"], "min_type": "highest"}, "min_type"),
        ({"type": "data_bar", "colors": ["#FF0000"], "min_value": 3}, "min_value"),
        ({"type": "data_bar", "colors": ["#FF0000"], "max_type": "percent"}, "needs max_value"),
        (
            {"type": "data_bar", "colors": ["#FF0000"], "max_type": "percent", "max_value": 120},
            "0 to 100",
        ),
        (
            {"type": "data_bar", "colors": ["#FF0000"], "negative_border_color": "000000"},
            "border_color",
        ),
        ({"type": "data_bar", "colors": ["#FF0000"], "icons": ["none"]}, "does not use icons"),
        ({"type": "icon_set", "icon_set": "3Arrows", "icons": ["none"]}, "needs 3"),
        (
            {"type": "icon_set", "icon_set": "3Arrows", "icons": ["none", "3Flags:4", "none"]},
            "Invalid icon",
        ),
        (
            {"type": "icon_set", "icon_set": "3Arrows", "icons": ["none"] * 3, "reverse": True},
            "remove reverse",
        ),
        (
            {
                "type": "data_bar",
                "colors": ["#FF0000"],
                "max_type": "formula",
                "max_value": '=WEBSERVICE("x")',
            },
            "not allowed",
        ),
    ],
)
async def test_invalid_extended_rules_are_rejected(
    call_error: ToolCall, sample: Path, rule: dict[str, Any], message: str
) -> None:
    before = sample.read_bytes()
    found = await call_error(
        "add_conditional_format", **BOOK, sheet="Data", range="C2:C5", rule=rule
    )
    assert message in found
    assert sample.read_bytes() == before


async def test_priorities_are_shared_and_extended_rules_keep_their_partner(
    call: ToolCall, sample: Path
) -> None:
    bar = {"type": "data_bar", "colors": ["#638EC6"]}
    await call("add_conditional_format", **BOOK, sheet="Data", range="C2:C5", rule=bar)
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="D2:D5",
        rule={"type": "icon_set", "icon_set": "3Stars"},
    )
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="A2:A5",
        rule={"type": "duplicate", "fill_color": "#FFC7CE", "priority": 1},
    )
    xml = _sheet_xml(sample)
    assert re.findall(r'<cfRule type="(\w+)" priority="(\d)"', xml) == [
        ("dataBar", "2"),
        ("duplicateValues", "1"),
    ]
    assert 'priority="3"' in _x14_rules(sample)[1]
    guid = re.search(r"<x14:id>(\{[^}]*\})</x14:id>", xml)[1]  # type: ignore[index]
    assert f'id="{guid}"' in _x14_rules(sample)[0]
    assert_package_is_consistent(read_parts(sample))


async def test_extended_rules_survive_edits_and_sheet_copies(call: ToolCall, sample: Path) -> None:
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="C2:C5",
        rule={
            "type": "data_bar",
            "colors": ["#638EC6"],
            "max_type": "formula",
            "max_value": "=Data!$F$1",
        },
    )
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="D2:D5",
        rule={"type": "icon_set", "icon_set": "5Boxes"},
    )
    await call("write_range", **BOOK, sheet="Data", start_cell="H1", rows=[[1]])
    await call("copy_sheet", **BOOK, sheet="Data", new_name="Copy")
    source, copied = _x14_rules(sample), _x14_rules(sample, "Copy")
    assert len(source) == len(copied) == 2
    assert "<xm:f>Data!$F$1</xm:f>" in source[0] and "<xm:f>'Copy'!$F$1</xm:f>" in copied[0]
    ids = lambda parts: set(re.findall(r"\{[0-9A-F-]{36}\}", "".join(parts)))  # noqa: E731
    assert ids(source).isdisjoint(ids(copied))
    assert re.findall(r"<x14:id>(\{[^}]*\})</x14:id>", _sheet_xml(sample, "Copy"))[0] in ids(copied)
    assert_package_is_consistent(read_parts(sample))


async def test_extended_rules_follow_inserted_rows(call: ToolCall, sample: Path) -> None:
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="C3:C5",
        rule={"type": "data_bar", "colors": ["#638EC6"]},
    )
    workbook = load_workbook(sample)
    package.capture(sample, workbook, 50_000_000)
    sheet = workbook["Data"]
    edit = _edit("Data", "rows", 2)
    rewrite_lines(workbook, edit, _formula(edit))
    state = state_of(workbook).sheet(sheet)
    assert list(state.rule_extensions) == [("C4:C6", "1")]
    assert "<xm:sqref>C4:C6</xm:sqref>" in state.extensions[ext.CONDITIONAL_FORMATS]
