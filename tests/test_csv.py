"""CSV import and export: field typing, dialects, formula safety, limits and paths."""

import datetime as dt
from collections.abc import AsyncIterator
from pathlib import Path

import pytest
from mcp import Client
from openpyxl import Workbook, load_workbook

from excel_mcp.config import Limits, Settings
from excel_mcp.server import create_server
from tests.conftest import ToolCall, error_text

pytestmark = pytest.mark.anyio


def _csv(files: Path, text: str, name: str = "in.csv", encoding: str = "utf-8") -> str:
    (files / name).write_bytes(text.encode(encoding))
    return name


async def _import(call: ToolCall, files: Path, text: str, **options: object) -> dict[str, object]:
    name = _csv(files, text, encoding=str(options.pop("file_encoding", "utf-8")))
    return await call("import_csv", path="sales.xlsx", csv_path=name, sheet="Imported", **options)


def _sheet(files: Path, name: str = "Imported"):
    return load_workbook(files / "sales.xlsx")[name]


async def test_fields_become_numbers_booleans_dates_and_text(
    call: ToolCall, sample: Path, files: Path
) -> None:
    result = await _import(
        call,
        files,
        'name,qty,price,on,when,at,pct,big,note\r\nAda,3,"1,234.50",true,2026-01-31,'
        '2026-01-31 09:30,12.5%,"1,000",""\r\n',
    )
    assert result == {"sheet": "Imported", "range": "A1:I2", "cells_written": 18}
    row = [cell.value for cell in _sheet(files)[2]]
    assert row == [
        "Ada", 3, 1234.5, True, dt.datetime(2026, 1, 31), dt.datetime(2026, 1, 31, 9, 30),
        0.125, 1000, None,
    ]  # fmt: skip
    sheet = _sheet(files)
    assert sheet["C2"].number_format == "#,##0.00"
    assert sheet["G2"].number_format == "0.0%"
    assert sheet["E2"].number_format == "yyyy-mm-dd"
    assert sheet["E2"].is_date and sheet["F2"].is_date


async def test_text_that_excel_would_alter_stays_text(
    call: ToolCall, sample: Path, files: Path
) -> None:
    fields = [
        "007",
        "0123.5",
        "12345678901234567",
        "31/12/2026",
        "01/02/2026",
        "1e999",
        "NaN",
        "+15551234",
    ]
    await _import(call, files, ",".join(fields) + "\r\n")
    assert [cell.value for cell in _sheet(files)[1]] == fields
    assert _sheet(files)["A1"].data_type == "s"


async def test_zero_and_decimals_below_one_are_numbers(
    call: ToolCall, sample: Path, files: Path
) -> None:
    await _import(call, files, "0,0.5,-0,-7,1e3\r\n")
    assert [cell.value for cell in _sheet(files)[1]] == [0, 0.5, 0, -7, 1000.0]


@pytest.mark.parametrize(
    ("order", "text", "expected"),
    [
        ("dmy", "31/12/2026", dt.datetime(2026, 12, 31)),
        ("dmy", "1.2.2026", dt.datetime(2026, 2, 1)),
        ("mdy", "12-31-2026", dt.datetime(2026, 12, 31)),
        ("mdy", "31/12/2026", "31/12/2026"),
        ("ymd", "31/12/2026", "31/12/2026"),
    ],
)
async def test_numeric_dates_only_in_the_order_given(
    call: ToolCall, sample: Path, files: Path, order: str, text: str, expected: object
) -> None:
    await _import(call, files, text + "\r\n", date_order=order)
    assert _sheet(files)["A1"].value == expected


async def test_decimal_comma_and_semicolon_delimiter(
    call: ToolCall, sample: Path, files: Path
) -> None:
    await _import(call, files, 'a;b;c\r\n"1.234,5";2,5;1.5\r\n', decimal=",")
    sheet = _sheet(files)
    assert [cell.value for cell in sheet[2]] == [1234.5, 2.5, "1.5"]
    assert sheet["A2"].number_format == "#,##0.0"


async def test_as_text_keeps_every_field_text(call: ToolCall, sample: Path, files: Path) -> None:
    await _import(call, files, "1,TRUE,2026-01-31\r\n", as_text=True)
    assert [cell.value for cell in _sheet(files)[1]] == ["1", "TRUE", "2026-01-31"]


async def test_formula_like_fields_are_text_never_formulas(
    call: ToolCall, sample: Path, files: Path
) -> None:
    hostile = ["=1+1", "+1+2", "@SUM(A1)", "-1+2", '=HYPERLINK("http://x.example")', "\tx"]
    await _import(
        call,
        files,
        "\r\n".join(f'"{field.replace(chr(34), chr(34) * 2)}",-5' for field in hostile) + "\r\n",
    )
    sheet = _sheet(files)
    for row, field in enumerate(hostile, start=1):
        cell = sheet.cell(row, 1)
        assert (cell.value, cell.data_type, cell.quotePrefix) == (field, "s", True)
        assert sheet.cell(row, 2).value == -5  # a negative number is a number
        assert not sheet.cell(row, 2).quotePrefix
    assert b"<f>" not in _sheet_xml(files / "sales.xlsx")


def _sheet_xml(path: Path) -> bytes:
    import zipfile

    with zipfile.ZipFile(path) as archive:
        return b"".join(
            archive.read(name) for name in archive.namelist() if name.startswith("xl/worksheets/")
        )


@pytest.mark.parametrize(
    ("encoding", "text"),
    [
        ("utf-8", "Zoë,Müller"),
        ("utf-8-sig", "Zoë,Müller"),
        ("cp1252", "Zoë,Müller"),
        ("utf-16", "Zoë,Müller"),
    ],
)
async def test_encodings(
    call: ToolCall, sample: Path, files: Path, encoding: str, text: str
) -> None:
    await _import(call, files, text + "\r\n", encoding=encoding, file_encoding=encoding)
    assert [cell.value for cell in _sheet(files)[1]] == ["Zoë", "Müller"]


async def test_utf8_import_accepts_a_byte_order_mark(
    call: ToolCall, sample: Path, files: Path
) -> None:
    await _import(call, files, "Zoë,x\r\n", file_encoding="utf-8-sig")
    assert _sheet(files)["A1"].value == "Zoë"


async def test_wrong_encoding_says_what_to_try(
    call_error: ToolCall, sample: Path, files: Path
) -> None:
    _csv(files, "Zoë\r\n", encoding="cp1252")
    message = await call_error("import_csv", path="sales.xlsx", csv_path="in.csv", sheet="N")
    assert "not valid utf-8" in message and "cp1252" in message


@pytest.mark.parametrize(
    ("header", "expected"),
    [("a;b;c", ";"), ("a\tb\tc", "\t"), ("a|b|c", "|"), ('"a,x,y";b;c', ";"), ("abc", ",")],
)
async def test_delimiter_is_detected_from_the_first_record(
    call: ToolCall, sample: Path, files: Path, header: str, expected: str
) -> None:
    await _import(call, files, header + "\r\n")
    assert len([cell.value for cell in _sheet(files)[1] if cell.value is not None]) == (
        1 if expected == "," else 3
    )


async def test_quotes_and_line_breaks_inside_fields(
    call: ToolCall, sample: Path, files: Path
) -> None:
    await _import(call, files, 'a,"line one\r\nline ""two"""\r\nx,y\r\n')
    sheet = _sheet(files)
    assert sheet["B1"].value == 'line one\nline "two"'
    assert sheet["A2"].value == "x"


async def test_import_into_an_existing_sheet_at_a_cell(
    call: ToolCall, sample: Path, files: Path
) -> None:
    _csv(files, "1,2\r\n\r\n3,\r\n")
    result = await call(
        "import_csv", path="sales.xlsx", csv_path="in.csv", sheet="Data", at="C2", delimiter="comma"
    )
    assert result["range"] == "C2:D4"
    sheet = load_workbook(files / "sales.xlsx")["Data"]
    assert (sheet["C2"].value, sheet["D2"].value, sheet["C4"].value) == (1, 2, 3)
    assert sheet["D4"].value is None  # an empty field empties the cell
    assert sheet["A2"].value == "North"
    assert load_workbook(files / "sales.xlsx").sheetnames == ["Data", "Report"]


async def test_dates_widen_their_column(call: ToolCall, sample: Path, files: Path) -> None:
    await _import(call, files, "2026-01-31,2026-01-31 09:30\r\n")
    sheet = _sheet(files)
    assert sheet["B1"].number_format == "yyyy-mm-dd hh:mm:ss"
    assert 10 < sheet.column_dimensions["A"].width < sheet.column_dimensions["B"].width


async def test_the_workbook_is_unchanged_when_the_import_fails(
    call_error: ToolCall, sample: Path, files: Path
) -> None:
    before = sample.read_bytes()
    _csv(files, "a,b\r\nx\x01y,2\r\n")
    message = await call_error("import_csv", path="sales.xlsx", csv_path="in.csv", sheet="N")
    assert "control characters" in message
    assert sample.read_bytes() == before


@pytest.mark.parametrize("text", ["", "\r\n\r\n"])
async def test_an_empty_file_is_refused(
    call_error: ToolCall, sample: Path, files: Path, text: str
) -> None:
    _csv(files, text)
    message = await call_error("import_csv", path="sales.xlsx", csv_path="in.csv", sheet="N")
    assert "no data" in message


async def test_bad_arguments_are_errors(call_error: ToolCall, sample: Path, files: Path) -> None:
    _csv(files, "a\r\n")
    base = {"path": "sales.xlsx", "csv_path": "in.csv"}
    assert "does not exist" in await call_error(
        "import_csv", **{**base, "csv_path": "nope.csv"}, sheet="N"
    )
    assert "1 to 31" in await call_error("import_csv", **base, sheet="x" * 40)
    assert "Invalid arguments" in await call_error("import_csv", **base, sheet="N", delimiter=";")
    _csv(files, "a,b\r\n")
    assert "does not fit" in await call_error("import_csv", **base, sheet="Data", at="XFD1")
    assert await call_error("import_csv", **base, sheet="N", at="nonsense")


@pytest.fixture
async def small(files: Path) -> AsyncIterator[Client]:
    settings = Settings(allowed_dirs=[files], limits=Limits(max_cells=6, max_file_bytes=20_000))
    async with Client(create_server(settings)) as client:
        yield client


async def test_cell_and_file_size_limits(small: Client, sample: Path, files: Path) -> None:
    _csv(files, "1,2,3\r\n4,5,6\r\n7,8,9\r\n")
    result = await small.call_tool(
        "import_csv", {"path": "sales.xlsx", "csv_path": "in.csv", "sheet": "N"}
    )
    assert result.is_error and "more than 6 cells" in error_text(result)
    _csv(files, "x" * 30_000)
    result = await small.call_tool(
        "import_csv", {"path": "sales.xlsx", "csv_path": "in.csv", "sheet": "N"}
    )
    assert result.is_error and "limit is 20,000" in error_text(result)
    assert "N" not in load_workbook(files / "sales.xlsx").sheetnames


async def test_export_is_limited_by_max_cells(small: Client, files: Path) -> None:
    workbook = Workbook()
    workbook.active.append([1, 2, 3])  # type: ignore[union-attr]
    workbook.active.append([4, 5, 6])  # type: ignore[union-attr]
    workbook.active.append([7, 8, 9])  # type: ignore[union-attr]
    workbook.save(files / "big.xlsx")
    result = await small.call_tool(
        "export_csv", {"path": "big.xlsx", "sheet": "Sheet", "csv_path": "out.csv"}
    )
    assert result.is_error and "more than 6 cells" in error_text(result)
    assert not (files / "out.csv").exists()


@pytest.mark.parametrize(
    "csv_path", ["../outside.csv", "sub/../../outside.csv", "notes.txt", "data.xlsx", "C:/x.csv"]
)
async def test_import_and_export_use_the_path_gate(
    call_error: ToolCall, sample: Path, files: Path, csv_path: str
) -> None:
    (files.parent / "outside.csv").write_bytes(b"secret\r\n")
    message = await call_error("import_csv", path="sales.xlsx", csv_path=csv_path, sheet="N")
    assert "outside" in message or "must point to a CSV" in message or "drive" in message
    message = await call_error("export_csv", path="sales.xlsx", sheet="Data", csv_path=csv_path)
    assert "outside" in message or "must point to a CSV" in message or "drive" in message
    assert "secret" not in message
    assert (files.parent / "outside.csv").read_bytes().decode() == "secret\r\n"


async def test_symlinks_out_of_the_root_are_refused(
    call_error: ToolCall, sample: Path, files: Path
) -> None:
    secret = files.parent / "secret.csv"
    secret.write_bytes(b"secret\r\n")
    try:
        (files / "link.csv").symlink_to(secret)
    except OSError:
        pytest.skip("symlinks are not available")
    message = await call_error("import_csv", path="sales.xlsx", csv_path="link.csv", sheet="N")
    assert "outside" in message
    message = await call_error(
        "export_csv", path="sales.xlsx", sheet="Data", csv_path="link.csv", overwrite=True
    )
    assert "outside" in message
    assert secret.read_bytes().decode() == "secret\r\n"


async def test_export_writes_raw_values_with_iso_dates(
    call: ToolCall, sample: Path, files: Path
) -> None:
    result = await call("export_csv", path="sales.xlsx", sheet="Data", csv_path="out.csv")
    assert result == {"path": "out.csv", "sheet": "Data", "range": "A1:D5"}
    lines = (files / "out.csv").read_bytes().decode().split("\r\n")
    assert lines[:2] == ["Region,Product,Units,Price", "North,Apples,10,1.5"]


@pytest.fixture
def rich(files: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Rich"
    sheet.append(["Name", "Amount", "Day", "Ok", "Note"])
    sheet.append(['Zoë, "Z"\nline', 1234.5, dt.date(2026, 1, 31), True, "=1+1"])
    sheet.append(["+15551234", -2, dt.datetime(2026, 1, 31, 9, 30), False, None])
    sheet.append(["x", "=B2*2", None, None, "@cmd"])
    sheet["B2"].number_format = "#,##0.00"
    sheet["C2"].number_format = "mmm d, yyyy"
    sheet["B3"].number_format = "0.0%"
    sheet["E2"].data_type = "s"
    path = files / "rich.xlsx"
    workbook.save(path)
    return path


async def test_export_quotes_and_neutralises_formula_text(
    call: ToolCall, rich: Path, files: Path
) -> None:
    await call("export_csv", path="rich.xlsx", sheet="Rich", csv_path="out.csv")
    text = (files / "out.csv").read_bytes().decode("utf-8")
    assert text == (
        "Name,Amount,Day,Ok,Note\r\n"
        '"Zoë, ""Z""\nline",1234.5,2026-01-31,TRUE,\'=1+1\r\n'
        "'+15551234,-2,2026-01-31T09:30:00,FALSE,\r\n"
        "x,2469,,,'@cmd\r\n"
    )  # a number is not prefixed; there is no BOM


async def test_export_without_escaping_writes_text_as_is(
    call: ToolCall, rich: Path, files: Path
) -> None:
    await call(
        "export_csv", path="rich.xlsx", sheet="Rich", csv_path="out.csv", escape_formulas=False
    )
    text = (files / "out.csv").read_bytes().decode("utf-8")
    assert ",=1+1" in text and "\r\n+15551234," in text


async def test_export_formatted_shows_what_excel_shows(
    call: ToolCall, rich: Path, files: Path
) -> None:
    await call(
        "export_csv", path="rich.xlsx", sheet="Rich", csv_path="out.csv", formatted=True,
        range="B2:C3",
    )  # fmt: skip
    assert (files / "out.csv").read_bytes().decode("utf-8") == (
        '"1,234.50","Jan 31, 2026"\r\n-200.0%,2026-01-31 9:30:00\r\n'
    )


async def test_export_computes_formulas_and_pads_ranges(
    call: ToolCall, rich: Path, files: Path
) -> None:
    result = await call(
        "export_csv", path="rich.xlsx", sheet="Rich", csv_path="out.csv", range="A4:C5"
    )
    assert result["range"] == "A4:C5"
    assert (files / "out.csv").read_bytes().decode("utf-8") == "x,2469,\r\n,,\r\n"


async def test_encodings_and_delimiters_on_export(call: ToolCall, rich: Path, files: Path) -> None:
    await call(
        "export_csv", path="rich.xlsx", sheet="Rich", csv_path="a.csv", range="A2:B2",
        encoding="utf-8-sig", delimiter="semicolon",
    )  # fmt: skip
    raw = (files / "a.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and raw.decode("utf-8-sig").endswith(";1234.5\r\n")
    await call(
        "export_csv", path="rich.xlsx", sheet="Rich", csv_path="b.tsv", range="A3:B3",
        encoding="cp1252", delimiter="tab",
    )  # fmt: skip
    assert (files / "b.tsv").read_bytes() == b"'+15551234\t-2\r\n"
    await call(
        "export_csv", path="rich.xlsx", sheet="Rich", csv_path="c.csv", range="A2:A2",
        encoding="utf-16",
    )  # fmt: skip
    assert (files / "c.csv").read_bytes().decode("utf-16").startswith('"Zoë')


async def test_export_refuses_what_the_encoding_cannot_hold(
    call_error: ToolCall, files: Path
) -> None:
    workbook = Workbook()
    workbook.active.append(["日本"])  # type: ignore[union-attr]
    workbook.save(files / "jp.xlsx")
    message = await call_error(
        "export_csv", path="jp.xlsx", sheet="Sheet", csv_path="o.csv", encoding="cp1252"
    )
    assert "cannot hold" in message
    assert not (files / "o.csv").exists()


async def test_export_will_not_replace_a_file_unless_asked(
    call: ToolCall, call_error: ToolCall, sample: Path, files: Path
) -> None:
    (files / "out.csv").write_bytes(b"keep\r\n")
    message = await call_error("export_csv", path="sales.xlsx", sheet="Data", csv_path="out.csv")
    assert "already exists" in message
    assert (files / "out.csv").read_bytes().decode() == "keep\r\n"
    await call("export_csv", path="sales.xlsx", sheet="Data", csv_path="out.csv", overwrite=True)
    assert (files / "out.csv").read_bytes().decode().startswith("Region,")
    assert not list(files.glob(".~*"))


async def test_export_of_an_unknown_sheet_or_workbook_is_an_error(
    call_error: ToolCall, sample: Path
) -> None:
    assert "Available sheets" in await call_error(
        "export_csv", path="sales.xlsx", sheet="Nope", csv_path="o.csv"
    )
    assert "does not exist" in await call_error(
        "export_csv", path="gone.xlsx", sheet="Data", csv_path="o.csv"
    )


async def test_export_then_import_round_trips(call: ToolCall, sample: Path, files: Path) -> None:
    await call("export_csv", path="sales.xlsx", sheet="Data", csv_path="out.csv")
    await call("import_csv", path="sales.xlsx", csv_path="out.csv", sheet="Copy")
    book = load_workbook(files / "sales.xlsx")
    assert [[c.value for c in row] for row in book["Copy"].iter_rows()] == [
        [c.value for c in row] for row in book["Data"].iter_rows()
    ]


async def test_csv_tools_are_listed_and_hidden_in_read_only_mode(
    client: Client, files: Path
) -> None:
    tools = {tool.name: tool for tool in (await client.list_tools()).tools}
    for name in ("import_csv", "export_csv"):
        annotations = tools[name].annotations
        assert annotations and not annotations.read_only_hint and annotations.destructive_hint
        assert annotations.open_world_hint is False
    assert tools["import_csv"].input_schema["required"] == ["path", "csv_path", "sheet"]
    assert tools["export_csv"].input_schema["properties"]["delimiter"]["enum"] == [
        "comma", "semicolon", "tab", "pipe",
    ]  # fmt: skip
    async with Client(create_server(Settings(allowed_dirs=[files], read_only=True))) as read_only:
        names = {tool.name for tool in (await read_only.list_tools()).tools}
    assert not {"import_csv", "export_csv"} & names
