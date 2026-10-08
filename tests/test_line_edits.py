from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.hyperlink import Hyperlink

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


async def insert(call: ToolCall, sheet: str, axis: str, at: int, count: int = 1) -> None:
    await call(
        "insert_rows_or_columns", path="sales.xlsx", sheet=sheet, axis=axis, at=at, count=count
    )


async def delete(call: ToolCall, sheet: str, axis: str, at: int, count: int = 1) -> None:
    await call(
        "delete_rows_or_columns", path="sales.xlsx", sheet=sheet, axis=axis, at=at, count=count
    )


async def test_formulas_everywhere_follow_the_move(call: ToolCall, sample: Path) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Data",
        start_cell="F1",
        rows=[["=B3+A5"], ["=SUM(C2:C5)"], ["=SUM(C:C)"], ["=$C$5*D$2"]],
    )
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="A1",
        rows=[["=Data!C3+Data!C5"], ["=SUM(Data!C2:C5)"], ["=A1+1"]],
    )
    await insert(call, "Data", "rows", 4, 3)
    data, report = load_workbook(sample)["Data"], load_workbook(sample)["Report"]
    assert [data[f"F{row}"].value for row in (1, 2, 3, 7)] == [
        "=B3+A8",
        "=SUM(C2:C8)",
        "=SUM(C:C)",
        "=$C$8*D$2",
    ]
    assert [report[f"A{row}"].value for row in (1, 2, 3)] == [
        "=Data!C3+Data!C8",
        "=SUM(Data!C2:C8)",
        "=A1+1",
    ]
    assert data["A7"].value == "North"
    assert data["A4"].value is None


async def test_deleted_cells_become_ref_errors(call: ToolCall, sample: Path) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="A1",
        rows=[["=Data!C3"], ["=SUM(Data!C2:C4)"], ["=SUM(Data!C3:C4)"], ["=Data!C5"]],
    )
    await delete(call, "Data", "rows", 3, 2)
    report = load_workbook(sample)["Report"]
    assert [report[f"A{row}"].value for row in (1, 2, 3, 4)] == [
        "=Data!#REF!",
        "=SUM(Data!C2:C2)",
        "=SUM(Data!#REF!)",
        "=Data!C3",
    ]


async def test_columns(call: ToolCall, sample: Path) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Data",
        start_cell="F2",
        rows=[["=C2+D2", "=SUM(B2:D2)", "=B2"]],
    )
    await delete(call, "Data", "columns", 2)
    data = load_workbook(sample)["Data"]
    assert [data[f"{column}2"].value for column in "EFG"] == ["=B2+C2", "=SUM(B2:C2)", "=#REF!"]
    await insert(call, "Data", "columns", 1, 2)
    moved = load_workbook(sample)["Data"]
    assert [moved[f"{column}2"].value for column in "GHI"] == ["=D2+E2", "=SUM(D2:E2)", "=#REF!"]


async def test_names_rules_merges_and_layout(call: ToolCall, sample: Path) -> None:
    await call("set_defined_name", path="sales.xlsx", name="Units", refers_to="Data!$C$2:$C$5")
    await call(
        "set_defined_name", path="sales.xlsx", name="Local", refers_to="Data!$A$3", sheet="Report"
    )
    await call(
        "add_conditional_format",
        path="sales.xlsx",
        sheet="Data",
        range="A3:D5",
        rule={"type": "formula", "formula": "=$C3>5", "fill_color": "FF0000"},
    )
    await call(
        "add_data_validation",
        path="sales.xlsx",
        sheet="Data",
        range="C3:C5",
        rule={"type": "whole", "operator": "lessThan", "minimum": "$C$2"},
    )
    await call("merge_cells", path="sales.xlsx", sheet="Data", range="F3:G4")
    await call(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Data",
        layout={
            "freeze_panes": "A3",
            "row_heights": {"3": 30},
            "print_setup": {"print_area": "A1:D5", "title_rows": "1:2"},
        },
    )
    await insert(call, "Data", "rows", 2)
    data = load_workbook(sample)["Data"]
    book = load_workbook(sample)
    assert book.defined_names["Units"].attr_text == "Data!$C$3:$C$6"
    assert book["Report"].defined_names["Local"].attr_text == "Data!$A$4"
    assert [str(entry.sqref) for entry in data.conditional_formatting] == ["A4:D6"]
    assert data.conditional_formatting._cf_rules  # pyright: ignore[reportAttributeAccessIssue]
    rule = next(iter(data.conditional_formatting)).rules[0]
    assert rule.formula == ["$C4>5"]
    validation = data.data_validations.dataValidation[0]
    assert (str(validation.sqref), validation.formula1) == ("C4:C6", "$C$3")
    assert [str(area) for area in data.merged_cells.ranges] == ["F4:G5"]
    assert data.freeze_panes == "A4"
    assert data.row_dimensions[4].height == 30
    assert data.print_area == "'Data'!$A$1:$D$6"
    assert data.print_title_rows == "$1:$3"


async def test_conditional_format_split_by_a_deleted_first_row(
    call: ToolCall, sample: Path
) -> None:
    await call(
        "add_conditional_format",
        path="sales.xlsx",
        sheet="Data",
        range="A3:C8",
        rule={"type": "formula", "formula": "=$A3>$B5", "fill_color": "FF0000"},
    )
    await delete(call, "Data", "rows", 3)
    entries = list(load_workbook(sample)["Data"].conditional_formatting)
    assert [(str(entry.sqref), entry.rules[0].formula) for entry in entries] == [
        ("A3:C7", ["$A3>$B5"])
    ]


async def test_conditional_format_split_in_the_middle(call: ToolCall, sample: Path) -> None:
    await call(
        "add_conditional_format",
        path="sales.xlsx",
        sheet="Data",
        range="A3:C8",
        rule={"type": "formula", "formula": "=$A3>$B5", "fill_color": "FF0000"},
    )
    await insert(call, "Data", "rows", 6)
    entries = list(load_workbook(sample)["Data"].conditional_formatting)
    found = {str(entry.sqref): entry.rules[0].formula for entry in entries}
    assert found == {"A3:C3 A7:C9": ["$A3>$B5"], "A4:C6": ["$A4>$B7"]}


async def test_chart_and_hyperlink_and_note_move(call: ToolCall, sample: Path) -> None:
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        data_range="Data!B1:C5",
        chart_type="column",
        anchor_cell="B2",
    )
    await call("set_note", path="sales.xlsx", sheet="Data", cell="A3", text="hello")
    book = load_workbook(sample)
    book["Data"]["A4"].hyperlink = Hyperlink(ref="A4", location="Report!A1")
    book.save(sample)
    await insert(call, "Data", "rows", 3)
    await insert(call, "Report", "rows", 1)
    book = load_workbook(sample)
    chart = book["Report"]._charts[0]  # pyright: ignore[reportAttributeAccessIssue]
    assert chart.series[0].val.numRef.f == "'Data'!$C$2:$C$6"
    assert chart.anchor._from.row == 2
    assert book["Data"]["A4"].comment is not None
    assert book["Data"]["A5"].hyperlink.ref == "A5"


async def test_table_columns_and_structured_references(call: ToolCall, sample: Path) -> None:
    await call("create_table", path="sales.xlsx", sheet="Data", range="A1:D5", name="Sales")
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="A1",
        rows=[["=SUM(Sales[Units])"], ["=SUM(Sales[Price])"]],
    )
    await insert(call, "Data", "columns", 2)
    table = load_workbook(sample)["Data"].tables["Sales"]
    assert table.ref == "A1:E5"
    assert [column.name for column in table.tableColumns] == [
        "Region",
        "Column1",
        "Product",
        "Units",
        "Price",
    ]
    assert load_workbook(sample)["Data"]["B1"].value == "Column1"
    await delete(call, "Data", "columns", 4)
    book = load_workbook(sample)
    assert [c.name for c in book["Data"].tables["Sales"].tableColumns] == [
        "Region",
        "Column1",
        "Product",
        "Price",
    ]
    assert [book["Report"]["A1"].value, book["Report"]["A2"].value] == [
        "=SUM(#REF!)",
        "=SUM(Sales[Price])",
    ]


async def test_table_rows(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    await call("create_table", path="sales.xlsx", sheet="Data", range="A1:D5", name="Sales")
    message = await call_error(
        "delete_rows_or_columns", path="sales.xlsx", sheet="Data", axis="rows", at=1, count=2
    )
    assert "header row of table 'Sales'" in message
    await insert(call, "Data", "rows", 3)
    assert load_workbook(sample)["Data"].tables["Sales"].ref == "A1:D6"
    await delete(call, "Data", "rows", 2, 5)
    data = load_workbook(sample)["Data"]
    assert data.tables["Sales"].ref == "A1:D2"
    assert data["A2"].value is None  # Excel keeps one empty data row
    await delete(call, "Data", "rows", 1, 2)
    assert not load_workbook(sample)["Data"].tables


async def test_array_formulas_move_but_cannot_be_split(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    book = load_workbook(sample)
    book["Report"]["A5"] = ArrayFormula("A5:A7", "=Data!C2:C4*2")
    book.save(sample)
    for tool in ("insert_rows_or_columns", "delete_rows_or_columns"):
        message = await call_error(tool, path="sales.xlsx", sheet="Report", axis="rows", at=6)
        assert "array formula in A5:A7" in message
    await insert(call, "Data", "rows", 2)
    await insert(call, "Report", "rows", 1)
    formula = load_workbook(sample)["Report"]["A6"].value
    assert isinstance(formula, ArrayFormula)
    assert (formula.ref, formula.text) == ("A6:A8", "=Data!C3:C5*2")


async def test_pivot_tables_move_and_cannot_be_split(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await call(
        "create_pivot_table",
        path="sales.xlsx",
        source_sheet="Data",
        source_range="A1:D5",
        rows=["Region"],
        values=[{"field": "Units"}],
        target_sheet="Report",
        target_cell="B3",
    )
    message = await call_error(
        "insert_rows_or_columns", path="sales.xlsx", sheet="Report", axis="rows", at=4
    )
    assert "PivotTable" in message
    await insert(call, "Report", "rows", 1)
    await insert(call, "Data", "rows", 3)
    pivot = load_workbook(sample)["Report"]._pivots[0]  # pyright: ignore[reportAttributeAccessIssue]
    assert pivot.location.ref == "B4:C7"
    assert pivot.cache.cacheSource.worksheetSource.ref == "A1:D6"


async def test_limits_and_paths(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    await call("write_range", path="sales.xlsx", sheet="Data", start_cell="A1048500", rows=[[1]])
    message = await call_error(
        "insert_rows_or_columns", path="sales.xlsx", sheet="Data", axis="rows", at=1, count=1000
    )
    assert "past the last" in message
    outside = await call_error(
        "delete_rows_or_columns", path="../outside.xlsx", sheet="Data", axis="rows", at=1
    )
    assert "outside" in outside.lower() or "not allowed" in outside.lower()


async def test_nothing_is_saved_when_an_edit_is_refused(call_error: ToolCall, sample: Path) -> None:
    book = Workbook()
    book.worksheets[0]["A1"] = ArrayFormula("A1:A3", "=B1:B3")
    path = sample.parent / "array.xlsx"
    book.save(path)
    before = path.read_bytes()
    await call_error("insert_rows_or_columns", path="array.xlsx", sheet="Sheet", axis="rows", at=2)
    assert path.read_bytes() == before


async def test_inserted_lines_inherit_sizes_and_rule_ranges(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Data",
        layout={
            "row_heights": {"3": 30},
            "column_widths": {"B": 20},
            "columns": [{"span": "D", "action": "hide"}],
            "rows": [{"span": "4", "action": "hide"}],
        },
    )
    await call(
        "add_data_validation",
        path="sales.xlsx",
        sheet="Data",
        range="A2:A5",
        rule={"type": "list", "options": ["a", "b"]},
    )
    await insert(call, "Data", "rows", 4)
    await insert(call, "Data", "rows", 6)
    await insert(call, "Data", "columns", 3)
    await insert(call, "Data", "columns", 6)
    data = load_workbook(sample)["Data"]
    heights = [data.row_dimensions[row].height for row in (3, 4, 5)]
    hidden = [bool(data.row_dimensions[row].hidden) for row in (4, 5, 6)]
    assert (heights, hidden) == ([30, 30, None], [False, True, False])
    assert data.column_dimensions["C"].width == 20 and not data.column_dimensions["C"].hidden
    assert data.column_dimensions["E"].hidden and not data.column_dimensions["F"].hidden
    assert str(data.data_validations.dataValidation[0].sqref) == "A2:A7"
