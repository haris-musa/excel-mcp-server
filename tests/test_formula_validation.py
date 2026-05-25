import os
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from excel_mcp.calculations import apply_formula  # noqa: E402
from excel_mcp.exceptions import CalculationError  # noqa: E402
from excel_mcp.validation import validate_formula  # noqa: E402


class TestFormulaSafetyValidation(unittest.TestCase):
    def test_rejects_unsafe_functions_case_insensitively(self):
        formulas = [
            '=WEBSERVICE("http://127.0.0.1:9/canary")',
            '=webservice("http://127.0.0.1:9/canary")',
            '=HyperLink("http://127.0.0.1:9/canary","open")',
            '=indirect("A1")',
            '=Dget(A1:B2,"field",C1:D2)',
            '=rTd("prog.id",,"topic")',
        ]

        for formula in formulas:
            with self.subTest(formula=formula):
                is_valid, message = validate_formula(formula)
                self.assertFalse(is_valid)
                self.assertIn("Unsafe function", message)

    def test_rejects_direct_external_workbook_references(self):
        formulas = [
            "='[external-secrets.xlsx]Sheet1'!A1",
            "='C:\\tmp\\[external-secrets.xlsx]Sheet1'!A1",
            "='https://example.com/[book.xlsx]Sheet1'!A1",
            '=SUM([external-secrets.xlsx]Sheet1!A1, 1)',
        ]

        for formula in formulas:
            with self.subTest(formula=formula):
                is_valid, message = validate_formula(formula)
                self.assertFalse(is_valid)
                self.assertEqual(message, "External workbook references are not allowed")

    def test_allows_safe_formula(self):
        self.assertEqual(validate_formula("=SUM(1, 2)"), (True, "Formula is valid"))

    def test_apply_formula_does_not_persist_rejected_formula(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name

        try:
            wb = Workbook()
            wb.active.title = "Sheet1"
            wb.save(path)

            with self.assertRaises(CalculationError):
                apply_formula(
                    path,
                    "Sheet1",
                    "A1",
                    '=webservice("http://127.0.0.1:9/canary")',
                )

            reloaded = load_workbook(path, data_only=False)
            self.assertIsNone(reloaded["Sheet1"]["A1"].value)
        finally:
            os.unlink(path)


if __name__ == "__main__":
    unittest.main()
