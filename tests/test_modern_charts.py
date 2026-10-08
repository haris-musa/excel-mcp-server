"""Waterfall, histogram, Pareto, box and whisker, treemap, sunburst and funnel charts.

The expected XML is what Microsoft Excel itself writes for the same chart (checked by
creating each chart in Excel on the same data and comparing the parts element by element).
"""

import re
import zipfile
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from excel_mcp import package
from excel_mcp.package import state_of
from excel_mcp.package.lines import LineEdit
from excel_mcp.package.references import rewrite_lines
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

LABELLED = {"values": "Data!C2:C5", "name": "Data!C1", "categories": "Data!A2:A5"}
CHARTS: dict[str, dict[str, Any]] = {
    "waterfall": {"series": [LABELLED]},
    "histogram": {"series": [{"values": "Data!C2:C5", "name": "Data!C1"}]},
    "pareto": {"series": [LABELLED]},
    "box_whisker": {"series": [LABELLED]},
    "treemap": {"series": [{**LABELLED, "categories": "Data!A2:B5"}]},
    "sunburst": {"series": [{**LABELLED, "categories": "Data!A2:B5"}]},
    "funnel": {"series": [LABELLED]},
}
LAYOUTS = {
    "waterfall": "waterfall",
    "histogram": "clusteredColumn",
    "pareto": "clusteredColumn",
    "box_whisker": "boxWhisker",
    "treemap": "treemap",
    "sunburst": "sunburst",
    "funnel": "funnel",
}


async def add(call: ToolCall, chart_type: str, anchor: str = "B2", **extra: Any) -> str:
    arguments = {**CHARTS[chart_type], **extra}
    return await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        chart_type=chart_type,
        anchor_cell=anchor,
        **arguments,
    )


def parts(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {n: archive.read(n).decode("utf-8") for n in archive.namelist()}


def chart_text(path: Path, number: int = 1) -> str:
    return parts(path)[f"xl/charts/chartEx{number}.xml"]


def hidden_names(path: Path) -> dict[str, str]:
    workbook = parts(path)["xl/workbook.xml"]
    found = re.findall(r'<definedName name="(_xlchart[^"]+)" hidden="1">([^<]*)<', workbook)
    return dict(found)


@pytest.mark.parametrize("chart_type", list(CHARTS))
async def test_each_type_writes_the_parts_excel_writes(
    call: ToolCall, sample: Path, chart_type: str
) -> None:
    await add(call, chart_type)

    saved = parts(sample)
    chart = saved["xl/charts/chartEx1.xml"]
    assert f'<cx:series layoutId="{LAYOUTS[chart_type]}"' in chart
    assert (chart_type == "pareto") == ('layoutId="paretoLine" ownerIdx="0"' in chart)
    for name in ("xl/charts/style1.xml", "xl/charts/colors1.xml"):
        assert name in saved
    types = saved["[Content_Types].xml"]
    for name, kind in (
        ("chartEx1", "application/vnd.ms-office.chartex+xml"),
        ("style1", "application/vnd.ms-office.chartstyle+xml"),
        ("colors1", "application/vnd.ms-office.chartcolorstyle+xml"),
    ):
        assert f'PartName="/xl/charts/{name}.xml" ContentType="{kind}"' in types
    assert "relationships/chartEx" in saved["xl/drawings/_rels/drawing1.xml.rels"]
    assert "relationships/chartStyle" in saved["xl/charts/_rels/chartEx1.xml.rels"]
    assert "relationships/chartColorStyle" in saved["xl/charts/_rels/chartEx1.xml.rels"]
    drawing = saved["xl/drawings/drawing1.xml"]
    required = "cx2" if chart_type == "funnel" else "cx1"
    assert f'<mc:Choice xmlns:{required}="' in drawing and f'Requires="{required}"' in drawing
    assert "<xdr:twoCellAnchor" in drawing and "<mc:Fallback>" in drawing


@pytest.mark.parametrize(
    ("chart_type", "version", "ranges"),
    [
        ("waterfall", 1, ["Data!$A$2:$A$5", "Data!$C$1", "Data!$C$2:$C$5"]),
        ("histogram", 1, ["Data!$C$1", "Data!$C$2:$C$5"]),
        ("treemap", 1, ["Data!$A$2:$B$5", "Data!$C$1", "Data!$C$2:$C$5"]),
        ("funnel", 2, ["Data!$A$2:$A$5", "Data!$C$1", "Data!$C$2:$C$5"]),
    ],
)
async def test_data_lives_in_hidden_names_numbered_like_excels(
    call: ToolCall, sample: Path, chart_type: str, version: int, ranges: list[str]
) -> None:
    await add(call, chart_type)

    names = hidden_names(sample)
    assert list(names) == [f"_xlchart.v{version}.{n}" for n in range(len(ranges))]
    assert list(names.values()) == ranges
    formulas = re.findall(r"<cx:f>([^<]*)</cx:f>", chart_text(sample))
    assert sorted(formulas) == sorted(names)
    assert "<cx:v>Units</cx:v>" in chart_text(sample)


async def test_a_second_chart_continues_the_numbering(call: ToolCall, sample: Path) -> None:
    await add(call, "waterfall")
    await add(call, "funnel", "B20")

    assert list(hidden_names(sample)) == [
        *(f"_xlchart.v1.{n}" for n in range(3)),
        *(f"_xlchart.v2.{n}" for n in range(3, 6)),
    ]
    assert "xl/charts/chartEx2.xml" in parts(sample)


async def test_title_legend_labels_and_axes(call: ToolCall, sample: Path) -> None:
    await add(
        call,
        "waterfall",
        options={
            "title": "Bridge <1>",
            "legend": "none",
            "data_labels": {
                "show": ["value", "category"],
                "position": "outside_end",
                "number_format": "0.0",
            },
            "y_axis": {"title": "Units", "min": 0, "max": 40, "major_unit": 10,
                       "number_format": "0", "major_gridlines": False},
            "x_axis": {"title": "Region"},
        },
    )  # fmt: skip

    chart = chart_text(sample)
    assert "<cx:v>Bridge &lt;1&gt;</cx:v>" in chart
    assert "<cx:legend" not in chart
    assert (
        '<cx:dataLabels pos="outEnd"><cx:numFmt formatCode="0.0" sourceLinked="0"/>'
        '<cx:visibility seriesName="0" categoryName="1" value="1"/></cx:dataLabels>'
    ) in chart
    assert '<cx:valScaling max="40" min="0" majorUnit="10"/>' in chart
    assert "<cx:majorGridlines/>" not in chart
    assert '<cx:numFmt formatCode="0" sourceLinked="0"/></cx:axis>' in chart
    assert chart.count("<cx:title><cx:tx><cx:txData><cx:v>") == 2


async def test_the_legend_sits_where_asked(call: ToolCall, sample: Path) -> None:
    await add(call, "treemap", options={"legend": "right"})

    assert '<cx:legend pos="r" align="ctr" overlay="0"/>' in chart_text(sample)


async def test_waterfall_totals_and_connectors(call: ToolCall, sample: Path) -> None:
    await add(call, "waterfall", options={"totals": [4, 1, 4], "connector_lines": False})

    chart = chart_text(sample)
    assert '<cx:visibility connectorLines="0"/>' in chart
    assert '<cx:subtotals><cx:idx val="0"/><cx:idx val="3"/></cx:subtotals>' in chart


async def test_a_plain_waterfall_has_no_totals(call: ToolCall, sample: Path) -> None:
    await add(call, "waterfall")

    assert "<cx:subtotals></cx:subtotals>" in chart_text(sample)
    assert "connectorLines" not in chart_text(sample)


@pytest.mark.parametrize(
    ("bins", "expected"),
    [
        ({}, '<cx:binning intervalClosed="r"></cx:binning>'),
        ({"width": 5}, '<cx:binning intervalClosed="r"><cx:binSize val="5"/></cx:binning>'),
        ({"count": 4}, '<cx:binning intervalClosed="r"><cx:binCount val="4"/></cx:binning>'),
        (
            {"width": 2.5, "underflow": 1, "overflow": 20},
            '<cx:binning intervalClosed="r" underflow="1" overflow="20">'
            '<cx:binSize val="2.5"/></cx:binning>',
        ),
    ],
)
async def test_histogram_bins(
    call: ToolCall, sample: Path, bins: dict[str, float], expected: str
) -> None:
    await add(call, "histogram", options={"bins": bins})

    assert expected in chart_text(sample)


async def test_pareto_has_its_cumulative_line_and_percentage_axis(
    call: ToolCall, sample: Path
) -> None:
    await add(call, "pareto")

    chart = chart_text(sample)
    assert "<cx:aggregation/>" in chart
    assert '<cx:axis id="2"><cx:valScaling max="1" min="0"/><cx:units unit="percentage"/>' in chart
    assert chart.count("<cx:axisId") == 2


async def test_box_defaults_are_excels_and_options_write_visibility(
    call: ToolCall, sample: Path
) -> None:
    await add(call, "box_whisker")
    assert '<cx:layoutPr><cx:statistics quartileMethod="exclusive"/></cx:layoutPr>' in (
        chart_text(sample)
    )

    await add(
        call, "box_whisker", "B30", index=1,
        options={"box": {"quartiles": "inclusive", "mean_line": True, "outliers": False}},
    )  # fmt: skip
    chart = chart_text(sample)
    assert (
        '<cx:visibility meanLine="1" meanMarker="1" nonoutliers="0" outliers="0"/>'
        '<cx:statistics quartileMethod="inclusive"/>'
    ) in chart


async def test_box_and_histogram_take_several_series(call: ToolCall, sample: Path) -> None:
    second = {"values": "Data!D2:D5", "name": "Data!D1", "categories": "Data!A2:A5"}
    await add(call, "box_whisker", series=[LABELLED, second])

    chart = chart_text(sample)
    assert chart.count('<cx:series layoutId="boxWhisker"') == 2
    assert chart.count("<cx:data id=") == 2
    assert chart.count("<cx:dataId val=") == 2
    assert len(hidden_names(sample)) == 5


async def test_treemap_parent_labels(call: ToolCall, sample: Path) -> None:
    await add(call, "treemap", options={"parent_labels": "banner"})

    assert '<cx:parentLabelLayout val="banner"/>' in chart_text(sample)


async def test_treemap_block_reads_levels_then_sizes(call: ToolCall, sample: Path) -> None:
    await call(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type="sunburst",
        anchor_cell="B2", data_range="Data!A1:C5",
    )  # fmt: skip

    assert list(hidden_names(sample).values()) == [
        "Data!$A$2:$B$5",
        "Data!$C$1",
        "Data!$C$2:$C$5",
    ]


async def test_blocks_for_labelled_and_value_only_charts(call: ToolCall, sample: Path) -> None:
    await call(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type="waterfall",
        anchor_cell="B2", data_range="Data!B1:C5",
    )  # fmt: skip
    await call(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type="histogram",
        anchor_cell="B20", data_range="Data!C1:D5",
    )  # fmt: skip

    names = list(hidden_names(sample).values())
    assert names[:3] == ["Data!$B$2:$B$5", "Data!$C$1", "Data!$C$2:$C$5"]
    assert names[3:] == ["Data!$C$1", "Data!$C$2:$C$5", "Data!$D$1", "Data!$D$2:$D$5"]


async def test_size_follows_the_options_and_the_sheet_geometry(
    call: ToolCall, sample: Path
) -> None:
    await add(call, "funnel", "C3", options={"width_cm": 10, "height_cm": 5})

    drawing = parts(sample)["xl/drawings/drawing1.xml"]
    assert "<xdr:from><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row>" in drawing
    assert '<a:ext cx="3600000" cy="1800000"/>' in drawing


async def test_charts_are_listed_after_the_classic_ones(call: ToolCall, sample: Path) -> None:
    await add(call, "waterfall", options={"title": "Bridge"})
    await call(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type="column",
        anchor_cell="B30", data_range="Data!B1:C5",
    )  # fmt: skip
    await add(call, "pareto", "B60")

    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")

    assert [(c["index"], c["type"], c["title"], c["anchor"]) for c in details["charts"]] == [
        (1, "column", None, "B30"),
        (2, "waterfall", "Bridge", "B2"),
        (3, "pareto", None, "B60"),
    ]
    assert details["charts"][1]["series"] == ["Data!$C$2:$C$5"]


async def test_delete_removes_the_chart_its_parts_and_its_names(
    call: ToolCall, sample: Path
) -> None:
    await add(call, "waterfall")
    await add(call, "funnel", "B30")

    message = await call("delete_chart", path="sales.xlsx", sheet="Report", index=1)

    assert "Deleted waterfall chart 1" in message
    saved = parts(sample)
    remaining = [n for n in saved if re.fullmatch(r"xl/charts/chartEx\d+\.xml", n)]
    assert len(remaining) == 1 and 'layoutId="funnel"' in saved[remaining[0]]
    assert list(hidden_names(sample)) == [f"_xlchart.v2.{n}" for n in (3, 4, 5)]
    await call("delete_chart", path="sales.xlsx", sheet="Report", index=1)
    saved = parts(sample)
    assert not [n for n in saved if "chart" in n.lower() or "drawing" in n.lower()]
    assert not hidden_names(sample)


async def test_replacing_a_chart_keeps_its_place(call: ToolCall, sample: Path) -> None:
    await add(call, "waterfall")
    await add(call, "funnel", "B30")

    message = await add(call, "treemap", "B60", index=1)

    assert message == "Replaced chart 1 of Report with a treemap chart at B60."
    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert [(c["index"], c["type"], c["anchor"]) for c in details["charts"]] == [
        (1, "treemap", "B60"),
        (2, "funnel", "B30"),
    ]
    assert len(hidden_names(sample)) == 6


async def test_replacing_across_kinds_moves_the_new_chart_last(
    call: ToolCall, sample: Path
) -> None:
    await call(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type="column",
        anchor_cell="B30", data_range="Data!B1:C5",
    )  # fmt: skip
    await add(call, "waterfall")

    message = await add(call, "funnel", "B60", index=1)
    assert "It is now chart 2." in message
    message = await call(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type="line",
        anchor_cell="B2", data_range="Data!B1:C5", index=2,
    )  # fmt: skip
    assert "It is now chart 1." in message

    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert [c["type"] for c in details["charts"]] == ["line", "waterfall"]
    assert len(hidden_names(sample)) == 3


async def test_copy_sheet_copies_the_charts_pointing_at_the_copy(
    call: ToolCall, sample: Path
) -> None:
    await call(
        "create_chart", path="sales.xlsx", sheet="Data", chart_type="funnel",
        anchor_cell="F2", **CHARTS["funnel"],
    )  # fmt: skip

    await call("copy_sheet", path="sales.xlsx", sheet="Data", new_name="Data (2)")

    saved = parts(sample)
    assert len([n for n in saved if re.fullmatch(r"xl/charts/chartEx\d+\.xml", n)]) == 2
    names = hidden_names(sample)
    assert sorted(names.values()) == sorted([
        "Data!$A$2:$A$5", "Data!$C$1", "Data!$C$2:$C$5",
        "'Data (2)'!$A$2:$A$5", "'Data (2)'!$C$1", "'Data (2)'!$C$2:$C$5",
    ])  # fmt: skip
    copied = await call("describe_sheet", path="sales.xlsx", sheet="Data (2)")
    assert copied["charts"][0]["series"] == ["'Data (2)'!$C$2:$C$5"]
    original = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert original["charts"][0]["series"] == ["Data!$C$2:$C$5"]
    for name in ("style", "colors"):
        assert sum(1 for n in saved if n.startswith(f"xl/charts/{name}")) >= 1


async def test_charts_survive_other_edits(call: ToolCall, sample: Path) -> None:
    await add(call, "sunburst")
    before = chart_text(sample)

    await call("write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[[1]])
    await call("create_sheet", path="sales.xlsx", sheet="Other")

    assert re.sub(r"uniqueId=\"[^\"]*\"", "", chart_text(sample)) == re.sub(
        r"uniqueId=\"[^\"]*\"", "", before
    )
    assert list(hidden_names(sample)) == [f"_xlchart.v1.{n}" for n in range(3)]


async def test_anchors_follow_inserted_rows_through_the_references_api(
    call: ToolCall, sample: Path
) -> None:
    await add(call, "waterfall", "B4")
    workbook = load_workbook(sample)
    package.capture(sample, workbook, 50_000_000)

    rewrite_lines(workbook, LineEdit("Report", "rows", 1, 2, False), lambda text, host: text)

    anchor = state_of(workbook).sheet(workbook["Report"]).anchors[0]
    assert re.search(r"<xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5<", anchor)


@pytest.mark.parametrize(
    ("chart_type", "arguments", "message"),
    [
        ("waterfall", {"options": {"totals": [9]}}, "totals are positions from 1 to 4"),
        ("waterfall", {"options": {"totals": [0]}}, "totals are positions"),
        ("column", {"options": {"totals": [1]}}, "totals does not apply to column charts"),
        ("funnel", {"options": {"bins": {"count": 3}}}, "bins does not apply to funnel"),
        ("histogram", {"options": {"bins": {"width": 1, "count": 2}}}, "width or a count"),
        ("histogram", {"options": {"bins": {"underflow": 5, "overflow": 1}}}, "underflow must"),
        ("waterfall", {"options": {"box": {"mean_line": True}}}, "box does not apply to waterfall"),
        ("pie", {"options": {"parent_labels": "banner"}}, "parent_labels does not apply to pie"),
        ("waterfall", {"options": {"colors": ["#FF0000"]}}, "colors does not apply"),
        ("waterfall", {"options": {"grouping": "stacked"}}, "grouping does not apply"),
        ("waterfall", {"options": {"style": 3}}, "style does not apply"),
        ("funnel", {"options": {"y_axis": {"title": "x"}}}, "y_axis.title does not apply"),
        ("treemap", {"options": {"x_axis": {"title": "x"}}}, "x_axis.title does not apply"),
        ("waterfall", {"options": {"x_axis": {"reverse": True}}}, "x_axis.reverse does not"),
        ("waterfall", {"options": {"y_axis": {"log": True}}}, "y_axis.log does not apply"),
        (
            "funnel",
            {"options": {"data_labels": {"position": "center"}}},
            "place their labels themselves",
        ),
        (
            "waterfall",
            {"options": {"data_labels": {"show": ["percent"]}}},
            "Percent labels fit pie",
        ),
        (
            "waterfall",
            {"options": {"data_labels": {"position": "above"}}},
            "does not fit",
        ),
        (
            "waterfall",
            {"series": [{**LABELLED, "color": "#FF0000"}]},
            "Series 1: color does not apply",
        ),
        (
            "waterfall",
            {"series": [{**LABELLED, "trendline": {"type": "linear"}}]},
            "trendline does not apply",
        ),
        ("waterfall", {"series": [LABELLED, LABELLED]}, "plots one series, got 2"),
        ("histogram", {"series": [LABELLED]}, "leave out categories"),
        ("pareto", {"series": [{"values": "Data!C2:C5"}]}, "pareto chart needs categories"),
        ("treemap", {"series": [{"values": "Data!C2:C5"}]}, "treemap chart needs categories"),
        ("waterfall", {"series_in": "rows"}, "read their data in columns"),
        ("waterfall", {"series": [], "data_range": "Data!A1:A5"}, "label column and a value"),
        ("histogram", {"series": [], "data_range": "Data!C1:C1"}, "header row"),
        ("waterfall", {"series": [], "data_range": None}, "either data_range"),
        ("waterfall", {"series": [{"values": "Data!B2:C5"}]}, "one row or one column"),
        ("waterfall", {"series": [{"values": "Nope!B2:B5"}]}, "not found"),
    ],
)
async def test_invalid_requests_are_errors_and_change_nothing(
    call_error: ToolCall, sample: Path, chart_type: str, arguments: dict[str, Any], message: str
) -> None:
    before = parts(sample)
    defaults = {**CHARTS.get(chart_type, {"data_range": "Data!B1:C5"})}
    if "data_range" in arguments or "series" in arguments:
        defaults.pop("series", None)
    result = await call_error(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type=chart_type,
        anchor_cell="B2", **{**defaults, **{k: v for k, v in arguments.items() if v is not None}},
    )  # fmt: skip

    assert message in result
    assert parts(sample) == before


async def test_modern_charts_need_a_cell_to_sit_at(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_chart", path="sales.xlsx", sheet="Fresh", chart_type="funnel", **CHARTS["funnel"]
    )

    assert "give anchor_cell" in message


async def test_index_must_exist(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    await add(call, "funnel")

    message = await call_error(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type="treemap",
        anchor_cell="B2", index=2, **CHARTS["treemap"],
    )  # fmt: skip
    assert "no chart 2" in message


async def test_stays_inside_the_allowed_folder(
    call_error: ToolCall, sample: Path, tmp_path: Path
) -> None:
    outside = tmp_path / "outside.xlsx"
    outside.write_bytes(sample.read_bytes())

    message = await call_error(
        "create_chart", path=str(outside), sheet="Report", chart_type="funnel",
        anchor_cell="B2", **CHARTS["funnel"],
    )  # fmt: skip

    assert "outside" in message


async def test_the_tool_schema_lists_the_new_types(client: Any) -> None:
    tools = {tool.name: tool for tool in (await client.list_tools()).tools}

    schema = str(tools["create_chart"].input_schema)

    for chart_type in CHARTS:
        assert f"'{chart_type}'" in schema
    assert "'totals'" in schema and "'bins'" in schema


@pytest.mark.parametrize(
    ("chart_type", "legend"),
    [
        ("waterfall", "t"), ("treemap", "t"), ("histogram", None), ("pareto", None),
        ("box_whisker", None), ("sunburst", None), ("funnel", None),
    ],
)  # fmt: skip
async def test_legend_default_is_excels_for_the_chart_type(
    call: ToolCall, sample: Path, chart_type: str, legend: str | None
) -> None:
    await add(call, chart_type)

    found = re.findall(r'<cx:legend pos="(\w)"', chart_text(sample))
    assert found == ([legend] if legend else [])


async def test_pareto_percentage_axis_takes_the_secondary_axis_options(
    call: ToolCall, sample: Path
) -> None:
    await add(
        call, "pareto",
        options={"secondary_y_axis": {"title": "Share", "max": 0.8}},
    )  # fmt: skip

    assert (
        '<cx:axis id="2"><cx:valScaling max="0.8" min="0"/><cx:title><cx:tx><cx:txData>'
        '<cx:v>Share</cx:v></cx:txData></cx:tx></cx:title><cx:units unit="percentage"/>'
        "<cx:tickLabels/></cx:axis>"
    ) in chart_text(sample)


async def test_the_secondary_axis_fits_pareto_only(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_chart", path="sales.xlsx", sheet="Report", chart_type="waterfall",
        anchor_cell="B2", options={"secondary_y_axis": {"title": "x"}}, **CHARTS["waterfall"],
    )  # fmt: skip

    assert "secondary_y_axis.title does not apply" in message
