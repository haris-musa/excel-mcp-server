import zipfile
from base64 import b64encode
from hashlib import sha512
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.protection import WorkbookProtection

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


async def settings(call: ToolCall, **values: object) -> str:
    return await call("set_workbook_settings", path="sales.xlsx", settings=values)


def app_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("docProps/app.xml").decode()


async def test_properties_including_company(call: ToolCall, sample: Path) -> None:
    await settings(
        call,
        doc_properties={
            "title": "Q3 Sales",
            "subject": "Revenue",
            "author": "Ada",
            "keywords": "sales, q3",
            "company": "Acme",
        },
    )
    properties = load_workbook(sample).properties
    assert (properties.title, properties.subject, properties.creator, properties.keywords) == (
        "Q3 Sales",
        "Revenue",
        "Ada",
        "sales, q3",
    )
    assert "<Company>Acme</Company>" in app_xml(sample)
    info = await call("describe_workbook", path="sales.xlsx")
    assert info["doc_properties"] == {
        "title": "Q3 Sales",
        "subject": "Revenue",
        "author": "Ada",
        "keywords": "sales, q3",
        "company": "Acme",
    }


async def test_company_survives_other_edits_and_can_be_cleared(
    call: ToolCall, sample: Path
) -> None:
    await settings(call, doc_properties={"company": "Acme", "title": "T"})
    await call("create_sheet", path="sales.xlsx", new_name="Extra")
    assert "Acme" in app_xml(sample)
    await settings(call, doc_properties={"company": ""})
    assert "Company" not in app_xml(sample)
    assert load_workbook(sample).properties.title == "T"
    assert load_workbook(sample).sheetnames == ["Data", "Report", "Extra"]


async def test_calculation_settings(call: ToolCall, sample: Path) -> None:
    await settings(
        call,
        calculation={
            "mode": "manual",
            "iterative": True,
            "max_iterations": 50,
            "max_change": 0.01,
            "full_calc_on_load": False,
        },
    )
    calc = load_workbook(sample).calculation
    assert (calc.calcMode, calc.iterate, calc.iterateCount, calc.iterateDelta) == (
        "manual",
        True,
        50,
        0.01,
    )
    assert calc.fullCalcOnLoad is False
    info = await call("describe_workbook", path="sales.xlsx")
    assert info["calculation"] == {
        "mode": "manual",
        "iterative": True,
        "max_iterations": 50,
        "max_change": 0.01,
    }
    await settings(call, calculation={"mode": "auto_except_tables"})
    assert load_workbook(sample).calculation.calcMode == "autoNoTable"


async def test_structure_protection_with_password(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await settings(call, structure_protection={"enabled": True, "password": "s3cret"})
    security = load_workbook(sample).security
    assert security.lockStructure is True
    assert security.workbookPassword
    assert (await call("describe_workbook", path="sales.xlsx"))["structure_protected"] is True
    assert "already protected" in await call_error(
        "set_workbook_settings",
        path="sales.xlsx",
        settings={"structure_protection": {"enabled": True}},
    )
    assert "right password" in await call_error(
        "set_workbook_settings",
        path="sales.xlsx",
        settings={"structure_protection": {"enabled": False, "password": "wrong"}},
    )
    await settings(call, structure_protection={"enabled": False, "password": "s3cret"})
    assert not load_workbook(sample).security.lockStructure


async def test_unprotect_structure_protected_by_excel_2013(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    salt = b"0123456789abcdef"
    digest = sha512(salt + "pw123".encode("utf-16-le")).digest()
    for index in range(100):
        digest = sha512(digest + index.to_bytes(4, "little")).digest()
    workbook = load_workbook(sample)
    workbook.security = WorkbookProtection(
        lockStructure=True,
        workbookAlgorithmName="SHA-512",
        workbookHashValue=b64encode(digest).decode(),
        workbookSaltValue=b64encode(salt).decode(),
        workbookSpinCount=100,
    )
    workbook.save(sample)
    wrong = {"structure_protection": {"enabled": False, "password": "no"}}
    assert "right password" in await call_error(
        "set_workbook_settings", path="sales.xlsx", settings=wrong
    )
    await settings(call, structure_protection={"enabled": False, "password": "pw123"})
    assert not load_workbook(sample).security.lockStructure


async def test_settings_reject_invalid_input(call_error: ToolCall, sample: Path) -> None:
    for bad in (
        {"calculation": {"mode": "sometimes"}},
        {"calculation": {"max_iterations": 0}},
        {"calculation": {"max_change": 0}},
        {"doc_properties": {"colour": "red"}},
        {"structure_protection": {"password": "x"}},
    ):
        await call_error("set_workbook_settings", path="sales.xlsx", settings=bad)


async def test_settings_stay_inside_the_sandbox(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "set_workbook_settings", path="../sales.xlsx", settings={"doc_properties": {"title": "x"}}
    )
    assert "outside" in message


# Sheet order, view and cell protection


def layout(**values: object) -> dict[str, object]:
    return {"path": "sales.xlsx", "sheet": "Data", "layout": values}


async def test_move_sheet_keeps_the_active_sheet(call: ToolCall, sample: Path) -> None:
    await call("create_sheet", path="sales.xlsx", new_name="Third")
    await call(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Report",
        layout={"view": {"active": True}},
    )
    await call("set_sheet_layout", **layout(position=3))
    assert load_workbook(sample).sheetnames == ["Report", "Third", "Data"]
    info = await call("describe_workbook", path="sales.xlsx")
    assert [sheet["name"] for sheet in info["sheets"] if sheet.get("active")] == ["Report"]
    await call("set_sheet_layout", **layout(position=1))
    assert load_workbook(sample).sheetnames == ["Data", "Report", "Third"]


async def test_move_sheet_beyond_the_last_is_rejected(call_error: ToolCall, sample: Path) -> None:
    assert "1 to 2" in await call_error("set_sheet_layout", **layout(position=3))


async def test_view_options_and_active_sheet(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Report",
        layout={
            "view": {
                "zoom": 150,
                "gridlines": False,
                "headings": False,
                "show_formulas": True,
                "right_to_left": True,
                "active": True,
                "selected_cell": "c4",
            }
        },
    )
    workbook = load_workbook(sample)
    view = workbook["Report"].sheet_view
    assert (view.zoomScale, view.showGridLines, view.showRowColHeaders) == (150, False, False)
    assert (view.showFormulas, view.rightToLeft) == (True, True)
    assert view.selection[0].activeCell == "C4"
    assert workbook.active is workbook["Report"]
    assert [sheet.sheet_view.tabSelected for sheet in workbook.worksheets] == [False, True]
    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert details["view"] == {
        "zoom": 150,
        "gridlines": False,
        "headings": False,
        "show_formulas": True,
        "right_to_left": True,
        "active": True,
        "selected_cell": "C4",
    }
    assert "view" not in await call("describe_sheet", path="sales.xlsx", sheet="Data")


async def test_view_rejects_bad_values(call_error: ToolCall, sample: Path) -> None:
    assert "zoom" in await call_error("set_sheet_layout", **layout(view={"zoom": 5}))
    assert "explicit rows" in await call_error(
        "set_sheet_layout", **layout(view={"selected_cell": "x"})
    )
    await call_error("set_sheet_layout", **layout(view={"active": False}))


async def test_hidden_sheet_cannot_be_activated(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Report",
        layout={"visibility": "hidden", "view": {"active": True}},
    )
    assert "hidden" in message


async def test_cell_protection_flags(call: ToolCall, sample: Path) -> None:
    await call(
        "format_range",
        path="sales.xlsx",
        sheet="Data",
        range="C2:C3",
        style={"locked": False, "formula_hidden": True},
    )
    sheet = load_workbook(sample)["Data"]
    assert (sheet["C2"].protection.locked, sheet["C2"].protection.hidden) == (False, True)
    assert sheet["C4"].protection.locked is True
    await call(
        "format_range", path="sales.xlsx", sheet="Data", range="C2", style={"formula_hidden": False}
    )
    sheet = load_workbook(sample)["Data"]
    assert (sheet["C2"].protection.locked, sheet["C2"].protection.hidden) == (False, False)
