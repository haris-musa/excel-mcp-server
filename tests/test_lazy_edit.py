"""Edits of a large sheet leave the rows they do not touch alone, and give the same workbook as
loading and saving the whole sheet does.

The sheets here are small: the size from which a sheet is edited this way, and the number of
rows in a block, are lowered so that every edit takes the large-sheet path.
"""

import datetime as dt
import re
import shutil
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from excel_mcp import lazy_edit, lazy_workbook
from tests.conftest import ToolCall
from tests.package_support import (
    FIXTURES,
    assert_namespaces_declared,
    assert_package_is_consistent,
    read_parts,
    sheet_part,
)

pytestmark = pytest.mark.anyio

ROWS = 40


@contextmanager
def lazily(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """Every sheet is edited as a large one, in blocks of four rows."""
    with monkeypatch.context() as context:
        context.setattr(lazy_edit, "LAZY_SHEET_BYTES", 0)
        context.setattr(lazy_edit, "BLOCK_ROWS", 4)
        context.setattr(lazy_workbook, "BLOCK_ROWS", 4)
        yield


def make_book(path: Path) -> None:
    """A sheet with something of everything that belongs to cells and rows."""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Id", "Name", "Date", "Amount", "Double", "Note"])
    for row in range(2, ROWS + 1):
        sheet.append(
            [
                row - 1,
                f"item {row}",
                dt.date(2026, 1, 1) + dt.timedelta(days=row),
                row * 1.5,
                f"=A{row}*2",
                "text" if row % 3 else None,
            ]
        )
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEEFF")
    for row in range(2, ROWS + 1):
        sheet.cell(row, 4).number_format = "0.00"
    for row in range(9, ROWS, 9):
        sheet.row_dimensions[row].height = 24
    sheet.row_dimensions[17].hidden = True
    sheet.column_dimensions["B"].width = 22
    sheet.merge_cells("B20:C21")
    sheet["B30"].hyperlink = "https://example.com/a"
    sheet["C35"].comment = Comment("hello", "me")
    validation = DataValidation(type="whole", operator="between", formula1="0", formula2="100")
    sheet.add_data_validation(validation)
    validation.add("A2:A40")
    red = PatternFill("solid", bgColor="FFC7CE")
    sheet.conditional_formatting.add(
        "D2:D40", CellIsRule(operator="greaterThan", formula=["50"], fill=red)
    )
    sheet.freeze_panes = "A2"
    keys = [["Key", "Count", "Share"]] + [[f"k{n}", n, n / 10] for n in range(12)]
    for row, values in enumerate(keys, start=1):
        for column, value in enumerate(values, start=8):
            sheet.cell(row, column, value)
    sheet.add_table(Table(displayName="Keys", ref="H1:J13"))
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "=SUM(Data!A2:A40)"
    notes["A2"] = "see Data"
    workbook.save(path)


def style_of(cell: Any) -> tuple[Any, ...]:
    return (
        cell.number_format,
        cell.font.b,
        cell.font.i,
        cell.font.name,
        cell.font.sz,
        cell.fill.patternType,
        cell.fill.fgColor.rgb,
        cell.alignment.horizontal,
        cell.alignment.wrap_text,
        cell.border.left.style,
        cell.protection.locked,
    )


def snapshot(path: Path) -> dict[str, Any]:
    """Everything about a workbook's sheets that a person would see or use."""
    book = load_workbook(path)
    found: dict[str, Any] = {}
    for sheet in book.worksheets:
        cells: dict[tuple[int, int], Any] = {}
        for (row, column), cell in sorted(sheet._cells.items()):  # pyright: ignore[reportPrivateUsage]
            if isinstance(cell, MergedCell):
                cells[row, column] = "merged"
                continue
            link = cell.hyperlink
            if cell.value is None and not (cell.has_style or link or cell.comment):
                continue  # as good as no cell: whole sheets do not write it, the file may have it
            cells[row, column] = (
                ("array", cell.value.ref, cell.value.text)
                if isinstance(cell.value, ArrayFormula)
                else cell.value,
                cell.data_type,
                style_of(cell),
                (link.target, link.location, link.tooltip) if link else None,
                cell.comment.text if cell.comment else None,
            )
        found[sheet.title] = {
            "cells": cells,
            "rows": {
                row: (dim.height, dim.hidden, dim.outlineLevel, dim.customHeight)
                for row, dim in sorted(sheet.row_dimensions.items())
            },
            "columns": {
                letter: (dim.width, dim.hidden)
                for letter, dim in sorted(sheet.column_dimensions.items())
            },
            "merged": sorted(str(merged) for merged in sheet.merged_cells.ranges),
            "tables": dict(sheet.tables.items()),
            "validations": [
                (str(rule.sqref), rule.type, rule.formula1)
                for rule in sheet.data_validations.dataValidation
            ],
            "formats": [
                (str(rules.sqref), [(rule.type, rule.formula) for rule in rules.rules])
                for rules in sheet.conditional_formatting
            ],
            "freeze": sheet.freeze_panes,
            "charts": len(sheet._charts),  # pyright: ignore[reportPrivateUsage]
        }
    return found


STEPS: list[tuple[str, dict[str, Any]]] = [
    ("write_range", {"sheet": "Data", "at": "D38", "rows": [[999, 1000]]}),
    ("write_range", {"sheet": "Data", "at": "A50", "rows": [[1, "beyond", 3]]}),
    (
        "format_range",
        {"sheet": "Data", "range": "B36:C38", "style": {"bold": True, "fill_color": "FFFF00"}},
    ),
    (
        "set_sheet_layout",
        {
            "sheet": "Data",
            "layout": {
                "row_heights_pt": {"33": 40, "9": 30},
                "rows": [{"span": "28:29", "action": "hide"}],
            },
        },
    ),
    ("merge_cells", {"sheet": "Data", "range": "A37:A38"}),
    (
        "write_range",
        {
            "sheet": "Data",
            "at": "F39",
            "rows": [["link"]],
            "links": [{"cell": "F39", "target": "https://example.org"}],
        },
    ),
    ("set_note", {"sheet": "Data", "cell": "A25", "text": "a note"}),
    ("clear_range", {"sheet": "Data", "range": "A3:B4", "clear": "all"}),
    ("copy_range", {"sheet": "Data", "range": "A2:F3", "at": "A44"}),
    ("write_range", {"sheet": "Data", "at": "B20", "rows": [["top left of a merge"]]}),
    ("merge_cells", {"sheet": "Data", "range": "B20:C21", "action": "unmerge"}),
    ("set_table", {"sheet": "Data", "name": "Keys", "options": {"style": "TableStyleMedium2"}}),
    ("write_range", {"sheet": "Data", "at": "L1", "rows": [["a", "b"], [1, 2], [3, 4], [5, 6]]}),
    ("set_table", {"sheet": "Data", "name": "Extra", "range": "L1:M4"}),
    ("create_chart", {"sheet": "Data", "chart_type": "column", "source": "Data!B1:D12", "at": "O2"}),
    (
        "add_conditional_format",
        {
            "sheet": "Data",
            "range": "A30:A40",
            "rule": {
                "type": "cell_value",
                "operator": "greaterThan",
                "values": ["5"],
                "fill_color": "#FFC7CE",
            },
        },
    ),
    ("replace_cells", {"query": "item 3", "replacement": "thing 3", "sheet": "Data"}),
    (
        "sort_range",
        {"sheet": "Data", "range": "H1:J13", "sort_by": [{"column": "Count", "order": "descending"}]},
    ),
    ("insert_rows_or_columns", {"sheet": "Data", "axis": "rows", "start": 10, "count": 3}),
    ("delete_rows_or_columns", {"sheet": "Data", "axis": "rows", "start": 30, "count": 2}),
    ("insert_rows_or_columns", {"sheet": "Data", "axis": "columns", "start": 3, "count": 1}),
    (
        "transform_range",
        {"sheet": "Data", "range": "D2:D36", "transform": {"operation": "fill", "direction": "down"}},
    ),
    ("write_range", {"sheet": "Notes", "at": "B1", "rows": [["other sheet"]]}),
]


async def test_every_edit_gives_what_loading_the_whole_sheet_gives(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_book(files / "lazy.xlsx")
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")

    for tool, arguments in STEPS:
        with lazily(monkeypatch):
            lazy_result = await call(tool, path="lazy.xlsx", **arguments)
        full_result = await call(tool, path="full.xlsx", **arguments)
        assert str(lazy_result).replace("lazy", "full") == str(full_result), tool
        assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx"), tool
    assert_package_is_consistent(read_parts(files / "lazy.xlsx"))
    assert_namespaces_declared(read_parts(files / "lazy.xlsx"))


async def test_the_values_read_back_are_the_same(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_book(files / "lazy.xlsx")
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")
    for tool, arguments in STEPS[:6]:
        with lazily(monkeypatch):
            await call(tool, path="lazy.xlsx", **arguments)
        await call(tool, path="full.xlsx", **arguments)

    lazy = await call("read_range", path="lazy.xlsx", sheet="Data")
    full = await call("read_range", path="full.xlsx", sheet="Data")

    assert lazy["values"] == full["values"] and lazy["range"] == full["range"]


@pytest.mark.parametrize("fixture", sorted(path.name for path in FIXTURES.glob("*.xlsx")))
async def test_workbooks_written_by_excel(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch, fixture: str
) -> None:
    shutil.copy(FIXTURES / fixture, files / "lazy.xlsx")
    shutil.copy(FIXTURES / fixture, files / "full.xlsx")
    sheet = load_workbook(files / "lazy.xlsx").worksheets[0].title
    arguments = {"sheet": sheet, "at": "XA9", "rows": [["edited", 5]]}

    with lazily(monkeypatch):
        await call("write_range", path="lazy.xlsx", **arguments)
    await call("write_range", path="full.xlsx", **arguments)

    assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx")
    assert_package_is_consistent(read_parts(files / "lazy.xlsx"))
    assert_namespaces_declared(read_parts(files / "lazy.xlsx"))


def test_an_edit_parses_the_blocks_it_touches_and_no_others(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    for row in range(1, 61):
        sheet.append([row, f"text {row}", f"=A{row}+1"])
    workbook.save(files / "plain.xlsx")

    with lazily(monkeypatch):
        book, edit = lazy_edit.load(files / "plain.xlsx")
        assert edit is not None
        cells = book.worksheets[0]._cells  # pyright: ignore[reportPrivateUsage]
        cells[50, 1].value = "changed"
        cells.get((7, 3))
        assert cells.loaded == {50 // 4, 7 // 4}


def test_hyperlinks_comments_and_merges_are_parsed_when_the_workbook_is_opened(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_book(files / "book.xlsx")

    with lazily(monkeypatch):
        book, edit = lazy_edit.load(files / "book.xlsx")
        assert edit is not None
        cells = book["Data"]._cells  # pyright: ignore[reportPrivateUsage]
        assert cells.loaded == {20 // 4, 30 // 4, 35 // 4}
        assert cells[30, 2].hyperlink.target == "https://example.com/a"
        assert cells[30, 2].value == "item 30"
        assert cells[35, 3].comment.text == "hello"
        assert isinstance(cells[21, 3], MergedCell)


def test_row_dimensions_are_all_there_and_using_one_parses_its_rows(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_book(files / "book.xlsx")

    with lazily(monkeypatch):
        book, _ = lazy_edit.load(files / "book.xlsx")
        sheet = book["Data"]
        cells = sheet._cells  # pyright: ignore[reportPrivateUsage]
        before = set(cells.loaded)
        assert dict.keys(sheet.row_dimensions) >= {9, 17, 18, 27, 36}
        assert cells.loaded == before
        assert sheet.row_dimensions[27].height == 24
        assert 27 // 4 in cells.loaded


async def test_rows_that_are_not_parsed_are_kept_with_their_results_dropped(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_book(files / "book.xlsx")
    parts = read_parts(files / "book.xlsx")
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    stored = re.sub(r'(<c r="E(\d+)"[^>]*>)<f>(.*?)</f><v ?/?>(</v>)?', r"\1<f>\3</f><v>\2</v>", sheet)
    assert len(re.findall(r"</f><v>\d", stored)) == ROWS - 1
    parts["xl/worksheets/sheet1.xml"] = stored.encode()
    with ZipFile(files / "book.xlsx", "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    with lazily(monkeypatch):
        await call("write_range", path="book.xlsx", sheet="Data", at="A2", rows=[[500]])

    saved = read_parts(files / "book.xlsx")["xl/worksheets/sheet1.xml"].decode()
    assert not re.search(r"</f><v>[^<]", saved)
    assert saved.count("<f>") == ROWS - 1
    data = await call("read_range", path="book.xlsx", sheet="Data", range="E2:E4")
    assert data["values"] == [[1000], [4], [6]]


async def test_the_table_of_strings_is_kept_for_the_rows_that_refer_to_it(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    shutil.copy(FIXTURES / "excel_slicers.xlsx", files / "book.xlsx")
    table = read_parts(files / "book.xlsx")["xl/sharedStrings.xml"]

    with lazily(monkeypatch):
        await call("write_range", path="book.xlsx", sheet="Data", at="H1", rows=[["written"]])

    parts = read_parts(files / "book.xlsx")
    assert parts["xl/sharedStrings.xml"] == table
    assert 't="s"' in parts[sheet_part(parts, "Data")].decode()
    assert load_workbook(files / "book.xlsx")["Data"]["A1"].value == "Region"
    assert_package_is_consistent(parts)


async def test_prefixes_that_the_kept_rows_use_are_declared(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_book(files / "book.xlsx")
    parts = read_parts(files / "book.xlsx")
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace(
        "<worksheet ",
        '<worksheet xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" '
        'mc:Ignorable="x14ac" ',
        1,
    )
    sheet = re.sub(r'<row r="(\d+)"', r'<row r="\1" spans="1:10" x14ac:dyDescent="0.25"', sheet)
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    with ZipFile(files / "book.xlsx", "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    with lazily(monkeypatch):
        await call("write_range", path="book.xlsx", sheet="Data", at="D40", rows=[[1]])

    saved = read_parts(files / "book.xlsx")
    assert_namespaces_declared(saved)
    text = saved["xl/worksheets/sheet1.xml"].decode()
    assert 'x14ac:dyDescent="0.25"' in text
    assert 'mc:Ignorable="x14ac"' in text


async def test_the_dimension_covers_the_rows_written_and_the_rows_kept(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_book(files / "book.xlsx")

    with lazily(monkeypatch):
        await call("write_range", path="book.xlsx", sheet="Data", at="Z60", rows=[[1]])

    saved = read_parts(files / "book.xlsx")["xl/worksheets/sheet1.xml"].decode()
    assert '<dimension ref="A1:Z60"/>' in saved


@pytest.mark.parametrize("marker", ['t="array"', 't="shared"', 'cm="1"'])
def test_sheets_whose_cells_depend_on_each_other_are_loaded_whole(
    files: Path, monkeypatch: pytest.MonkeyPatch, marker: str
) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    for row in range(1, 11):
        sheet.append([row, f"=A{row}*2"])
    workbook.save(files / "book.xlsx")
    parts = read_parts(files / "book.xlsx")
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
        b'<c r="B2"', b'<c r="B2" ' + marker.encode(), 1
    )
    with ZipFile(files / "book.xlsx", "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    with lazily(monkeypatch):
        _, edit = lazy_edit.load(files / "book.xlsx")

    assert edit is None


def test_small_sheets_are_loaded_whole(files: Path) -> None:
    make_book(files / "book.xlsx")

    assert not lazy_edit.has_large_sheet(files / "book.xlsx")
