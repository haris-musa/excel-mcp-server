"""Formulas without stored results are calculated from the rows they use, not the whole sheet."""

import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from excel_mcp.lazy_workbook import LazyCells, load_lazy, split_sheet
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

ROWS = 400  # several blocks of rows


def make_running_totals(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in range(1, ROWS + 1):
        sheet.cell(row, 1, row)
        sheet.cell(row, 2, f"=A{row}*2")
        sheet.cell(row, 3, f"=C{row - 1}+B{row}" if row > 1 else "=B1")
    sheet["E1"] = "=SUM(A:A)"
    sheet["E2"] = "=SUMPRODUCT(A1:A400,B1:B400)"
    sheet["E3"] = "=SUBTOTAL(109,A1:A10)"
    sheet["E4"] = "=H1+I2"
    sheet.row_dimensions[5].hidden = True
    sheet["H1"] = 5
    sheet.merge_cells("H1:I2")
    workbook.save(path)


async def test_formulas_far_apart_are_calculated_from_the_rows_they_use(
    call: ToolCall, files: Path
) -> None:
    make_running_totals(files / "totals.xlsx")

    tail = await call("read_range", path="totals.xlsx", sheet="Data", range="A399:C400")
    top = await call("read_range", path="totals.xlsx", sheet="Data", range="E1:E4")

    assert tail["values"] == [[399, 798, 399 * 400], [400, 800, 400 * 401]]
    assert top["values"] == [
        [ROWS * (ROWS + 1) // 2],
        [2 * sum(r * r for r in range(1, ROWS + 1))],
        [50],
        [5],
    ]
    assert "uncalculated" not in tail and "uncalculated" not in top


def test_rows_are_parsed_only_when_asked_for(files: Path) -> None:
    make_running_totals(files / "totals.xlsx")

    book = load_lazy(files / "totals.xlsx", {}, data_only=False)
    cells = book["Data"]._cells  # pyright: ignore[reportPrivateUsage]

    assert isinstance(cells, LazyCells) and cells.loaded == set()
    assert cells.get((300, 2)).value == "=A300*2"  # pyright: ignore[reportOptionalMemberAccess]
    assert cells.loaded == {300 // 128}


async def test_shared_formulas_are_read_whole(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in range(1, 7):
        sheet.cell(row, 1, row)
        sheet.cell(row, 2, f"=A{row}*2")
    workbook.save(files / "shared.xlsx")
    with ZipFile(files / "shared.xlsx") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    xml = parts["xl/worksheets/sheet1.xml"].decode()
    xml = xml.replace("<f>A1*2</f>", '<f t="shared" ref="B1:B6" si="0">A1*2</f>')
    xml = re.sub(r"<f>A\d\*2</f>", '<f t="shared" si="0"/>', xml)
    parts["xl/worksheets/sheet1.xml"] = xml.encode()
    with ZipFile(files / "shared.xlsx", "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    data = await call("read_range", path="shared.xlsx", sheet="Data")

    assert data["values"][5] == [6, 12]
    assert split_sheet(parts["xl/worksheets/sheet1.xml"]) is None
