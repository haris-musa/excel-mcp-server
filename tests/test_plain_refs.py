"""Formulas of plain references are handled as text, with the results the tokenizer gives."""

import random

import pytest
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils import get_column_letter

from excel_mcp import plain_refs
from excel_mcp.errors import InvalidFormulaError, UnsafeFormulaError
from excel_mcp.formulas import check_formula, storable_formula
from excel_mcp.lazy_shared import SharedMaster, _Translator
from excel_mcp.operations.shifting import Shifter
from excel_mcp.package.lines import LineEdit
from excel_mcp.rewrite import ReferenceRewriter
from excel_mcp.spill import show_spills, store_spills
from excel_mcp.xlfn import add_prefixes, display_formula, remove_prefixes

COLUMNS = ["A", "B", "C", "AA", "XFD", "ZZZ", "a", "Bc"]
ROWS = [1, 2, 3, 4, 5, 6, 10, 100, 1048576, 1048577, 0]
FUNCTIONS = ["SUM", "IF", "MAX", "LOG10", "ATAN2", "INDEX", "ROUND"]
NAMES = ["Tax", "Q1_sales", "TRUE", "R1C1x", "Sheet1", "my.name"]


class Formulas:
    """Random formulas, many of them with something that is not a plain reference."""

    def __init__(self, seed: int) -> None:
        self.rng = random.Random(seed)

    def cell(self, row: int) -> str:
        rng = self.rng
        return f"{rng.choice(['', '$'])}{rng.choice(COLUMNS)}{rng.choice(['', '$'])}{row}"

    def reference(self) -> str:
        rng = self.rng
        cell = (
            f"{rng.choice(['', '$'])}{rng.choice(COLUMNS)}{rng.choice(['', '$'])}{rng.choice(ROWS)}"
        )
        kind = rng.random()
        if kind < 0.55:
            return cell
        if kind < 0.8:
            return f"{cell}:{self.cell(rng.choice([1, 3, 9, 100]))}"
        if kind < 0.85:
            return f"{rng.choice(COLUMNS)}:{rng.choice(COLUMNS)}"
        if kind < 0.9:
            return f"{rng.randint(1, 9)}:{rng.randint(1, 9)}"
        if kind < 0.93:
            return f"Sheet2!{cell}"
        if kind < 0.95:
            return f"'S'!{cell}"
        if kind < 0.97:
            return f"{cell}#"
        return f"Tbl[{rng.choice(['Col', '[Col]'])}]"

    def operand(self, depth: int = 0) -> str:
        rng = self.rng
        kind = rng.random()
        if kind < 0.5:
            return self.reference()
        if kind < 0.6:
            return str(rng.choice([1, 2.5, 100, "1E5", "0.5%"]))
        if kind < 0.7:
            return rng.choice(['"A1"', '"x ""B2"" y"', '"a!b"', '""'])
        if kind < 0.75:
            return rng.choice(NAMES)
        if depth < 3:
            arguments = ",".join(self.operand(depth + 1) for _ in range(rng.randint(1, 3)))
            return f"{rng.choice(FUNCTIONS)}({arguments})"
        return self.reference()

    def formula(self) -> str:
        parts = [self.operand()]
        for _ in range(self.rng.randint(0, 3)):
            parts += [
                self.rng.choice(["+", "-", "*", "/", "&", ">", "<=", " ", "^"]),
                self.operand(),
            ]
        return "=" + "".join(parts)


def test_moving_plain_references_as_text_gives_what_the_tokenizer_gives() -> None:
    formulas = Formulas(7)
    plain = 0
    for _ in range(4000):
        formula = formulas.formula()
        rng = formulas.rng
        edit = LineEdit(
            "S",
            rng.choice(["rows", "columns"]),
            rng.randint(1, 6),
            rng.randint(1, 3),
            rng.random() < 0.4,
        )
        shifter = Shifter(edit, {})
        host = rng.choice(["S", "S", "Other"])
        try:
            expected = ReferenceRewriter.formula(shifter, formula, host)
        except TokenizerError:
            continue
        plain += plain_refs.shape(formula) is not None
        assert shifter.formula(formula, host) == expected, (formula, edit, host)
    assert plain > 1000


def test_translating_plain_references_as_text_gives_what_openpyxl_gives() -> None:
    formulas = Formulas(11)
    for _ in range(4000):
        formula = formulas.formula()
        rng = formulas.rng
        row, column = rng.randint(1, 9), rng.randint(1, 5)
        origin = f"{get_column_letter(column)}{row}"
        dest_column = max(1, column + rng.randint(-4, 6))
        dest = f"{get_column_letter(dest_column)}{max(1, row + rng.randint(-8, 40))}"
        if dest == origin:
            continue
        try:
            expected = Translator(formula, origin).translate_formula(dest)
        except TranslatorError:
            expected = "error"
        except TokenizerError:
            continue
        master = SharedMaster("0", row, column, origin, formula[1:])
        try:
            got = _Translator(master).translate_formula(dest)
        except TranslatorError:
            got = "error"
        assert got == expected, (formula, origin, dest)


def test_showing_formulas_without_the_tokenizer_gives_what_it_gives() -> None:
    formulas = Formulas(5)
    for _ in range(3000):
        formula = formulas.formula()
        if formulas.rng.random() < 0.1:
            formula = formula.replace("SUM(", "_xlfn.SUM(").replace("A1", "_xlfn.ANCHORARRAY(A1)")
        if formulas.rng.random() < 0.05:
            formula = formula + formulas.rng.choice([" ", "\n", "  "])
        try:
            expected = show_spills(remove_prefixes(formula))
        except TokenizerError:
            expected = formula
        assert display_formula(formula) == expected, formula


@pytest.mark.parametrize(
    "formula",
    [
        "=SUM(A1:B5)+C3",
        '=IF(D2>50,"big","small")',
        "=XLOOKUP(A1,B:B,C:C)",
        "=_xlfn.XLOOKUP(A1,B1:B5,C1:C5)",
        '=HYPERLINK("https://example.com/"&A1,"x")',
        '=HYPERLINK("#Sheet1!A1","x")',
        '=WEBSERVICE("https://x.example/"&A1)',
        "=webservice(A1)",
        "=_xlfn.WEBSERVICE(A1)",
        "=CALL(A1)",
        "=FILES(A1)",
        "=SUM(A1,",
        '=IF(A1,"x)',
        "=A1+",
        "A1+B1",
        "=LET(x,A1,x+1)",
    ],
)
def test_the_check_of_a_formula_does_not_depend_on_which_cells_it_uses(formula: str) -> None:
    """Formulas of one shape are checked once; the answer is the one checking each gives."""
    names = ["Sheet1"]

    def checked(text: str) -> str:
        try:
            return storable_formula(text, names)
        except (InvalidFormulaError, UnsafeFormulaError) as error:
            return type(error).__name__

    for moved in (formula, formula.replace("A1", "Z99").replace("B5", "C17")):
        try:
            check_formula(moved, names)
            expected = add_prefixes(store_spills(moved))
        except (InvalidFormulaError, UnsafeFormulaError) as error:
            expected = type(error).__name__
        assert checked(moved) == expected
        assert checked(moved) == expected  # and again, from the cache
