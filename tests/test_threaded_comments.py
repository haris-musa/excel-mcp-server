"""Threaded comments written, read, resolved, deleted and kept through other edits."""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.conftest import ToolCall
from tests.package_support import (
    assert_package_is_consistent,
    read_parts,
    relationships,
    sheet_part,
    text,
)

pytestmark = pytest.mark.anyio

BOOK = {"path": "sales.xlsx"}
THREADS = "xl/threadedComments/threadedComment1.xml"
GUID = r"\{[0-9A-F]{8}(?:-[0-9A-F]{4}){3}-[0-9A-F]{12}\}"


async def _comment(call: ToolCall, cell: str, text: str, author: str = "Ann") -> str:
    return await call(
        "set_note", **BOOK, sheet="Data", cell=cell, text=text, author=author, threaded=True
    )


async def test_thread_with_replies_is_written_as_excel_does(call: ToolCall, sample: Path) -> None:
    assert "a thread" in await _comment(call, "B2", "Is this right?")
    assert "a reply" in await _comment(call, "B2", "Yes & checked", "Bob")
    await _comment(call, "C3", "Second")

    parts = read_parts(sample)
    threads = text(parts, THREADS)
    entries = re.findall(r"<threadedComment [^>]*>", threads)
    assert len(entries) == 3
    root_id = re.search(
        rf'<threadedComment ref="B2" dT="\d{{4}}-\d\d-\d\dT[\d:]+\.\d\d" '
        rf'personId="{GUID}" id="({GUID})" done="0">',
        threads,
    )
    assert root_id and f'parentId="{root_id[1]}" done="0">' in entries[1]
    assert "<text>Yes &amp; checked</text>" in threads
    persons = text(parts, "xl/persons/person.xml")
    assert (
        re.findall(r'displayName="(\w+)" id="(\{[^}]*\})" userId="\w+" providerId="None"', persons)[
            0
        ][0]
        == "Ann"
    )
    assert persons.count("<person ") == 2
    # the stand-in note for older versions of Excel carries the thread id
    note = load_workbook(sample)["Data"]["B2"].comment
    assert note is not None and note.author == f"tc={root_id[1]}"
    assert note.text.startswith("[Threaded comment]") and note.text.endswith(
        "Comment:\n    Is this right?\nReply:\n    Yes & checked"
    )
    kinds = {k for k, _ in relationships(parts, sheet_part(parts, "Data")).values()}
    assert "threadedComment" in kinds
    assert_package_is_consistent(parts)


async def test_describe_sheet_lists_threads_not_their_stand_in_notes(
    call: ToolCall, sample: Path
) -> None:
    await call("set_note", **BOOK, sheet="Data", cell="A1", text="Plain")
    await _comment(call, "B2", "Question", "Ann")
    await _comment(call, "B2", "Answer", "Bob")
    details = await call("describe_sheet", **BOOK, sheet="Data")
    assert [n["cell"] for n in details["notes"]] == ["A1"]
    [thread] = details["comments"]
    assert thread["cell"] == "B2" and not thread.get("resolved")
    assert [(c["author"], c["text"]) for c in thread["comments"]] == [
        ("Ann", "Question"),
        ("Bob", "Answer"),
    ]
    assert re.fullmatch(r"\d{4}-\d\d-\d\dT[\d:]+\.\d\d", thread["comments"][0]["date"])


async def test_resolve_and_reopen(call: ToolCall, sample: Path) -> None:
    await _comment(call, "B2", "Question")
    await call("resolve_comment", **BOOK, sheet="Data", cell="B2")
    assert 'done="1"' in text(read_parts(sample), THREADS)
    details = await call("describe_sheet", **BOOK, sheet="Data")
    assert details["comments"][0]["resolved"] is True
    await call("resolve_comment", **BOOK, sheet="Data", cell="B2", resolved=False)
    assert 'done="0"' in text(read_parts(sample), THREADS)


async def test_delete_reply_then_thread(call: ToolCall, sample: Path) -> None:
    for note in ("one", "two", "three"):
        await _comment(call, "B2", note)
    await call("delete_note", **BOOK, sheet="Data", cell="B2", reply=1)
    details = await call("describe_sheet", **BOOK, sheet="Data")
    assert [c["text"] for c in details["comments"][0]["comments"]] == ["one", "three"]
    note = load_workbook(sample)["Data"]["B2"].comment
    assert note is not None and note.text.endswith("Comment:\n    one\nReply:\n    three")

    await call("delete_note", **BOOK, sheet="Data", cell="B2")
    parts = read_parts(sample)
    assert not [n for n in parts if "threadedComments/" in n]
    assert load_workbook(sample)["Data"]["B2"].comment is None
    assert_package_is_consistent(parts)


async def test_threads_survive_edits_and_follow_their_cells(call: ToolCall, sample: Path) -> None:
    await _comment(call, "B2", "Stays with its cell")
    await call("insert_rows_or_columns", **BOOK, sheet="Data", axis="rows", at=1, count=2)
    details = await call("describe_sheet", **BOOK, sheet="Data")
    assert details["comments"][0]["cell"] == "B4"
    assert load_workbook(sample)["Data"]["B4"].comment is not None


async def test_copy_sheet_gives_the_copy_its_own_thread(call: ToolCall, sample: Path) -> None:
    await _comment(call, "B2", "Question")
    await _comment(call, "B2", "Answer", "Bob")
    await call("copy_sheet", **BOOK, sheet="Data", new_name="Copy")
    parts = read_parts(sample)
    ids = re.findall(
        rf'\bid="({GUID})"', "".join(text(parts, n) for n in parts if "threadedC" in n)
    )
    assert len(ids) == 4 and len(set(ids)) == 4
    details = await call("describe_sheet", **BOOK, sheet="Copy")
    assert [c["text"] for c in details["comments"][0]["comments"]] == ["Question", "Answer"]
    note = load_workbook(sample)["Copy"]["B2"].comment
    assert note is not None and note.author in {f"tc={i}" for i in ids}
    assert_package_is_consistent(parts)


async def test_notes_and_threads_do_not_mix(
    call_error: ToolCall, call: ToolCall, sample: Path
) -> None:
    await call("set_note", **BOOK, sheet="Data", cell="A1", text="Plain")
    assert "has a note" in await call_error(
        "set_note", **BOOK, sheet="Data", cell="A1", text="x", threaded=True
    )
    await _comment(call, "B2", "Thread")
    assert "has a threaded comment" in await call_error(
        "set_note", **BOOK, sheet="Data", cell="B2", text="x"
    )


async def test_invalid_input(call_error: ToolCall, call: ToolCall, sample: Path) -> None:
    assert "no threaded comment" in await call_error(
        "resolve_comment", **BOOK, sheet="Data", cell="B2"
    )
    assert "no threaded comment" in await call_error(
        "delete_note", **BOOK, sheet="Data", cell="B2", reply=1
    )
    await _comment(call, "B2", "Thread")
    assert "has 0 replies" in await call_error(
        "delete_note", **BOOK, sheet="Data", cell="B2", reply=1
    )
    assert "Sheet 'X' not found" in await call_error(
        "set_note", **BOOK, sheet="X", cell="B2", text="x", threaded=True
    )
    assert "past XFD" in await call_error(
        "set_note", **BOOK, sheet="Data", cell="XFE1", text="x", threaded=True
    )
