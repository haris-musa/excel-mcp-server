from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


async def test_format_range_changes_only_given_fields(call: ToolCall, sample: Path) -> None:
    await call(
        "format_range",
        path="sales.xlsx",
        sheet="Data",
        range="A1:D1",
        style={"font_size": 14, "font_color": "#1F4E78"},
    )
    await call(
        "format_range",
        path="sales.xlsx",
        sheet="Data",
        range="A1:D1",
        style={"bold": True, "fill_color": "FFF2CC", "border_style": "thin"},
    )
    cell = load_workbook(sample)["Data"]["B1"]
    assert cell.font.bold is True
    assert cell.font.size == 14
    assert cell.font.color.rgb == "FF1F4E78"
    assert cell.fill.start_color.rgb == "FFFFF2CC"
    assert cell.border.left.style == "thin"


async def test_invalid_color(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "format_range", path="sales.xlsx", sheet="Data", range="A1", style={"fill_color": "red"}
    )
    assert "Invalid color" in message


async def test_merge_and_unmerge(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    await call("merge_cells", path="sales.xlsx", sheet="Report", range="A1:C1")
    assert [str(r) for r in load_workbook(sample)["Report"].merged_cells.ranges] == ["A1:C1"]
    await call("merge_cells", path="sales.xlsx", sheet="Report", range="A1:C1", action="unmerge")
    assert not load_workbook(sample)["Report"].merged_cells.ranges
    assert "not a merged" in await call_error(
        "merge_cells", path="sales.xlsx", sheet="Report", range="A1:C1", action="unmerge"
    )


async def test_set_sheet_layout(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Data",
        layout={
            "column_widths_chars": {"a": 18},
            "row_heights_pt": {"1": 24},
            "autofit_columns": ["B"],
            "freeze_panes": "A2",
            "auto_filter": {"range": "A1:D5"},
            "tab_color": "#00B050",
        },
    )
    sheet = load_workbook(sample)["Data"]
    assert sheet.column_dimensions["A"].width == 18
    assert sheet.column_dimensions["B"].width == 9
    assert sheet.row_dimensions[1].height == 24
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:D5"
    tab_color = sheet.sheet_properties.tabColor
    assert tab_color is not None
    assert tab_color.rgb == "FF00B050"


async def test_row_height_outside_the_sheet_is_rejected(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Data",
        layout={"row_heights_pt": {"0": 24}},
    )
    assert "Row 0" in message


async def test_conditional_formats(call: ToolCall, sample: Path) -> None:
    for rule in [
        {"type": "color_scale", "colors": ["#F8696B", "#63BE7B"]},
        {"type": "data_bar", "colors": ["#638EC6"]},
        {"type": "cell_value", "operator": "greaterThan", "values": ["5"], "fill_color": "#FFC7CE"},
        {"type": "formula", "formula": "=$C2>5", "font_color": "#9C0006"},
    ]:
        await call(
            "add_conditional_format", path="sales.xlsx", sheet="Data", range="C2:C5", rule=rule
        )
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert [entry["type"] for entry in details["conditional_formats"]] == [
        "colorScale",
        "dataBar",
        "cellIs",
        "expression",
    ]


async def test_rule_formulas_are_checked(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "add_conditional_format",
        path="sales.xlsx",
        sheet="Data",
        range="C2:C5",
        rule={"type": "formula", "formula": '=WEBSERVICE("https://x")<>""'},
    )
    assert "not allowed" in message
    message = await call_error(
        "add_data_validation",
        path="sales.xlsx",
        sheet="Data",
        range="C2:C5",
        rule={"type": "whole", "operator": "greaterThan", "minimum": 'INDIRECT("A1")'},
    )
    assert "not allowed" in message


async def test_data_validation(call: ToolCall, sample: Path) -> None:
    await call(
        "add_data_validation",
        path="sales.xlsx",
        sheet="Data",
        range="A2:A5",
        rule={"type": "list", "options": ["North", "South"], "error_message": "Pick a region"},
    )
    await call(
        "add_data_validation",
        path="sales.xlsx",
        sheet="Data",
        range="C2:C5",
        rule={"type": "whole", "operator": "between", "minimum": "0", "maximum": "1000"},
    )
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert details["data_validations"][0]["formula1"] == '"North,South"'
    assert details["data_validations"][1]["operator"] == "between"


async def test_font_name_replaces_the_theme_font(call: ToolCall, sample: Path) -> None:
    await call(
        "format_range", path="sales.xlsx", sheet="Data", range="A1", style={"font_name": "Arial"}
    )
    await call("format_range", path="sales.xlsx", sheet="Data", range="B1", style={"bold": True})
    sheet = load_workbook(sample)["Data"]
    assert (sheet["A1"].font.name, sheet["A1"].font.scheme) == ("Arial", None)
    assert (sheet["B1"].font.name, sheet["B1"].font.scheme) == ("Calibri", "minor")
