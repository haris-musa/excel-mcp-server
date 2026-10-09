from pathlib import Path
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image

from excel_mcp.operations.images import list_images
from excel_mcp.operations.pivot_index import sheet_pivots
from tests.conftest import ToolCall
from tests.package_support import (
    assert_namespaces_declared,
    assert_package_is_consistent,
    copy_fixture,
    read_parts,
    relationships,
    sheet_part,
    text,
)

pytestmark = pytest.mark.anyio

BOOK = {"path": "sales.xlsx"}


@pytest.fixture
def picture(files: Path) -> Path:
    path = files / "logo.png"
    Image.new("RGB", (200, 100), "red").save(path)
    return path


def _sheet(path: Path, name: str) -> Worksheet:
    sheet = load_workbook(path)[name]
    assert isinstance(sheet, Worksheet)
    return sheet


async def _decorate_data_sheet(call: ToolCall) -> None:
    await call("write_range", **BOOK, sheet="Data", at="C6", rows=[["=SUM(Sales[Units])"]])
    await call("write_range", **BOOK, sheet="Data", at="D6", rows=[["=SUM(Data!D2:D5)"]])
    await call("set_table", **BOOK, sheet="Data", range="A1:D5", name="Sales")
    await call(
        "add_data_validation",
        **BOOK,
        sheet="Data",
        range="B2:B5",
        rule={"type": "list", "options": ["Apples", "Pears"]},
    )
    await call(
        "add_conditional_format",
        **BOOK,
        sheet="Data",
        range="C2:C5",
        rule={
            "type": "cell_value",
            "operator": "greaterThan",
            "values": ["6"],
            "fill_color": "#FFC7CE",
        },
    )
    await call("set_note", **BOOK, sheet="Data", cell="A2", text="check")
    await call("insert_image", **BOOK, sheet="Data", image_path="logo.png", at="G2")
    await call(
        "create_chart",
        **BOOK,
        sheet="Data",
        source="A1:C5",
        chart_type="column",
        at="G10",
    )
    await call(
        "create_chart",
        **BOOK,
        sheet="Data",
        source="Report!A1:B3",
        chart_type="line",
        at="G30",
    )
    await call("set_defined_name", **BOOK, name="Local", refers_to="Data!$A$1:$A$3", sheet="Data")
    await call("merge_cells", **BOOK, sheet="Data", range="A8:C8")
    await call(
        "set_sheet_layout",
        **BOOK,
        sheet="Data",
        layout={
            "column_widths_chars": {"A": 22},
            "freeze_panes": "A2",
            "rows": [{"span": "3", "action": "hide"}],
            "print_setup": {"orientation": "landscape", "print_area": "A1:D6", "title_rows": "1:1"},
            "protection": {"enabled": True},
        },
    )
    await call(
        "create_pivot_table",
        **BOOK,
        source="Data!A1:D5",
        row_fields=["Region"],
        value_fields=[{"field": "Units"}],
        sheet="Data",
        at="N2",
    )


async def test_copy_includes_everything_the_server_supports(
    call: ToolCall, sample: Path, picture: Path
) -> None:
    await call("write_range", **BOOK, sheet="Report", at="A1", rows=[["k", "v"], ["a", 1]])
    await _decorate_data_sheet(call)
    await call("copy_sheet", **BOOK, sheet="Data", new_name="Copy")

    assert load_workbook(sample).sheetnames == ["Data", "Report", "Copy"]
    original, copied = _sheet(sample, "Data"), _sheet(sample, "Copy")
    assert copied["A2"].value == "North"
    assert copied["A2"].comment is not None
    assert copied["C6"].value == "=SUM(Sales2[Units])"
    assert copied["D6"].value == "=SUM('Copy'!D2:D5)"
    assert original["C6"].value == "=SUM(Sales[Units])"
    assert list(copied.tables) == ["Sales2"]
    assert list(original.tables) == ["Sales"]
    assert [str(v.sqref) for v in copied.data_validations.dataValidation] == ["B2:B5"]
    assert [str(entry.sqref) for entry in copied.conditional_formatting] == ["C2:C5"]
    assert len(list_images(copied)) == 1
    assert [str(r) for r in copied.merged_cells.ranges] == ["A8:C8"]
    assert copied.column_dimensions["A"].width == 22
    assert copied.row_dimensions[3].hidden
    assert copied.freeze_panes == "A2"
    assert copied.protection.sheet
    assert copied.page_setup.orientation == "landscape"
    assert copied.print_area == "'Copy'!$A$1:$D$6"
    assert copied.print_title_rows == "$1:$1"
    assert copied.defined_names["Local"].attr_text == "'Copy'!$A$1:$A$3"
    assert not any(view.tabSelected for view in copied.views.sheetView)
    assert [pivot.name for pivot in sheet_pivots(copied)] == [
        pivot.name for pivot in sheet_pivots(original)
    ]


async def test_copied_charts_point_at_the_copy_only_when_their_data_is_on_the_source(
    call: ToolCall, sample: Path, picture: Path
) -> None:
    await call("write_range", **BOOK, sheet="Report", at="A1", rows=[["k", "v"], ["a", 1]])
    await _decorate_data_sheet(call)
    await call("copy_sheet", **BOOK, sheet="Data", new_name="Copy")

    own, other = _sheet(sample, "Copy")._charts  # pyright: ignore[reportAttributeAccessIssue]
    own_refs = {s.val.numRef.f for s in own.series} | {s.cat.numRef.f for s in own.series}
    assert own_refs and all(ref.startswith("'Copy'!") for ref in own_refs)
    assert {s.val.numRef.f for s in other.series} == {"'Report'!$B$2:$B$3"}
    original = _sheet(sample, "Data")._charts[0]  # pyright: ignore[reportAttributeAccessIssue]
    assert all(s.val.numRef.f.startswith("'Data'!") for s in original.series)


async def test_copy_keeps_the_source_untouched_and_names_unique(
    call: ToolCall, sample: Path, picture: Path
) -> None:
    await _decorate_data_sheet(call)
    await call("copy_sheet", **BOOK, sheet="Data", new_name="Copy")
    await call("copy_sheet", **BOOK, sheet="Data", new_name="Copy 2")
    assert list(_sheet(sample, "Copy").tables) == ["Sales2"]
    assert list(_sheet(sample, "Copy 2").tables) == ["Sales3"]
    assert _sheet(sample, "Copy 2")["C6"].value == "=SUM(Sales3[Units])"


async def test_copy_refuses_unsafe_formulas(call_error: ToolCall, files: Path) -> None:
    workbook = Workbook()
    workbook.worksheets[0]["A1"] = '=INDIRECT("B1")'
    workbook.save(files / "unsafe.xlsx")
    message = await call_error("copy_sheet", path="unsafe.xlsx", sheet="Sheet", new_name="Copy")
    assert "INDIRECT" in message


@pytest.mark.parametrize(
    "fixture",
    [
        "excel_slicers.xlsx",
        "excel_chartex.xlsx",
        "excel_controls.xlsx",
        "excel_sparklines.xlsx",
        "excel_pivot_ext.xlsx",
        "excel_dynamic_arrays.xlsx",
    ],
)
async def test_copies_of_excel_authored_sheets_declare_every_namespace_they_use(
    call: ToolCall, files: Path, fixture: str
) -> None:
    path = copy_fixture(files, fixture)
    for sheet in load_workbook(path).sheetnames:
        await call("copy_sheet", path="book.xlsx", sheet=sheet, new_name=f"{sheet} copy")
    parts = read_parts(path)
    assert_namespaces_declared(parts)
    assert_package_is_consistent(parts)


async def test_copied_slicers_are_numbered_from_the_name_without_its_number(
    call: ToolCall, files: Path
) -> None:
    copy_fixture(files, "excel_slicers.xlsx")
    await call("copy_sheet", path="book.xlsx", sheet="Data", new_name="Data copy")
    for sheet, expected in (
        ("Data", ["Region", "Region 1"]),
        ("Data copy", ["Region 2", "Region 3"]),
    ):
        details = await call("describe_sheet", path="book.xlsx", sheet=sheet)
        assert sorted(s["name"] for s in details["slicers"]) == expected


XDR = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"


def _drawing_of(parts: dict[str, bytes], sheet: str) -> tuple[str, ElementTree.Element]:
    related = relationships(parts, sheet_part(parts, sheet))
    name = next(target for kind, target in related.values() if kind == "drawing")
    return name, ElementTree.fromstring(parts[name])


def _shape_ids(drawing: ElementTree.Element) -> dict[str, str]:
    """Shape ids by name, for every shape in the drawing, groups included."""
    return {e.get("name", ""): e.get("id", "") for e in drawing.iter(f"{XDR}cNvPr")}


async def test_copy_keeps_shapes_text_boxes_connectors_and_groups(
    call: ToolCall, files: Path, picture: Path
) -> None:
    copy_fixture(files, "excel_shapes.xlsx")
    for cell in ("K2", "K12"):  # their ids are the ones the shapes have
        await call("insert_image", path="book.xlsx", sheet="Data", image_path="logo.png", at=cell)
    await call("copy_sheet", path="book.xlsx", sheet="Data", new_name="Copy")
    parts = read_parts(files / "book.xlsx")

    (_, source), (name, drawing) = (_drawing_of(parts, sheet) for sheet in ("Data", "Copy"))
    ids = _shape_ids(drawing)
    assert set(ids) == set(_shape_ids(source))
    assert set(ids) >= {"Box", "Oval", "Note", "Link", "Group", "G1", "G2", "Line"}
    assert len(set(ids.values())) == len(ids)
    connected = {
        e.tag.rpartition("}")[2]: e.get("id") for e in drawing.iter() if e.tag.endswith("Cxn")
    }
    assert connected == {"stCxn": ids["Box"], "endCxn": ids["Oval"]}
    assert drawing.find(f".//{XDR}grpSp") is not None
    assert [k for k, _ in relationships(parts, name).values()] == ["image"] * 3
    assert_namespaces_declared(parts)
    assert_package_is_consistent(parts)


async def test_copy_keeps_form_controls_with_their_own_properties(
    call: ToolCall, files: Path
) -> None:
    copy_fixture(files, "excel_controls.xlsx")
    await call("copy_sheet", path="book.xlsx", sheet="Data", new_name="Copy")
    parts = read_parts(files / "book.xlsx")

    related = relationships(parts, sheet_part(parts, "Copy"))
    properties = sorted(target for kind, target in related.values() if kind == "ctrlProp")
    original = sorted(
        target
        for kind, target in relationships(parts, sheet_part(parts, "Data")).values()
        if kind == "ctrlProp"
    )
    assert len(properties) == len(original) == 2
    assert not set(properties) & set(original)
    assert {parts[p] for p in properties} == {parts[p] for p in original}
    _, drawing = _drawing_of(parts, "Copy")
    assert sorted(_shape_ids(drawing)) == ["Check Box 1", "Drop Down 1"]
    vml = next(target for kind, target in related.values() if kind == "vmlDrawing")
    assert parts[vml].count(b"<x:ClientData ObjectType=") >= 2
    assert_namespaces_declared(parts)
    assert_package_is_consistent(parts)


async def test_copy_keeps_protected_ranges_and_embedded_objects_as_excel_does(
    call: ToolCall, files: Path
) -> None:
    copy_fixture(files, "excel_objects.xlsx")
    result = await call("copy_sheet", path="book.xlsx", sheet="Data", new_name="Copy")
    assert "note" not in result
    parts = read_parts(files / "book.xlsx")

    original, copied = (text(parts, sheet_part(parts, s)) for s in ("Data", "Copy"))
    assert '<protectedRange sqref="B1:B3" name="Editable"/>' in copied
    assert "ignoredErrors" in original and "ignoredErrors" not in copied
    assert 'shapeId="2049"' in copied and 'shapeId="1025"' in original
    embedded = [
        target
        for sheet in ("Data", "Copy")
        for kind, target in relationships(parts, sheet_part(parts, sheet)).values()
        if kind == "oleObject"
    ]
    assert len(set(embedded)) == 2
    assert_namespaces_declared(parts)
    assert_package_is_consistent(parts)
