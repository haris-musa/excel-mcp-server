"""Every place a formula is stored gets the storage prefixes of newer functions."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


@pytest.fixture
def book(files: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row, value in enumerate([3, 1, 2], start=1):
        sheet.cell(row, 1, value)
        sheet.cell(row, 2, f"=A{row}*2")
    # A formula as an older version of this server stored it, without the prefix.
    sheet["C1"] = "=IFS(A1>2,1,TRUE,0)"
    path = files / "book.xlsx"
    workbook.save(path)
    return path


async def test_conditional_format_formulas(call: ToolCall, book: Path) -> None:
    await call(
        "add_conditional_format",
        path="book.xlsx",
        sheet="Data",
        range="A1:A3",
        rule={"type": "formula", "formula": "=IFS($A1>2,TRUE,TRUE,FALSE)", "fill_color": "#FFC7CE"},
    )
    await call(
        "add_conditional_format",
        path="book.xlsx",
        sheet="Data",
        range="A1:A3",
        rule={
            "type": "cell_value",
            "operator": "greaterThan",
            "values": ['MAXIFS(A1:A3,A1:A3,"<3")'],
            "fill_color": "#FFC7CE",
        },
    )

    rules = [
        formula
        for ranges in load_workbook(book)["Data"].conditional_formatting
        for rule in ranges.rules
        for formula in rule.formula
    ]
    assert rules == ["_xlfn.IFS($A1>2,TRUE,TRUE,FALSE)", '_xlfn.MAXIFS(A1:A3,A1:A3,"<3")']


async def test_data_validation_formulas(call: ToolCall, book: Path) -> None:
    await call(
        "add_data_validation",
        path="book.xlsx",
        sheet="Data",
        range="A1:A3",
        rule={"type": "custom", "formula": "=IFS(A1>0,TRUE,TRUE,FALSE)"},
    )
    await call(
        "add_data_validation",
        path="book.xlsx",
        sheet="Data",
        range="B1:B3",
        rule={
            "type": "decimal",
            "operator": "between",
            "minimum": "0",
            "maximum": 'MAXIFS(A1:A3,A1:A3,"<3")',
        },
    )

    validations = load_workbook(book)["Data"].data_validations.dataValidation
    assert validations[0].formula1 == "_xlfn.IFS(A1>0,TRUE,TRUE,FALSE)"
    assert validations[1].formula2 == '_xlfn.MAXIFS(A1:A3,A1:A3,"<3")'


async def test_copy_range_prefixes_formulas(call: ToolCall, book: Path) -> None:
    await call("copy_range", path="book.xlsx", sheet="Data", range="C1", target_cell="D1")

    assert load_workbook(book)["Data"]["D1"].value == "=_xlfn.IFS(B1>2,1,TRUE,0)"


async def test_sort_range_prefixes_formulas(call: ToolCall, book: Path) -> None:
    await call(
        "sort_range",
        path="book.xlsx",
        sheet="Data",
        range="A1:C3",
        sort_by=[{"column": "A"}],
        has_header=False,
    )

    sheet = load_workbook(book)["Data"]
    assert [sheet[f"A{row}"].value for row in (1, 2, 3)] == [1, 2, 3]
    assert sheet["C3"].value == "=_xlfn.IFS(A3>2,1,TRUE,0)"


async def test_defined_names_are_prefixed(call: ToolCall, book: Path) -> None:
    await call(
        "set_defined_name",
        name="Peak",
        refers_to='=MAXIFS(Data!$A$1:$A$3,Data!$A$1:$A$3,"<3")',
        path="book.xlsx",
    )

    names = load_workbook(book).defined_names
    assert names["Peak"].attr_text == '_xlfn.MAXIFS(Data!$A$1:$A$3,Data!$A$1:$A$3,"<3")'
