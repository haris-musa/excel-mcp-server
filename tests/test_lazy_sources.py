# pyright: reportAttributeAccessIssue=false, reportOptionalMemberAccess=false
# The tests look at openpyxl internals and at what the lazy cells hold.
"""A large sheet is copied and used as a PivotTable source without parsing its rows, with the
results that loading it whole gives."""

import datetime as dt
import re
import shutil
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from excel_mcp import lazy_edit
from excel_mcp.cell_budget import CellBudget
from excel_mcp.operations import sheet_copy
from tests.conftest import ToolCall, lazily
from tests.package_support import read_parts
from tests.test_lazy_edit import snapshot
from tests.test_lazy_lines import make_linked_book

pytestmark = pytest.mark.anyio

RECORDS = 60


def make_sales(path: Path, *, stored: bool) -> None:
    """Sales with a formula column; ``stored`` gives the formulas the results Excel stores."""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Region", "Product", "Date", "Units", "Price", "Amount", "Note"])
    for row in range(2, RECORDS + 2):
        sheet.append(
            [
                ["North", "South", "East"][row % 3],
                ["Apples", "Pears"][row % 2],
                dt.date(2026, 1, 1) + dt.timedelta(days=row),
                row % 11,
                row * 0.5,
                f"=D{row}*E{row}",
                "late" if row % 5 == 0 else None,
            ]
        )
    sheet.append(["West", "Plums", dt.date(2026, 3, 1), 4, 2.5, "=D62*E62", "=1/0"])
    for row in range(2, RECORDS + 2):
        sheet.cell(row, 4).number_format = "0"
    workbook.create_sheet("Report")
    workbook.save(path)
    if stored:
        with ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        xml = parts["xl/worksheets/sheet1.xml"].decode()
        xml = re.sub(
            r'(<c r="F(\d+)"[^>]*>)<f>D\d+\*E\d+</f><v ?/>',
            lambda m: f"{m[1]}<f>D{m[2]}*E{m[2]}</f><v>{int(m[2]) % 11 * int(m[2]) * 0.5}</v>",
            xml,
        )
        xml = re.sub(r'(<c r="F62"[^>]*>)<f>D62\*E62</f><v ?/>', r"\1<f>D62*E62</f><v>10</v>", xml)
        xml = re.sub(
            r'<c r="G62"([^>]*)>(<f>1/0</f>)<v ?/>', r'<c r="G62"\1 t="e">\2<v>#DIV/0!</v>', xml
        )
        parts["xl/worksheets/sheet1.xml"] = xml.encode()
        with ZipFile(path, "w") as archive:
            for name, content in parts.items():
                archive.writestr(name, content)


PIVOT = {
    "sheet": "Report",
    "at": "A1",
    "source": "Data!A1:G62",
    "row_fields": ["Region", "Product"],
    "value_fields": [{"field": "Units"}, {"field": "Amount", "function": "average"}],
}


@pytest.mark.parametrize("stored", [True, False], ids=["stored results", "no results"])
async def test_a_pivot_table_of_a_large_sheet_is_what_one_of_the_whole_sheet_is(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch, stored: bool
) -> None:
    make_sales(files / "lazy.xlsx", stored=stored)
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")

    with lazily(monkeypatch):
        await call("create_pivot_table", path="lazy.xlsx", **PIVOT)
    await call("create_pivot_table", path="full.xlsx", **PIVOT)

    assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx")
    lazy, full = read_parts(files / "lazy.xlsx"), read_parts(files / "full.xlsx")
    for part in ("xl/pivotCache/pivotCacheRecords1.xml", "xl/pivotTables/pivotTable1.xml"):
        assert lazy[part] == full[part], part
    for definition in (lazy, full):
        definition["xl/pivotCache/pivotCacheDefinition1.xml"] = re.sub(
            rb'refreshedDate="[^"]*"', b"", definition["xl/pivotCache/pivotCacheDefinition1.xml"]
        )
    assert (
        lazy["xl/pivotCache/pivotCacheDefinition1.xml"]
        == full["xl/pivotCache/pivotCacheDefinition1.xml"]
    )


def test_the_source_is_read_without_parsing_it_when_every_formula_has_a_result(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_sales(files / "stored.xlsx", stored=True)
    make_sales(files / "calculated.xlsx", stored=False)

    with lazily(monkeypatch):
        book, _ = lazy_edit.load(files / "stored.xlsx", CellBudget(3_000_000))
        assert book["Data"]._cells.stored_rows(1, RECORDS + 2) is not None
        assert book["Data"]._cells.loaded == set()
        book, _ = lazy_edit.load(files / "calculated.xlsx", CellBudget(3_000_000))
        assert book["Data"]._cells.stored_rows(1, RECORDS + 2) is None


async def test_a_copy_of_a_large_sheet_is_the_copy_of_the_whole_sheet(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_linked_book(files / "lazy.xlsx")
    await call("write_range", path="lazy.xlsx", sheet="Data", at="H1", rows=[["=SUM(Data!A2:A9)"]])
    await call("write_range", path="lazy.xlsx", sheet="Data", at="H2", rows=[["=Data!B3*2+A2"]])
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")

    with lazily(monkeypatch):
        await call("copy_sheet", path="lazy.xlsx", sheet="Data", new_name="Copy")
    await call("copy_sheet", path="full.xlsx", sheet="Data", new_name="Copy")

    assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx")


def test_a_copy_shares_the_rows_of_the_original_that_were_not_parsed(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_linked_book(files / "book.xlsx")

    with lazily(monkeypatch):
        book, _ = lazy_edit.load(files / "book.xlsx", CellBudget(3_000_000))
        original = book["Data"]._cells
        before = set(original.loaded)
        sheet_copy.copy_sheet(book, "Data", "Copy", 1_000_000)
        copied = book["Copy"]._cells

        assert original.loaded == before
        assert copied.loaded == before
        assert copied.source.data != original.source.data  # its formulas are its own
        assert copied.count() == original.count()


async def test_a_pivot_table_beside_its_source_checks_only_the_rows_it_covers(
    call: ToolCall, call_error: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_sales(files / "lazy.xlsx", stored=True)
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")
    beside = {**PIVOT, "sheet": "Data", "at": "J1"}

    with lazily(monkeypatch):
        await call("create_pivot_table", path="lazy.xlsx", **beside)
    await call("create_pivot_table", path="full.xlsx", **beside)
    assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx")

    for name in ("lazy", "full"):
        await call("write_range", path=f"{name}.xlsx", sheet="Data", at="N40", rows=[["x"]])
    taken = {**PIVOT, "sheet": "Data", "at": "N38"}
    with lazily(monkeypatch):
        refused = await call_error("create_pivot_table", path="lazy.xlsx", **taken)
    assert "N40 already holds a value" in refused
    assert refused == await call_error("create_pivot_table", path="full.xlsx", **taken)


async def test_renaming_a_sheet_rewrites_the_references_of_a_large_sheet_to_it(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_linked_book(files / "lazy.xlsx")
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")

    with lazily(monkeypatch):
        await call("rename_sheet", path="lazy.xlsx", sheet="Data", new_name="Facts and Figures")
    await call("rename_sheet", path="full.xlsx", sheet="Data", new_name="Facts and Figures")

    assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx")
