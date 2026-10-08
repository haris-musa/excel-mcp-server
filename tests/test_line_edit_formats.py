from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.table import TableFormula

from excel_mcp.rewrite import chart_references
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


async def insert(call: ToolCall, sheet: str, axis: str, at: int, count: int = 1) -> None:
    await call(
        "insert_rows_or_columns", path="sales.xlsx", sheet=sheet, axis=axis, at=at, count=count
    )


async def test_inserted_cells_are_formatted_like_their_neighbours(
    call: ToolCall, sample: Path
) -> None:
    await call(
        "format_range",
        path="sales.xlsx",
        sheet="Data",
        range="A2:D2",
        style={"bold": True, "number_format": "0.00", "fill_color": "#FFFF00"},
    )
    await insert(call, "Data", "rows", 3, 2)
    await insert(call, "Data", "columns", 5)
    data = load_workbook(sample)["Data"]
    for ref in ("A3", "B4", "D3", "E2"):
        assert data[ref].font.b and data[ref].number_format == "0.00", ref
        assert data[ref].fill.fgColor.rgb.endswith("FFFF00"), ref
    assert data["B4"].value is None
    assert not data["A5"].font.b


async def test_nothing_is_copied_into_the_first_line(call: ToolCall, sample: Path) -> None:
    await call("format_range", path="sales.xlsx", sheet="Data", range="A1:D1", style={"bold": True})
    await insert(call, "Data", "rows", 1)
    await insert(call, "Data", "columns", 1)
    data = load_workbook(sample)["Data"]
    assert not data["A1"].font.b and not data["B1"].font.b and not data["A2"].font.b
    assert data["B2"].font.b


async def test_inserted_cells_keep_only_borders_shared_with_the_far_side(
    call: ToolCall, sample: Path
) -> None:
    style = {"border_style": "thin"}
    await call("format_range", path="sales.xlsx", sheet="Data", range="A1:A3", style=style)
    await insert(call, "Data", "rows", 2)
    await insert(call, "Data", "rows", 5)
    data = load_workbook(sample)["Data"]
    inside, below = data["A2"].border, data["A5"].border
    assert (inside.left.style, inside.top.style, inside.bottom.style) == ("thin",) * 3
    assert [side is None or side.style is None for side in (below.left, below.top)] == [True] * 2


async def test_rows_inserted_into_a_table_get_calculated_columns(
    call: ToolCall, sample: Path
) -> None:
    await call("create_table", path="sales.xlsx", sheet="Data", range="A1:D5", name="Sales")
    book = load_workbook(sample)
    book["Data"].tables["Sales"].tableColumns[3].calculatedColumnFormula = TableFormula(
        attr_text="Sales[[#This Row],[Units]]*2"
    )
    book.save(sample)
    await insert(call, "Data", "rows", 3, 2)
    data = load_workbook(sample)["Data"]
    assert data.tables["Sales"].ref == "A1:D7"
    assert [data[f"D{row}"].value for row in (3, 4)] == ["=Sales[[#This Row],[Units]]*2"] * 2
    assert data["A3"].value is None
    await insert(call, "Data", "rows", 8)
    assert load_workbook(sample)["Data"]["D8"].value is None


async def test_two_tables_cannot_be_cut_at_once(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await call("create_table", path="sales.xlsx", sheet="Data", range="A1:B5", name="Left")
    await call("create_table", path="sales.xlsx", sheet="Data", range="D1:D3", name="Right")
    message = await call_error(
        "insert_rows_or_columns", path="sales.xlsx", sheet="Data", axis="rows", at=2
    )
    assert "Left, Right" in message
    await insert(call, "Data", "rows", 4)


async def test_rename_sheet_rewrites_references(call: ToolCall, sample: Path) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="A1",
        rows=[["=Data!C2+Data!C3"], ["=SUM(Data!C2:C5)"], ["='Data'!$C$2"], ['="Data!C2"']],
    )
    await call("set_defined_name", path="sales.xlsx", name="Units", refers_to="Data!$C$2:$C$5")
    await call(
        "add_conditional_format",
        path="sales.xlsx",
        sheet="Report",
        range="A1:A4",
        rule={"type": "formula", "formula": "=A1>Data!$C$2", "fill_color": "FF0000"},
    )
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        data_range="Data!B1:C5",
        chart_type="column",
        anchor_cell="D2",
    )
    await call("rename_sheet", path="sales.xlsx", sheet="Data", new_name="Q1 Sales")
    book = load_workbook(sample)
    report = book["Report"]
    assert [report[f"A{row}"].value for row in (1, 2, 3, 4)] == [
        "='Q1 Sales'!C2+'Q1 Sales'!C3",
        "=SUM('Q1 Sales'!C2:C5)",
        "='Q1 Sales'!$C$2",
        '="Data!C2"',
    ]
    assert book.defined_names["Units"].attr_text == "'Q1 Sales'!$C$2:$C$5"
    rule = next(iter(report.conditional_formatting)).rules[0]
    assert rule.formula == ["A1>'Q1 Sales'!$C$2"]
    chart = report._charts[0]  # pyright: ignore[reportAttributeAccessIssue]
    assert chart.series[0].val.numRef.f == "'Q1 Sales'!$C$2:$C$5"
    await call("rename_sheet", path="sales.xlsx", sheet="Q1 Sales", new_name="Sales")
    assert load_workbook(sample)["Report"]["A1"].value == "=Sales!C2+Sales!C3"


async def test_deleting_the_totals_row_removes_it_from_the_table(
    call: ToolCall, sample: Path
) -> None:
    await call("create_table", path="sales.xlsx", sheet="Data", range="A1:D5", name="Sales")
    book = load_workbook(sample)
    table = book["Data"].tables["Sales"]
    table.ref, table.totalsRowCount = "A1:D5", 1
    table.tableColumns[2].totalsRowFunction = "sum"
    book.save(sample)
    await insert(call, "Data", "rows", 5)
    table = load_workbook(sample)["Data"].tables["Sales"]
    assert (table.ref, table.totalsRowCount) == ("A1:D6", 1)
    await call(
        "delete_rows_or_columns", path="sales.xlsx", sheet="Data", axis="rows", at=5, count=2
    )
    table = load_workbook(sample)["Data"].tables["Sales"]
    assert (table.ref, table.totalsRowCount) == ("A1:D4", None)
    assert table.tableColumns[2].totalsRowFunction is None


async def test_chart_series_and_chart_sheets_follow_edits_and_renames(
    call: ToolCall, sample: Path
) -> None:
    series = [{"values": "Data!C2:C5", "name": "Data!C1"}]
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Sheet Chart",
        chart_type="line",
        categories="Data!B2:B5",
        series=series,
    )
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        chart_type="scatter",
        anchor_cell="E2",
        categories="Data!C2:C5",
        series=[{"values": "Data!D2:D5"}],
    )
    await insert(call, "Data", "rows", 3)
    await call("rename_sheet", path="sales.xlsx", sheet="Data", new_name="Q1")
    book = load_workbook(sample)
    refs = [
        str(reference.f)
        for sheet in (book["Sheet Chart"], book["Report"])
        for chart in sheet._charts  # pyright: ignore[reportAttributeAccessIssue]
        for plot in chart._charts
        for reference in chart_references(plot)
    ]
    assert "'Q1'!$C$2:$C$6" in refs and "'Q1'!$B$2:$B$6" in refs and "'Q1'!$D$2:$D$6" in refs
    assert not any("Data" in ref for ref in refs)


async def test_filter_columns_and_sort_follow_column_edits(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Data",
        layout={
            "auto_filter": {
                "range": "A1:D5",
                "filters": [{"column": "Units", "type": "values", "values": ["10", "5"]}],
            }
        },
    )
    await insert(call, "Data", "columns", 2)
    saved = load_workbook(sample)["Data"].auto_filter
    assert (saved.ref, saved.filterColumn[0].colId) == ("A1:E5", 3)
    await call("delete_rows_or_columns", path="sales.xlsx", sheet="Data", axis="columns", at=1)
    saved = load_workbook(sample)["Data"].auto_filter
    assert (saved.ref, saved.filterColumn[0].colId) == ("A1:D5", 2)
    await call("delete_rows_or_columns", path="sales.xlsx", sheet="Data", axis="columns", at=3)
    saved = load_workbook(sample)["Data"]
    assert saved.auto_filter.filterColumn == []
    assert not any(dimension.hidden for dimension in saved.row_dimensions.values())


async def test_deleting_one_of_several_filtered_columns_is_refused(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    filters = [
        {"column": "Units", "type": "values", "values": ["10"]},
        {"column": "Region", "type": "values", "values": ["North"]},
    ]
    await call(
        "set_sheet_layout",
        path="sales.xlsx",
        sheet="Data",
        layout={"auto_filter": {"range": "A1:D5", "filters": filters}},
    )
    message = await call_error(
        "delete_rows_or_columns", path="sales.xlsx", sheet="Data", axis="columns", at=3
    )
    assert "remove the filter first" in message


async def test_an_intersection_with_a_deleted_reference_is_kept_like_excel(
    call: ToolCall, sample: Path
) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="A1",
        rows=[["=SUM(Data!C2:D2 Data!C3:D3)"]],
    )
    await call("delete_rows_or_columns", path="sales.xlsx", sheet="Data", axis="rows", at=3)
    assert load_workbook(sample)["Report"]["A1"].value == "=SUM(Data!C2:D2 Data!#REF!)"
