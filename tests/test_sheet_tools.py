from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


async def test_create_rename_copy_delete(call: ToolCall, sample: Path) -> None:
    await call("create_sheet", path="sales.xlsx", new_name="Notes", position=1)
    await call("rename_sheet", path="sales.xlsx", sheet="Notes", new_name="Readme")
    await call("copy_sheet", path="sales.xlsx", sheet="Data", new_name="Data copy")
    await call("delete_sheet", path="sales.xlsx", sheet="Report")

    workbook = load_workbook(sample)
    assert workbook.sheetnames == ["Readme", "Data", "Data copy"]
    assert workbook["Data copy"]["C2"].value == 10


async def test_sheet_name_rules(call_error: ToolCall, sample: Path) -> None:
    assert "already exists" in await call_error("create_sheet", path="sales.xlsx", new_name="data")
    assert "cannot contain" in await call_error("create_sheet", path="sales.xlsx", new_name="a/b")
    assert "31" in await call_error("create_sheet", path="sales.xlsx", new_name="x" * 32)


async def test_unknown_sheet_lists_available_sheets(call_error: ToolCall, sample: Path) -> None:
    message = await call_error("read_range", path="sales.xlsx", sheet="Nope")
    assert "'Data', 'Report'" in message


async def test_cannot_delete_last_sheet(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    await call("delete_sheet", path="sales.xlsx", sheet="Report")
    assert "at least one" in await call_error("delete_sheet", path="sales.xlsx", sheet="Data")


async def test_insert_and_delete_rows_and_columns(call: ToolCall, sample: Path) -> None:
    await call("insert_rows_or_columns", path="sales.xlsx", sheet="Data", axis="rows", start=2)
    await call(
        "delete_rows_or_columns", path="sales.xlsx", sheet="Data", axis="columns", start=2, count=2
    )
    sheet = load_workbook(sample)["Data"]
    assert [cell.value for cell in sheet[1]] == ["Region", "Price"]
    assert sheet["A2"].value is None
    assert sheet["A3"].value == "North"


async def test_describe_sheet(call: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    sheet = workbook["Data"]
    sheet.merge_cells("F1:G1")
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["B"].width = 20
    workbook.save(sample)

    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert details["freeze_panes"] == "A2"
    assert details["merged_ranges"] == ["F1:G1"]
    assert details["column_widths_chars"] == [{"column": "B", "width": 20}]
