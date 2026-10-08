"""AutoFilter criteria. Expected hidden rows are what Excel itself hid for the same data."""

import datetime as dt
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

NAMES = [
    "Ann",
    "bob",
    "Cy",
    None,
    "ann",
    "Dee",
    "Ed",
    "Fay",
    "Gus",
    "Hal",
    "Ivy",
    "Jo",
    "Kim",
    "Lee",
    "Max",
    "Ned",
    "Opal",
    "Pam",
    "Quin",
    "Rae",
]
NUMBERS = [5, 10, 10, 3, 8, 7, 7, 7, 1, 2, 9, 4, 6, 12, 11, 15, 15, 15, 20, 0]
WORDS = ["apple", "banana pie", "cherry", "Apple pie", "date"]
RANGE = "A1:F21"


def rows(*spans: tuple[int, int]) -> set[int]:
    return {row for first, last in spans for row in range(first, last + 1)}


@pytest.fixture
def book(files: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Name", "Num", "Date", "Txt", "Pct", "Color"])
    for index, (name, number) in enumerate(zip(NAMES, NUMBERS, strict=True)):
        sheet.append(
            [name, number, dt.date(2026, 1, 1 + index % 5), WORDS[index % 5], number / 100]
        )
        sheet.cell(index + 2, 5).number_format = "0.0%"
    for row, text in ((4, "x"), (6, "y")):
        sheet.cell(row, 6, text).fill = PatternFill("solid", fgColor="FFFF0000")
    sheet["F7"] = "z"
    sheet["F7"].fill = PatternFill("solid", fgColor="FF00FF00")
    path = files / "filter.xlsx"
    workbook.save(path)
    return path


def hidden_rows(path: Path) -> set[int]:
    sheet: Worksheet = load_workbook(path)["Data"]
    return {row for row, dimension in sheet.row_dimensions.items() if dimension.hidden}


async def apply(call: ToolCall, filters: list[dict[str, Any]], range: str = RANGE) -> str:
    return await call(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={"auto_filter": {"range": range, "filters": filters}},
    )


CASES = [
    (
        [{"column": "Name", "type": "values", "values": ["Ann", "Cy", ""]}],
        rows((3, 3), (7, 21)),
    ),
    ([{"column": "Num", "type": "values", "values": ["10", "7"]}], rows((2, 2), (5, 6), (10, 21))),
    (
        [
            {
                "column": "Num",
                "type": "compare",
                "conditions": [
                    {"operator": "greater_or_equal", "value": 5},
                    {"operator": "less_or_equal", "value": 10},
                ],
            }
        ],
        {5, 10, 11, 13, *range(15, 22)},
    ),
    (
        [
            {
                "column": "B",
                "type": "compare",
                "combine": "or",
                "conditions": [
                    {"operator": "less_than", "value": 3},
                    {"operator": "greater_than", "value": 15},
                ],
            }
        ],
        rows((2, 9), (12, 19)),
    ),
    (
        [
            {
                "column": "Txt",
                "type": "compare",
                "conditions": [{"operator": "begins_with", "value": "a"}],
            }
        ],
        {3, 4, 6, 8, 9, 11, 13, 14, 16, 18, 19, 21},
    ),
    (
        [
            {
                "column": "Txt",
                "type": "compare",
                "conditions": [{"operator": "contains", "value": "pie"}],
            }
        ],
        {2, 4, 6, 7, 9, 11, 12, 14, 16, 17, 19, 21},
    ),
    (
        [
            {
                "column": "Txt",
                "type": "compare",
                "conditions": [{"operator": "not_contains", "value": "pie"}],
            }
        ],
        {3, 5, 8, 10, 13, 15, 18, 20},
    ),
    ([{"column": "Num", "type": "top", "count": 3}], rows((2, 16), (21, 21))),
    ([{"column": "Num", "type": "bottom", "count": 3}], rows((2, 9), (12, 20))),
    ([{"column": "Num", "type": "top", "count": 10, "percent": True}], rows((2, 16), (21, 21))),
    ([{"column": "Num", "type": "top", "count": 25, "percent": True}], {*range(2, 15), 16, 21}),
    (
        [{"column": "Num", "type": "bottom", "count": 30, "percent": True}],
        {3, 4, 6, 7, 8, 9, 12, 14, 15, 16, 17, 18, 19, 20},
    ),
    ([{"column": "Num", "type": "above_average"}], {2, 5, 6, 7, 8, 9, 10, 11, 13, 14, 21}),
    ([{"column": "Num", "type": "below_average"}], {3, 4, 12, 15, 16, 17, 18, 19, 20}),
    ([{"column": "Color", "type": "color", "color": "#FF0000"}], rows((2, 3), (5, 5), (7, 21))),
    ([{"column": "Name", "type": "values", "values": [""]}], rows((2, 4), (6, 21))),
    (
        [
            {
                "column": "Num",
                "type": "compare",
                "conditions": [{"operator": "greater_than", "value": 5}],
            },
            {
                "column": "Txt",
                "type": "compare",
                "conditions": [{"operator": "contains", "value": "pie"}],
            },
        ],
        {2, 4, 5, 6, 7, 9, 10, 11, 12, 13, 14, 16, 17, 19, 21},
    ),
]


@pytest.mark.parametrize(("filters", "expected"), CASES)
async def test_criteria_hide_the_rows_excel_hides(
    call: ToolCall, book: Path, filters: list[dict[str, Any]], expected: set[int]
) -> None:
    await apply(call, filters)
    assert hidden_rows(book) == expected
    saved = load_workbook(book)["Data"].auto_filter
    assert saved.ref == RANGE
    assert len(saved.filterColumn) == len(filters)


async def test_stored_criteria(call: ToolCall, book: Path) -> None:
    await apply(call, [{"column": "Num", "type": "top", "count": 3}])
    top = load_workbook(book)["Data"].auto_filter.filterColumn[0].top10
    assert (top.val, top.filterVal, top.percent) == (3, 15, None)

    await apply(call, [{"column": "Num", "type": "bottom", "count": 30, "percent": True}])
    bottom = load_workbook(book)["Data"].auto_filter.filterColumn[0].top10
    assert (bottom.top, bottom.percent, bottom.filterVal) == (False, True, 5)

    await apply(call, [{"column": "Num", "type": "above_average"}])
    dynamic = load_workbook(book)["Data"].auto_filter.filterColumn[0].dynamicFilter
    assert (dynamic.type, dynamic.val) == ("aboveAverage", 8.35)

    await apply(
        call,
        [
            {
                "column": "Num",
                "type": "compare",
                "conditions": [
                    {"operator": "less_than", "value": 3},
                    {"operator": "not_contains", "value": "5*"},
                ],
            }
        ],
    )
    custom = load_workbook(book)["Data"].auto_filter.filterColumn[0].customFilters
    assert [(c.operator, c.val) for c in custom.customFilter] == [
        ("lessThan", "3"),
        ("notEqual", "*5~**"),
    ]
    assert custom._and is True


async def test_values_filter_stores_text_and_blank(call: ToolCall, book: Path) -> None:
    await apply(call, [{"column": "A", "type": "values", "values": ["Ann", "Cy", ""]}])
    filters = load_workbook(book)["Data"].auto_filter.filterColumn[0].filters
    assert filters.blank is True
    assert list(filters.filter) == ["Ann", "Cy"]


async def test_dates_are_filtered_as_date_groups(call: ToolCall, book: Path) -> None:
    await apply(call, [{"column": "Date", "type": "values", "values": ["2026-01-02"]}])
    assert hidden_rows(book) == set(range(2, 22)) - {3, 8, 13, 18}
    (item,) = load_workbook(book)["Data"].auto_filter.filterColumn[0].filters.dateGroupItem
    assert (item.year, item.month, item.day, item.dateTimeGrouping) == (2026, 1, 2, "day")

    await apply(
        call,
        [
            {
                "column": "Date",
                "type": "compare",
                "conditions": [{"operator": "greater_or_equal", "value": "2026-01-04"}],
            }
        ],
    )
    assert hidden_rows(book) == {r for r in range(2, 22) if (r - 2) % 5 < 3}


async def test_formula_results_are_filtered(call: ToolCall, book: Path) -> None:
    await call(
        "write_range",
        path="filter.xlsx",
        sheet="Data",
        start_cell="G1",
        rows=[["Double"], *[[f"=B{row}*2"] for row in range(2, 22)]],
    )
    await apply(
        call,
        [
            {
                "column": "Double",
                "type": "compare",
                "conditions": [{"operator": "greater_than", "value": 20}],
            }
        ],
        range="A1:G21",
    )
    assert hidden_rows(book) == {r for r in range(2, 22) if NUMBERS[r - 2] * 2 <= 20}


async def test_new_criteria_replace_the_old_and_remove_shows_every_row(
    call: ToolCall, book: Path
) -> None:
    await apply(call, [{"column": "Num", "type": "top", "count": 3}])
    await apply(call, [{"column": "Txt", "type": "values", "values": ["date"]}])
    assert hidden_rows(book) == {r for r in range(2, 22) if (r - 2) % 5 != 4}
    assert len(load_workbook(book)["Data"].auto_filter.filterColumn) == 1

    await apply(call, [])
    assert hidden_rows(book) == set()
    assert load_workbook(book)["Data"].auto_filter.ref == RANGE

    await apply(call, [{"column": "Num", "type": "top", "count": 1}])
    await call(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={"auto_filter": {"range": RANGE, "remove": True}},
    )
    assert hidden_rows(book) == set()
    assert not load_workbook(book)["Data"].auto_filter.ref


async def test_moving_the_filter_shows_the_old_rows_and_keeps_other_hidden_rows(
    call: ToolCall, book: Path
) -> None:
    await call(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={"rows": [{"span": "30", "action": "hide"}]},
    )
    await apply(call, [{"column": "Num", "type": "top", "count": 1}])
    await apply(call, [], range="A1:D10")
    assert hidden_rows(book) == {30}
    assert load_workbook(book)["Data"].auto_filter.ref == "A1:D10"


async def test_tables_can_be_filtered(call: ToolCall, book: Path) -> None:
    await call("create_table", path="filter.xlsx", sheet="Data", range=RANGE, name="People")
    await call(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={
            "auto_filter": {
                "range": "people",
                "filters": [{"column": "Num", "type": "top", "count": 3}],
            }
        },
    )
    sheet = load_workbook(book)["Data"]
    assert hidden_rows(book) == rows((2, 16), (21, 21))
    assert not sheet.auto_filter.ref
    table = sheet.tables["People"]
    assert table.autoFilter.ref == RANGE
    assert table.autoFilter.filterColumn[0].top10.filterVal == 15

    await call(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={"auto_filter": {"range": "People", "remove": True}},
    )
    assert hidden_rows(book) == set()
    assert load_workbook(book)["Data"].tables["People"].autoFilter is None


async def test_a_range_inside_a_table_points_to_the_table(
    call: ToolCall, call_error: ToolCall, book: Path
) -> None:
    await call("create_table", path="filter.xlsx", sheet="Data", range=RANGE, name="People")
    message = await call_error(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={"auto_filter": {"range": RANGE}},
    )
    assert "table's name" in message and "'People'" in message


@pytest.mark.parametrize(
    ("filters", "message"),
    [
        ([{"column": "Nope", "type": "top", "count": 1}], "Column 'Nope' is not in"),
        ([{"column": "Num", "type": "top"}], "need a count"),
        ([{"column": "Num", "type": "top", "count": 101, "percent": True}], "need a count"),
        ([{"column": "Num", "type": "values"}], "needs values"),
        ([{"column": "Num", "type": "compare"}], "needs conditions"),
        ([{"column": "Num", "type": "color"}], "needs a color"),
        ([{"column": "Num", "type": "top", "count": 1, "values": ["1"]}], "does not use values"),
        ([{"column": "Txt", "type": "top", "count": 1}], "no numbers"),
        ([{"column": "Pct", "type": "values", "values": ["0.05"]}], "compare filter"),
        (
            [
                {
                    "column": "Num",
                    "type": "compare",
                    "conditions": [{"operator": "equals", "value": ""}],
                }
            ],
            "values filter",
        ),
    ],
)
async def test_invalid_filters_are_rejected(
    call_error: ToolCall, book: Path, filters: list[dict[str, Any]], message: str
) -> None:
    before = book.read_bytes()
    result = await call_error(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={"auto_filter": {"range": RANGE, "filters": filters}},
    )
    assert message in result
    assert book.read_bytes() == before


async def test_filters_respect_the_cell_limit(call_error: ToolCall, book: Path) -> None:
    message = await call_error(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={
            "auto_filter": {
                "range": "A1:F1048576",
                "filters": [{"column": "Num", "type": "top", "count": 1}],
            }
        },
    )
    assert "at most" in message


async def test_filter_needs_a_header_and_data(call_error: ToolCall, book: Path) -> None:
    message = await call_error(
        "set_sheet_layout",
        path="filter.xlsx",
        sheet="Data",
        layout={"auto_filter": {"range": "A1:F1"}},
    )
    assert "header row and at least one data row" in message


async def test_describe_sheet_reports_the_filter_and_hidden_rows(
    call: ToolCall, book: Path
) -> None:
    await apply(call, [{"column": "Num", "type": "top", "count": 3}])
    details = await call("describe_sheet", path="filter.xlsx", sheet="Data")
    assert details["auto_filter"] == RANGE
    assert details["hidden_rows"] == ["2:16", "21:21"]
