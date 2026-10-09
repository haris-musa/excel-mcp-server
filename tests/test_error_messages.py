from pathlib import Path

import pytest
from openpyxl import load_workbook

from excel_mcp.errors import ExcelMCPError, InvalidArgumentError, InvalidFormulaError
from excel_mcp.formulas import check_formula, storable_formula
from excel_mcp.refs import CellRange, parse_clamped_range, parse_range
from excel_mcp.spill import show_spills
from excel_mcp.text import quoted
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

SHEETS = ["Sheet1", "Sheet 2"]


def test_quoted_lists_names_the_same_way_everywhere() -> None:
    assert quoted([]) == ""
    assert quoted(["a"]) == "'a'"
    assert quoted(["a", "b", "c"]) == "'a', 'b', 'c'"
    assert quoted(["values", "formulas"], "or") == "'values' or 'formulas'"
    assert quoted(["a", "b", "c"], "or") == "'a', 'b' or 'c'"


async def test_validation_error_is_one_clean_line(call_error: ToolCall) -> None:
    message = await call_error("read_range", path="a.xlsx", sheet="S", range="A1", mode="formula")
    assert message.endswith(
        "Invalid arguments for read_range: mode: 'formula' is not valid; "
        "use 'values' or 'formulas'."
    )
    assert "pydantic" not in message
    assert "input_value" not in message


async def test_each_problem_has_its_own_line(call_error: ToolCall) -> None:
    message = await call_error("read_range", path="a.xlsx", mode="x", max_cells=0)
    assert "- sheet: required" in message
    assert "- mode: 'x' is not valid; use 'values' or 'formulas'." in message
    assert "- max_cells: must be at least 1" in message


async def test_unknown_field_suggests_the_close_match(call_error: ToolCall) -> None:
    message = await call_error("read_range", path="a.xlsx", sheet="S", rnage="A1")
    assert message.endswith("rnage: unknown field; did you mean 'range'?")


async def test_unknown_field_without_a_close_match_lists_valid_fields(
    call_error: ToolCall,
) -> None:
    message = await call_error("read_range", path="a.xlsx", sheet="S", zzz=1)
    assert "zzz: unknown field; valid fields: 'path', 'sheet', 'range', 'mode'" in message


async def test_nested_paths_and_limits(call_error: ToolCall) -> None:
    message = await call_error(
        "set_sheet_layout",
        path="a.xlsx",
        sheet="S",
        layout={"print_setup": {"scale": 900, "oriantation": "portrait"}},
    )
    assert "layout.print_setup.oriantation: unknown field; did you mean 'orientation'?" in message
    message = await call_error(
        "set_sheet_layout", path="a.xlsx", sheet="S", layout={"print_setup": {"scale": 900}}
    )
    assert message.endswith("layout.print_setup.scale: must be at most 400")


async def test_union_failures_collapse_to_one_line(call_error: ToolCall) -> None:
    message = await call_error(
        "write_range", path="a.xlsx", sheet="S", at="A1", rows=[[1, {"a": 1}]]
    )
    assert message.endswith("rows[0][1]: expected text, number or boolean; got object")


async def test_input_values_are_never_echoed_in_full(call_error: ToolCall) -> None:
    message = await call_error("read_range", path="a.xlsx", sheet="S", mode="secret-" * 50)
    assert "secret-" * 10 not in message
    assert len(message) < 300


def test_range_errors_name_the_problem() -> None:
    with pytest.raises(InvalidArgumentError, match=r"Column ZZZ is past XFD, the last column\."):
        parse_range("A1:ZZZ5")
    with pytest.raises(InvalidArgumentError, match=r"Row 2000000 is past 1048576, the last row\."):
        parse_range("A1:B2000000")
    with pytest.raises(InvalidArgumentError, match=r"rows start at 1"):
        parse_range("A0")


def test_clamped_ranges_accept_whole_columns_and_rows() -> None:
    used = CellRange(2, 1, 10, 4)
    assert str(parse_clamped_range("B:B", lambda: used)) == "B2:B10"
    assert str(parse_clamped_range("$B:$c", lambda: used)) == "B2:C10"
    assert str(parse_clamped_range("3:4", lambda: used)) == "A3:D4"
    assert str(parse_clamped_range("C5:D6", lambda: used)) == "C5:D6"
    with pytest.raises(InvalidArgumentError, match="Column ZZZ is past XFD"):
        parse_clamped_range("A:ZZZ", lambda: used)


async def test_read_range_accepts_whole_columns_and_rows(call: ToolCall, sample: Path) -> None:
    column = await call("read_range", path="sales.xlsx", sheet="Data", range="C:C")
    assert column["range"] == "C1:C5"
    assert column["values"] == [["Units"], [10], [5], [7], [3]]
    rows = await call("read_range", path="sales.xlsx", sheet="Data", range="2:3", mode="formulas")
    assert rows["range"] == "A2:D3"
    assert rows["values"][0] == ["North", "Apples", 10, 1.5]


async def test_insert_and_delete_results_say_what_changed(call: ToolCall, sample: Path) -> None:
    arguments = {"path": "sales.xlsx", "sheet": "Data"}
    inserted = await call("insert_rows_or_columns", **arguments, axis="rows", start=2)
    assert inserted == {"sheet": "Data", "range": "2:2"}
    inserted = await call("insert_rows_or_columns", **arguments, axis="columns", start=3, count=3)
    assert inserted == {"sheet": "Data", "range": "C:E"}
    deleted = await call("delete_rows_or_columns", **arguments, axis="rows", start=5, count=2)
    assert deleted == {"sheet": "Data", "range": "5:6"}
    deleted = await call("delete_rows_or_columns", **arguments, axis="columns", start=1)
    assert deleted == {"sheet": "Data", "range": "A:A"}


async def test_empty_object_lists_say_so(call_error: ToolCall, sample: Path) -> None:
    arguments = {"path": "sales.xlsx", "sheet": "Data", "name": "Chart 1"}
    images = await call_error("delete_image", **arguments)
    charts = await call_error("delete_chart", **arguments)
    assert images.endswith("Sheet 'Data' has no images.")
    assert charts.endswith("Sheet 'Data' has no charts.")


async def test_pivot_and_field_lists_are_quoted(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_pivot_table",
        path="sales.xlsx",
        source="Data!A1:D5",
        sheet="Report",
        at="A1",
        row_fields=["Nope"],
        value_fields=[{"field": "Units"}],
    )
    assert "Available fields: 'Region', 'Product', 'Units', 'Price'." in message


VALID = [
    "=1",
    "=-A1",
    "=+A1",
    "=A1%",
    "=2^-1",
    "=SUM(A1:A3)",
    "=SUM(A1,)",
    "=SUM(,)",
    "=IF(1,,3)",
    "=VLOOKUP(A1,B1:C3,2,)",
    '="a(b"',
    '="a""b"',
    "=A1 B1",
    "=SUM(A1:B2 B2:C3)",
    "=(A1,B1)",
    "=SUM((A1,B1))",
    "=A1,B1",
    "=1 + 2",
    "= 1",
    "={1,2;3,4}",
    "={-1,+2}",
    '={"a","b"}',
    "=SUM({1,2,3})",
    "=A:A",
    "=$A:$C",
    "=1:3",
    "=Sheet1!A1:Sheet1!B2",
    "='Sheet 2'!A1:B2",
    "=MyName",
    "=Table1[Col]",
    "=Table1[[#This Row],[Col]]",
    "=[@Col]",
    "=LET(x,1,x+1)",
    "=@A1",
    "=1+@A1",
    "=@SUM(A1:A3)",
    "=A1:INDEX(A1:A3,2)",
    "=INDEX(A1:A3,1):A3",
    "=1%%",
    "=#REF!",
    "=1E+5",
    "=FOO(1)",
]
INVALID = [
    ("=", "empty"),
    ("=SUM(A1:", "unclosed '('"),
    ("=SUM(A1:A3", "unclosed '('"),
    ("=(1+2", "unclosed '('"),
    ("={1,2", "unclosed '{'"),
    ("=SUM(A1:A3))", "unmatched ')'"),
    ("=1+2)", "unmatched ')'"),
    ("=)", "unmatched ')'"),
    ("=1+", "'+' needs a value after it"),
    ("=A1*", "'*' needs a value after it"),
    ("=SUM(1+)", "'+' needs a value after it"),
    ("=*1", "'*' needs a value before it"),
    ("==1", "'=' needs a value before it"),
    ("=1&&2", "'&' needs a value before it"),
    ("=()", "an expression is missing before ')'"),
    ("=1 2", "an operator is missing before '2'"),
    ("=(1)(2)", "an operator is missing before '('"),
    ("=1(2)", "not a valid function name"),
    ('="a" 1', "an operator is missing"),
    ('="abc', "unterminated text"),
    ('="a""', "unterminated text"),
    ("=@", "'@' needs a value after it"),
    ("=1+@", "'@' needs a value after it"),
    ("=A1@", "not a valid reference"),
    ("=A1 #", "Invalid error code"),
    ("=#REF", "Invalid error code"),
    ("=12abc", "not a valid reference"),
    ("=1.2.3", "not a valid reference"),
    ("=$A", "not a valid reference"),
    ("=A1:", "not a valid reference"),
    ("=Sheet1!", "not a valid reference"),
    ("=1,2", "can only join references"),
    ("=A1;B1", "unexpected ';'"),
    ("=SUM(A1;B1)", "unexpected ';'"),
    ("={1,,2}", "cannot have empty elements"),
    ("={1,2;3}", "same number of elements"),
    ("={A1}", "array constants can only hold"),
    ("={1+1}", "array constants can only hold"),
    ("={{1}}", "array constants can only hold"),
    ("=Table1[Col", "unmatched '['"),
    ("=Table1[[#Foo],[Col]]", "not a valid reference"),
    ("=A-", "'-' needs a value after it"),
    ("=--", "'-' needs a value after it"),
]


@pytest.mark.parametrize("formula", VALID)
def test_valid_formulas_are_accepted(formula: str) -> None:
    check_formula(formula, SHEETS)


@pytest.mark.parametrize(("formula", "reason"), INVALID)
def test_invalid_formulas_are_rejected_with_the_reason(formula: str, reason: str) -> None:
    with pytest.raises(InvalidFormulaError) as raised:
        check_formula(formula, SHEETS)
    message = str(raised.value)
    assert message.startswith(f"Formula {formula!r} is not valid: ")
    assert reason in message


def test_message_for_unclosed_function() -> None:
    with pytest.raises(InvalidFormulaError) as raised:
        check_formula("=SUM(A1:", SHEETS)
    assert str(raised.value) == "Formula '=SUM(A1:' is not valid: unclosed '('."


def test_long_formulas_are_truncated_in_the_message() -> None:
    with pytest.raises(InvalidFormulaError) as raised:
        check_formula("=SUM(" + "A1," * 100, SHEETS)
    assert len(str(raised.value)) < 120


async def test_every_formula_sink_rejects_invalid_formulas(
    call_error: ToolCall, sample: Path
) -> None:
    expected = "Formula '=SUM(A1:' is not valid: unclosed '('."
    arguments = {"path": "sales.xlsx", "sheet": "Data"}
    assert expected in await call_error("write_range", **arguments, at="F1", rows=[["=SUM(A1:"]])
    assert expected in await call_error(
        "add_conditional_format",
        **arguments,
        range="A1:A5",
        rule={"type": "formula", "formula": "=SUM(A1:", "fill_color": "FF0000"},
    )
    assert expected in await call_error(
        "set_defined_name", path="sales.xlsx", name="Total", refers_to="=SUM(A1:"
    )
    assert load_workbook(sample)["Data"]["F1"].value is None


@pytest.mark.parametrize(
    ("formula", "reason"),
    [
        ("=SUM()", "SUM needs at least 1 argument, got 0."),
        ("=IF(1)", "IF needs at least 2 arguments, got 1."),
        ("=IF(1,2,3,)", "IF takes at most 3 arguments, got 4."),
        ("=PI(1)", "PI takes at most 0 arguments, got 1."),
        ("=IFS(A1,1,A1)", "IFS takes 2, 4, 6... arguments, got 3."),
        ("=_xlfn.XLOOKUP(1,A1)", "XLOOKUP needs at least 3 arguments, got 2."),
    ],
)
def test_argument_counts_are_checked(formula: str, reason: str) -> None:
    with pytest.raises(InvalidFormulaError) as raised:
        check_formula(formula, SHEETS)
    assert reason in str(raised.value)


@pytest.mark.parametrize(
    "formula",
    [
        "=PI()",
        "=NOW()",
        "=SUM(A1)",
        "=SUM(,)",
        "=IF(1,2)",
        "=IFS(A1,1)",
        "=MyUdf()",
        "=MyUdf(1,2,3)",
        "=LET(x,1,x+1)",
        "=LAMBDA(x,x)",
        "=FILTER(A1:A3,A1:A3>1)",
        "=ROW()",
        "=ROW(A1)",
    ],
)
def test_valid_argument_counts_and_unknown_functions_pass(formula: str) -> None:
    check_formula(formula, SHEETS)


@pytest.mark.parametrize(
    ("typed", "stored"),
    [
        ("=SUM(A1#)", "=SUM(_xlfn.ANCHORARRAY(A1))"),
        ("=Sheet1!A1#", "=_xlfn.ANCHORARRAY(Sheet1!A1)"),
        ("='Sheet 2'!$A$1#+1", "=_xlfn.ANCHORARRAY('Sheet 2'!$A$1)+1"),
        ("=MyName#", "=_xlfn.ANCHORARRAY(MyName)"),
        ('=IF(A1#="#",A1#,"x#")', '=IF(_xlfn.ANCHORARRAY(A1)="#",_xlfn.ANCHORARRAY(A1),"x#")'),
        ("=Table1[[#This Row],[Col]]", "=Table1[[#This Row],[Col]]"),
        ("=IFERROR(A1,#N/A)", "=IFERROR(A1,#N/A)"),
    ],
)
def test_spill_references_are_stored_as_excel_stores_them(typed: str, stored: str) -> None:
    assert storable_formula(typed, SHEETS) == stored
    assert show_spills(stored) == typed


@pytest.mark.parametrize("formula", ["=A1 #", "=@+1", "=Other!A1#", "=#", "=A1##"])
def test_stray_hashes_and_other_workbooks_stay_rejected(formula: str) -> None:
    with pytest.raises(ExcelMCPError):
        check_formula(formula, SHEETS)


async def test_spill_reference_round_trips_through_the_tools(call: ToolCall, sample: Path) -> None:
    await call("write_range", path="sales.xlsx", sheet="Data", at="F1", rows=[["=SUM(C2#)"]])
    assert load_workbook(sample)["Data"]["F1"].value == "=SUM(_xlfn.ANCHORARRAY(C2))"
    read = await call("read_range", path="sales.xlsx", sheet="Data", range="F1", mode="formulas")
    assert read["values"] == [["=SUM(C2#)"]]
