"""Making a PivotTable again from its own definition, for every case of the Excel golden fixture."""

import json
import re
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import tostring

from openpyxl import Workbook
from openpyxl.pivot.table import TableDefinition
from pivot_cases import SOURCE_RANGE
from pivot_workbook import MAX_CELLS, build_workbook, stored_grids

from excel_mcp.operations import pivot
from excel_mcp.operations.pivot_index import sheet_pivots
from excel_mcp.operations.pivot_rebuild import request_of

FIXTURE = Path(__file__).parent / "fixtures" / "pivot_golden.json"


def _xml(definition: TableDefinition) -> str:
    text = tostring(definition.to_tree(), encoding="unicode") + tostring(
        definition.cache.to_tree(), encoding="unicode"
    )
    return re.sub(r'refreshedDate="[^"]*"', "", text)


def test_a_pivot_made_again_from_its_definition_is_the_same() -> None:
    cases = json.loads(FIXTURE.read_text(encoding="utf-8"))["cases"]
    original = build_workbook(
        [{key: value for key, value in case.items() if key != "expected"} for case in cases]
    )
    again = Workbook()
    again.worksheets[0].title = "Data"
    for row in original["Data"].iter_rows(values_only=True):
        again["Data"].append(row)
    for sheet in original.worksheets[1:]:
        request, (_, reference) = request_of(original, sheet, sheet_pivots(sheet)[0])
        pivot.create_pivot(
            again,
            again["Data"],
            reference,
            again.create_sheet(sheet.title),
            "A1",
            request,
            MAX_CELLS,
        )
    differing: list[Any] = []
    first, second = stored_grids(original), stored_grids(again)
    for sheet in original.worksheets[1:]:
        copy = again[sheet.title]
        if first[sheet.title] != second[sheet.title] or _xml(sheet_pivots(sheet)[0]) != _xml(
            sheet_pivots(copy)[0]
        ):
            differing.append(sheet.title)
    assert SOURCE_RANGE
    assert not differing, f"{len(differing)} of {len(cases)} differ: {differing[:10]}"
