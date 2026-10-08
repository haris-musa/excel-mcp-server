"""Chart options, new chart types, chart listing and deletion, checked in the saved file."""

import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

ALL_TYPES = ["column", "bar", "line", "area", "pie", "scatter", "doughnut", "radar"]


async def add_chart(
    call: ToolCall, chart_type: str = "column", data_range: str = "B1:D5", **options: Any
) -> None:
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        data_sheet="Data",
        data_range=data_range,
        chart_type=chart_type,
        anchor_cell="B2",
        options=options,
    )


def chart_xml(path: Path, number: int = 1) -> ElementTree.Element:
    with zipfile.ZipFile(path) as archive:
        return ElementTree.fromstring(archive.read(f"xl/charts/chart{number}.xml"))


def saved_chart_count(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        return len([n for n in archive.namelist() if n.startswith("xl/charts/chart")])


def elements(root: ElementTree.Element, tag: str) -> list[ElementTree.Element]:
    return [node for node in root.iter() if node.tag.rpartition("}")[2] == tag]


def values(root: ElementTree.Element, tag: str) -> list[str | None]:
    return [node.get("val") for node in elements(root, tag)]


async def test_legend_position(call: ToolCall, sample: Path) -> None:
    await add_chart(call, legend_position="bottom")
    assert values(chart_xml(sample), "legendPos") == ["b"]


async def test_legend_position_needs_a_legend(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        data_range="A1:B2",
        chart_type="column",
        anchor_cell="A1",
        options={"show_legend": False, "legend_position": "top"},
    )
    assert "show_legend" in message


async def test_data_labels(call: ToolCall, sample: Path) -> None:
    await add_chart(call, data_labels=True)
    root = chart_xml(sample)
    assert values(root, "showVal") == ["1"]
    assert set(values(root, "showPercent")) == {"0"}


async def test_data_labels_are_off_by_default(call: ToolCall, sample: Path) -> None:
    await add_chart(call)
    assert elements(chart_xml(sample), "dLbls") == []


@pytest.mark.parametrize(
    ("chart_type", "grouping", "expected"),
    [
        ("column", "stacked", "stacked"),
        ("bar", "percent_stacked", "percentStacked"),
        ("area", "stacked", "stacked"),
        ("line", "percent_stacked", "percentStacked"),
        ("column", "standard", "clustered"),
    ],
)
async def test_grouping(
    call: ToolCall, sample: Path, chart_type: str, grouping: str, expected: str
) -> None:
    await add_chart(call, chart_type, grouping=grouping)
    root = chart_xml(sample)
    assert expected in values(root, "grouping")
    stacked = grouping != "standard" and chart_type in ("column", "bar")
    assert (values(root, "overlap") == ["100"]) is stacked


async def test_colors_fill_series_in_order(call: ToolCall, sample: Path) -> None:
    await add_chart(call, colors=["#1F4E78", "c00000"])
    series = elements(chart_xml(sample), "ser")
    fills = [values(elements(item, "spPr")[0], "srgbClr") for item in series]
    assert fills == [["1F4E78"], ["C00000"]]


async def test_fewer_colors_leave_other_series_alone(call: ToolCall, sample: Path) -> None:
    await add_chart(call, colors=["#1F4E78"])
    series = elements(chart_xml(sample), "ser")
    assert values(series[0], "srgbClr") == ["1F4E78"]
    assert values(series[1], "srgbClr") == []


async def test_line_colors_color_the_line(call: ToolCall, sample: Path) -> None:
    await add_chart(call, "line", colors=["#112233", "#445566"], markers=True)
    series = elements(chart_xml(sample), "ser")
    line_colors = [values(elements(item, "ln")[0], "srgbClr") for item in series]
    assert line_colors == [["112233"], ["445566"]]
    assert values(series[0], "symbol") == ["circle"]


@pytest.mark.parametrize("chart_type", ["pie", "doughnut"])
async def test_round_charts_color_each_slice(call: ToolCall, sample: Path, chart_type: str) -> None:
    await add_chart(call, chart_type, data_range="B1:C5", colors=["#111111", "#222222", "#333333"])
    points = elements(chart_xml(sample), "dPt")
    assert [values(point, "srgbClr") for point in points] == [["111111"], ["222222"], ["333333"]]


async def test_color_errors(call_error: ToolCall, sample: Path) -> None:
    assert "Invalid color" in await call_error(
        "create_chart",
        path="sales.xlsx",
        sheet="Data",
        data_range="B1:C5",
        chart_type="column",
        anchor_cell="G1",
        options={"colors": ["red"]},
    )
    assert "3 colors given for 1 series" in await call_error(
        "create_chart",
        path="sales.xlsx",
        sheet="Data",
        data_range="B1:C5",
        chart_type="column",
        anchor_cell="G1",
        options={"colors": ["#111111", "#222222", "#333333"]},
    )
    assert "4 slices" in await call_error(
        "create_chart",
        path="sales.xlsx",
        sheet="Data",
        data_range="B1:C5",
        chart_type="pie",
        anchor_cell="G1",
        options={"colors": ["#111111"] * 5},
    )


@pytest.mark.parametrize("chart_type", ["line", "scatter"])
async def test_markers_and_smooth(call: ToolCall, sample: Path, chart_type: str) -> None:
    await add_chart(call, chart_type, data_range="C1:D5", markers=False, smooth=True)
    root = chart_xml(sample)
    assert set(values(root, "symbol")) == {"none"}
    assert set(values(root, "smooth")) == {"1"}


async def test_axis_range_and_number_format(call: ToolCall, sample: Path) -> None:
    await add_chart(call, y_axis_min=0, y_axis_max=1.5, y_axis_number_format="#,##0")
    root = chart_xml(sample)
    assert [float(value or "") for value in values(root, "min")] == [0]
    assert [float(value or "") for value in values(root, "max")] == [1.5]
    number_format = elements(elements(root, "valAx")[0], "numFmt")[0]
    assert number_format.get("formatCode") == "#,##0"
    assert number_format.get("sourceLinked") == "0"


@pytest.mark.parametrize("chart_type", ["column", "bar"])
async def test_secondary_axis_line_combo(call: ToolCall, sample: Path, chart_type: str) -> None:
    await add_chart(
        call,
        chart_type,
        secondary_line_columns=["Price"],
        colors=["#111111", "#222222"],
        data_labels=True,
        y_axis_title="Units",
    )
    root = chart_xml(sample)
    (bars,) = elements(root, "barChart")
    (lines,) = elements(root, "lineChart")
    assert [f.text for f in elements(elements(bars, "tx")[0], "f")] == ["'Data'!C1"]
    assert [f.text for f in elements(elements(lines, "tx")[0], "f")] == ["'Data'!D1"]
    assert values(bars, "srgbClr") == ["111111"]
    assert "222222" in values(lines, "srgbClr")
    assert values(bars, "showVal") == ["1"] and values(lines, "showVal") == ["1"]
    assert len(elements(root, "valAx")) == 2
    assert "max" in values(root, "crosses")
    assert values(root, "delete") == ["0", "0", "0"]


async def test_combo_survives_later_edits(call: ToolCall, sample: Path) -> None:
    await add_chart(call, secondary_line_columns=["Price"])
    await call("write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[["x"]])
    root = chart_xml(sample)
    assert len(elements(root, "barChart")) == len(elements(root, "lineChart")) == 1
    assert values(root, "delete") == ["0", "0", "0"]


async def test_secondary_column_errors(call_error: ToolCall, sample: Path) -> None:
    base = {
        "path": "sales.xlsx",
        "sheet": "Data",
        "data_range": "B1:D5",
        "chart_type": "column",
        "anchor_cell": "G1",
    }
    missing = await call_error("create_chart", **base, options={"secondary_line_columns": ["Cost"]})
    assert "'Units', 'Price'" in missing
    everything = await call_error(
        "create_chart", **base, options={"secondary_line_columns": ["Units", "Price"]}
    )
    assert "at least one" in everything
    empty = await call_error("create_chart", **base, options={"secondary_line_columns": []})
    assert "at least one column" in empty


@pytest.mark.parametrize(
    ("chart_type", "options", "fragment"),
    [
        ("pie", {"grouping": "stacked"}, "grouping does not apply to pie"),
        ("scatter", {"grouping": "stacked"}, "grouping does not apply to scatter"),
        ("column", {"markers": True}, "markers does not apply to column"),
        ("area", {"smooth": True}, "smooth does not apply to area"),
        ("line", {"secondary_line_columns": ["Units"]}, "does not apply to line"),
        ("doughnut", {"y_axis_min": 0}, "y_axis_min does not apply to doughnut"),
        ("pie", {"y_axis_number_format": "0%"}, "y_axis_number_format does not apply"),
        ("column", {"y_axis_min": 5, "y_axis_max": 5}, "must be below"),
    ],
)
async def test_options_that_do_not_fit_are_rejected(
    call_error: ToolCall, sample: Path, chart_type: str, options: dict[str, Any], fragment: str
) -> None:
    message = await call_error(
        "create_chart",
        path="sales.xlsx",
        sheet="Data",
        data_range="B1:D5",
        chart_type=chart_type,
        anchor_cell="G1",
        options=options,
    )
    assert fragment in message
    assert saved_chart_count(sample) == 0


@pytest.mark.parametrize(
    ("chart_type", "tag"), [("doughnut", "doughnutChart"), ("radar", "radarChart")]
)
async def test_new_chart_types(call: ToolCall, sample: Path, chart_type: str, tag: str) -> None:
    await add_chart(call, chart_type, title="T")
    assert len(elements(chart_xml(sample), tag)) == 1
    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert details["charts"][0]["type"] == chart_type


@pytest.mark.parametrize("chart_type", ALL_TYPES)
async def test_axes_stay_visible_after_later_edits(
    call: ToolCall, sample: Path, chart_type: str
) -> None:
    await add_chart(call, chart_type, data_range="B1:C5")
    await call("write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[["x"]])
    expected = ["0", "0"] if chart_type not in ("pie", "doughnut") else []
    assert values(chart_xml(sample), "delete") == expected


async def test_describe_sheet_lists_charts(call: ToolCall, sample: Path) -> None:
    await add_chart(call, "bar", title="Horizontal")
    await add_chart(call, "column", secondary_line_columns=["Price"])
    await add_chart(call, "pie", data_range="B1:C5")
    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert details["chart_count"] == 3
    assert details["charts"] == [
        {"index": 1, "type": "bar", "title": "Horizontal", "anchor": "B2"},
        {"index": 2, "type": "column", "title": None, "anchor": "B2"},
        {"index": 3, "type": "pie", "title": None, "anchor": "B2"},
    ]
    assert (await call("describe_sheet", path="sales.xlsx", sheet="Data"))["charts"] == []


async def test_delete_chart(call: ToolCall, sample: Path) -> None:
    await add_chart(call, "bar", title="First")
    await add_chart(call, "line", title="Second")
    message = await call("delete_chart", path="sales.xlsx", sheet="Report", index=1)
    assert "bar chart 1 'First'" in message
    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert [(item["index"], item["type"], item["title"]) for item in details["charts"]] == [
        (1, "line", "Second")
    ]
    assert saved_chart_count(sample) == 1


async def test_delete_chart_rejects_bad_indices(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    arguments = {"path": "sales.xlsx", "sheet": "Report"}
    assert "has no charts" in await call_error("delete_chart", **arguments, index=1)
    await add_chart(call)
    await add_chart(call, "line")
    for index in (0, 3, -1):
        message = await call_error("delete_chart", **arguments, index=index)
        assert f"no chart {index}" in message and "1 to 2" in message
    assert saved_chart_count(sample) == 2
    assert "not found" in await call_error("delete_chart", path="sales.xlsx", sheet="Nope", index=1)


async def test_delete_chart_stays_inside_the_allowed_folder(
    call: ToolCall, call_error: ToolCall, sample: Path, tmp_path: Path
) -> None:
    outside = tmp_path / "outside.xlsx"
    workbook = Workbook()
    workbook.save(outside)
    message = await call_error("delete_chart", path=str(outside), sheet="Sheet", index=1)
    assert "outside" in message
