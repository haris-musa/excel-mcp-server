# pyright: reportAttributeAccessIssue=false, reportOptionalMemberAccess=false
# The tests look at openpyxl internals and at what the lazy cells hold.
"""Rows and columns inserted into or deleted from a large sheet are moved in its XML, and the
result is what loading the whole sheet gives.

Blocks are four rows here, so every group of shared formulas and every edit crosses blocks.
"""

import re
import shutil
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from excel_mcp import lazy_edit
from excel_mcp.operations import line_edits
from tests.conftest import ToolCall, lazily
from tests.package_support import read_parts, sheet_part
from tests.test_lazy_edit import snapshot
from tests.test_lazy_shared import assert_groups_are_whole, groups, share

pytestmark = pytest.mark.anyio

ROWS = 40
GROUP_END = 20  # the formulas are filled down to here, and written one by one below


def make_linked_book(path: Path) -> None:
    """Formulas filled down on 'Data', and on 'Other', which refers to 'Data'."""
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Id", "Qty", "Price", "Amount", "Label", "Gx"])
    for row in range(2, ROWS + 1):
        sheet.append(
            [
                row - 1,
                row % 7,
                row * 1.5,
                f"=B{row}*C{row}",
                f'=IF(D{row}>10,"big & bold","small")',
                f"=$A{row}*B$1+Other!$A$1",
            ]
        )
    for row in range(5, 11):
        sheet.cell(row, 2).border = Border(
            left=Side(), right=Side("thin"), top=Side(), bottom=Side("thin")
        )
        sheet.cell(row, 3).fill = PatternFill("solid", fgColor="FFFF00")
    sheet.cell(1, 1).font = Font(bold=True)
    other = workbook.create_sheet("Other")
    other["A1"] = "=SUM(Data!D2:D40)"
    other["A2"] = "=Data!B30+Data!B3"
    for row in range(3, 30):
        other.cell(row, 1, f"=Data!B{row}*2")
        other.cell(row, 2, f"=A{row}+Data!$A$1")
        other.cell(row, 3, f"=SUM(Data!B2:B{row})")
    workbook.save(path)
    for number, columns, last in ((0, "D", GROUP_END), (1, "E", GROUP_END), (2, "F", GROUP_END)):
        share(path, number, columns, 2, last)
    share(path, 3, "A", 3, 29, "Other")
    share(path, 4, "B", 3, 29, "Other")
    share(path, 5, "C", 3, 29, "Other")


@pytest.fixture(scope="module")
def linked_book(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("linked") / "linked.xlsx"
    make_linked_book(path)
    return path


async def assert_same_as_whole(
    call: ToolCall,
    files: Path,
    monkeypatch: pytest.MonkeyPatch,
    linked_book: Path,
    tool: str,
    **arguments: object,
) -> None:
    shutil.copy(linked_book, files / "lazy.xlsx")
    shutil.copy(linked_book, files / "full.xlsx")
    with lazily(monkeypatch):
        await call(tool, path="lazy.xlsx", **arguments)
    await call(tool, path="full.xlsx", **arguments)
    assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx")
    assert_groups_are_whole(files / "lazy.xlsx")


@pytest.mark.parametrize("delete", [False, True])
@pytest.mark.parametrize("at", [1, 3, 5, 12, 21, 30, 41])
@pytest.mark.parametrize("count", [1, 4])
async def test_rows_move_as_loading_the_whole_sheet_moves_them(
    call: ToolCall,
    files: Path,
    monkeypatch: pytest.MonkeyPatch,
    linked_book: Path,
    delete: bool,
    at: int,
    count: int,
) -> None:
    tool = "delete_rows_or_columns" if delete else "insert_rows_or_columns"
    await assert_same_as_whole(
        call,
        files,
        monkeypatch,
        linked_book,
        tool,
        sheet="Data",
        axis="rows",
        start=at,
        count=count,
    )


@pytest.mark.parametrize("delete", [False, True])
@pytest.mark.parametrize("at", [1, 3, 4, 6, 8])
@pytest.mark.parametrize("count", [1, 2])
async def test_columns_move_as_loading_the_whole_sheet_moves_them(
    call: ToolCall,
    files: Path,
    monkeypatch: pytest.MonkeyPatch,
    linked_book: Path,
    delete: bool,
    at: int,
    count: int,
) -> None:
    tool = "delete_rows_or_columns" if delete else "insert_rows_or_columns"
    await assert_same_as_whole(
        call,
        files,
        monkeypatch,
        linked_book,
        tool,
        sheet="Data",
        axis="columns",
        start=at,
        count=count,
    )


async def test_an_edit_of_the_other_sheet_moves_the_references_to_it_the_same_way(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch, linked_book: Path
) -> None:
    await assert_same_as_whole(
        call,
        files,
        monkeypatch,
        linked_book,
        "insert_rows_or_columns",
        sheet="Other",
        axis="rows",
        start=1,
        count=2,
    )


def test_inserting_below_the_filled_formulas_parses_only_the_rows_around_the_edit(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_linked_book(files / "book.xlsx")

    with lazily(monkeypatch):
        book, _ = lazy_edit.load(files / "book.xlsx")
        sheet = book["Data"]
        line_edits.edit_lines(book, sheet, "rows", 30, 3, delete=False, formula_values=lambda _: {})
        cells = sheet._cells

        assert len(cells.loaded) <= 4
        assert cells.loaded >= {29 // 4, 33 // 4}
        assert 1 not in cells.loaded and 3 not in cells.loaded
        # Rows 33 on are the rows 30 on of the file, renumbered but not parsed.
        assert cells[33, 1].value == 29
        assert cells[33, 4].value == "=B33*C33"


async def test_filled_formulas_that_move_alike_stay_filled_formulas(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_linked_book(files / "book.xlsx")

    with lazily(monkeypatch):
        await call(
            "insert_rows_or_columns", path="book.xlsx", sheet="Data", axis="rows", start=30, count=3
        )

    found = groups(files / "book.xlsx")
    assert found["0"]["master"] == ("D2", f"D2:D{GROUP_END}")
    assert found["2"]["master"] == ("F2", f"F2:F{GROUP_END}")
    assert groups(files / "book.xlsx", "Other")["3"]["master"] == ("A3", "A3:A29")


@pytest.mark.parametrize(("at", "moves", "first"), [(1, True, 4), (5, False, 2), (31, True, 2)])
async def test_a_group_the_edit_cuts_is_split_and_one_it_misses_is_moved(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch, at: int, moves: bool, first: int
) -> None:
    make_linked_book(files / "book.xlsx")

    with lazily(monkeypatch):
        await call(
            "insert_rows_or_columns", path="book.xlsx", sheet="Data", axis="rows", start=at, count=2
        )

    # Inserted at row 5 the edit cuts D2:D20, and the formulas become ordinary ones.
    kept = "0" in groups(files / "book.xlsx")
    assert kept is moves
    saved = load_workbook(files / "book.xlsx")["Data"]
    assert saved.cell(first, 4).value == f"=B{first}*C{first}"


async def test_inserting_past_the_last_row_is_refused_as_it_is_for_a_whole_sheet(
    call: ToolCall, call_error: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_linked_book(files / "book.xlsx")
    await call("write_range", path="book.xlsx", sheet="Data", at="A1048576", rows=[[1]])

    with lazily(monkeypatch):
        message = await call_error(
            "insert_rows_or_columns", path="book.xlsx", sheet="Data", axis="rows", start=2
        )

    assert "past the last" in message


async def test_a_sheet_whose_cells_have_no_address_is_loaded_whole(
    call: ToolCall, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = Workbook()
    for row in range(1, 11):
        workbook.worksheets[0].append([row, row * 2, f"=A{row}+B{row}"])
    workbook.save(files / "lazy.xlsx")
    parts = read_parts(files / "lazy.xlsx")
    name = sheet_part(parts, "Sheet")
    parts[name] = re.sub(rb'(<c) r="[A-Z]+\d+"', rb"\1", parts[name])
    with ZipFile(files / "lazy.xlsx", "w") as archive:
        for part, content in parts.items():
            archive.writestr(part, content)
    shutil.copy(files / "lazy.xlsx", files / "full.xlsx")

    with lazily(monkeypatch):
        await call(
            "insert_rows_or_columns", path="lazy.xlsx", sheet="Sheet", axis="columns", start=2
        )
    await call("insert_rows_or_columns", path="full.xlsx", sheet="Sheet", axis="columns", start=2)

    assert snapshot(files / "lazy.xlsx") == snapshot(files / "full.xlsx")
