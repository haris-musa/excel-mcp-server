"""Series selection, scatter and bubble charts, trendlines, error bars, axes, combos, formatting,
chart sheets and replacing charts, checked in the saved file."""

import zipfile
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from tests.conftest import ToolCall
from tests.test_charts import chart_xml, elements, saved_chart_count, values

pytestmark = pytest.mark.anyio

CATEGORIES = "Data!B2:B5"


async def chart(
    call: ToolCall,
    chart_type: str = "column",
    series: list[dict[str, Any]] | None = None,
    **arguments: Any,
) -> str:
    """Create a chart on 'Report' from explicit series (Units by default)."""
    arguments.setdefault("options", {})
    arguments.setdefault("categories", CATEGORIES)
    return await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        chart_type=chart_type,
        anchor_cell="B2",
        series=series or [{"values": "Data!C2:C5"}],
        **arguments,
    )


async def chart_error(call_error: ToolCall, chart_type: str = "column", **arguments: Any) -> str:
    arguments.setdefault("options", {})
    arguments.setdefault("series", [{"values": "Data!C2:C5"}])
    return await call_error(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        chart_type=chart_type,
        anchor_cell="B2",
        **arguments,
    )


def formulas(root: Any, tag: str) -> list[str | None]:
    return [node.text for item in elements(root, tag) for node in elements(item, "f")]


async def test_series_need_not_be_adjacent_and_take_names(call: ToolCall, sample: Path) -> None:
    await chart(
        call,
        series=[
            {"values": "Data!D2:D5", "name": "Data!D1"},
            {"values": "Data!C2:C5", "name": "Items sold"},
        ],
    )
    root = chart_xml(sample)
    assert formulas(root, "val") == ["'Data'!$D$2:$D$5", "'Data'!$C$2:$C$5"]
    assert formulas(root, "cat") == ["'Data'!$B$2:$B$5"] * 2
    assert formulas(root, "tx") == ["'Data'!$D$1"]
    assert [node.text for node in elements(root, "v")] == ["Items sold"]


async def test_series_and_categories_come_from_any_sheet(call: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook.create_sheet("My Sheet").append(["Q1", 4])
    workbook.save(sample)
    await chart(
        call,
        categories="'My Sheet'!A1:A1",
        series=[{"values": "'My Sheet'!B1:B1"}, {"values": "D2:D2", "categories": "Data!B2:B2"}],
    )
    root = chart_xml(sample)
    assert formulas(root, "val") == ["'My Sheet'!$B$1", "'Report'!$D$2"]
    assert formulas(root, "cat") == ["'My Sheet'!$A$1", "'Data'!$B$2"]


async def test_name_that_is_not_a_cell_reference_is_text(call: ToolCall, sample: Path) -> None:
    await chart(call, series=[{"values": "Data!C2:C5", "name": "Q1"}])
    root = chart_xml(sample)
    assert formulas(root, "tx") == []
    assert [node.text for node in elements(root, "v")] == ["Q1"]


async def test_series_in_rows(call: ToolCall, sample: Path) -> None:
    await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="A10",
        rows=[["", "Q1", "Q2"], ["North", 1, 2], ["South", 3, 4]],
    )
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        chart_type="line",
        anchor_cell="E2",
        data_range="A10:C12",
        series_in="rows",
    )
    root = chart_xml(sample)
    assert formulas(root, "val") == ["'Report'!$B$11:$C$11", "'Report'!$B$12:$C$12"]
    assert formulas(root, "cat") == ["'Report'!$B$10:$C$10"] * 2
    assert formulas(root, "tx") == ["'Report'!$A$11", "'Report'!$A$12"]


@pytest.mark.parametrize(
    ("arguments", "fragment"),
    [
        ({"data_range": "Data!A1:C5"}, "either data_range"),
        ({"series": [], "data_range": None}, "either data_range"),
        ({"series": [{"values": "Data!C2:D5"}]}, "one row or one column"),
        ({"series": [{"values": "Nope!C2:C5"}]}, "'Nope' not found"),
        ({"series": [{"values": "C2:C5", "extra": 1}]}, "unknown field"),
        ({"categories": "Data!A2:B5"}, "one row or one column"),
        ({"series": [{"values": "Data!C2:C5", "sizes": "Data!D2:D5"}]}, "bubble charts"),
    ],
)
async def test_series_errors(
    call_error: ToolCall, sample: Path, arguments: dict[str, Any], fragment: str
) -> None:
    message = await chart_error(call_error, **arguments)
    assert fragment in message
    assert saved_chart_count(sample) == 0


@pytest.mark.parametrize(
    ("style", "smooth", "line", "marker"),
    [
        ("markers", "0", False, "circle"),
        ("lines_markers", "0", True, "circle"),
        ("lines", "0", True, "none"),
        ("smooth_markers", "1", True, "circle"),
        ("smooth", "1", True, "none"),
    ],
)
async def test_scatter_subtypes(
    call: ToolCall, sample: Path, style: str, smooth: str, line: bool, marker: str
) -> None:
    await chart(call, "scatter", options={"scatter_style": style})
    root = chart_xml(sample)
    series = elements(root, "ser")[0]
    assert values(series, "smooth") == [smooth]
    assert values(series, "symbol") == [marker]
    assert bool(elements(elements(series, "ln")[0], "noFill")) is not line if line else True
    assert values(root, "scatterStyle") == ["smoothMarker" if smooth == "1" else "lineMarker"]


async def test_bubble_chart_from_a_block_and_from_series(call: ToolCall, sample: Path) -> None:
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        chart_type="bubble",
        anchor_cell="B2",
        data_range="Data!C1:E5",
    )
    await chart(
        call,
        "bubble",
        categories="Data!C2:C5",
        series=[{"values": "Data!D2:D5", "sizes": "Data!C2:C5", "name": "Bubbles"}],
    )
    first, second = chart_xml(sample, 1), chart_xml(sample, 2)
    assert len(elements(first, "bubbleChart")) == len(elements(second, "bubbleChart")) == 1
    assert formulas(first, "bubbleSize") == ["'Data'!$E$2:$E$5"]
    assert formulas(second, "bubbleSize") == ["'Data'!$C$2:$C$5"]
    assert values(second, "bubbleScale") == ["100"]


async def test_bubble_needs_sizes(call_error: ToolCall, sample: Path) -> None:
    message = await chart_error(call_error, "bubble")
    assert "sizes are required for bubble" in message


@pytest.mark.parametrize(
    ("kind", "extra", "xml_type"),
    [
        ("linear", {}, "linear"),
        ("exponential", {}, "exp"),
        ("logarithmic", {}, "log"),
        ("power", {}, "power"),
        ("polynomial", {"order": 3}, "poly"),
        ("moving_average", {"period": 2}, "movingAvg"),
    ],
)
async def test_trendlines(
    call: ToolCall, sample: Path, kind: str, extra: dict[str, int], xml_type: str
) -> None:
    await chart(
        call, "line", series=[{"values": "Data!C2:C5", "trendline": {"type": kind, **extra}}]
    )
    trend = elements(chart_xml(sample), "trendline")[0]
    assert values(trend, "trendlineType") == [xml_type]
    assert values(trend, "order") == ([str(extra["order"])] if "order" in extra else [])
    assert values(trend, "period") == ([str(extra["period"])] if "period" in extra else [])


async def test_trendline_equation_and_r_squared(call: ToolCall, sample: Path) -> None:
    trendline = {"type": "linear", "equation": True, "r_squared": True}
    await chart(call, "scatter", series=[{"values": "Data!C2:C5", "trendline": trendline}])
    trend = elements(chart_xml(sample), "trendline")[0]
    assert values(trend, "dispEq") == ["1"] and values(trend, "dispRSqr") == ["1"]
    assert values(elements(trend, "spPr")[0], "prstDash") == ["sysDot"]
    assert len(elements(trend, "trendlineLbl")) == 1


@pytest.mark.parametrize(
    ("chart_type", "trendline", "fragment"),
    [
        ("line", {"type": "linear", "order": 2}, "order applies to polynomial"),
        ("line", {"type": "linear", "period": 2}, "period applies to moving_average"),
        ("line", {"type": "moving_average"}, "needs a period"),
        ("line", {"type": "moving_average", "period": 2, "equation": True}, "no equation"),
        ("pie", {"type": "linear"}, "trendline does not apply to pie"),
        ("radar", {"type": "linear"}, "trendline does not apply to radar"),
    ],
)
async def test_trendline_errors(
    call_error: ToolCall, sample: Path, chart_type: str, trendline: dict[str, Any], fragment: str
) -> None:
    message = await chart_error(
        call_error, chart_type, series=[{"values": "Data!C2:C5", "trendline": trendline}]
    )
    assert fragment in message and message.count("Series 1:") == 1


async def test_stacked_series_take_no_trendline(call_error: ToolCall, sample: Path) -> None:
    message = await chart_error(
        call_error,
        series=[{"values": "Data!C2:C5", "trendline": {"type": "linear"}}],
        options={"grouping": "stacked"},
    )
    assert "stacked" in message


@pytest.mark.parametrize(
    ("bars", "expected"),
    [
        ({"kind": "fixed", "value": 2}, ("fixedVal", "both", "2", "0")),
        ({"kind": "percent", "value": 10, "direction": "plus"}, ("percentage", "plus", "10", "0")),
        ({"kind": "std_dev", "value": 1, "direction": "minus"}, ("stdDev", "minus", "1", "0")),
        ({"kind": "std_error", "end_cap": False}, ("stdErr", "both", None, "1")),
    ],
)
async def test_error_bars(
    call: ToolCall, sample: Path, bars: dict[str, Any], expected: tuple[str | None, ...]
) -> None:
    await chart(call, series=[{"values": "Data!C2:C5", "error_bars": bars}])
    node = elements(chart_xml(sample), "errBars")[0]
    kind, direction, value, no_cap = expected
    assert values(node, "errValType") == [kind]
    assert values(node, "errBarType") == [direction]
    assert values(node, "val") == ([value] if value else [])
    assert values(node, "noEndCap") == [no_cap]
    assert values(node, "errDir") == ["y"]


async def test_scatter_error_bars_in_x(call: ToolCall, sample: Path) -> None:
    bars = {"kind": "fixed", "value": 1, "axis": "x"}
    await chart(call, "scatter", series=[{"values": "Data!C2:C5", "error_bars": bars}])
    assert values(chart_xml(sample), "errDir") == ["x"]


@pytest.mark.parametrize(
    ("chart_type", "bars", "fragment"),
    [
        ("column", {"kind": "fixed"}, "need a value"),
        ("column", {"kind": "std_error", "value": 1}, "take no value"),
        ("column", {"kind": "fixed", "value": 1, "axis": "x"}, "scatter and bubble"),
        ("pie", {"kind": "fixed", "value": 1}, "error_bars does not apply to pie"),
        ("column", {"kind": "fixed", "value": 0}, "greater than 0"),
    ],
)
async def test_error_bar_errors(
    call_error: ToolCall, sample: Path, chart_type: str, bars: dict[str, Any], fragment: str
) -> None:
    message = await chart_error(
        call_error, chart_type, series=[{"values": "Data!C2:C5", "error_bars": bars}]
    )
    assert fragment in message


async def test_axis_scale_options(call: ToolCall, sample: Path) -> None:
    axis = {"min": 1, "max": 100, "major_unit": 10, "log": True, "reverse": True}
    await chart(call, options={"y_axis": axis | {"number_format": "0.0%"}})
    value_axis = elements(chart_xml(sample), "valAx")[0]
    assert values(value_axis, "logBase") == ["10"]
    assert values(value_axis, "orientation") == ["maxMin"]
    assert values(value_axis, "majorUnit") == ["10"]
    assert elements(value_axis, "numFmt")[0].get("formatCode") == "0.0%"


async def test_scatter_x_axis_is_a_value_axis(call: ToolCall, sample: Path) -> None:
    await chart(call, "scatter", options={"x_axis": {"min": 0, "max": 10, "log": False}})
    assert values(elements(chart_xml(sample), "valAx")[0], "max") == ["10"]


@pytest.mark.parametrize(
    ("chart_type", "axis", "settings", "expected"),
    [
        ("column", "y_axis", {}, [True, False]),
        ("column", "y_axis", {"major_gridlines": False}, [False, False]),
        ("column", "y_axis", {"minor_gridlines": True}, [True, True]),
        ("column", "x_axis", {"major_gridlines": True}, [True, False]),
        ("scatter", "x_axis", {}, [True, False]),
        ("scatter", "x_axis", {"major_gridlines": False}, [False, False]),
    ],
)
async def test_gridlines(
    call: ToolCall,
    sample: Path,
    chart_type: str,
    axis: str,
    settings: dict[str, Any],
    expected: list[bool],
) -> None:
    await chart(call, chart_type, options={axis: settings})
    root = chart_xml(sample)
    index = 0 if axis == "x_axis" else 1 if chart_type == "scatter" else 0
    target = elements(root, "catAx" if chart_type != "scatter" and axis == "x_axis" else "valAx")
    node = target[index if chart_type == "scatter" else 0]
    assert [bool(elements(node, tag)) for tag in ("majorGridlines", "minorGridlines")] == expected


async def test_tick_labels_position(call: ToolCall, sample: Path) -> None:
    await chart(call, options={"x_axis": {"labels": "low"}, "y_axis": {"labels": "high"}})
    root = chart_xml(sample)
    assert values(elements(root, "catAx")[0], "tickLblPos") == ["low"]
    assert values(elements(root, "valAx")[0], "tickLblPos") == ["high"]


@pytest.mark.parametrize(
    ("reverse", "orientation", "crosses"), [(False, "maxMin", "max"), (True, "minMax", "autoZero")]
)
async def test_bar_chart_reverse_gives_excels_own_order(
    call: ToolCall, sample: Path, reverse: bool, orientation: str, crosses: str
) -> None:
    await chart(call, "bar", options={"x_axis": {"reverse": reverse}})
    root = chart_xml(sample)
    assert values(elements(root, "catAx")[0], "orientation") == [orientation]
    assert values(elements(root, "valAx")[0], "crosses") == [crosses]


@pytest.mark.parametrize(
    ("chart_type", "series_type", "primary", "second"),
    [
        ("column", "line", "barChart", "lineChart"),
        ("line", "column", "barChart", "lineChart"),
        ("area", "column", "areaChart", "barChart"),
    ],
)
async def test_combo_charts_in_both_directions(
    call: ToolCall, sample: Path, chart_type: str, series_type: str, primary: str, second: str
) -> None:
    await chart(
        call,
        chart_type,
        series=[{"values": "Data!C2:C5"}, {"values": "Data!D2:D5", "type": series_type}],
    )
    root = chart_xml(sample)
    assert len(elements(root, primary)) == len(elements(root, second)) == 1
    assert len(elements(root, "valAx")) == 1


async def test_secondary_axis_for_a_series(call: ToolCall, sample: Path) -> None:
    await chart(
        call,
        "line",
        series=[{"values": "Data!C2:C5"}, {"values": "Data!D2:D5", "secondary_axis": True}],
        options={"secondary_y_axis": {"min": 0, "max": 9, "title": "Price"}},
    )
    root = chart_xml(sample)
    primary, secondary = elements(root, "lineChart")
    assert len(values(primary, "axId")) == len(values(secondary, "axId")) == 2
    assert set(values(primary, "axId")).isdisjoint(values(secondary, "axId"))
    first, second = elements(root, "valAx")
    assert values(first, "crosses") == ["autoZero"] and values(second, "crosses") == ["max"]
    assert values(second, "axPos") == ["r"] and values(second, "max") == ["9"]
    assert len(elements(second, "title")) == 1 and len(elements(second, "majorGridlines")) == 0
    hidden = elements(root, "catAx")[1]
    assert values(hidden, "delete") == ["1"]


async def test_secondary_scatter_axes(call: ToolCall, sample: Path) -> None:
    await chart(
        call,
        "scatter",
        series=[{"values": "Data!C2:C5"}, {"values": "Data!D2:D5", "secondary_axis": True}],
    )
    root = chart_xml(sample)
    assert len(elements(root, "valAx")) == 4
    assert len(elements(root, "scatterChart")) == 2
    assert not [
        node
        for node in elements(root, "valAx")
        if values(node, "axId") == ["500"] and elements(node, "majorGridlines")
    ]


async def test_grouping_applies_to_the_charts_own_type(call: ToolCall, sample: Path) -> None:
    await chart(
        call,
        series=[{"values": "Data!C2:C5"}, {"values": "Data!D2:D5", "type": "line"}],
        options={"grouping": "stacked"},
    )
    root = chart_xml(sample)
    assert values(elements(root, "barChart")[0], "grouping") == ["stacked"]
    assert values(elements(root, "lineChart")[0], "grouping") == ["standard"]


@pytest.mark.parametrize(
    ("chart_type", "series", "fragment"),
    [
        (
            "pie",
            [{"values": "Data!C2:C5", "type": "line"}],
            "needs chart_type column, line or area",
        ),
        ("column", [{"values": "Data!C2:C5", "secondary_axis": True}], "primary axis"),
        (
            "pie",
            [{"values": "Data!C2:C5", "secondary_axis": True}, {"values": "Data!D2:D5"}],
            "secondary_axis does not apply to pie",
        ),
        (
            "column",
            [{"values": "Data!C2:C5", "marker": "circle"}],
            "marker does not apply to column",
        ),
        (
            "column",
            [{"values": "Data!C2:C5", "line_width": 2}],
            "line_width does not apply to column",
        ),
        ("column", [{"values": "Data!C2:C5", "color": "red"}], "Invalid color"),
    ],
)
async def test_combo_and_series_format_errors(
    call_error: ToolCall, sample: Path, chart_type: str, series: list[dict[str, Any]], fragment: str
) -> None:
    assert fragment in await chart_error(call_error, chart_type, series=series)


async def test_series_formatting(call: ToolCall, sample: Path) -> None:
    await chart(
        call,
        "line",
        series=[
            {
                "values": "Data!C2:C5",
                "color": "#C00000",
                "line_width": 3,
                "marker": "diamond",
                "marker_size": 9,
            }
        ],
    )
    series = elements(chart_xml(sample), "ser")[0]
    assert values(elements(series, "ln")[0], "srgbClr") == ["C00000"]
    assert elements(series, "ln")[0].get("w") == "38100"
    assert values(series, "symbol") == ["diamond"] and values(series, "size") == ["9"]


async def test_series_color_beats_the_color_list(call: ToolCall, sample: Path) -> None:
    await chart(
        call,
        series=[{"values": "Data!C2:C5", "color": "#111111"}, {"values": "Data!D2:D5"}],
        options={"colors": ["#222222", "#333333"]},
    )
    assert [values(item, "srgbClr")[0] for item in elements(chart_xml(sample), "ser")] == [
        "111111",
        "333333",
    ]


async def test_data_label_position_and_format(call: ToolCall, sample: Path) -> None:
    labels = {"position": "outside_end", "number_format": "0.0", "show": ["value", "category"]}
    await chart(call, series=[{"values": "Data!C2:C5", "data_labels": labels}])
    node = elements(chart_xml(sample), "dLbls")[0]
    assert values(node, "dLblPos") == ["outEnd"]
    assert (
        elements(node, "numFmt")[0].get("formatCode"),
        elements(node, "numFmt")[0].get("sourceLinked"),
    ) == ("0.0", "0")
    assert values(node, "showVal") == ["1"] and values(node, "showCatName") == ["1"]


async def test_series_labels_win_over_chart_labels(call: ToolCall, sample: Path) -> None:
    await chart(
        call,
        series=[
            {"values": "Data!C2:C5", "data_labels": {"position": "center"}},
            {"values": "Data!D2:D5"},
        ],
        options={"data_labels": {"position": "inside_end"}},
    )
    assert values(chart_xml(sample), "dLblPos") == ["ctr", "inEnd"]


@pytest.mark.parametrize(
    ("chart_type", "options", "fragment"),
    [
        ("column", {"position": "above"}, "Valid positions: center, inside_end"),
        ("line", {"position": "outside_end"}, "Valid positions: center, above"),
        ("area", {"position": "center"}, "Valid positions: none"),
        ("column", {"show": ["percent"]}, "pie and doughnut"),
    ],
)
async def test_data_label_errors(
    call_error: ToolCall, sample: Path, chart_type: str, options: dict[str, Any], fragment: str
) -> None:
    message = await chart_error(
        call_error, chart_type, series=[{"values": "Data!C2:C5", "data_labels": options}]
    )
    assert fragment in message


async def test_stacked_columns_take_no_outside_labels(call_error: ToolCall, sample: Path) -> None:
    message = await chart_error(
        call_error,
        options={"grouping": "stacked", "data_labels": {"position": "outside_end"}},
    )
    assert "stacked column" in message


async def test_pie_percent_labels(call: ToolCall, sample: Path) -> None:
    await chart(call, "pie", options={"data_labels": {"show": ["percent"], "position": "best_fit"}})
    root = chart_xml(sample)
    assert values(root, "showPercent") == ["1"] and values(root, "showVal") == ["0"]
    assert values(root, "dLblPos") == ["bestFit"]


async def test_chart_style_title_size_and_plot_color(call: ToolCall, sample: Path) -> None:
    await chart(call, options={"title": "T", "title_size": 20, "plot_color": "#F2F2F2", "style": 4})
    root = chart_xml(sample)
    assert values(root, "style") == ["4"]
    assert elements(elements(root, "title")[0], "defRPr")[0].get("sz") == "2000"
    plot_fill = elements(elements(root, "plotArea")[0], "spPr")[-1]
    assert values(plot_fill, "srgbClr") == ["F2F2F2"]


async def test_default_look_matches_excel(call: ToolCall, sample: Path) -> None:
    await chart(call)
    root = chart_xml(sample)
    assert values(root, "roundedCorners") == ["0"]
    assert values(root, "varyColors") == ["0"]
    assert values(root, "invertIfNegative") == ["0"]
    assert values(root, "gapWidth") == ["219"] and values(root, "overlap") == ["-27"]
    assert values(root, "legendPos") == ["b"]
    assert set(values(root, "majorTickMark")) == {"none"}
    assert "style" not in {node.tag.rpartition("}")[2] for node in root}


async def test_chart_details_survive_later_edits(call: ToolCall, sample: Path) -> None:
    await chart(
        call,
        series=[{"values": "Data!C2:C5", "data_labels": {"number_format": "0.0"}}],
        options={"style": 5, "plot_color": "#F2F2F2", "x_axis": {"title": "Axis"}},
    )
    await chart(call, options={"title": "Title"})
    await call("write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[["x"]])
    root = chart_xml(sample, 1)
    assert values(root, "style") == ["5"] and values(root, "roundedCorners") == ["0"]
    assert values(elements(root, "plotArea")[0], "srgbClr")[-1] == "F2F2F2"
    assert [
        node.get("sourceLinked") for node in elements(elements(root, "dLbls")[0], "numFmt")
    ] == ["0"]
    assert "None" not in [node.text for node in elements(root, "t")]
    assert len(elements(root, "t")) == 1


async def test_chart_on_its_own_sheet(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    message = await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Units chart",
        chart_type="column",
        series=[{"values": "Data!C2:C5"}],
        options={"title": "Units"},
    )
    assert "chart sheet 'Units chart'" in message
    workbook = load_workbook(sample)
    assert workbook.sheetnames == ["Data", "Report", "Units chart"]
    assert [sheet.title for sheet in workbook.chartsheets] == ["Units chart"]
    with zipfile.ZipFile(sample) as archive:
        assert "xl/chartsheets/sheet1.xml" in archive.namelist()
    info = await call("describe_workbook", path="sales.xlsx")
    assert info["chart_sheets"] == ["Units chart"]
    await call("write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[["x"]])
    assert saved_chart_count(sample) == 1
    await call("delete_sheet", path="sales.xlsx", sheet="Units chart")
    assert load_workbook(sample).sheetnames == ["Data", "Report"]


@pytest.mark.parametrize(
    ("arguments", "fragment"),
    [
        ({"sheet": "Data"}, "already exists"),
        ({"sheet": "Bad/name"}, "cannot contain"),
        ({"sheet": "New", "index": 1}, "holds one chart"),
        ({"sheet": "Report", "index": 1}, "holds one chart"),
    ],
)
async def test_chart_sheet_errors(
    call_error: ToolCall, sample: Path, arguments: dict[str, Any], fragment: str
) -> None:
    message = await call_error(
        "create_chart",
        path="sales.xlsx",
        chart_type="column",
        series=[{"values": "Data!C2:C5"}],
        **arguments,
    )
    assert fragment in message
    assert saved_chart_count(sample) == 0


async def test_replace_a_chart_in_place(call: ToolCall, sample: Path) -> None:
    await chart(call, "bar", options={"title": "First"})
    await chart(call, "line", options={"title": "Second"})
    message = await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        chart_type="area",
        anchor_cell="H9",
        index=1,
        series=[{"values": "Data!D2:D5"}],
        options={"title": "Replaced"},
    )
    assert "Replaced chart 1" in message
    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert [
        (item["index"], item["type"], item["title"], item["anchor"]) for item in details["charts"]
    ] == [
        (1, "area", "Replaced", "H9"),
        (2, "line", "Second", "B2"),
    ]
    assert saved_chart_count(sample) == 2


async def test_replace_rejects_bad_indices(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await chart(call)
    for index in (0, 2):
        message = await call_error(
            "create_chart",
            path="sales.xlsx",
            sheet="Report",
            chart_type="line",
            anchor_cell="B2",
            index=index,
            series=[{"values": "Data!C2:C5"}],
        )
        assert f"no chart {index}" in message and "1 to 1" in message
    assert saved_chart_count(sample) == 1


async def test_create_chart_stays_inside_the_allowed_folder(
    call_error: ToolCall, sample: Path, tmp_path: Path
) -> None:
    outside = tmp_path / "outside.xlsx"
    load_workbook(sample).save(outside)
    message = await call_error(
        "create_chart",
        path=str(outside),
        sheet="Report",
        chart_type="column",
        anchor_cell="B2",
        series=[{"values": "Data!C2:C5"}],
    )
    assert "outside" in message
