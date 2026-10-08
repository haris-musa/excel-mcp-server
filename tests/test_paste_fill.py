"""Paste Special options for copy_range and Fill Down, Right and Series."""

import datetime as dt
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio

BOOK = {"path": "paste.xlsx", "sheet": "Data"}
RED = PatternFill("solid", fgColor="FFFF0000")
GREEN = PatternFill("solid", fgColor="FF00FF00")


@pytest.fixture
def book(files: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet["A1"] = dt.date(2026, 1, 31)
    sheet["A1"].number_format = "yyyy-mm-dd"
    sheet["A1"].fill = RED
    sheet["A2"] = "=C1*2"
    sheet["A2"].number_format = "0.00"
    sheet["A3"] = 5
    sheet["A3"].font = Font(bold=True)
    sheet["C1"] = 21
    path = files / "paste.xlsx"
    workbook.save(path)
    return path


def data(path: Path) -> Worksheet:
    return load_workbook(path)["Data"]


async def copy(call: ToolCall, **arguments: Any) -> str:
    return await call("copy_range", **BOOK, **arguments)


async def test_paste_values_gives_results_without_formatting(call: ToolCall, book: Path) -> None:
    await copy(call, range="A1:A3", target_cell="E1", paste="values")
    sheet = data(book)
    assert [sheet.cell(row, 5).value for row in (1, 2, 3)] == [46053, 42, 5]
    assert all(sheet.cell(row, 5).number_format == "General" for row in (1, 2, 3))
    assert sheet["E1"].fill.fill_type is None


async def test_paste_formulas_keeps_formulas_and_leaves_formatting(
    call: ToolCall, book: Path
) -> None:
    await copy(call, range="A2", target_cell="B2", paste="formulas")
    sheet = data(book)
    assert sheet["B2"].value == "=D1*2"
    assert sheet["B2"].number_format == "General"


async def test_paste_formats_copies_only_styles(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, start_cell="E1", rows=[["keep"]])
    await copy(call, range="A1:A3", target_cell="E1", paste="formats")
    sheet = data(book)
    assert sheet["E1"].value == "keep"
    assert sheet["E1"].fill.start_color.rgb == "FFFF0000"
    assert sheet["E1"].number_format == "yyyy-mm-dd"
    assert sheet["E3"].font.bold is True
    assert sheet["E2"].value is None


async def test_paste_all_still_copies_values_and_formats(call: ToolCall, book: Path) -> None:
    await copy(call, range="A1:A3", target_cell="E1")
    sheet = data(book)
    assert sheet["E1"].value == dt.datetime(2026, 1, 31)
    assert sheet["E1"].fill.start_color.rgb == "FFFF0000"
    assert sheet["E2"].value == "=G1*2"


async def test_skip_blanks_leaves_destination_cells_alone(call: ToolCall, book: Path) -> None:
    workbook = load_workbook(book)
    sheet = workbook["Data"]
    sheet["A2"] = None
    sheet["A2"].fill = RED
    for row in (1, 2, 3):
        sheet.cell(row, 5, "old").fill = GREEN
    workbook.save(book)
    await copy(call, range="A1:A3", target_cell="E1", skip_blanks=True)
    sheet = data(book)
    assert [sheet.cell(row, 5).value for row in (1, 2, 3)] == [dt.datetime(2026, 1, 31), "old", 5]
    assert sheet["E2"].fill.start_color.rgb == "FF00FF00"
    assert sheet["E1"].fill.start_color.rgb == "FFFF0000"

    await copy(call, range="A1:A3", target_cell="E1")
    assert data(book)["E2"].value is None


async def test_transpose_swaps_rows_and_columns_and_formula_offsets(
    call: ToolCall, book: Path
) -> None:
    workbook = load_workbook(book)
    sheet = workbook["Data"]
    sheet["H1"], sheet["I1"], sheet["J1"] = 1, 2, 3
    sheet["H2"], sheet["I2"], sheet["J2"] = "=H1*2", "=I1*2", "=I2+J1"
    workbook.save(book)
    message = await copy(call, range="H1:J2", target_cell="L5", transpose=True)
    assert message == "Copied Data!H1:J2 to Data!L5:M7."
    sheet = data(book)
    # Excel's result: a reference's offset (rows, columns) becomes (columns, rows).
    assert [[sheet.cell(row, col).value for col in (12, 13)] for row in (5, 6, 7)] == [
        [1, "=L5*2"],
        [2, "=L6*2"],
        [3, "=M6+L7"],
    ]


async def test_transposed_references_follow_excel(call: ToolCall, book: Path) -> None:
    await call("create_sheet", path="paste.xlsx", sheet="Other")
    formula = "=$C$1+C1+$C1+C$1+Other!B1+Other!$B2+SUM(C1:D2)+SUM($C1:D$2)"
    await call("write_range", **BOOK, start_cell="M10", rows=[[formula]])
    await copy(call, range="M10", target_cell="O20", transpose=True)
    # Recorded from Excel: relative references swap their offsets, any "$" reference stays.
    assert data(book)["O20"].value == (
        "=$C$1+F10+$C1+C$1+Other!F9+Other!$B2+SUM(F10:G11)+SUM($C1:D$2)"
    )


async def test_transposing_formulas_with_whole_lines_or_off_the_sheet_is_refused(
    call: ToolCall, call_error: ToolCall, book: Path
) -> None:
    await call("write_range", **BOOK, start_cell="H1", rows=[["=SUM(C:C)"], ["=A1"]])
    message = await call_error("copy_range", **BOOK, range="H1", target_cell="H4", transpose=True)
    assert "whole row or column" in message
    message = await call_error("copy_range", **BOOK, range="H2", target_cell="A9", transpose=True)
    assert "off the sheet" in message


async def test_paste_values_of_uncalculable_formulas_is_refused(
    call: ToolCall, call_error: ToolCall, book: Path
) -> None:
    await call("write_range", **BOOK, start_cell="H1", rows=[["=NOSUCHFUNCTION(1)"]])
    message = await call_error("copy_range", **BOOK, range="H1", target_cell="H4", paste="values")
    assert "H1" in message and "recalculates" in message


async def test_paste_options_keep_the_array_formula_guard(call_error: ToolCall, book: Path) -> None:
    from openpyxl.worksheet.formula import ArrayFormula

    workbook = load_workbook(book)
    workbook["Data"]["H1"] = ArrayFormula("H1", "=C1*2")
    workbook.save(book)
    message = await call_error("copy_range", **BOOK, range="H1", target_cell="H4")
    assert "array formula" in message


# -- fill ------------------------------------------------------------------------------------


async def fill(call: ToolCall, range: str, **transform: Any) -> str:
    return await call(
        "transform_range", **BOOK, range=range, transform={"operation": "fill", **transform}
    )


async def test_fill_down_repeats_the_first_row_with_formats(call: ToolCall, book: Path) -> None:
    workbook = load_workbook(book)
    sheet = workbook["Data"]
    sheet["J1"], sheet["K1"] = 5, "=J1*2"
    sheet["J1"].fill = sheet["K1"].fill = RED
    sheet["J3"] = "overwritten"
    workbook.save(book)
    assert await fill(call, "J1:K4", direction="down") == "Filled J1:K4."
    sheet = data(book)
    assert [(sheet.cell(row, 10).value, sheet.cell(row, 11).value) for row in (2, 3, 4)] == [
        (5, "=J2*2"),
        (5, "=J3*2"),
        (5, "=J4*2"),
    ]
    assert sheet["K4"].fill.start_color.rgb == "FFFF0000"


async def test_fill_right_repeats_the_first_column(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, start_cell="H1", rows=[[1], ["=H1+C1"]])
    await fill(call, "H1:K2", direction="right")
    sheet = data(book)
    assert [sheet.cell(1, col).value for col in range(8, 12)] == [1, 1, 1, 1]
    assert [sheet.cell(2, col).value for col in range(8, 12)] == [
        "=H1+C1",
        "=I1+D1",
        "=J1+E1",
        "=K1+F1",
    ]


@pytest.mark.parametrize(
    ("seed", "options", "expected"),
    [
        (1, {"series": "linear"}, [1, 2, 3, 4, 5]),
        (1, {"series": "linear", "step": 0.1}, [1, 1.1, 1.2, 1.3, 1.4]),
        (10, {"series": "linear", "step": -3, "stop": 0}, [10, 7, 4, 1]),
        (1, {"series": "linear", "step": 2, "stop": 7}, [1, 3, 5, 7]),
        (2, {"series": "growth", "step": 3, "stop": 100}, [2, 6, 18, 54]),
        (100, {"series": "growth", "step": 0.5}, [100, 50, 25, 12.5, 6.25]),
    ],
)
async def test_number_series(
    call: ToolCall, book: Path, seed: float, options: dict[str, Any], expected: list[float]
) -> None:
    await call("write_range", **BOOK, start_cell="H1", rows=[[seed]])
    await fill(call, "H1:H5", direction="down", **options)
    sheet = data(book)
    assert [sheet.cell(row, 8).value for row in range(1, 6)] == expected + [None] * (
        5 - len(expected)
    )


@pytest.mark.parametrize(
    ("seed", "options", "expected"),
    [
        ("2026-01-31", {"unit": "month"}, ["2026-01-31", "2026-02-28", "2026-03-31", "2026-04-30"]),
        (
            "2026-01-31",
            {"unit": "month", "step": -1},
            ["2026-01-31", "2025-12-31", "2025-11-30", "2025-10-31"],
        ),
        (
            "2026-01-31",
            {"unit": "day", "step": 7},
            ["2026-01-31", "2026-02-07", "2026-02-14", "2026-02-21"],
        ),
        ("2024-02-29", {"unit": "year"}, ["2024-02-29", "2025-02-28", "2026-02-28", "2027-02-28"]),
        (
            "2026-02-26",
            {"unit": "weekday"},
            ["2026-02-26", "2026-02-27", "2026-03-02", "2026-03-03"],
        ),
        (
            "2026-01-31",
            {"unit": "month", "stop": "2026-03-31"},
            ["2026-01-31", "2026-02-28", "2026-03-31", None],
        ),
    ],
)
async def test_date_series(
    call: ToolCall, book: Path, seed: str, options: dict[str, Any], expected: list[str | None]
) -> None:
    await call("write_range", **BOOK, start_cell="H1", rows=[[seed]])
    await fill(call, "H1:H4", direction="down", series="date", **options)
    sheet = data(book)
    found = [sheet.cell(row, 8).value for row in range(1, 5)]
    assert [v.date().isoformat() if isinstance(v, dt.datetime) else v for v in found] == expected
    assert sheet["H1"].number_format == "yyyy-mm-dd"


async def test_series_copy_the_seed_formatting_and_run_per_seed_cell(
    call: ToolCall, book: Path
) -> None:
    workbook = load_workbook(book)
    sheet = workbook["Data"]
    sheet["H1"], sheet["I1"] = 1, 100
    sheet["H1"].fill = RED
    sheet["H1"].number_format = "0.00"
    sheet["H2"].number_format = "0%"
    workbook.save(book)
    await fill(call, "H1:I3", direction="down", series="linear", step=5)
    sheet = data(book)
    assert [sheet.cell(row, 8).value for row in (1, 2, 3)] == [1, 6, 11]
    assert [sheet.cell(row, 9).value for row in (1, 2, 3)] == [100, 105, 110]
    assert sheet["H2"].number_format == "0.00"
    assert sheet["H3"].fill.start_color.rgb == "FFFF0000"


async def test_series_right_and_single_cell_with_stop(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, start_cell="H1", rows=[[10]])
    assert (
        await fill(call, "H1", direction="right", series="linear", step=2, stop=17)
        == "Filled H1:K1."
    )
    assert [data(book).cell(1, col).value for col in range(8, 13)] == [10, 12, 14, 16, None]


@pytest.mark.parametrize(
    ("range", "options", "message"),
    [
        ("H2:H5", {"direction": "down", "series": "linear"}, "number cell"),
        ("H1", {"direction": "down", "series": "linear"}, "needs a range of cells or a stop"),
        ("H1:H1", {"direction": "down"}, "at least two rows"),
        ("H1:H3", {"series": "linear"}, "needs a direction"),
        ("H1:H3", {"direction": "down", "series": "date"}, "date cell"),
        ("H1:H3", {"direction": "down", "series": "growth", "step": 0}, "other than 0"),
        ("H1:H3", {"direction": "down", "series": "linear", "stop": "x"}, "stop is a number"),
        ("H1:J1", {"direction": "right", "series": "date", "stop": "soon"}, "stop is a date"),
        ("H1:H3", {"direction": "down", "text_qualifier": "'"}, "does not use text_qualifier"),
    ],
)
async def test_invalid_fills_are_rejected(
    call: ToolCall,
    call_error: ToolCall,
    book: Path,
    range: str,
    options: dict[str, Any],
    message: str,
) -> None:
    await call("write_range", **BOOK, start_cell="H1", rows=[[1]])
    before = book.read_bytes()
    result = await call_error(
        "transform_range", **BOOK, range=range, transform={"operation": "fill", **options}
    )
    assert message in result
    assert book.read_bytes() == before
