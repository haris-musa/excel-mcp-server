"""A workbook may create only so many cells in memory, however small its file is."""

import io
import zipfile
from pathlib import Path
from typing import cast

import pytest
from mcp import Client
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from excel_mcp import cell_budget, lazy_edit
from excel_mcp.cell_budget import CellBudget
from excel_mcp.config import Limits, Settings
from excel_mcp.errors import LimitExceededError
from excel_mcp.lazy_source import split_sheet
from excel_mcp.lazy_workbook import LazyCells
from excel_mcp.server import create_server
from tests.conftest import error_text, lazily

pytestmark = pytest.mark.anyio

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ROWS, COLUMNS = 40, 10
LIMIT = 100  # Below the 400 cells and 40 rows of the sheet, above any one block of it.

_PACKAGE = "http://schemas.openxmlformats.org/package/2006"
_OFFICE = "http://schemas.openxmlformats.org/officeDocument/2006"
_SHEET_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml"
_CONTENT_TYPES = (
    f'<Types xmlns="{_PACKAGE}/content-types">'
    '<Default Extension="rels"'
    ' ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    f'<Override PartName="/xl/workbook.xml" ContentType="{_SHEET_TYPE}.sheet.main+xml"/>'
    f'<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{_SHEET_TYPE}.worksheet+xml"/>'
    "</Types>"
)
_RELS = (
    f'<Relationships xmlns="{_PACKAGE}/relationships">'
    f'<Relationship Id="rId1" Type="{_OFFICE}/relationships/officeDocument"'
    ' Target="xl/workbook.xml"/>'
    "</Relationships>"
)
_WORKBOOK = (
    f'<workbook xmlns="{MAIN}" xmlns:r="{_OFFICE}/relationships">'
    '<sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>'
)
_WORKBOOK_RELS = (
    f'<Relationships xmlns="{_PACKAGE}/relationships">'
    f'<Relationship Id="rId1" Type="{_OFFICE}/relationships/worksheet"'
    ' Target="worksheets/sheet1.xml"/>'
    "</Relationships>"
)


def write_workbook(path: Path, sheet_xml: str, part: str = "xl/worksheets/sheet1.xml") -> None:
    relative = part.removeprefix("xl/")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml", _CONTENT_TYPES.replace("xl/worksheets/sheet1.xml", part)
        )
        archive.writestr("_rels/.rels", _RELS)
        archive.writestr("xl/workbook.xml", _WORKBOOK)
        archive.writestr(
            "xl/_rels/workbook.xml.rels", _WORKBOOK_RELS.replace("worksheets/sheet1.xml", relative)
        )
        archive.writestr(part, sheet_xml)


def sheet_xml(*, addressed: bool, tail: str = "") -> str:
    """Rows of numbers; without the addresses, a file is a fraction of the size."""
    rows = []
    for row in range(1, ROWS + 1):
        cells = "".join(
            f'<c r="{chr(64 + column)}{row}"><v>{row * column}</v></c>'
            if addressed
            else f"<c><v>{row * column}</v></c>"
            for column in range(1, COLUMNS + 1)
        )
        rows.append(f'<row r="{row}">{cells}</row>' if addressed else f"<row>{cells}</row>")
    return f'<worksheet xmlns="{MAIN}"><sheetData>{"".join(rows)}</sheetData>{tail}</worksheet>'


@pytest.fixture(params=[True, False], ids=["addressed", "unaddressed"])
def big(request: pytest.FixtureRequest, files: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A sheet that is parsed whole, as every sheet of this size is."""
    monkeypatch.setattr(lazy_edit, "LAZY_SHEET_BYTES", 1 << 30)
    path = files / "big.xlsx"
    write_workbook(path, sheet_xml(addressed=request.param))
    return path


async def call(files: Path, tool: str, limit: int, **arguments: object) -> tuple[bool, str]:
    settings = Settings(allowed_dirs=[files], limits=Limits(max_loaded_cells=limit))
    async with Client(create_server(settings)) as client:
        result = await client.call_tool(tool, arguments)
    return result.is_error, error_text(result)


@pytest.mark.parametrize(
    ("tool", "arguments"),
    [
        ("describe_sheet", {}),
        ("write_range", {"at": "A1", "rows": [[1]]}),
        ("clear_range", {"range": "A1"}),
        ("sort_range", {"range": "A1:B5", "sort_by": [{"column": "A"}], "has_header": False}),
        ("copy_sheet", {"new_name": "Copy"}),
        ("set_note", {"cell": "A1", "text": "x"}),
    ],
)
async def test_tools_that_parse_the_whole_sheet_refuse_one_over_the_limit(
    big: Path, files: Path, tool: str, arguments: dict[str, object]
) -> None:
    before = big.read_bytes()
    failed, message = await call(files, tool, LIMIT, path="big.xlsx", sheet="Data", **arguments)
    assert failed
    assert f"more than {LIMIT:,} cells and rows" in message
    assert "--max-loaded-cells" in message
    assert big.read_bytes() == before


async def test_a_sheet_within_the_limit_is_opened(big: Path, files: Path) -> None:
    failed, message = await call(
        files, "describe_sheet", ROWS * (COLUMNS + 1), path="big.xlsx", sheet="Data"
    )
    assert not failed, message


async def test_streaming_reads_do_not_need_the_cells_in_memory(big: Path, files: Path) -> None:
    failed, message = await call(
        files, "read_range", LIMIT, path="big.xlsx", sheet="Data", range="A1:B2"
    )
    assert not failed, message


@pytest.mark.parametrize(
    "tail",
    [
        '<mergeCells count="1"><mergeCell ref="A1:Z1000"/></mergeCells>',
        '<hyperlinks><hyperlink ref="A1:Z1000" location="Data!A1"/></hyperlinks>',
    ],
    ids=["merged", "hyperlinked"],
)
async def test_one_element_that_names_a_big_range_is_counted_as_its_cells(
    files: Path, tail: str
) -> None:
    write_workbook(
        files / "range.xlsx", f'<worksheet xmlns="{MAIN}"><sheetData/>{tail}</worksheet>'
    )
    failed, message = await call(files, "describe_sheet", 1000, path="range.xlsx", sheet="Data")
    assert failed
    assert "more than 1,000 cells and rows" in message


async def test_a_sheet_is_counted_wherever_the_package_keeps_it(files: Path) -> None:
    write_workbook(files / "odd.xlsx", sheet_xml(addressed=False), part="xl/elsewhere/data.xml")
    failed, message = await call(files, "describe_sheet", LIMIT, path="odd.xlsx", sheet="Data")
    assert failed
    assert "--max-loaded-cells" in message


async def test_a_sheet_that_cannot_be_split_is_counted_before_it_is_parsed(files: Path) -> None:
    # Rows out of order cannot be taken apart, so a lazy read has to parse the whole sheet.
    rows = "".join(f'<row r="{row}"><c r="A{row}"><v>1</v></c></row>' for row in (3, 2, 1))
    rows = rows.replace('<v>1</v></c></row><row r="2"', '<f>A2</f></c></row><row r="2"', 1)
    write_workbook(
        files / "order.xlsx", f'<worksheet xmlns="{MAIN}"><sheetData>{rows}</sheetData></worksheet>'
    )
    arguments = {"path": "order.xlsx", "sheet": "Data", "range": "A3", "direction": "precedents"}
    failed, message = await call(files, "trace_formula", 4, **arguments)
    assert failed
    assert "more than 4 cells and rows" in message
    failed, message = await call(files, "trace_formula", 12, **arguments)  # Two copies are read.
    assert not failed, message


async def test_an_edit_that_parses_only_some_rows_needs_only_those_cells(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    write_workbook(files / "big.xlsx", sheet_xml(addressed=False))
    with lazily(monkeypatch):
        failed, message = await call(
            files, "write_range", LIMIT, path="big.xlsx", sheet="Data", at="B2", rows=[["x"]]
        )
        assert not failed, message
        failed, message = await call(
            files, "clear_range", LIMIT, path="big.xlsx", sheet="Data", range="A1:J40"
        )
        assert failed
        assert "--max-loaded-cells" in message
    sheet = load_workbook(files / "big.xlsx")["Data"]
    assert sheet["B2"].value == "x"
    assert sheet["A2"].value == 2
    assert sheet["J40"].value == 400


async def test_a_sheet_without_addresses_gets_them_when_it_is_edited_lazily(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    write_workbook(files / "big.xlsx", sheet_xml(addressed=False))
    expected = {
        (row, column): row * column for row in range(1, ROWS + 1) for column in range(1, 11)
    }
    with lazily(monkeypatch):
        failed, message = await call(
            files,
            "format_range",
            10_000,
            path="big.xlsx",
            sheet="Data",
            range="A1",
            style={"bold": True},
        )
    assert not failed, message
    sheet = load_workbook(files / "big.xlsx")["Data"]
    assert {(c.row, c.column): c.value for row in sheet.iter_rows() for c in row} == expected
    assert sheet["A1"].font.bold


def test_cells_and_rows_are_numbered_by_position_when_the_file_gives_no_number() -> None:
    xml = (
        f'<worksheet xmlns="{MAIN}"><sheetData>'
        "<row><c><v>1</v></c><c><v>2</v></c></row>"
        '<row r="5"><c r="C5"><v>3</v></c><c><v>4</v></c><c s="1"/></row>'
        "<row/>"
        '<row><c t="s"><v>0</v></c></row>'
        "</sheetData></worksheet>"
    ).encode()
    split = split_sheet(xml)
    assert split
    _, source = split
    assert source.row_numbers == [1, 5, 6, 7]
    rows = [source.data[start:end].decode() for start, end in source.spans]
    assert rows[0] == '<row r="1"><c r="A1"><v>1</v></c><c r="B1"><v>2</v></c></row>'
    assert rows[1] == (
        '<row r="5"><c r="C5"><v>3</v></c><c r="D5"><v>4</v></c><c r="E5" s="1"/></row>'
    )
    assert rows[2] == '<row r="6"/>'
    assert rows[3] == '<row r="7"><c r="A7" t="s"><v>0</v></c></row>'


def test_a_sheet_is_counted_with_the_namespace_prefix_it_uses() -> None:
    xml = (
        f'<x:worksheet xmlns:x="{MAIN}"><x:sheetData><x:row><x:c/><x:c/></x:row></x:sheetData>'
        "</x:worksheet>"
    ).encode()
    budget = CellBudget(100)
    cell_budget.spend_sheet(io.BytesIO(xml), budget)
    assert budget.used == 3


def test_the_budget_is_shared_by_all_sheets_of_a_workbook() -> None:
    budget = CellBudget(5)
    sheet = f'<worksheet xmlns="{MAIN}"><sheetData><row><c/><c/></row></sheetData></worksheet>'
    cell_budget.spend_sheet(io.BytesIO(sheet.encode()), budget)
    with pytest.raises(LimitExceededError):
        cell_budget.spend_sheet(io.BytesIO(sheet.encode()), budget)


def test_a_lazy_load_that_parses_every_row_is_refused(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    write_workbook(files / "big.xlsx", sheet_xml(addressed=True))
    with lazily(monkeypatch):
        workbook, _ = lazy_edit.load(files / "big.xlsx", CellBudget(LIMIT))
        cells = cast(LazyCells, workbook["Data"]._cells)
        cells.load(1)
        with pytest.raises(LimitExceededError):
            cells.load_all()


def test_a_cell_with_a_link_is_counted_as_several() -> None:
    xml = f'<worksheet xmlns="{MAIN}"><hyperlinks><hyperlink ref="A1:B2"/></hyperlinks></worksheet>'
    budget = CellBudget(100)
    cell_budget.spend_sheet(io.BytesIO(xml.encode()), budget)
    assert budget.used == 12


async def test_describe_sheet_of_a_large_sheet_does_not_parse_all_its_cells(
    files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    for row in range(1, ROWS + 1):
        sheet.append([row * column for column in range(1, COLUMNS + 1)])
    sheet["B3"].comment = Comment("check", "Ann")
    sheet["C30"].hyperlink = "https://example.com/"
    sheet.merge_cells("A38:B39")
    sheet.row_dimensions[20].hidden = True
    sheet["Z2"] = "far"
    workbook.save(files / "notes.xlsx")
    arguments = {"path": "notes.xlsx", "sheet": "Data"}
    failed, whole = await call(files, "describe_sheet", 10_000, **arguments)
    assert not failed, whole
    with lazily(monkeypatch):
        failed, lazy = await call(files, "describe_sheet", 200, **arguments)
    assert not failed, lazy
    assert lazy == whole
    assert '"check"' in lazy and "example.com" in lazy and "A38:B39" in lazy and "20" in lazy
    assert "Z40" in lazy


async def test_formulas_calculated_during_an_edit_share_its_budget(files: Path) -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Region", "Units", "Double"])
    for row in range(2, 31):
        sheet.append(["N" if row % 2 else "S", row, f"=B{row}*2"])
    workbook.save(files / "formulas.xlsx")
    arguments = {
        "path": "formulas.xlsx",
        "sheet": "Data",
        "at": "F1",
        "source": "Data!A1:C30",
        "row_fields": ["Region"],
        "value_fields": [{"field": "Double"}],
    }
    cells = 30 * 3 + 30
    failed, message = await call(files, "create_pivot_table", 5 * cells, **arguments)
    assert not failed, message
    failed, message = await call(files, "create_pivot_table", 2 * cells, **arguments)
    assert failed
    assert "Open the workbook in Excel and save it" in message
