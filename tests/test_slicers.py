"""Slicers and timelines: what is written, what they filter, and how they are listed."""

import re
from collections.abc import Callable, Coroutine
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from tests.package_support import relationships, sheet_part, text
from tests.slicer_support import (
    PIVOT,
    TABLE,
    ToolCall,
    add,
    build_book,
    named,
    parts_of,
    pivot_cells,
)

pytestmark = pytest.mark.anyio


@pytest.fixture
async def book(call: ToolCall, files: Path) -> str:
    return await build_book(call, files)


async def test_a_pivot_slicer_is_written_as_excel_writes_it(call: ToolCall, book: str) -> None:
    await add(call, book, field="Product", columns=2, style="SlicerStyleDark2", caption="What")
    parts = parts_of(book)
    cache = named(parts, "xl/slicerCaches/")[0]
    assert 'name="Slicer_Product" sourceName="Product"' in cache
    assert '<pivotTable tabId="2" name="PivotSales"/>' in cache
    assert '<items count="3"><i x="0" s="1"/><i x="2" s="1"/><i x="1" s="1"/></items>' in cache
    slicer = named(parts, "xl/slicers/")[0]
    assert (
        '<slicer name="Product" cache="Slicer_Product" caption="What" columnCount="2" '
        'style="SlicerStyleDark2" rowHeight="241300"/>' in slicer
    )
    workbook = text(parts, "xl/workbook.xml")
    assert '<definedName name="Slicer_Product">#N/A</definedName>' in workbook
    assert "{BBE1A952-AA13-448e-AADC-164F8A28A991}" in workbook
    assert "{79F54976-1DA5-4618-B147-4CDE4B953A38}" in workbook
    sheet = text(parts, sheet_part(parts, "Pivot"))
    assert "{A8765BA9-456A-4dab-B4F3-ACF838C121DE}" in sheet
    kinds = {kind for kind, _ in relationships(parts, sheet_part(parts, "Pivot")).values()}
    assert {"slicer", "drawing", "pivotTable"} <= kinds
    drawing = next(text(parts, n) for n in parts if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n))
    assert 'Requires="a14"' in drawing and 'name="Product"' in drawing
    assert "<xdr:from><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row>" in drawing
    # Excel keeps the id of the pivot cache in the cache, and the slicer cache names it
    definition = text(parts, "xl/pivotCache/pivotCacheDefinition1.xml")
    cache_id = re.search(r'pivotCacheId="(\d+)"', cache)
    assert cache_id and f'x14:pivotCacheDefinition pivotCacheId="{cache_id[1]}"' in definition


async def test_selected_items_hide_the_others_in_the_pivot_table(call: ToolCall, book: str) -> None:
    await add(call, book, field="Region", selected_items=["North", "South"], sort="descending")
    assert pivot_cells(book)[:4] == [
        ["Product", "Sum of Amount"],
        ["Apples", 30],
        ["Pears", 70],
        ["Grand Total", 100],
    ]
    parts = parts_of(book)
    cache = named(parts, "xl/slicerCaches/")[0]
    # items in the order the slicer lists them (descending); unselected ones have no s
    assert '<items count="4"><i x="3"/><i x="1" s="1"/><i x="0" s="1"/><i x="2"/></items>' in cache
    assert 'sortOrder="descending"' in cache
    pivot = text(parts, "xl/pivotTables/pivotTable1.xml")
    assert pivot.count('h="1"') == 2
    info = await call("describe_sheet", path=book, sheet="Pivot")
    assert info["slicers"] == [
        {
            "name": "Region",
            "kind": "pivot",
            "caption": "Region",
            "target": "Pivot!PivotSales",
            "field": "Region",
            "range": "E3:G16",
            "selected_items": ["North", "South"],
        }
    ]


async def test_slicers_on_a_row_field_and_on_another_field_combine(
    call: ToolCall, book: str
) -> None:
    await add(call, book, field="Product", selected_items=["Apples", "Pears"])
    await add(call, book, field="Region", at="E20", selected_items=["North"])
    assert pivot_cells(book)[:4] == [
        ["Product", "Sum of Amount"],
        ["Apples", 10],
        ["Pears", 30],
        ["Grand Total", 40],
    ]
    info = await call("describe_sheet", path=book, sheet="Pivot")
    assert [s["name"] for s in info["slicers"]] == ["Product", "Region"]
    assert info["pivot_tables"][0]["range"] == "A3:B6"


async def test_a_timeline_limits_the_period(call: ToolCall, book: str) -> None:
    name = await add(
        call,
        book,
        field="Date",
        timeline={"level": "quarters", "start": "2025-02-01", "end": "2025-05-31"},
    )
    assert name == {"sheet": "Pivot", "range": "E3:J10", "name": "Date"}
    assert pivot_cells(book)[:5] == [
        ["Product", "Sum of Amount"],
        ["Apples", 20],
        ["Kiwis", 50],
        ["Pears", 70],
        ["Grand Total", 140],
    ]
    parts = parts_of(book)
    cache = named(parts, "xl/timelineCaches/")[0]
    assert 'filterType="dateBetween"' in cache
    assert '<selection startDate="2025-02-01T00:00:00" endDate="2025-05-31T00:00:00"/>' in cache
    assert '<bounds startDate="2025-01-01T00:00:00" endDate="2026-01-01T00:00:00"/>' in cache
    timeline = named(parts, "xl/timelines/")[0]
    assert 'name="Date" cache="NativeTimeline_Date" caption="Date" level="1"' in timeline
    pivot = text(parts, "xl/pivotTables/pivotTable1.xml")
    assert 'type="dateBetween"' in pivot and 'operator="greaterThanOrEqual"' in pivot
    drawing = next(text(parts, n) for n in parts if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n))
    assert 'Requires="tsle"' in drawing
    info = await call("describe_sheet", path=book, sheet="Pivot")
    assert info["slicers"][0] | {"range": None} == {
        "name": "Date",
        "kind": "timeline",
        "caption": "Date",
        "target": "Pivot!PivotSales",
        "field": "Date",
        "range": None,
        "level": "quarters",
        "start": "2025-02-01",
        "end": "2025-05-31",
    }


async def test_a_timeline_without_a_period_shows_everything(call: ToolCall, book: str) -> None:
    await add(call, book, field="Date", timeline={"level": "years"}, style="TimeSlicerStyleDark2")
    assert pivot_cells(book)[4][1] == 360
    cache = named(parts_of(book), "xl/timelineCaches/")[0]
    assert 'filterType="unknown"' in cache and "<selection" not in cache


async def test_a_table_slicer_filters_the_table(call: ToolCall, book: str) -> None:
    await add(
        call, book, sheet="Data", target=TABLE, field="Product", at="G2", selected_items=["Pears"]
    )
    sheet = load_workbook(book)["Data"]
    assert [r for r in range(2, 10) if sheet.row_dimensions[r].hidden] == [2, 3, 6, 7, 8]
    table = text(parts_of(book), "xl/tables/table1.xml")
    assert re.search(r'<filterColumn colId="1"[^>]*>\s*<filters>\s*<filter val="Pears"', table)
    parts = parts_of(book)
    cache = named(parts, "xl/slicerCaches/")[0]
    assert '<x15:tableSlicerCache tableId="1" column="2"/>' in cache
    assert "{3A4CF648-6AED-40f4-86FF-DC5316D8AED3}" in text(parts, sheet_part(parts, "Data"))
    drawing = next(text(parts, n) for n in parts if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n))
    assert 'Requires="sle15"' in drawing and 'editAs="absolute"' in drawing
    info = await call("describe_sheet", path=book, sheet="Data")
    assert info["slicers"][0]["selected_items"] == ["Pears"]
    assert info["slicers"][0]["kind"] == "table"


async def test_options_of_the_cache_follow_excel(call: ToolCall, book: str) -> None:
    await add(call, book, field="Product", sort="descending", hide_empty_items=True)
    await add(
        call, book, sheet="Data", target=TABLE, field="Region", sort="descending",
        hide_empty_items=True,
    )  # fmt: skip
    caches = named(parts_of(book), "xl/slicerCaches/")
    assert 'sortOrder="descending"' in caches[0] and "slicerCacheHideItemsWithNoData" in caches[0]
    assert '<x15:tableSlicerCache tableId="1" column="1" sortOrder="descending"/>' in caches[1]
    assert "slicerCacheHideItemsWithNoData" in caches[1]


async def test_names_are_made_unique_as_excel_does(call: ToolCall, book: str) -> None:
    await add(call, book, field="Product")
    second = await add(call, book, sheet="Data", target=TABLE, field="Product", at="G2")
    assert second["name"] == "Product 1"
    workbook = text(parts_of(book), "xl/workbook.xml")
    assert "Slicer_Product1" in workbook


@pytest.mark.parametrize(
    ("arguments", "message"),
    [
        ({"field": "Nope"}, "Field 'Nope' not found. Fields: 'Region', 'Product'"),
        (
            {"field": "Region", "selected_items": ["Mars"]},
            "Field 'Region' has no item 'Mars'. Items: ['East', 'North', 'South', 'West']",
        ),
        ({"field": "Product", "style": "Fancy"}, "style: string should match pattern"),
        (
            {"field": "Date", "timeline": {"level": "days"}, "style": "SlicerStyleDark1"},
            "must be a built-in timeline style",
        ),
        (
            {"field": "Date", "timeline": {"start": "2025-01-01"}},
            "Give both timeline.start and timeline.end",
        ),
        (
            {"field": "Date", "timeline": {"start": "2025-03-01", "end": "2025-01-01"}},
            "timeline.start is after timeline.end",
        ),
        (
            {"field": "Date", "timeline": {"start": "2030-01-01", "end": "2030-02-01"}},
            "leave no data",
        ),
        (
            {"field": "Product", "timeline": {"level": "days"}, "columns": 2},
            "A timeline takes no selected_items, columns, sort or hide_empty_items",
        ),
        (
            {"field": "Region", "connect": [{"sheet": "Pivot", "name": "PivotSales"}]},
            "is the PivotTable the slicer is for",
        ),
        (
            {"field": "Product", "target": {"sheet": "Pivot", "name": "Missing"}},
            "no table or Pivot",
        ),
        ({"field": "Product", "target": {"sheet": "Nope", "name": "x"}}, "Nope"),
        ({"field": "Product", "at": "A0"}, "Row 0 is not valid"),
    ],
)
async def test_invalid_requests_are_errors(
    call: ToolCall, call_error: Callable[..., Coroutine[Any, Any, str]], book: str,
    arguments: dict[str, Any], message: str,
) -> None:  # fmt: skip
    error = await call_error("add_slicer", path=book, **{**_defaults(), **arguments})
    assert message in error


def _defaults() -> dict[str, Any]:
    return {"sheet": "Pivot", "target": PIVOT, "at": "E3"}


async def test_a_field_has_one_slicer_and_one_timeline(
    call: ToolCall, call_error: Callable[..., Coroutine[Any, Any, str]], book: str
) -> None:
    await add(call, book, field="Product")
    error = await call_error("add_slicer", path=book, **_defaults(), field="Product")
    assert "already has the pivot 'Slicer_Product'" in error
    await add(call, book, field="Date", timeline={})
    error = await call_error("add_slicer", path=book, **_defaults(), field="Date", timeline={})
    assert "already has the timeline" in error


async def test_table_requests_that_do_not_apply_are_errors(
    call: ToolCall, call_error: Callable[..., Coroutine[Any, Any, str]], book: str
) -> None:
    error = await call_error(
        "add_slicer", path=book, sheet="Data", target=TABLE, field="Date", at="G2", timeline={}
    )
    assert "tables have no timelines" in error
    error = await call_error(
        "add_slicer", path=book, sheet="Data", target=TABLE, field="Nope", at="G2"
    )
    assert "Table 'Sales' has no column 'Nope'" in error
    await add(call, book, sheet="Data", target=TABLE, field="Region", selected_items=["East"])
    error = await call_error(
        "add_slicer", path=book, sheet="Data", target=TABLE, field="Region", at="G20"
    )
    assert "already has the slicer cache 'Slicer_Region'" in error


async def test_slicers_stay_when_the_workbook_is_edited(call: ToolCall, book: str) -> None:
    await add(call, book, field="Region", selected_items=["East"])
    await call("write_range", path=book, sheet="Data", at="F1", rows=[["note"]])
    await call("create_sheet", path=book, new_name="Other")
    info = await call("describe_sheet", path=book, sheet="Pivot")
    assert info["slicers"][0]["selected_items"] == ["East"]
    assert pivot_cells(book)[3][1] == 120


async def test_the_path_must_be_inside_the_workbook_folder(
    call_error: Callable[..., Coroutine[Any, Any, str]],
) -> None:
    error = await call_error(
        "add_slicer",
        path="../outside.xlsx",
        sheet="Pivot",
        target=PIVOT,
        field="Product",
        at="E3",
    )
    assert "outside" in error.lower() or "not allowed" in error.lower() or "folder" in error.lower()
