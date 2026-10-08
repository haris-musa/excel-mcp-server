from base64 import b64encode
from hashlib import sha512
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.header_footer import HeaderFooterItem
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image

from excel_mcp.operations import images
from excel_mcp.workspace import get_sheet
from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


def layout_args(**layout: object) -> dict[str, object]:
    return {"path": "sales.xlsx", "sheet": "Data", "layout": layout}


def open_sheet(path: Path, name: str) -> Worksheet:
    return get_sheet(load_workbook(path), name)


def fits_to_page(sheet: Worksheet) -> bool | None:
    properties = sheet.sheet_properties.pageSetUpPr
    assert properties is not None
    return properties.fitToPage


def header_part(item: HeaderFooterItem | None) -> HeaderFooterItem:
    assert item is not None
    return item


def image_count(sheet: Worksheet) -> int:
    return len(sheet._images)  # pyright: ignore[reportAttributeAccessIssue]


@pytest.fixture
def picture(files: Path) -> Path:
    path = files / "logo.png"
    Image.new("RGB", (200, 100), "red").save(path)
    return path


# Hidden and grouped lines


async def test_hide_show_and_group_lines(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        **layout_args(
            rows=[{"span": "2:3", "action": "hide"}, {"span": "5", "action": "group"}],
            columns=[{"span": "b:c", "action": "hide"}, {"span": "D", "action": "group"}],
        ),
    )
    sheet = open_sheet(sample, "Data")
    assert [sheet.row_dimensions[row].hidden for row in (1, 2, 3, 4)] == [False, True, True, False]
    assert sheet.row_dimensions[5].outlineLevel == 1
    assert sheet.sheet_format.outlineLevelRow == 1
    assert sheet.column_dimensions["B"].hidden and sheet.column_dimensions["C"].hidden
    assert sheet.column_dimensions["D"].outlineLevel == 1
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert details["hidden_rows"] == ["2:3"]
    assert details["hidden_columns"] == ["B:C"]

    await call(
        "set_sheet_layout",
        **layout_args(
            rows=[{"span": "2:3", "action": "show"}, {"span": "5", "action": "ungroup"}],
            columns=[{"span": "B:C", "action": "show"}],
        ),
    )
    sheet = open_sheet(sample, "Data")
    assert not sheet.row_dimensions[2].hidden
    assert sheet.row_dimensions[5].outlineLevel == 0
    assert not sheet.column_dimensions["B"].hidden


async def test_changing_one_column_of_a_shared_definition(call: ToolCall, files: Path) -> None:
    workbook = Workbook()
    columns = workbook.worksheets[0].column_dimensions
    columns.group("A", "D", hidden=False)
    columns["A"].width = 12
    workbook.save(files / "wide.xlsx")
    await call(
        "set_sheet_layout",
        path="wide.xlsx",
        sheet="Sheet",
        layout={
            "column_widths": {"B": 30},
            "columns": [{"span": "C", "action": "hide"}],
        },
    )
    dimensions = open_sheet(files / "wide.xlsx", "Sheet").column_dimensions
    assert {letter: (dim.min, dim.max) for letter, dim in dimensions.items()} == {
        "A": (1, 1),
        "B": (2, 2),
        "C": (3, 3),
        "D": (4, 4),
    }
    assert dimensions["B"].width == 30
    assert dimensions["A"].width == 12
    assert dimensions["C"].hidden and not dimensions["D"].hidden


@pytest.mark.parametrize("span", ["", "x1", "0", "1:20000", "2:", "B:D"])
async def test_invalid_row_span(call_error: ToolCall, sample: Path, span: str) -> None:
    message = await call_error(
        "set_sheet_layout", **layout_args(rows=[{"span": span, "action": "hide"}])
    )
    assert "span" in message


async def test_invalid_column_span(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "set_sheet_layout", **layout_args(columns=[{"span": "3", "action": "hide"}])
    )
    assert "Invalid columns span" in message


# Sheet visibility


async def test_hide_sheet_moves_the_active_tab(call: ToolCall, sample: Path) -> None:
    await call("set_sheet_layout", **layout_args(visibility="hidden"))
    workbook = load_workbook(sample)
    assert workbook["Data"].sheet_state == "hidden"
    assert workbook.active is workbook["Report"]
    assert not workbook["Data"].sheet_view.tabSelected
    await call("set_sheet_layout", **layout_args(visibility="visible"))
    assert open_sheet(sample, "Data").sheet_state == "visible"


async def test_last_visible_sheet_cannot_be_hidden(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await call("set_sheet_layout", **layout_args(visibility="hidden"))
    message = await call_error(
        "set_sheet_layout", path="sales.xlsx", sheet="Report", layout={"visibility": "hidden"}
    )
    assert "at least one visible" in message
    assert open_sheet(sample, "Report").sheet_state == "visible"


# Print setup


async def test_print_setup(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        **layout_args(
            print_setup={
                "orientation": "landscape",
                "paper_size": "a4",
                "fit_to_pages": {"wide": 1, "tall": 0},
                "margins_cm": {"left": 2.54, "top": 1.27},
                "print_area": "a1:d5",
                "title_rows": "1",
                "title_columns": "A:B",
                "center_horizontally": True,
                "gridlines": True,
                "header": {"center": "Sales &A"},
                "footer": {"left": "Page &P of &N", "right": "&D"},
            }
        ),
    )
    sheet = open_sheet(sample, "Data")
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.paperSize == 9
    assert fits_to_page(sheet) is True
    assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 0)
    assert (sheet.page_margins.left, sheet.page_margins.top) == (1.0, 0.5)
    assert sheet.page_margins.right == 0.75
    assert sheet.print_area == "'Data'!$A$1:$D$5"
    assert sheet.print_title_rows == "$1:$1"
    assert sheet.print_title_cols == "$A:$B"
    assert sheet.print_options.horizontalCentered is True
    assert sheet.print_options.gridLines is True
    assert header_part(sheet.oddHeader).center.text == "Sales &A"
    assert header_part(sheet.oddFooter).left.text == "Page &P of &N"
    assert header_part(sheet.oddFooter).right.text == "&D"
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert details["print_area"] == "'Data'!$A$1:$D$5"


async def test_print_setup_scale_and_clearing(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        **layout_args(
            print_setup={
                "fit_to_pages": {},
                "print_area": "A1:B2",
                "title_rows": "1:2",
                "header": {"left": "Draft"},
            }
        ),
    )
    await call(
        "set_sheet_layout",
        **layout_args(
            print_setup={
                "scale": 80,
                "print_area": "",
                "title_rows": "",
                "header": {"left": ""},
            }
        ),
    )
    sheet = open_sheet(sample, "Data")
    assert fits_to_page(sheet) is False
    assert sheet.page_setup.scale == 80
    assert not sheet.print_area
    assert sheet.print_title_rows is None
    assert not header_part(sheet.oddHeader).left.text


@pytest.mark.parametrize(
    ("setup", "expected"),
    [
        ({"scale": 50, "fit_to_pages": {}}, "either scale or fit_to_pages"),
        ({"print_area": "nonsense"}, "Invalid range"),
        ({"title_rows": "A:B"}, "Invalid rows span"),
        ({"title_columns": "1:2"}, "Invalid columns span"),
        ({"paper_size": "a0"}, "paper_size"),
        ({"margins_cm": {"left": -1}}, "margins_cm"),
    ],
)
async def test_invalid_print_setup(
    call_error: ToolCall, sample: Path, setup: dict[str, object], expected: str
) -> None:
    assert expected in await call_error("set_sheet_layout", **layout_args(print_setup=setup))


# Protection


async def test_protect_with_password_and_allowed_actions(call: ToolCall, sample: Path) -> None:
    await call(
        "set_sheet_layout",
        **layout_args(
            protection={
                "enabled": True,
                "password": "secret",
                "allow": ["select_locked_cells", "format_cells", "sort"],
            }
        ),
    )
    protection = open_sheet(sample, "Data").protection
    assert protection.sheet is True
    assert protection.password == "DAA7"
    assert (protection.selectLockedCells, protection.selectUnlockedCells) == (False, True)
    assert (protection.formatCells, protection.sort, protection.insertRows) == (False, False, True)
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert details["protected"] is True


async def test_unprotect_needs_the_password(
    call: ToolCall, call_error: ToolCall, sample: Path
) -> None:
    await call("set_sheet_layout", **layout_args(protection={"enabled": True, "password": "pw"}))
    assert "password protected" in await call_error(
        "set_sheet_layout", **layout_args(protection={"enabled": False})
    )
    assert "right password" in await call_error(
        "set_sheet_layout", **layout_args(protection={"enabled": False, "password": "nope"})
    )
    assert "Unprotect it first" in await call_error(
        "set_sheet_layout", **layout_args(protection={"enabled": True})
    )
    assert open_sheet(sample, "Data").protection.sheet is True
    await call("set_sheet_layout", **layout_args(protection={"enabled": False, "password": "pw"}))
    protection = open_sheet(sample, "Data").protection
    assert not protection.sheet
    assert protection.password is None


async def test_protect_without_password_is_removed_without_one(
    call: ToolCall, sample: Path
) -> None:
    await call("set_sheet_layout", **layout_args(protection={"enabled": True}))
    assert open_sheet(sample, "Data").protection.sheet is True
    await call("set_sheet_layout", **layout_args(protection={"enabled": False}))
    assert not open_sheet(sample, "Data").protection.sheet


# Images


async def test_insert_list_and_delete_images(call: ToolCall, sample: Path, picture: Path) -> None:
    await call("insert_image", path="sales.xlsx", sheet="Report", image_path="logo.png", cell="B2")
    await call(
        "insert_image",
        path="sales.xlsx",
        sheet="Report",
        image_path="logo.png",
        cell="E5",
        width_cm=4,
    )
    await call(
        "insert_image",
        path="sales.xlsx",
        sheet="Report",
        image_path="logo.png",
        cell="H9",
        height_cm=3,
    )
    await call(
        "insert_image",
        path="sales.xlsx",
        sheet="Report",
        image_path="logo.png",
        cell="K1",
        width_cm=2,
        height_cm=6,
    )
    listed = (await call("describe_sheet", path="sales.xlsx", sheet="Report"))["images"]
    assert [(entry["index"], entry["anchor"]) for entry in listed] == [
        (1, "B2"),
        (2, "E5"),
        (3, "H9"),
        (4, "K1"),
    ]
    sizes = [entry[key] for entry in listed for key in ("width_cm", "height_cm")]
    assert sizes == pytest.approx([5.29, 2.65, 4, 2, 6, 3, 2, 6], abs=0.02)

    await call("delete_image", path="sales.xlsx", sheet="Report", index=2)
    remaining = (await call("describe_sheet", path="sales.xlsx", sheet="Report"))["images"]
    assert [entry["anchor"] for entry in remaining] == ["B2", "H9", "K1"]
    assert image_count(open_sheet(sample, "Report")) == 3


async def test_images_survive_other_edits(call: ToolCall, sample: Path, picture: Path) -> None:
    await call("insert_image", path="sales.xlsx", sheet="Data", image_path="logo.png", cell="F1")
    await call("write_range", path="sales.xlsx", sheet="Data", start_cell="A10", rows=[["x"]])
    assert image_count(open_sheet(sample, "Data")) == 1


async def test_delete_image_out_of_range(call_error: ToolCall, sample: Path) -> None:
    message = await call_error("delete_image", path="sales.xlsx", sheet="Data", index=1)
    assert "no images" in message


async def test_insert_image_rejects_bad_files(
    call_error: ToolCall, files: Path, sample: Path, picture: Path
) -> None:
    (files / "fake.png").write_text("not an image")
    Image.new("RGB", (4, 4)).save(files / "other.png", format="GIF")
    (files / "notes.txt").write_text("hi")
    (files / "big.jpg").write_bytes(b"\xff" * (11 * 1024 * 1024))
    arguments = {"path": "sales.xlsx", "sheet": "Data", "cell": "A1"}
    assert "not a valid PNG or JPEG" in await call_error(
        "insert_image", image_path="fake.png", **arguments
    )
    assert "only PNG and JPEG" in await call_error(
        "insert_image", image_path="other.png", **arguments
    )
    assert "must point to an image" in await call_error(
        "insert_image", image_path="notes.txt", **arguments
    )
    assert "limit" in await call_error("insert_image", image_path="big.jpg", **arguments)
    assert "does not exist" in await call_error(
        "insert_image", image_path="missing.png", **arguments
    )
    assert "Invalid range" in await call_error(
        "insert_image", image_path="logo.png", **{**arguments, "cell": "nope"}
    )
    assert not image_count(open_sheet(sample, "Data"))


async def test_insert_image_rejects_too_many_pixels(
    call_error: ToolCall, sample: Path, picture: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(images, "MAX_IMAGE_PIXELS", 1000)
    message = await call_error(
        "insert_image", path="sales.xlsx", sheet="Data", image_path="logo.png", cell="A1"
    )
    assert "200x100 pixels" in message


async def test_images_stay_inside_the_allowed_folder(
    call_error: ToolCall, tmp_path: Path, sample: Path
) -> None:
    outside = tmp_path / "outside.png"
    Image.new("RGB", (4, 4)).save(outside)
    for image_path in (str(outside), "../outside.png"):
        message = await call_error(
            "insert_image", path="sales.xlsx", sheet="Data", image_path=image_path, cell="A1"
        )
        assert "outside the allowed directories" in message


async def test_unprotect_a_sheet_protected_by_excel(
    call: ToolCall, call_error: ToolCall, files: Path
) -> None:
    salt = b"0123456789abcdef"
    digest = sha512(salt + "pw123".encode("utf-16-le")).digest()
    for index in range(100):
        digest = sha512(digest + index.to_bytes(4, "little")).digest()
    workbook = Workbook()
    workbook.worksheets[0].protection = SheetProtection(
        sheet=True,
        algorithmName="SHA-512",
        hashValue=b64encode(digest).decode(),
        saltValue=b64encode(salt).decode(),
        spinCount=100,
    )
    workbook.save(files / "excel.xlsx")
    arguments = {"path": "excel.xlsx", "sheet": "Sheet"}
    wrong = {"protection": {"enabled": False, "password": "no"}}
    assert "right password" in await call_error("set_sheet_layout", layout=wrong, **arguments)
    right = {"protection": {"enabled": False, "password": "pw123"}}
    await call("set_sheet_layout", layout=right, **arguments)
    assert not open_sheet(files / "excel.xlsx", "Sheet").protection.sheet
