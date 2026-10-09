from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


async def test_create_table(call: ToolCall, call_error: ToolCall, sample: Path) -> None:
    created = await call("set_table", path="sales.xlsx", sheet="Data", range="A1:D5")
    assert created == {"sheet": "Data", "name": "Table1", "range": "A1:D5"}
    assert load_workbook(sample)["Data"].tables["Table1"].ref == "A1:D5"
    assert "already used" in await call_error(
        "set_table", path="sales.xlsx", sheet="Report", range="A1:A2", name="table1"
    )


async def test_create_table_requires_text_headers(call_error: ToolCall, sample: Path) -> None:
    message = await call_error("set_table", path="sales.xlsx", sheet="Data", range="C2:D5")
    assert "header" in message


@pytest.mark.parametrize("chart_type", ["column", "bar", "line", "area", "pie", "scatter"])
async def test_create_chart(call: ToolCall, sample: Path, chart_type: str) -> None:
    await call(
        "create_chart",
        path="sales.xlsx",
        sheet="Report",
        source="Data!B1:C5",
        chart_type=chart_type,
        at="B2",
        options={"title": "Units"},
    )
    details = await call("describe_sheet", path="sales.xlsx", sheet="Report")
    assert len(details["charts"]) == 1


async def test_chart_needs_labels_and_a_series(call_error: ToolCall, sample: Path) -> None:
    message = await call_error(
        "create_chart",
        path="sales.xlsx",
        sheet="Data",
        source="C1:C5",
        chart_type="line",
        at="F2",
    )
    assert "header row" in message
