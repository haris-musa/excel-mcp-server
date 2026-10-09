"""Cached dynamic array results follow the data, and rows can be inserted through them."""

import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from tests.conftest import ToolCall
from tests.package_support import read_parts, sheet_part, text

pytestmark = pytest.mark.anyio

BOOK = {"path": "book.xlsx"}


@pytest.fixture
def book(files: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "S"
    for row in [[1, 10], [2, 20], [3, 30], [4, 40], [5, 50]]:
        sheet.append(row)
    path = files / "book.xlsx"
    workbook.save(path)
    return path


async def _values(call: ToolCall, area: str) -> list[list[object]]:
    return (await call("read_range", **BOOK, sheet="S", range=area))["values"]


async def test_spilled_values_follow_changes_to_the_data(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, sheet="S", at="D1", rows=[["=FILTER(A1:B5,A1:A5>2)"]])
    assert await _values(call, "D1:E4") == [[3, 30], [4, 40], [5, 50]]

    await call("write_range", **BOOK, sheet="S", at="B2", rows=[[999]])
    await call("write_range", **BOOK, sheet="S", at="A5", rows=[[1]])
    assert await _values(call, "D1:E4") == [[3, 30], [4, 40]]
    sheet = load_workbook(book)["S"]
    assert sheet["D1"].value.ref == "D1:E2"  # type: ignore[union-attr]
    assert sheet["D3"].value is None  # the spill range shrank

    await call("write_range", **BOOK, sheet="S", at="A1", rows=[[9], [9], [9], [9], [9]])
    assert len(await _values(call, "D1:E6")) == 5


async def test_a_spill_that_gets_blocked_shows_spill_error(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, sheet="S", at="D1", rows=[["=FILTER(A1:B5,A1:A5>2)"]])
    await call("write_range", **BOOK, sheet="S", at="E2", rows=[["in the way"]])

    assert (await _values(call, "D1:D1")) == [["#SPILL!"]]
    assert load_workbook(book)["S"]["E2"].value == "in the way"


async def test_a_result_the_calculator_cannot_reproduce_is_dropped(
    call: ToolCall, book: Path
) -> None:
    await call("write_range", **BOOK, sheet="S", at="D1", rows=[["=SEQUENCE(3)"]])
    sheet = load_workbook(book)["S"]
    sheet["D1"] = ArrayFormula("D1:D3", '=TEXTSPLIT("a,b",",")')
    sheet.parent.save(book)  # type: ignore[union-attr]
    await call("write_range", **BOOK, sheet="S", at="G1", rows=[[1]])

    anchor = text(read_parts(book), sheet_part(read_parts(book), "S"))
    assert "<v>" not in anchor.split('r="D1"')[1].split("</c>")[0]


async def test_edits_ask_excel_to_recalculate_when_it_opens_the_file(
    call: ToolCall, book: Path
) -> None:
    await call("write_range", **BOOK, sheet="S", at="A1", rows=[[5]])
    with zipfile.ZipFile(book) as archive:
        assert 'fullCalcOnLoad="1"' in archive.read("xl/workbook.xml").decode()


async def test_rows_can_be_inserted_through_a_spill_range(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, sheet="S", at="D1", rows=[["=SEQUENCE(4)"]])
    await call("insert_rows_or_columns", **BOOK, sheet="S", axis="rows", start=3, count=2)
    assert [row[0] for row in await _values(call, "D1:D6")] == [1, 2, 3, 4]
    assert load_workbook(book)["S"]["D1"].value.ref == "D1:D4"  # type: ignore[union-attr]

    await call("delete_rows_or_columns", **BOOK, sheet="S", axis="rows", start=2, count=2)
    assert [row[0] for row in await _values(call, "D1:D6")] == [1, 2, 3, 4]
    await call("delete_rows_or_columns", **BOOK, sheet="S", axis="rows", start=1, count=1)
    assert await _values(call, "D1:D6") == []


async def test_columns_can_be_inserted_through_a_spill_range(call: ToolCall, book: Path) -> None:
    await call("write_range", **BOOK, sheet="S", at="D1", rows=[["=SEQUENCE(1,3)"]])
    await call("insert_rows_or_columns", **BOOK, sheet="S", axis="columns", start=5, count=1)
    assert await _values(call, "D1:H1") == [[1, 2, 3]]


async def test_legacy_array_formulas_still_cannot_be_split(
    call: ToolCall, call_error: ToolCall, book: Path
) -> None:
    workbook = load_workbook(book)
    workbook["S"]["D1"] = ArrayFormula("D1:D3", "=A1:A3*2")
    workbook.save(book)
    message = await call_error(
        "insert_rows_or_columns", **BOOK, sheet="S", axis="rows", start=2, count=1
    )
    assert "array formula in D1:D3" in message
