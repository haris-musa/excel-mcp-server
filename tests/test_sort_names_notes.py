from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


def column(path: Path, sheet: str, letter: str, last: int) -> list[object]:
    worksheet = load_workbook(path)[sheet]
    return [worksheet[f"{letter}{row}"].value for row in range(1, last + 1)]


@pytest.fixture
def mixed(files: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Data"
    sheet.append(["Name", "Score", "Double"])
    rows = [("pear", 5), ("Apple", None), ("banana", 2), (True, 9), ("apple", 5), (3, 1), (None, 7)]
    for index, (name, score) in enumerate(rows, start=2):
        sheet.append([name, score, f"=B{index}*2"])
    sheet["A3"].font = Font(bold=True)
    sheet["A3"].comment = Comment("first apple", "me")
    path = files / "mixed.xlsx"
    workbook.save(path)
    return path


async def test_sort_orders_like_excel(call: ToolCall, mixed: Path) -> None:
    await call(
        "sort_range",
        path="mixed.xlsx",
        sheet="Data",
        range="A1:C8",
        sort_by=[{"column": "Name"}],
    )
    # Numbers, text (ignoring case, stable among equals), booleans, then blanks last.
    assert column(mixed, "Data", "A", 8) == [
        "Name",
        3,
        "Apple",
        "apple",
        "banana",
        "pear",
        True,
        None,
    ]


async def test_sort_descending_keeps_blanks_last_and_rows_together(
    call: ToolCall, mixed: Path
) -> None:
    await call(
        "sort_range",
        path="mixed.xlsx",
        sheet="Data",
        range="A1:C8",
        sort_by=[{"column": "B", "order": "descending"}, {"column": "Name"}],
    )
    sheet = load_workbook(mixed)["Data"]
    assert [sheet[f"B{row}"].value for row in range(2, 9)] == [9, 7, 5, 5, 2, 1, None]
    assert [sheet[f"A{row}"].value for row in (4, 5)] == ["apple", "pear"]
    assert sheet["A8"].value == "Apple"


async def test_sort_moves_formulas_styles_and_notes_with_their_row(
    call: ToolCall, mixed: Path
) -> None:
    await call(
        "sort_range", path="mixed.xlsx", sheet="Data", range="A1:C8", sort_by=[{"column": "Score"}]
    )
    sheet = load_workbook(mixed)["Data"]
    assert [sheet[f"C{row}"].value for row in range(2, 5)] == ["=B2*2", "=B3*2", "=B4*2"]
    assert sheet["B2"].value == 1
    # "Apple" (blank score) is last, and its bold font and note went with it.
    assert sheet["A8"].value == "Apple" and sheet["A8"].font.bold
    assert sheet["A8"].comment is not None and sheet["A8"].comment.text == "first apple"
    assert sheet["A3"].comment is None


async def test_sort_without_header_row(call: ToolCall, mixed: Path) -> None:
    await call(
        "sort_range",
        path="mixed.xlsx",
        sheet="Data",
        range="B1:B8",
        sort_by=[{"column": "B"}],
        has_header=False,
    )
    assert column(mixed, "Data", "B", 8) == [1, 2, 5, 5, 7, 9, "Score", None]


async def test_sort_rejects_bad_requests(call_error: ToolCall, mixed: Path) -> None:
    message = await call_error(
        "sort_range", path="mixed.xlsx", sheet="Data", range="A1:C8", sort_by=[{"column": "Nope"}]
    )
    assert "Name" in message and "column letter" in message
    message = await call_error(
        "sort_range", path="mixed.xlsx", sheet="Data", range="A1:C8", sort_by=[{"column": "C"}]
    )
    assert "C2 is a formula" in message
    message = await call_error(
        "sort_range", path="mixed.xlsx", sheet="Data", range="A1:C1", sort_by=[{"column": "A"}]
    )
    assert "fewer than two rows" in message
    workbook = load_workbook(mixed)
    workbook["Data"].merge_cells("A9:B9")
    workbook.save(mixed)
    message = await call_error(
        "sort_range", path="mixed.xlsx", sheet="Data", range="A1:C9", sort_by=[{"column": "A"}]
    )
    assert "merged" in message


async def test_sort_refuses_paths_outside_the_sandbox(call_error: ToolCall) -> None:
    message = await call_error(
        "sort_range", path="../x.xlsx", sheet="Data", range="A1:B3", sort_by=[{"column": "A"}]
    )
    assert "outside" in message or "not allowed" in message


async def test_defined_names_roundtrip(call: ToolCall, sample: Path) -> None:
    await call("set_defined_name", path="sales.xlsx", name="Units", refers_to="=Data!$C$2:$C$5")
    await call("set_defined_name", path="sales.xlsx", name="Tax", refers_to="0.075")
    await call("set_defined_name", path="sales.xlsx", name="Tax", refers_to="0.1", sheet="Report")
    result = await call("set_defined_name", path="sales.xlsx", name="Tax", refers_to="0.2")
    assert "Updated" in result
    names = (await call("describe_workbook", path="sales.xlsx"))["defined_names"]
    assert names == [
        {"name": "Tax", "refers_to": "0.2", "sheet": None},
        {"name": "Units", "refers_to": "Data!$C$2:$C$5", "sheet": None},
        {"name": "Tax", "refers_to": "0.1", "sheet": "Report"},
    ]
    await call("delete_defined_name", path="sales.xlsx", name="Tax", sheet="Report")
    await call("delete_defined_name", path="sales.xlsx", name="Tax")
    workbook = load_workbook(sample)
    assert list(workbook.defined_names) == ["Units"]
    assert not workbook["Report"].defined_names


@pytest.mark.parametrize(
    ("name", "refers_to", "expected"),
    [
        ("A1", "Data!A1", "Invalid name"),
        ("my name", "Data!A1", "Invalid name"),
        ("ok", 'WEBSERVICE("http://x")', "WEBSERVICE"),
        ("ok", "[1]Sheet1!A1", "other workbooks"),
        ("ok", "Missing!A1", "not a sheet"),
        ("ok", 'INDIRECT("A1")', "INDIRECT"),
        ("ok", " ", "empty"),
    ],
)
async def test_defined_name_rejects_unsafe_input(
    call_error: ToolCall, sample: Path, name: str, refers_to: str, expected: str
) -> None:
    before = sample.read_bytes()
    message = await call_error(
        "set_defined_name", path="sales.xlsx", name=name, refers_to=refers_to
    )
    assert expected in message
    assert sample.read_bytes() == before


async def test_delete_unknown_name_is_an_error(call_error: ToolCall, sample: Path) -> None:
    message = await call_error("delete_defined_name", path="sales.xlsx", name="Nope")
    assert "No name 'Nope'" in message


async def test_notes_roundtrip(call: ToolCall, sample: Path) -> None:
    await call(
        "set_note", path="sales.xlsx", sheet="Data", cell="C2", text="Check this", author="Ann"
    )
    await call("set_note", path="sales.xlsx", sheet="Data", cell="C2", text="Checked")
    await call("set_note", path="sales.xlsx", sheet="Data", cell="A1", text="Header")
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert details["notes"] == [
        {"cell": "A1", "text": "Header", "author": "Claude"},
        {"cell": "C2", "text": "Checked", "author": "Claude"},
    ]
    saved = load_workbook(sample)["Data"]
    assert saved["C2"].comment is not None and saved["C2"].comment.text == "Checked"
    await call("delete_note", path="sales.xlsx", sheet="Data", cell="C2")
    details = await call("describe_sheet", path="sales.xlsx", sheet="Data")
    assert [note["cell"] for note in details["notes"]] == ["A1"]


async def test_note_errors(call_error: ToolCall, sample: Path) -> None:
    assert "has no note" in await call_error(
        "delete_note", path="sales.xlsx", sheet="Data", cell="B2"
    )
    assert "Sheet 'X' not found" in await call_error(
        "set_note", path="sales.xlsx", sheet="X", cell="B2", text="hi"
    )
    assert "empty" in await call_error(
        "set_note", path="sales.xlsx", sheet="Data", cell="B2", text="  "
    )


async def test_text_order_matches_excel(call: ToolCall, files: Path) -> None:
    words = [
        "x-y",
        "Zz",
        "éclair",
        "xy",
        "eclair",
        "ECLAIR",
        "x y",
        "coop",
        "#z",
        "co-op",
        "zz",
        "1a",
    ]
    workbook = Workbook()
    workbook.worksheets[0].title = "Data"
    for word in words:
        workbook.worksheets[0].append([word])
    workbook.save(files / "words.xlsx")
    await call(
        "sort_range",
        path="words.xlsx",
        sheet="Data",
        range="A1:A12",
        sort_by=[{"column": "A"}],
        has_header=False,
    )
    # The order Excel's own Range.Sort gives: case and hyphens ignored first, then accents.
    assert column(files / "words.xlsx", "Data", "A", 12) == [
        "#z",
        "1a",
        "coop",
        "co-op",
        "eclair",
        "ECLAIR",
        "éclair",
        "x y",
        "xy",
        "x-y",
        "Zz",
        "zz",
    ]
