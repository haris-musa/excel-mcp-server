# pyright: reportAttributeAccessIssue=false, reportPrivateUsage=false
# The tests look at openpyxl internals: the cells and styles a workbook holds.
"""Rows read and written as text are the rows openpyxl reads and writes."""

import datetime as dt
import random
import re
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.reader.excel import ExcelReader
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet._reader import WorkSheetParser
from openpyxl.worksheet.formula import ArrayFormula

from excel_mcp import lazy_rows
from excel_mcp.lazy_shared import Translators
from excel_mcp.lazy_source import split_sheet

NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
TEXTS = [
    "plain",
    " lead",
    "trail ",
    "a&b<c>d",
    "line\nbreak",
    "ünï",
    '"q"',
    "x" * 300,
    "=text",
    "1",
]
ROWS = 120
COLUMNS = 10


def make_corpus(path: Path) -> None:
    """Cells of every kind openpyxl writes, in a sheet with row heights, a merge, a link, a note
    and an array formula."""
    rng = random.Random(3)
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "S"
    for row in range(1, ROWS + 1):
        for column in range(1, COLUMNS + 1):
            cell, kind = sheet.cell(row, column), rng.random()
            if kind < 0.12:
                cell.value = rng.randint(-1000, 10**9)
            elif kind < 0.25:
                cell.value = rng.choice([0.1, 1e-7, 1.5e20, 3.14159265358979, -0.5, 1 / 3, 2.0])
            elif kind < 0.4:
                cell.value = rng.choice(TEXTS)
            elif kind < 0.45:
                cell.value = rng.choice([True, False])
            elif kind < 0.5:
                cell.value, cell.data_type = rng.choice(["#DIV/0!", "#N/A"]), "e"
            elif kind < 0.6:
                cell.value = dt.datetime(2020, 1, 1, 12, 30) + dt.timedelta(
                    days=rng.randint(0, 5000)
                )
            elif kind < 0.65:
                cell.value = rng.choice([dt.date(2021, 5, 17), dt.time(13, 45, 10)])
            elif kind < 0.68:
                cell.value = dt.timedelta(hours=rng.randint(1, 100))
            elif kind < 0.78:
                cell.value = f'=A{row}+{rng.randint(1, 9)}&"x & y"'
            elif kind < 0.84:
                pass
            elif kind < 0.9:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="FFFF00")
            else:
                cell.value = rng.choice(TEXTS)
                cell.number_format = rng.choice(["0.00", "General", "@", "yyyy-mm-dd"])
    sheet.row_dimensions[5].height = 30
    sheet.row_dimensions[6].hidden = True
    sheet.row_dimensions[7].outlineLevel = 1
    sheet.row_dimensions[ROWS + 10].height = 12
    sheet.merge_cells("B10:C11")
    sheet["A20"].hyperlink = "https://example.com"
    sheet["B21"].comment = Comment("note", "me")
    sheet["C22"] = ArrayFormula("C22:C23", "=SUM(A1:A5)")
    workbook.save(path)


def canonical(element: ET.Element) -> tuple[object, ...]:
    text = (element.text or "").strip() if len(element) == 0 else element.text
    return (
        element.tag,
        tuple(sorted(element.attrib.items())),
        text,
        tuple(map(canonical, element)),
    )


def test_cells_are_read_as_openpyxl_reads_them(tmp_path: Path) -> None:
    make_corpus(tmp_path / "book.xlsx")
    reader = ExcelReader(tmp_path / "book.xlsx")
    reader.read_manifest()
    reader.read_strings()
    book = load_workbook(tmp_path / "book.xlsx")
    xml = ZipFile(tmp_path / "book.xlsx").read("xl/worksheets/sheet1.xml")
    _, source = split_sheet(xml)  # type: ignore[misc]

    for data_only in (False, True):
        parser = WorkSheetParser(
            BytesIO(xml),
            reader.shared_strings,
            data_only,
            book.epoch,
            book._date_formats,
            book._timedelta_formats,
        )
        expected = {
            (item["row"], item["column"]): (item["style_id"], item["value"], item["data_type"])
            for _, cells in parser.parse()
            for item in cells
        }
        reading = lazy_rows.Reading(
            reader.shared_strings,
            Translators(source.masters),
            book._date_formats,
            book._timedelta_formats,
            book.epoch,
            data_only,
        )
        read: dict[tuple[int, int], tuple[object, ...]] = {}
        unusual = 0
        for number, (start, stop) in zip(source.row_numbers, source.spans, strict=True):
            row = lazy_rows.read(source.data[start:stop], reading)
            if row is None:
                unusual += 1
                continue
            read |= {
                (number, column): (style, value, kind) for column, style, value, kind in row[1]
            }

        assert read.items() <= expected.items()
        # Only the row of the array formula is left to openpyxl, when formulas are read.
        assert unusual == (0 if data_only else 1)
        assert len(read) >= len(expected) - COLUMNS


def test_rows_are_written_as_openpyxl_writes_them(tmp_path: Path) -> None:
    make_corpus(tmp_path / "book.xlsx")
    book = load_workbook(tmp_path / "book.xlsx")
    sheet = book.worksheets[0]
    rendered, left = lazy_rows.render(dict(sheet._cells), sheet.row_dimensions)

    saved = BytesIO()
    book.save(saved)
    xml = ZipFile(saved).read("xl/worksheets/sheet1.xml").decode()
    written = {
        int(re.search(r'\br="(\d+)"', row[0])[1]): row[0]  # type: ignore[index]
        for row in re.finditer(r"<row\b[^>]*?(?:/>|>.*?</row>)", xml, re.S)
    }
    # The rows with a link, a note or an array formula are written by openpyxl.
    assert {row for row, _ in left} == {20, 21, 22}
    assert rendered.keys() | {row for row, _ in left} == written.keys()
    for number, text in rendered.items():
        ours = ET.fromstring(text.decode().replace("<row", f'<row xmlns="{NAMESPACE}"', 1))
        theirs = ET.fromstring(written[number].replace("<row", f'<row xmlns="{NAMESPACE}"', 1))
        assert canonical(ours) == canonical(theirs), number
