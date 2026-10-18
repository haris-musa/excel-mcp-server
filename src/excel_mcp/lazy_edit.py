# pyright: reportAttributeAccessIssue=false
# openpyxl's stubs lack the private members of worksheets that this module replaces.

"""Editing large sheets without parsing or writing the rows an edit leaves alone.

Loading and saving a sheet with a million cells takes a minute, and most edits touch a
handful of them. A large sheet is loaded as `lazy_workbook` loads it, row blocks being parsed
when something asks for them, and saved by putting the rows of the blocks that were parsed
(which may have changed) together with the rows of the others, as the file has them.

A block that was not parsed is as in the file, so nothing that needs a cell may skip it. That
is why the sheets this handles hold no array or data-table formulas or cell metadata (their
cells depend on each other: arrays spill into cells that are cleared or kept with them), why
cells that carry a hyperlink, a comment or a merged range are parsed when the workbook is
opened, and why using a row's dimensions parses its block. Shared formulas, which depend on
their master cell, are looked after by `lazy_shared`.
Anything that looks at all cells parses all of them: slower, never wrong.
"""

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet._reader import WorksheetReader
from openpyxl.worksheet.dimensions import DimensionHolder
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.constants import SHARED_STRINGS

from excel_mcp import lazy_rows, lazy_shared
from excel_mcp.lazy_source import SheetSource, split_sheet
from excel_mcp.lazy_workbook import BLOCK_ROWS, LazyCells, Parts, lazy_sheets, open_lazily
from excel_mcp.package.opc import REL_BASE, Package, Rel, free_id
from excel_mcp.package.restore import MAIN, sheet_parts
from excel_mcp.package.scan import declare_in_root, root_scope

# Sheets smaller than this parse and write faster than they can be taken apart.
LAZY_SHEET_BYTES = 1_000_000

_STRINGS = "xl/sharedStrings.xml"
_SHEET_DATA = re.compile(rb"<sheetData\s*/>|<sheetData>.*</sheetData>", re.S)
_ROW = re.compile(rb"<row\b[^>]*?/>|<row\b[^>]*>.*?</row>", re.S)
_ROW_NUMBER = re.compile(rb'\br="(\d+)"')
_DIMENSION = re.compile(rb'<dimension\s+ref="[^"]*"\s*/>')
_CELL_COLUMN = re.compile(rb'<c\b[^>]*?\br="([A-Z]+)\d')
_UNADDRESSED_CELL = re.compile(rb"<c\b(?![^>]*\br=)")
# A row with attributes of its own (a height, a style, hidden), not just its number.
_ATTRIBUTED_ROW = re.compile(rb"<row\b[^>]*\s(?!r=|spans=|[\w.-]+:)[\w.-]+=[^>]*>")
# What a formula stored as its result when Excel last calculated it, which an edit makes stale.
_RESULT = re.compile(rb"(</f>|<f\b[^>]*/>)(?:<v>[^<]*</v>|<v\s*/>)")


def has_large_sheet(path: Path) -> bool:
    with ZipFile(path) as archive:
        return any(
            info.filename.startswith("xl/worksheets/") and info.file_size >= LAZY_SHEET_BYTES
            for info in archive.infolist()
        )


def _split(data: bytes) -> tuple[bytes, SheetSource] | None:
    cross_cell = (b't="array"', b't="dataTable"', b' cm="', b' vm="')
    if len(data) < LAZY_SHEET_BYTES or any(marker in data for marker in cross_cell):
        return None
    if _UNADDRESSED_CELL.search(data):  # Its column is not known until the row is read.
        return None
    split = split_sheet(data)
    return split if split and split[1].row_numbers else None


class RowDimensions(DimensionHolder):
    """The row dimensions of a lazily loaded sheet: using one parses the rows it belongs to.

    A row's dimension is kept as the file has it unless its block is parsed, so a dimension
    that is read or changed has to be one of a parsed block.
    """

    def __init__(self, sheet: Worksheet) -> None:
        super().__init__(worksheet=sheet, default_factory=sheet._add_row)
        dict.update(self, dict.items(sheet.row_dimensions))

    def _load(self, row: Any) -> None:
        cells = self.worksheet._cells
        if isinstance(cells, LazyCells) and isinstance(row, int):
            cells.load(row)

    def _load_listed(self) -> None:
        for row in dict.keys(self):
            self._load(row)

    def __getitem__(self, row: Any) -> Any:
        self._load(row)
        return super().__getitem__(row)

    def __setitem__(self, row: Any, dimension: Any) -> None:
        self._load(row)
        super().__setitem__(row, dimension)

    def __delitem__(self, row: Any) -> None:
        self._load(row)
        super().__delitem__(row)

    def __contains__(self, row: object) -> bool:
        self._load(row)
        return super().__contains__(row)

    def get(self, row: Any, default: Any = None) -> Any:
        self._load(row)
        return super().get(row, default)

    def pop(self, row: Any, *default: Any) -> Any:
        self._load(row)
        return super().pop(row, *default)

    def __iter__(self) -> Any:
        self._load_listed()
        return super().__iter__()

    def __len__(self) -> int:
        self._load_listed()
        return super().__len__()

    def keys(self) -> Any:
        self._load_listed()
        return super().keys()

    def values(self) -> Any:
        self._load_listed()
        return super().values()

    def items(self) -> Any:
        self._load_listed()
        return super().items()


@dataclass
class LazyEdit:
    """The sheets of a workbook that are loaded lazily, until they are saved."""

    path: Path
    strings_part: str | None
    frozen: dict[Worksheet, tuple[LazyCells, dict[int, bytes]]] = field(default_factory=dict)

    def freeze(self, workbook: Workbook) -> None:
        """Prepare each sheet for saving, with the rows that were parsed written as text.

        openpyxl writes what a sheet holds, which is only the rows that need it (those with a
        link, a note or a value it writes in a form of its own). The other rows of blocks that
        were parsed are written by `lazy_rows`, and the rows of blocks that were not parsed are
        put in from the file; `merge` puts them all together.
        """
        for sheet in workbook.worksheets:
            cells = sheet._cells
            if not isinstance(cells, LazyCells):
                continue
            lazy_shared.settle(cells)
            dimensions = DimensionHolder(worksheet=sheet, default_factory=sheet._add_row)
            for row, dimension in dict.items(sheet.row_dimensions):
                if cells.is_loaded(row):
                    dimensions[row] = dimension
            rendered, sheet._cells = lazy_rows.render(cells.held(), dimensions)
            kept = {row for row, _ in sheet._cells}
            sheet.row_dimensions = DimensionHolder(worksheet=sheet, default_factory=sheet._add_row)
            for row in kept & dimensions.keys():
                sheet.row_dimensions[row] = dimensions[row]
            self.frozen[sheet] = (cells, rendered)

    def merge(self, workbook: Workbook, content: bytes) -> bytes:
        """The saved workbook with the rows that were not parsed put back, as the file has them."""
        spliced = {
            sheet: cells for sheet, cells in self.frozen.items() if sheet in workbook.worksheets
        }
        if not spliced:
            return content
        package = Package(content)
        parts = sheet_parts(package)
        for sheet, (cells, rendered) in spliced.items():
            xml = _splice(package.read(parts[sheet.title]), cells, rendered)
            package.write(parts[sheet.title], xml)
        if self.strings_part:
            self._restore_strings(package)
        return package.to_bytes()

    def _restore_strings(self, package: Package) -> None:
        """Add the file's table of strings, which the rows kept from it refer to by position.

        openpyxl writes every string in the cell, so it writes no table.
        """
        with ZipFile(self.path) as original:
            table = original.read(self.strings_part or "")
        name = package.unique_name(_STRINGS)
        package.write(name, table)
        package.content_types.declare(name, SHARED_STRINGS)
        rels = package.rels(MAIN)
        rels.append(Rel(free_id(rels, ""), REL_BASE + "sharedStrings", name))


def load(path: Path) -> tuple[Workbook, LazyEdit | None]:
    """Load a workbook for editing, its large sheets without their rows."""
    parts: Parts = {}
    reader = open_lazily(path, parts, _split)
    reader.read()
    strings = reader.package.find(SHARED_STRINGS)
    edit = LazyEdit(path, strings.PartName.lstrip("/") if strings is not None else None)
    lazy = list(lazy_sheets(reader, parts))
    for sheet, source in lazy:
        _install(sheet, source, reader.shared_strings)
    return reader.wb, edit if lazy else None


def _install(sheet: Worksheet, source: SheetSource, strings: list[str]) -> None:
    carried = sheet._cells
    cells = LazyCells(sheet, source, strings)
    sheet._cells = cells
    _bind_row_dimensions(sheet, source, strings)
    for (row, column), carried_cell in carried.items():
        # Cells that merged ranges, hyperlinks and comments were bound to while the sheet had
        # no rows. Their rows are parsed so that the bound cell is the one that holds the value.
        cells.load(row)
        parsed = cells.get((row, column))
        if parsed is None or isinstance(carried_cell, MergedCell):
            cells[row, column] = carried_cell
            continue
        if carried_cell.hyperlink:
            parsed.hyperlink = carried_cell.hyperlink
        if carried_cell.comment:
            parsed.comment = carried_cell.comment
    sheet._current_row = cells.extent()[0] if cells else 0
    sheet.row_dimensions = RowDimensions(sheet)


def _bind_row_dimensions(sheet: Worksheet, source: SheetSource, strings: list[str]) -> None:
    """Give the sheet the dimensions of all the rows that have any, as loading it whole does."""
    tags_ = [_self_closed(tag.group()) for tag in _ATTRIBUTED_ROW.finditer(source.data)]
    if not tags_:
        return
    document = source.root + b"<sheetData>" + b"".join(tags_) + b"</sheetData></worksheet>"
    book: Any = sheet.parent
    reader = WorksheetReader(sheet, io.BytesIO(document), strings, book._data_only, False)
    for _ in reader.parser.parse():
        pass
    reader.bind_row_dimensions()


def _self_closed(tag: bytes) -> bytes:
    return tag if tag.endswith(b"/>") else tag[:-1] + b"/>"


def _splice(xml: bytes, cells: LazyCells, rendered: dict[int, bytes]) -> bytes:
    """The sheet openpyxl wrote, holding the rows of the parsed blocks, with the others added."""
    source = cells.source
    data = _SHEET_DATA.search(xml)
    if data is None:
        raise ValueError("A saved worksheet has no sheetData.")
    produced = {_row_number(row[0]): row[0] for row in _ROW.finditer(data[0])}
    produced.update(rendered)
    lazy_shared.restore(cells, produced)
    pieces = _rows(cells, produced)
    head, tail = xml[: data.start()], xml[data.end() :]
    head = _DIMENSION.sub(b'<dimension ref="%b"/>' % _dimension(source, produced), head, 1)
    declarations, ignorable = root_scope(source.root.decode("utf-8"))
    head = declare_in_root(head.decode("utf-8"), declarations, ignorable).encode("utf-8")
    return head + b"<sheetData>" + b"".join(pieces) + b"</sheetData>" + tail


def _row_number(row: bytes) -> int:
    number = _ROW_NUMBER.search(row[: row.index(b">")])
    assert number is not None
    return int(number[1])


def _rows(cells: LazyCells, produced: dict[int, bytes]) -> list[bytes]:
    """The rows in order: those written for parsed blocks, and the file's rows for the rest."""
    source = cells.source
    pieces: list[bytes] = []
    by_block: dict[int, list[bytes]] = {}
    for number, row in sorted(produced.items()):
        by_block.setdefault(number // BLOCK_ROWS, []).append(row)
    if not set(by_block) <= cells.loaded:
        raise ValueError("A row was written for a block that was never parsed.")
    done = 0
    for block in sorted(cells.loaded):
        first = source.between(block * BLOCK_ROWS, (block + 1) * BLOCK_ROWS - 1)
        pieces.append(_unchanged(source, done, first.start))
        pieces.extend(by_block.get(block, []))
        done = first.stop
    pieces.append(_unchanged(source, done, len(source.row_numbers)))
    return pieces


def _unchanged(source: SheetSource, first: int, stop: int) -> bytes:
    """The file's rows with positions ``first`` up to ``stop``, without their stored results."""
    if stop <= first:
        return b""
    rows = source.data[source.spans[first][0] : source.spans[stop - 1][1]]
    return _RESULT.sub(rb"\1", rows)


def _dimension(source: SheetSource, produced: dict[int, bytes]) -> bytes:
    """The range the cells of the file's rows and the written rows cover."""
    boxes: list[tuple[int, int, int, int]] = []
    if source.has_cells:
        boxes.append((source.min_row, source.max_row, *source.columns))
    written = b"".join(produced.values())
    columns = {column_index_from_string(name.decode()) for name in _CELL_COLUMN.findall(written)}
    rows = [number for number, row in produced.items() if _CELL_COLUMN.search(row)]
    if rows:
        boxes.append((min(rows), max(rows), min(columns), max(columns)))
    if not boxes:
        return b"A1:A1"
    top, bottom = min(box[0] for box in boxes), max(box[1] for box in boxes)
    left, right = min(box[2] for box in boxes), max(box[3] for box in boxes)
    return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}".encode()
