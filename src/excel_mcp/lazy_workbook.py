"""Workbooks whose cells are parsed, row by row, when something asks for them.

Calculating a page of a large sheet needs the cells of that page and of what its formulas
use, not the other million. The workbook is loaded without its cell data, and each sheet
keeps its XML so rows can be parsed on demand with openpyxl's own cell parser. Only the
cells and the hidden state of rows are loaded this way: that is all the calculator reads.
`lazy_edit` builds on this to edit large sheets without parsing the rows an edit leaves alone.
"""

import re
from collections.abc import Callable, Iterator
from io import BytesIO
from pathlib import Path
from typing import IO, Any
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.cell import Cell, MergedCell
from openpyxl.reader.excel import ExcelReader
from openpyxl.worksheet._reader import WorkSheetParser
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.constants import WORKSHEET_TYPE

from excel_mcp import lazy_rows, lazy_shared
from excel_mcp.lazy_source import SheetSource, split_sheet

BLOCK_ROWS = 128
# A formula that is not followed by a result.
_UNCALCULATED = re.compile(rb"<f\b[^>]*(?:/>|>[^<]*</f>)(?!<v>[^<])")

CellKey = tuple[int, int]


class LazyCells(dict[CellKey, Any]):
    """The cells of a worksheet, parsed from its rows the first time one is asked for.

    Asking for one cell parses the block of rows around it; asking for all of them (iterating,
    counting) parses the sheet. Changing a cell parses its block first, so a block that was
    parsed holds all of its rows: what is not in ``loaded`` is as the file has it.
    """

    def __init__(self, sheet: Worksheet, source: SheetSource, strings: list[str]) -> None:
        super().__init__()
        self.sheet = sheet
        self.source = source
        self.strings = strings
        self.loaded: set[int] = set()
        self.columns: dict[int, list[int]] = {}
        self.last_row = self.last_column = 0

    def keep(self, cell: Any) -> None:
        dict.__setitem__(self, (cell.row, cell.column), cell)
        self.columns.setdefault(cell.row, []).append(cell.column)

    def held(self) -> dict[CellKey, Any]:
        """The cells parsed or put in so far."""
        return dict(dict.items(self))

    def is_loaded(self, row: int) -> bool:
        return row // BLOCK_ROWS in self.loaded

    def all_loaded(self, first: int, last: int) -> bool:
        """Whether the blocks that hold rows ``first`` to ``last`` are parsed."""
        return all(
            block in self.loaded for block in range(first // BLOCK_ROWS, last // BLOCK_ROWS + 1)
        )

    def load_between(self, first: int, last: int) -> None:
        for block in range(first // BLOCK_ROWS, last // BLOCK_ROWS + 1):
            self.load(block * BLOCK_ROWS)

    def stored_rows(self, first: int, last: int) -> list[dict[int, tuple[int, Any, str]]] | None:
        """Rows ``first`` to ``last`` as the file stores them: the last results of formulas, not
        the formulas. Each row is its cells' style, value and type by column. Nothing is parsed
        into cells, and the sheet is not changed.

        None where that is not what the sheet holds or means: a block of these rows was parsed
        (it may have changed), a formula has no result, the first row has a formula, durations
        are formatted (a result is not told from its format), or a row is in an unusual form.
        """
        book: Any = self.sheet.parent
        source = self.source
        if book._timedelta_formats or any(
            block in self.loaded for block in range(first // BLOCK_ROWS, last // BLOCK_ROWS + 1)
        ):
            return None
        reading = lazy_rows.Reading(self.strings, {}, book._date_formats, set(), book.epoch, True)
        rows: list[dict[int, tuple[int, Any, str]]] = [{} for _ in range(last - first + 1)]
        found = source.between(first, last)
        for number, (start, stop) in zip(
            source.row_numbers[found], source.spans[found], strict=True
        ):
            xml = source.data[start:stop]
            if _UNCALCULATED.search(xml) or (number == first and b"<f" in xml):
                return None
            parsed = lazy_rows.read(xml, reading)
            if parsed is None:
                return None
            rows[number - first] = {c: (s, v, t) for c, s, v, t in parsed[1]}
        return rows

    def count(self) -> int:
        """The number of cells, without parsing any."""
        source = self.source
        unparsed = (
            source.data.count(b"<c ", start, stop)
            for number, (start, stop) in zip(source.row_numbers, source.spans, strict=True)
            if not self.is_loaded(number)
        )
        return dict.__len__(self) + sum(unparsed)

    def parsed_rows(self) -> list[tuple[int, int]]:
        """The first and last row of each block that is parsed."""
        return [(block * BLOCK_ROWS, (block + 1) * BLOCK_ROWS - 1) for block in sorted(self.loaded)]

    def in_row(self, row: int) -> list[tuple[CellKey, Any]]:
        """The cells of one row, which is parsed."""
        self.load(row)
        return [
            ((row, column), dict.__getitem__(self, (row, column)))
            for column in self.columns.get(row, ())
        ]

    def replace(
        self, cells: dict[CellKey, Any], source: SheetSource, parsed: list[tuple[int, int]]
    ) -> None:
        """Hold ``cells`` (those that were parsed, moved) over ``source`` (the rows that were not
        parsed, moved), and parse the blocks that ``parsed`` rows are in now, so that a block is
        either parsed whole or not at all."""
        dict.clear(self)
        self.columns = {}
        self.last_row = self.last_column = 0
        self.source = source
        self.loaded = set()
        for key, cell in cells.items():
            dict.__setitem__(self, key, cell)
            self.columns.setdefault(key[0], []).append(key[1])
            self.last_row, self.last_column = (
                max(self.last_row, key[0]),
                max(self.last_column, key[1]),
            )
        for first, last in parsed:
            self.load_between(first, last)

    def extent(self) -> tuple[int, int]:
        """The last row and column that hold a cell, without parsing any."""
        return max(self.source.max_row, self.last_row), max(
            self.source.max_column, self.last_column
        )

    def get(self, key: CellKey, default: Any = None) -> Any:
        self.load(key[0])
        return super().get(key, default)

    def __getitem__(self, key: CellKey) -> Any:
        self.load(key[0])
        return super().__getitem__(key)

    def __contains__(self, key: object) -> bool:
        self.load(key[0])  # pyright: ignore[reportIndexIssue]
        return super().__contains__(key)

    def __setitem__(self, key: CellKey, cell: Any) -> None:
        self.load(key[0])
        if not super().__contains__(key):
            self.columns.setdefault(key[0], []).append(key[1])
        super().__setitem__(key, cell)
        self.last_row, self.last_column = max(self.last_row, key[0]), max(self.last_column, key[1])

    def __delitem__(self, key: CellKey) -> None:
        self.load(key[0])
        super().__delitem__(key)
        self.columns[key[0]].remove(key[1])

    def pop(self, key: CellKey, *default: Any) -> Any:
        self.load(key[0])
        if super().__contains__(key):
            self.columns[key[0]].remove(key[1])
        return super().pop(key, *default)

    def setdefault(self, key: CellKey, default: Any = None) -> Any:
        if key not in self:
            self[key] = default
        return super().__getitem__(key)

    def __bool__(self) -> bool:
        return dict.__len__(self) > 0 or self.source.has_cells

    # Everything that looks at all the cells parses them all.

    def __iter__(self) -> Iterator[CellKey]:
        self.load_all()
        return super().__iter__()

    def __reversed__(self) -> Iterator[CellKey]:
        self.load_all()
        return super().__reversed__()

    def __len__(self) -> int:
        self.load_all()
        return super().__len__()

    def keys(self) -> Any:
        self.load_all()
        return super().keys()

    def values(self) -> Any:
        self.load_all()
        return super().values()

    def items(self) -> Any:
        self.load_all()
        return super().items()

    def copy(self) -> dict[CellKey, Any]:
        self.load_all()
        return super().copy()

    def stored_in(
        self, top: int, left: int, bottom: int, right: int
    ) -> Iterator[tuple[CellKey, Any]]:
        for row in self.source.row_numbers[self.source.between(top, bottom)]:
            self.load(row)
            for col in sorted(self.columns.get(row, ())):
                if left <= col <= right:
                    yield (row, col), super().__getitem__((row, col))

    def load_all(self) -> None:
        source = self.source
        spans = [
            span
            for number, span in zip(source.row_numbers, source.spans, strict=True)
            if number // BLOCK_ROWS not in self.loaded
        ]
        self.loaded.update(number // BLOCK_ROWS for number in source.row_numbers)
        self._parse(spans)

    def load(self, row: int) -> None:
        """Parse the block of rows that holds ``row``, unless that is done."""
        block = row // BLOCK_ROWS
        if block in self.loaded:
            return
        self.loaded.add(block)
        first, last = block * BLOCK_ROWS, (block + 1) * BLOCK_ROWS - 1
        self._parse(self.source.spans[self.source.between(first, last)])

    def _parse(self, spans: list[tuple[int, int]]) -> None:
        """Parse rows: those in the usual forms as text, the others with openpyxl's parser."""
        book: Any = self.sheet.parent
        shared = lazy_shared.Translators(self.source.masters)
        reading = lazy_rows.Reading(
            self.strings,
            shared,
            book._date_formats,
            book._timedelta_formats,
            book.epoch,
            book._data_only,
        )
        unusual = []
        for start, end in spans:
            row = lazy_rows.read(self.source.data[start:end], reading)
            if row is None:
                unusual.append((start, end))
                continue
            for column, style, value, data_type in row[1]:
                self._parsed(row[0], column, style, value, data_type)
        if unusual:
            self._parse_unusual(unusual, shared)

    def _parse_unusual(self, spans: list[tuple[int, int]], shared: lazy_shared.Translators) -> None:
        rows = b"".join(self.source.data[start:end] for start, end in spans)
        document = self.source.root + b"<sheetData>" + rows + b"</sheetData></worksheet>"
        book: Any = self.sheet.parent
        parser = WorkSheetParser(
            BytesIO(document),
            self.strings,
            book._data_only,
            book.epoch,
            book._date_formats,
            book._timedelta_formats,
        )
        parser.shared_formulae = shared
        for _, parsed in parser.parse():
            for item in parsed:
                self._parsed(
                    item["row"],
                    item["column"],
                    item["style_id"],
                    item["value"],
                    item["data_type"],
                )

    def _parsed(self, row: int, column: int, style: int, value: Any, data_type: str) -> None:
        if super().__contains__((row, column)):
            return
        book: Any = self.sheet.parent
        cell = Cell(self.sheet, row=row, column=column, style_array=book._cell_styles[style])
        cell._value = value
        cell.data_type = data_type
        self.keep(cell)


Parts = dict[str, tuple[bytes, SheetSource | None]]
Splitter = Callable[[bytes], tuple[bytes, SheetSource] | None]


class _SheetArchive(ZipFile):
    """An archive that serves each worksheet without its rows, which are kept in ``parts``."""

    def __init__(self, path: Path, parts: Parts, sheet_parts: set[str], split: Splitter) -> None:
        super().__init__(path)
        self.parts = parts
        self.sheet_parts = sheet_parts
        self.split = split
        self.worksheets: list[str] = []

    def open(self, name: Any, mode: Any = "r", pwd: Any = None, **options: Any) -> IO[bytes]:
        name = getattr(name, "filename", name)
        if name not in self.parts:
            with super().open(name, mode, pwd, **options) as part:
                data = part.read()
            if name not in self.sheet_parts:
                return BytesIO(data)
            split = self.split(data)
            self.parts[name] = split if split else (data, None)
        self.worksheets.append(name)
        return BytesIO(self.parts[name][0])


def open_lazily(path: Path, parts: Parts, split: Splitter, **options: Any) -> ExcelReader:
    """A reader that has not read the workbook yet, whose sheets are served by ``split``."""
    reader = ExcelReader(path, **options)
    reader.read_manifest()
    sheet_parts = {part.PartName.lstrip("/") for part in reader.package.findall(WORKSHEET_TYPE)}
    reader.archive.close()
    reader.archive = _SheetArchive(path, parts, sheet_parts, split)
    return reader


def lazy_sheets(reader: ExcelReader, parts: Parts) -> Iterator[tuple[Worksheet, SheetSource]]:
    """The worksheets of a workbook read by ``reader`` whose rows were split off."""
    archive: Any = reader.archive
    sheets = [sheet for sheet in reader.wb.worksheets if isinstance(sheet, Worksheet)]
    for sheet, name in zip(sheets, archive.worksheets, strict=True):
        if (source := parts[name][1]) is not None:
            yield sheet, source


def load_lazy(path: Path, parts: Parts, *, data_only: bool) -> Workbook:
    """Load a workbook the way ``load_workbook`` does, leaving the cells to be read on demand.

    A sheet this cannot split into rows is loaded whole. ``parts`` is shared by the loads of
    one file, so that each sheet is read from the file once.
    """
    reader = open_lazily(path, parts, split_sheet, data_only=data_only, keep_links=False)
    reader.read()
    for sheet, source in lazy_sheets(reader, parts):
        cells = LazyCells(sheet, source, reader.shared_strings)
        for cell in sheet._cells.values():
            if isinstance(cell, MergedCell):
                cells.keep(cell)
        sheet._cells = cells
        for row in source.hidden_rows:
            sheet.row_dimensions[row].hidden = True
    return reader.wb


def held_cells(sheet: Worksheet) -> dict[CellKey, Any]:
    """The cells a sheet holds in memory: all of them, but for the rows a lazy sheet has not parsed.

    For looking for what can only be in cells that were parsed because it matters to them
    (hyperlinks, comments, array formulas) or put in by an edit.
    """
    cells = sheet._cells
    return cells.held() if isinstance(cells, LazyCells) else cells


def cell_count(sheet: Worksheet) -> int:
    """The number of cells a sheet has, which loading a lazy sheet's rows must not."""
    cells = sheet._cells
    return cells.count() if isinstance(cells, LazyCells) else len(cells)


def extent(sheet: Worksheet) -> tuple[int, int]:
    """The last row and column that hold a cell, which loading a lazy sheet's rows must not."""
    cells = sheet._cells
    if isinstance(cells, LazyCells):
        return cells.extent()
    return sheet.max_row, sheet.max_column
