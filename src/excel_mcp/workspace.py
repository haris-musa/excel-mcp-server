"""Opening and saving workbooks safely."""

import io
import os
import tempfile
import threading
import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import TypeVar, cast

from openpyxl import Workbook, load_workbook
from openpyxl.chart import AreaChart
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp import chart_roundtrip, macros, package
from excel_mcp.config import Limits
from excel_mcp.errors import (
    InvalidArgumentError,
    LimitExceededError,
    SheetNotFoundError,
    WorkbookError,
    WorkbookExistsError,
    WorkbookNotFoundError,
)
from excel_mcp.lazy_workbook import Parts, load_lazy
from excel_mcp.paths import MACRO_SUFFIXES, TEMPLATE_SUFFIXES, PathPolicy

MAX_IMAGE_BYTES = 10 * 1024 * 1024
Sheet = TypeVar("Sheet")


class Workspace:
    """Resolves client paths and gives access to workbooks within the limits.

    Changes are serialized: tools run in parallel threads, and each edit loads
    and saves the whole file, so concurrent edits would otherwise lose data.
    """

    def __init__(self, paths: PathPolicy, limits: Limits, allow_macro_workbooks: bool) -> None:
        chart_roundtrip.install()
        self.paths = paths
        self.limits = limits
        self.allow_macro_workbooks = allow_macro_workbooks
        self._write_lock = threading.Lock()

    def resolve(self, raw_path: str) -> Path:
        return self.paths.resolve(raw_path)

    def resolve_existing(self, raw_path: str) -> Path:
        path = self.resolve(raw_path)
        if not path.is_file():
            raise WorkbookNotFoundError(f"Workbook {self.display(path)} does not exist.")
        size = path.stat().st_size
        if size > self.limits.max_file_bytes:
            raise LimitExceededError(
                f"Workbook is {size:,} bytes; the limit is {self.limits.max_file_bytes:,}."
            )
        return path

    def read_image(self, raw_path: str) -> bytes:
        """Read an image file confined like workbooks; its content is checked by the caller."""
        path = self.paths.resolve_image(raw_path)
        if not path.is_file():
            raise InvalidArgumentError(f"Image {self.display(path)} does not exist.")
        size = path.stat().st_size
        if size > MAX_IMAGE_BYTES:
            raise LimitExceededError(
                f"Image is {size:,} bytes; the limit is {MAX_IMAGE_BYTES:,}. Use a smaller image."
            )
        return path.read_bytes()

    def resolve_directory(self, raw_path: str) -> Path:
        return self.paths.resolve_directory(raw_path)

    def display(self, path: Path) -> str:
        return self.paths.display(path)

    @contextmanager
    def read(
        self, raw_path: str, *, data_only: bool = False, with_package: bool = False
    ) -> Iterator[Workbook]:
        """Open a workbook for reading; changes are never saved.

        ``with_package`` also loads what openpyxl drops, as `edit` does, for readers of it.
        """
        path = self.resolve_existing(raw_path)
        workbook = self._load(path, data_only=data_only)
        if with_package:
            package.capture(path, workbook, self.limits.max_file_bytes)
        try:
            yield workbook
        finally:
            close_workbook(workbook)

    @contextmanager
    def read_lazy(self, raw_path: str) -> Iterator[tuple[Workbook, Workbook]]:
        """Open a workbook as its stored results and its formulas, parsing cells when needed.

        For calculating some of a large sheet's formulas: only the rows that are asked for are
        read, so the cells they use are the only ones that cost anything. Changes are never saved.
        """
        path = self.resolve_existing(raw_path)
        self._check_zip(path)
        parts: Parts = {}
        try:
            stored = load_lazy(path, parts, data_only=True)
            formulas = load_lazy(path, parts, data_only=False)
        except Exception as error:
            raise self._unreadable(path, error) from None
        try:
            yield stored, formulas
        finally:
            close_workbook(stored)
            close_workbook(formulas)

    @contextmanager
    def stream(self, raw_path: str, *, data_only: bool = False) -> Iterator[Workbook]:
        """Open a workbook for streaming reads, which use constant memory on huge sheets.

        Its worksheets are openpyxl read-only views: they hold only cell values (no
        formatting, merged ranges, tables or charts) and re-parse the sheet on each pass.
        """
        workbook = self._load(self.resolve_existing(raw_path), data_only=data_only, stream=True)
        try:
            yield workbook
        finally:
            close_workbook(workbook)

    @contextmanager
    def edit(self, raw_path: str) -> Iterator[Workbook]:
        """Open a workbook and save it only if the block finishes without error."""
        with self._write_lock:
            path = self.resolve_existing(raw_path)
            workbook = self._load(path, data_only=False)
            _restore_area_chart_axes(workbook)
            package.capture(path, workbook, self.limits.max_file_bytes)
            workbook.calculation = workbook.calculation or CalcProperties()
            workbook.calculation.fullCalcOnLoad = True
            try:
                yield workbook
                # Imported here because the calculator imports this module.
                from excel_mcp.operations.spill import refresh_spills

                refresh_spills(workbook, self.limits.max_cells)
                _pin_hyperlinks(workbook)
                save_atomically(workbook, path, self.limits.max_file_bytes)
            finally:
                close_workbook(workbook)

    def create(self, raw_path: str, sheets: list[str], *, overwrite: bool) -> Path:
        path = self.resolve(raw_path)
        macro_enabled = path.suffix.lower() in MACRO_SUFFIXES
        if macro_enabled and not self.allow_macro_workbooks:
            raise InvalidArgumentError(
                "New workbooks cannot contain macros unless the server allows VBA writing. "
                "Create an .xlsx or .xltx file instead."
            )
        workbook = Workbook()
        workbook.properties.creator = None
        workbook.worksheets[0].title = sheets[0]
        for name in sheets[1:]:
            workbook.create_sheet(name)
        workbook.template = path.suffix.lower() in TEMPLATE_SUFFIXES
        if macro_enabled:
            macros.store_project(workbook, macros.new_project(workbook))
        with self._write_lock:
            try:
                self._check_overwrite(path, overwrite)
                save_atomically(workbook, path, self.limits.max_file_bytes)
            finally:
                close_workbook(workbook)
        return path

    def store(self, raw_path: str, content: bytes, *, overwrite: bool) -> Path:
        """Save uploaded workbook bytes as a file, without macros unless the file may hold them."""
        path = self.resolve(raw_path)
        if path.suffix.lower() not in MACRO_SUFFIXES:
            content = macros.without_macros(content)
        with self._write_lock:
            self._check_overwrite(path, overwrite)
            write_atomically(path, content)
        return path

    def _check_overwrite(self, path: Path, overwrite: bool) -> None:
        if path.exists() and not overwrite:
            raise WorkbookExistsError(
                f"{self.display(path)} already exists. Pass overwrite=true to replace it."
            )

    def _load(self, path: Path, *, data_only: bool, stream: bool = False) -> Workbook:
        self._check_zip(path)
        try:
            return load_workbook(
                path,
                data_only=data_only,
                read_only=stream,
                keep_vba=not stream and path.suffix.lower() in MACRO_SUFFIXES,
            )
        except Exception as error:
            raise self._unreadable(path, error) from None

    def _check_zip(self, path: Path) -> None:
        if not zipfile.is_zipfile(path):
            raise WorkbookError(
                f"{self.display(path)} is not an Excel workbook. Only .xlsx/.xlsm files are "
                "supported; legacy .xls and CSV files must be converted first."
            )

    def _unreadable(self, path: Path, error: Exception) -> WorkbookError:
        # openpyxl raises many different exception types for damaged files.
        return WorkbookError(
            f"{self.display(path)} could not be opened as an Excel workbook ({error})."
        )


def save_atomically(workbook: Workbook, path: Path, max_bytes: int) -> None:
    """Serialize the workbook in memory, then write it atomically.

    A result larger than the size limit is refused, because the server could not open it again.
    """
    buffer = io.BytesIO()
    try:
        package.prepare(workbook)
        workbook.save(buffer)
    except Exception as error:
        # openpyxl reports invalid workbook states with many exception types.
        raise WorkbookError(f"Could not save the workbook: {error}") from None
    content = package.apply(workbook, buffer.getvalue())
    if len(content) > max_bytes:
        raise LimitExceededError(
            f"The saved workbook would be {len(content):,} bytes; the limit is "
            f"{max_bytes:,}. The file was not changed."
        )
    write_atomically(path, content)


def write_atomically(path: Path, content: bytes) -> None:
    """Write to a new temporary file next to the target, then swap it in.

    mkstemp creates the file exclusively, so a symlink planted at the temporary
    name cannot redirect the write, and a failed write leaves the target intact.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(dir=path.parent, prefix=".~", suffix=".tmp")
    temp_path = Path(temp_name)
    try:
        with os.fdopen(handle, "wb") as file:
            file.write(content)
        temp_path.replace(path)
    except OSError as error:
        raise WorkbookError(f"Could not save the workbook: {error.strerror}.") from None
    finally:
        temp_path.unlink(missing_ok=True)


def close_workbook(workbook: Workbook) -> None:
    workbook.close()
    # openpyxl keeps macro parts in an in-memory zip that it never closes.
    if workbook.vba_archive:
        workbook.vba_archive.close()


def _restore_area_chart_axes(workbook: Workbook) -> None:
    """Undo an openpyxl bug that hides the axes of area charts.

    openpyxl drops the axes' ``delete`` flag when it reads an area chart, and
    Excel treats the missing flag as hidden, so without this every edit would
    remove the axes of area charts already in the workbook.
    """
    for sheet in worksheets(workbook):
        for chart in sheet._charts:  # pyright: ignore[reportAttributeAccessIssue]
            if isinstance(chart, AreaChart):
                for axis in (chart.x_axis, chart.y_axis):
                    if axis.delete is None:
                        axis.delete = False


def _pin_hyperlinks(workbook: Workbook) -> None:
    """Point each link at the cell that holds it now.

    openpyxl keeps the address a link had when it was read, so inserting or deleting rows
    or columns would leave links on the cells that took their place.
    """
    for sheet in worksheets(workbook):
        for cell in sheet._cells.values():
            if cell.hyperlink:
                cell.hyperlink.ref = cell.coordinate


def get_sheet(workbook: Workbook, name: str) -> Worksheet:
    return _find_sheet(workbook, name, Worksheet)


def get_streamed_sheet(workbook: Workbook, name: str) -> ReadOnlyWorksheet:
    return _find_sheet(workbook, name, ReadOnlyWorksheet)


def worksheets(workbook: Workbook) -> list[Worksheet]:
    """The workbook's worksheets, leaving out chart sheets."""
    return [sheet for sheet in workbook.worksheets if isinstance(sheet, Worksheet)]


def streamed_worksheets(workbook: Workbook) -> list[ReadOnlyWorksheet]:
    return [sheet for sheet in workbook.worksheets if isinstance(sheet, ReadOnlyWorksheet)]


def _find_sheet(workbook: Workbook, name: str, kind: type[Sheet]) -> Sheet:
    if name not in workbook.sheetnames:
        raise SheetNotFoundError(name, workbook.sheetnames)
    sheet = workbook[name]
    if not isinstance(sheet, kind):
        available = [sheet.title for sheet in workbook.worksheets if isinstance(sheet, kind)]
        raise SheetNotFoundError(name, available)
    return sheet


def sheet_names(sheet: Worksheet) -> list[str]:
    """The names of all sheets in the workbook that holds ``sheet``."""
    # openpyxl types parent as optional, but every sheet it loads or creates has one.
    return cast(Workbook, sheet.parent).sheetnames
