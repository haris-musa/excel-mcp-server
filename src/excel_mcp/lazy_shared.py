"""Shared formulas on lazily loaded sheets.

Excel writes a formula that was filled down once, as the master (the formula and the range it
covers, `ref`) and, in each other cell of the range, only the group number (`si`). openpyxl
turns a child into a formula by translating its master's, so a block of rows that holds
children needs the master wherever it is: `index` finds the masters when the sheet is split,
and `Translators` gives the parser the one a child asks for.

Saving rewrites the blocks that were parsed, which turns their cells into ordinary formulas,
and keeps the other rows as the file has them, whose children still refer to their master.
So a master in a parsed block that has children in rows that were not parsed is written as
a master again, as long as it is as it was. When the edit changed it, its children would
be left without one, so every block its range covers is parsed and they all become ordinary
formulas. A long fill-down is parsed whole only if its first cell is overwritten.
"""

import re
from dataclasses import dataclass
from functools import cached_property
from html import unescape
from typing import TYPE_CHECKING

from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter

from excel_mcp import plain_refs
from excel_mcp.refs import parse_range

if TYPE_CHECKING:
    from excel_mcp.lazy_workbook import LazyCells

# Only a master has the text of the formula; a child is `<f t="shared" si="3"/>`.
_MASTER = re.compile(rb'<f\b([^>]*\bt="shared"[^>]*)>([^<]+)</f>')
_COORDINATE = re.compile(rb'\br="([A-Z]+)(\d+)"')
_REF = re.compile(rb'\bref="([^"]+)"')
_GROUP = re.compile(rb'\bsi="(\d+)"')


@dataclass(frozen=True)
class SharedMaster:
    group: str
    row: int
    column: int
    ref: str
    text: str

    @property
    def coordinate(self) -> str:
        return f"{get_column_letter(self.column)}{self.row}"

    @property
    def rows(self) -> tuple[int, int]:
        """The first and last row of the range the formula covers."""
        area = parse_range(self.ref)
        return area.min_row, area.max_row

    @property
    def wide(self) -> bool:
        """Whether the range the formula covers has more than one column."""
        area = parse_range(self.ref)
        return area.min_col < area.max_col


def index(data: bytes, start: int, end: int) -> dict[str, SharedMaster]:
    """The masters of the shared formulas in ``data[start:end]``, by group number."""
    masters: dict[str, SharedMaster] = {}
    for found in _MASTER.finditer(data, start, end):
        group, ref = _GROUP.search(found[1]), _REF.search(found[1])
        cell = _COORDINATE.search(data, data.rfind(b"<c ", start, found.start()))
        if group is None or ref is None or cell is None:
            continue
        name = group[1].decode()
        masters.setdefault(
            name,
            SharedMaster(
                name,
                int(cell[2]),
                column_index_from_string(cell[1].decode()),
                ref[1].decode(),
                unescape(found[2].decode()),
            ),
        )
    return masters


class _Translator(Translator):
    """A master's formula, translated to the cells that follow it.

    A formula of plain references is translated as text. The tokenizer renders runs of spaces
    and line breaks in its own way, so a formula with those is translated through it.
    """

    def __init__(self, master: SharedMaster) -> None:
        self.master = master
        self.formula = "=" + master.text
        self.row, self.col = master.row, master.column
        self.plain = (
            self.formula.isprintable()
            and "  " not in self.formula
            and plain_refs.shape(self.formula) is not None
        )

    @cached_property
    def tokenizer(self) -> Tokenizer:  # pyright: ignore[reportIncompatibleVariableOverride]
        return Tokenizer(self.formula)

    def translate_formula(
        self, dest: str | None = None, row_delta: int = 0, col_delta: int = 0
    ) -> str:
        # The master is read as it is written, not as the tokens of it come out again.
        if dest == self.master.coordinate:
            return self.formula
        if dest:
            row, column = coordinate_to_tuple(dest)
            row_delta, col_delta = row - self.row, column - self.col
        if self.plain:
            moved = plain_refs.sub(self.formula, lambda m: _translated(m, row_delta, col_delta))
            assert moved is not None
            return moved
        return super().translate_formula(dest, row_delta, col_delta)


def _translated(match: plain_refs.Match, rows: int, columns: int) -> str:
    """A reference moved as openpyxl's `Translator.translate_range` moves it."""
    first = _moved(match[1] + match[2], match[3] + match[4], rows, columns)
    if match[6] is None:
        return first
    return f"{first}:{_moved(match[5] + match[6], match[7] + match[8], rows, columns)}"


def _moved(column: str, row: str, rows: int, columns: int) -> str:
    if Translator.CELL_REF_RE.match(column + row) is None:  # a name, not a cell
        return column + row
    return Translator.translate_col(column, columns) + Translator.translate_row(row, rows)


class Translators(dict[str, Translator]):
    """What the parser keeps its shared formulas in, with the masters of the whole sheet."""

    def __init__(self, masters: dict[str, SharedMaster]) -> None:
        super().__init__()
        self.masters = masters

    def __contains__(self, group: object) -> bool:
        return group in self.masters

    def __missing__(self, group: str) -> Translator:
        translator = self[group] = _Translator(self.masters[group])
        return translator


def _has_raw_children(cells: "LazyCells", master: SharedMaster) -> bool:
    return not cells.all_loaded(*master.rows)


def _unchanged(cells: "LazyCells", master: SharedMaster) -> bool:
    cell = cells.get((master.row, master.column))
    return cell is not None and cell.value == "=" + master.text


def settle(cells: "LazyCells") -> None:
    """Parse the range of every shared formula whose master an edit changed, before saving."""
    for master in cells.source.masters.values():
        if (
            cells.is_loaded(master.row)
            and _has_raw_children(cells, master)
            and not _unchanged(cells, master)
        ):
            cells.load_between(*master.rows)


def restore(cells: "LazyCells", produced: dict[int, bytes]) -> None:
    """Write the masters that children in rows kept from the file refer to as masters again."""
    for master in cells.source.masters.values():
        row = produced.get(master.row)
        if row is None or not cells.is_loaded(master.row) or not _has_raw_children(cells, master):
            continue
        cell = re.compile(rb'(<c r="%b"[^>]*><f)>' % master.coordinate.encode())
        shared = b' t="shared" ref="%b" si="%b">' % (master.ref.encode(), master.group.encode())
        produced[master.row] = cell.sub(rb"\1" + shared, row, 1)
