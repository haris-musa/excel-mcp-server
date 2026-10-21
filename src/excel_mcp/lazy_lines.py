"""Rows inserted or deleted, and references moved, in a sheet whose rows are not parsed.

A large sheet is held as the XML of its rows (`lazy_workbook`). Inserting rows above most of
them renumbers the rows and rewrites the references in their formulas, which needs no cells:
both are done on the XML, row by row, and only the rows around the edit are parsed.

Formulas filled down are stored as one master and the cells that follow it (`lazy_shared`).
They stay that way when every cell of the group moves alike: rewriting the master then
rewrites them all. A group whose cells or references fall on both sides of the edit (such a
cell now refers to another one than its master does) is parsed first and becomes ordinary
formulas, as Excel splits it.
"""

import re
from collections.abc import Callable
from dataclasses import replace
from functools import lru_cache
from html import unescape
from typing import Any
from xml.sax.saxutils import escape

from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.formulas import storable_formula
from excel_mcp.lazy_shared import SharedMaster
from excel_mcp.lazy_source import SheetSource
from excel_mcp.lazy_workbook import CellKey, LazyCells
from excel_mcp.package.lines import LineEdit
from excel_mcp.refs import parse_range
from excel_mcp.rewrite import ReferenceRewriter

_FORMULA = re.compile(rb"<f\b([^>]*)>([^<]*)</f>")
_ROW_NUMBERS = re.compile(rb'(<row\b[^>]*?\br="|<c\b[^>]*?\br="[A-Z]+)(\d+)(")')
_REF = re.compile(rb' ref="([^"]*)"')
_CELLS = re.compile(rb"<c\b[^>]*?(?:/>|>.*?</c>)", re.S)
_CELL_REF = re.compile(rb'<c\b[^>]*?\br="([A-Z]+)(\d+)"')
_STYLE = re.compile(rb' s="(\d+)"')
_SPANS = re.compile(rb' spans="[^"]*"')

Row = tuple[int, bytes]


def unshare(cells: LazyCells, host: str, rewriter: ReferenceRewriter, edit: LineEdit) -> None:
    """Parse the groups of shared formulas that ``edit`` would not leave alike."""
    for master in cells.source.masters.values():
        if not _moves_alike(master, host, rewriter, edit):
            cells.load_between(*master.rows)


def _moves_alike(
    master: SharedMaster, host: str, rewriter: ReferenceRewriter, edit: LineEdit
) -> bool:
    """Whether every cell of the group ends up with the formula its master gives it.

    The group moves alike if the master's rewritten formula, copied to the farthest cell of the
    group, is what rewriting the formula of that cell gives. The shift of a reference is the
    same for the cells in between, because it changes only once, at the edit.
    """
    area = parse_range(master.ref)
    if host.casefold() == edit.sheet.casefold():
        first, last = (
            (area.min_row, area.max_row) if edit.axis == "rows" else (area.min_col, area.max_col)
        )
        if edit.cuts(first, last):
            return False
    far = f"{get_column_letter(area.max_col)}{area.max_row}"
    formula = "=" + master.text
    copied = Translator(formula, master.coordinate).translate_formula(far)
    moved = Translator(rewriter.formula(formula, host), master.coordinate).translate_formula(far)
    return moved == rewriter.formula(copied, host)


def rewrite_formulas(
    cells: LazyCells,
    host: str,
    rewriter: ReferenceRewriter,
    names: list[str],
    *,
    check_all: bool = False,
) -> None:
    """Rewrite the references of the formulas in the rows that are not parsed.

    A formula that changes is checked as every written formula is; ``check_all`` also checks
    those that stay as they are (a copy of a sheet writes them all).
    """
    source = cells.source
    cache: dict[str, str] = {}

    def rewrite(text: str) -> str:
        if text not in cache:
            formula = "=" + text
            moved = rewriter.formula(formula, host)
            unchanged = moved == formula and not check_all
            cache[text] = text if unchanged else storable_formula(moved, names)[1:]
        return cache[text]

    def replace_formula(match: re.Match[bytes]) -> bytes:
        text = unescape(match[2].decode())
        moved = rewrite(text)
        if moved == text:
            return match[0]
        return b"<f%b>%b</f>" % (match[1], escape(moved).encode())

    rows: list[Row] = []
    for number, (start, stop) in zip(source.row_numbers, source.spans, strict=True):
        row = source.data[start:stop]
        if b"<f" in row and not cells.is_loaded(number):
            row = _FORMULA.sub(replace_formula, row)
        rows.append((number, row))
    masters = {group: replace(m, text=rewrite(m.text)) for group, m in source.masters.items()}
    cells.source = _assembled(source, rows, masters, source.hidden_rows)


def move_lines(
    cells: LazyCells,
    edit: LineEdit,
    moved: dict[CellKey, Any],
    inherit: Callable[[int, int | None], int] | None,
) -> None:
    """Insert or delete rows or columns: move the rows that are not parsed, in their XML, and
    hold ``moved`` (the cells that are parsed, already moved) with them.

    ``inherit`` gives the style of a cell that is inserted next to one: from the style of that
    cell and of the one on its other side.
    """
    source = cells.source
    rows: list[Row] = []
    hidden = source.hidden_rows
    parsed = cells.parsed_rows()
    if edit.axis == "rows":
        move_row = _row_mover(edit)
        for number, (start, stop) in zip(source.row_numbers, source.spans, strict=True):
            if (index := edit.index(number)) is not None and not cells.is_loaded(number):
                rows.append((index, move_row(source.data[start:stop], index - number)))
        hidden = [index for row in hidden if (index := edit.index(row)) is not None]
        parsed = [span for first, last in parsed if (span := edit.span(first, last))]
    else:
        move_cells = _column_mover(edit, inherit)
        for number, (start, stop) in zip(source.row_numbers, source.spans, strict=True):
            if not cells.is_loaded(number):
                rows.append((number, move_cells(source.data[start:stop])))
    cells.replace(moved, _assembled(source, rows, _moved_masters(source, edit), hidden), parsed)


def copied(cells: LazyCells, sheet: Worksheet) -> LazyCells:
    """The cells of a copy of a sheet: the rows that are not parsed are shared with the original,
    which keeps the cells that are parsed."""
    source = cells.source
    rows = [
        (number, source.data[start:stop])
        for number, (start, stop) in zip(source.row_numbers, source.spans, strict=True)
        if not cells.is_loaded(number)
    ]
    masters = dict(source.masters)
    return LazyCells(
        sheet, _assembled(source, rows, masters, source.hidden_rows), cells.strings, cells.budget
    )


def _moved_masters(source: SheetSource, edit: LineEdit) -> dict[str, SharedMaster]:
    masters = {}
    rows = edit.axis == "rows"
    for group, master in source.masters.items():
        area = edit.range(parse_range(master.ref))
        line = edit.index(master.row if rows else master.column)
        if line is not None and area is not None:
            position = {"row": line} if rows else {"column": line}
            masters[group] = replace(master, ref=str(area), **position)
    return masters


def _ref_mover(edit: LineEdit) -> Callable[[re.Match[bytes]], bytes]:
    """What a shared formula's ``ref`` becomes."""

    def move(match: re.Match[bytes]) -> bytes:
        area = edit.range(parse_range(match[1].decode()))
        return b' ref="%b"' % str(area).encode()

    return move


def _row_mover(edit: LineEdit) -> Callable[[bytes, int], bytes]:
    """A function that moves a row of XML down by some rows: its number, its cells, and the
    range of a shared formula it holds."""
    move_ref = _ref_mover(edit)

    def move(row: bytes, delta: int) -> bytes:
        if not delta:
            return row
        row = _ROW_NUMBERS.sub(lambda m: m[1] + b"%d" % (int(m[2]) + delta) + m[3], row)
        return _REF.sub(move_ref, row) if b" ref=" in row else row

    return move


def _column_mover(
    edit: LineEdit, inherit: Callable[[int, int | None], int] | None
) -> Callable[[bytes], bytes]:
    """A function that moves the cells of a row of XML sideways: it drops those in deleted
    columns, renames the others and gives the cells inserted next to one its format."""
    move_ref = _ref_mover(edit)

    def move(row: bytes) -> bytes:
        found = list(_CELLS.finditer(row))
        if not found:
            return row
        cells = []
        for match in found:
            ref = _CELL_REF.match(match[0])
            assert ref is not None  # `split_sheet` leaves out sheets with cells without an address.
            cells.append((_column(ref[1]), ref, match[0]))
        styles = {column: _style(cell) for column, _, cell in cells} if inherit else {}
        kept = []
        for column, ref, cell in cells:
            line = edit.index(column)
            if line is None:
                continue
            if line != column:
                cell = cell[: ref.start(1)] + _letters(line) + cell[ref.end(1) :]
                cell = _REF.sub(move_ref, cell) if b" ref=" in cell else cell
            kept.append(cell)
            if inherit and column == edit.at - 1:
                style = inherit(styles[column], styles.get(edit.at))
                kept.extend(
                    b'<c r="%b%b" s="%d"/>' % (_letters(new), ref[2], style)
                    for new in range(edit.at, edit.at + edit.count)
                    if style
                )
        head = _SPANS.sub(b"", row[: found[0].start()])
        return head + b"".join(kept) + row[found[-1].end() :]

    return move


@lru_cache(maxsize=1024)
def _column(letters: bytes) -> int:
    return column_index_from_string(letters.decode())


@lru_cache(maxsize=1024)
def _letters(column: int) -> bytes:
    return get_column_letter(column).encode()


def _style(cell: bytes) -> int:
    found = _STYLE.search(cell, 0, cell.index(b">"))
    return int(found[1]) if found else 0


def _assembled(
    source: SheetSource, rows: list[Row], masters: dict[str, SharedMaster], hidden: list[int]
) -> SheetSource:
    spans = []
    offset = 0
    for _, row in rows:
        spans.append((offset, offset + len(row)))
        offset += len(row)
    data = b"".join(row for _, row in rows)
    return SheetSource(data, source.root, [number for number, _ in rows], spans, hidden, masters)
