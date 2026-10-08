"""Parsing and formatting of A1-style cell references."""

import re
from collections.abc import Callable
from dataclasses import dataclass

from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import CellRange as SheetRange

from excel_mcp.errors import InvalidArgumentError, LimitExceededError

MAX_ROW = 1_048_576
MAX_COLUMN = 16_384

_COLUMN_SPAN = re.compile(r"([A-Za-z]{1,3}):([A-Za-z]{1,3})")
_ROW_SPAN = re.compile(r"(\d+):(\d+)")


@dataclass(frozen=True)
class CellRange:
    """A rectangular block of cells, 1-based and inclusive."""

    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def rows(self) -> int:
        return self.max_row - self.min_row + 1

    @property
    def cols(self) -> int:
        return self.max_col - self.min_col + 1

    @property
    def size(self) -> int:
        return self.rows * self.cols

    @property
    def top_left(self) -> str:
        return cell_name(self.min_row, self.min_col)

    @classmethod
    def of(cls, area: SheetRange) -> "CellRange":
        return cls(area.min_row, area.min_col, area.max_row, area.max_col)

    def overlaps(self, other: "CellRange") -> bool:
        return not (
            self.max_row < other.min_row
            or other.max_row < self.min_row
            or self.max_col < other.min_col
            or other.max_col < self.min_col
        )

    def within(self, max_cells: int) -> "CellRange":
        """Return the range, or raise if it has more than ``max_cells`` cells."""
        if self.size > max_cells:
            raise LimitExceededError(
                f"Range {self} has {self.size:,} cells; at most {max_cells:,} can be "
                "processed per call. Use a smaller range."
            )
        return self

    def __str__(self) -> str:
        start = self.top_left
        end = cell_name(self.max_row, self.max_col)
        return start if start == end else f"{start}:{end}"


def cell_name(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def parse_range(ref: str) -> CellRange:
    """Parse a bounded range such as ``B2`` or ``A1:C10`` (``$`` markers allowed)."""
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref.strip())
    except (ValueError, TypeError):
        raise InvalidArgumentError(
            f"Invalid range {ref!r}. Use A1 notation such as 'B2' or 'A1:C10'."
        ) from None
    if min_col is None or min_row is None or max_col is None or max_row is None:
        raise InvalidArgumentError(
            f"Range {ref!r} must have explicit rows and columns, such as 'A1:C10'."
        )
    cell_range = CellRange(
        min_row=min(min_row, max_row),
        min_col=min(min_col, max_col),
        max_row=max(min_row, max_row),
        max_col=max(min_col, max_col),
    )
    _check_rows(cell_range.min_row, cell_range.max_row)
    if cell_range.max_col > MAX_COLUMN:
        raise InvalidArgumentError(
            f"Column {get_column_letter(cell_range.max_col)} is past "
            f"{get_column_letter(MAX_COLUMN)}, the last column."
        )
    return cell_range


def parse_clamped_range(ref: str, used: Callable[[], CellRange]) -> CellRange:
    """`parse_range`, also accepting whole columns ('B:D') or rows ('2:3') limited to the
    used range, which is only looked up for those."""
    text = ref.strip().replace("$", "")
    column_span = _COLUMN_SPAN.fullmatch(text)
    if column_span:
        first, last = sorted(_column_number(letters) for letters in column_span.groups())
        area = used()
        return CellRange(area.min_row, first, area.max_row, last)
    row_span = _ROW_SPAN.fullmatch(text)
    if row_span:
        first, last = sorted(int(number) for number in row_span.groups())
        _check_rows(first, last)
        area = used()
        return CellRange(first, area.min_col, last, area.max_col)
    return parse_range(ref)


def _check_rows(first: int, last: int) -> None:
    if first < 1:
        raise InvalidArgumentError("Row 0 is not valid; rows start at 1.")
    if last > MAX_ROW:
        raise InvalidArgumentError(f"Row {last} is past {MAX_ROW}, the last row.")


def _column_number(letters: str) -> int:
    number = column_index_from_string(letters.upper())
    if number > MAX_COLUMN:
        raise InvalidArgumentError(
            f"Column {letters.upper()} is past {get_column_letter(MAX_COLUMN)}, the last column."
        )
    return number


def parse_cell(ref: str) -> tuple[int, int]:
    """Parse a single cell such as ``B2`` into ``(row, column)``."""
    cell_range = parse_range(ref)
    if cell_range.size != 1:
        raise InvalidArgumentError(f"Expected a single cell such as 'B2', got {ref!r}.")
    return cell_range.min_row, cell_range.min_col
