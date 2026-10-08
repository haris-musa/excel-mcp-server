"""Structured references (``Sales[Price]``, ``[@Price]``) as the cells they stand for."""

from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.calc.parser import (
    ArrayLiteral,
    Binary,
    Call,
    ErrorLiteral,
    Node,
    Percent,
    Ref,
    TableRef,
    Unary,
)
from excel_mcp.calc.values import REF, VALUE
from excel_mcp.refs import parse_range
from excel_mcp.structured import StructuredRef
from excel_mcp.workspace import worksheets


def bind(node: Node, workbook: Workbook, sheet: Worksheet, row: int, col: int) -> Node:
    """The tree with each table reference replaced by the cells it stands for, as seen from
    the cell at (row, col) of ``sheet``."""

    def walk(item: Node) -> Node:
        match item:
            case TableRef():
                return _cells(item.ref, workbook, sheet, row, col)
            case Unary(op, operand):
                return Unary(op, walk(operand))
            case Percent(operand):
                return Percent(walk(operand))
            case Binary(op, left, right):
                return Binary(op, walk(left), walk(right))
            case Call(name, args, prefixed):
                return Call(name, tuple(walk(a) for a in args), prefixed)
            case ArrayLiteral(rows):
                return ArrayLiteral(tuple(tuple(walk(i) for i in line) for line in rows))
        return item

    return walk(node)


def _cells(
    found: StructuredRef, workbook: Workbook, sheet: Worksheet, row: int, col: int
) -> Ref | ErrorLiteral:
    located = _find(found.table, workbook, sheet, row, col)
    if located is None:
        return ErrorLiteral(REF.code)
    host, table = located
    area = parse_range(table.ref)
    header = area.min_row if table.headerRowCount != 0 else None
    totals = area.max_row if table.totalsRowCount else None
    first_data = area.min_row + (header is not None)
    last_data = area.max_row - (totals is not None)
    spans: list[tuple[int, int]] = []
    for item in found.items or ("#Data",):
        match item:
            case "#All":
                spans.append((area.min_row, area.max_row))
            case "#Data":
                spans.append((first_data, last_data))
            case "#Headers" if header is not None:
                spans.append((header, header))
            case "#Totals" if totals is not None:
                spans.append((totals, totals))
            case "#This Row":
                if not first_data <= row <= last_data or host is not sheet:
                    return ErrorLiteral(VALUE.code)
                spans.append((row, row))
            case _:
                return ErrorLiteral(REF.code)
    names = [column.name.casefold() for column in table.tableColumns]
    left, right = 0, len(names) - 1
    if found.first is not None:
        try:
            left = right = names.index(found.first.casefold())
            if found.last is not None:
                right = names.index(found.last.casefold())
        except ValueError:
            return ErrorLiteral(REF.code)
    left, right = sorted((left, right))
    top, bottom = min(s[0] for s in spans), max(s[1] for s in spans)
    return Ref((host.title,), top, area.min_col + left, bottom, area.min_col + right, False)


def _find(
    name: str | None, workbook: Workbook, sheet: Worksheet, row: int, col: int
) -> tuple[Worksheet, Table] | None:
    for host in worksheets(workbook):
        for table in host.tables.values():
            if name is not None and table.displayName.casefold() == name.casefold():
                return host, table
            if name is None and host is sheet:
                area = parse_range(table.ref)
                if area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col:
                    return host, table
    return None
