"""Functions that look at where cells are: INDIRECT, ROW, COLUMN, CELL and INFO."""

from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.calc.parser import Call, Name, Node, Ref, parse_reference, split_sheet
from excel_mcp.calc.r1c1 import to_a1
from excel_mcp.calc.registry import function
from excel_mcp.calc.values import (
    REF,
    VALUE,
    ExcelError,
    FormulaError,
    Grid,
    RefGrid,
    Scalar,
    UncalculableError,
    Value,
    to_bool,
    to_text,
)

if TYPE_CHECKING:
    from excel_mcp.calc.engine import Engine

_DEFAULT_COLUMN_WIDTH = 8
_ALIGNMENT_PREFIX = {
    None: "'",
    "general": "'",
    "left": "'",
    "right": '"',
    "center": "^",
    "fill": "\\",
}


def _text(engine: "Engine", node: Node) -> str:
    value = engine.scalar(node)
    if isinstance(value, ExcelError):
        raise FormulaError(value)
    return to_text(value)


@function("INDIRECT", kind="lazy")
def indirect(engine: "Engine", ref_text: Node, a1: Node | None = None) -> Value:
    """A reference to a cell or range of this workbook, given as text."""
    text = _text(engine, ref_text)
    qualifier, body = split_sheet(text)
    if a1 is not None and not to_bool(engine.scalar(a1)):
        body = to_a1(body, engine.here.row, engine.here.col)
    target = parse_reference(body if qualifier is None else f"{qualifier}!{body}")
    if isinstance(target, Name) and not engine.is_defined(target):
        raise FormulaError(REF)
    result = engine.eval(target)
    # A name that holds a constant is not a reference.
    if isinstance(target, Name) and not isinstance(result, RefGrid):
        raise FormulaError(REF)
    return result


@function("CELL", kind="lazy")
def cell(engine: "Engine", info_type: Node, reference: Node | None = None) -> Value:
    """Information about the top-left cell of a reference."""
    kind = _text(engine, info_type).lower()
    if not isinstance(reference, Ref) or reference.top is None or reference.left is None:
        raise UncalculableError("CELL without a cell reference")
    sheets = engine.sheets_of(reference)
    if len(sheets) != 1:
        raise UncalculableError("CELL of a 3-D reference")
    sheet, row, col = sheets[0], reference.top, reference.left
    stored = sheet._cells.get((row, col))
    stored = stored if isinstance(stored, Cell) else None
    match kind:
        case "address":
            if sheet is not engine.here.sheet:
                raise UncalculableError("CELL address on another sheet")
            return f"${get_column_letter(col)}${row}"
        case "row":
            return float(row)
        case "col":
            return float(col)
        case "contents":
            value = engine.cell_value(sheet, row, col)
            return 0.0 if value is None else value
        case "type":
            return _type(engine.cell_value(sheet, row, col), stored is None)
        case "filename":
            return _filename(engine, sheet.title)
        case "protect":
            return 0.0 if stored is not None and not stored.protection.locked else 1.0
        case "prefix":
            return _prefix(engine.cell_value(sheet, row, col), stored)
        case "width":
            return _width(sheet, col)
        case "format":
            return _plain_format(stored, "G")
        case "color" | "parentheses":
            return _plain_format(stored, 0.0)
    raise FormulaError(VALUE)


def _type(value: Scalar, empty: bool) -> str:
    if value is None or empty:
        return "b"
    return "l" if isinstance(value, str) else "v"


def _prefix(value: Scalar, stored: Cell | None) -> str:
    if not isinstance(value, str):
        return ""
    alignment = stored.alignment.horizontal if stored is not None else None
    if alignment not in _ALIGNMENT_PREFIX:
        raise UncalculableError("CELL prefix of text with this alignment")
    return _ALIGNMENT_PREFIX[alignment]


def _width(sheet: Worksheet, col: int) -> float:
    dimension = sheet.column_dimensions.get(get_column_letter(col))
    if dimension is not None and (dimension.width or dimension.hidden):
        raise UncalculableError("CELL width of a resized column")
    return float(_DEFAULT_COLUMN_WIDTH)


def _plain_format(stored: Cell | None, answer: Scalar) -> Scalar:
    """The answer for a cell without a number format; other formats are not reproduced."""
    if stored is not None and stored.number_format != "General":
        raise UncalculableError("CELL with a number format")
    return answer


def _filename(engine: "Engine", sheet: str) -> str:
    if engine.filename is None:
        raise UncalculableError("CELL filename without a file")
    name = Path(engine.filename).name
    # The folder is shown as the server shows paths, so the host's own folders stay private.
    return f"{engine.filename.removesuffix(name)}[{name}]{sheet}"


@function("INFO", kind="lazy")
def info(engine: "Engine", info_type: Node) -> Value:
    """Information about the workbook's calculation mode; host details are not reproduced."""
    kind = _text(engine, info_type).lower()
    if kind == "recalc":
        mode = engine.workbook.calculation.calcMode or "auto"
        if mode == "manual":
            return "Manual"
        if mode == "auto":
            return "Automatic"
        raise UncalculableError("INFO recalc with automatic recalculation except tables")
    if kind in {"directory", "numfile", "origin", "osversion", "release", "system"}:
        raise UncalculableError(f"INFO {kind} depends on the host")
    raise FormulaError(VALUE)


def _origin(engine: "Engine", node: Node | None, *, rows: bool) -> Ref:
    if node is None:
        here = engine.here
        return Ref((), here.row, here.col, here.row, here.col, False)
    if isinstance(node, Name):
        raise UncalculableError("ROW or COLUMN of a name")
    if isinstance(node, Call) and node.name in ("INDIRECT", "OFFSET"):
        return _located(engine.eval(node), rows=rows)
    if not isinstance(node, Ref) or node.top is None or node.left is None:
        raise UncalculableError("ROW or COLUMN of a computed range")
    return node


def _located(value: Value, *, rows: bool) -> Ref:
    """Where a reference returned by INDIRECT or OFFSET lies; whole rows or columns are cut off."""
    if isinstance(value, ExcelError):
        raise FormulaError(value)
    if not isinstance(value, RefGrid) or (value.clipped_rows if rows else value.clipped_cols):
        raise UncalculableError("ROW or COLUMN of a computed range")
    bottom, right = value.top + value.height - 1, value.left + value.width - 1
    return Ref((), value.top, value.left, bottom, right, False)


@function("ROW", kind="lazy")
def row(engine: "Engine", reference: Node | None = None) -> Value:
    ref = _origin(engine, reference, rows=True)
    assert ref.top is not None and ref.bottom is not None
    numbers = [[float(r)] for r in range(ref.top, ref.bottom + 1)]
    return numbers[0][0] if len(numbers) == 1 else Grid(numbers)  # pyright: ignore[reportReturnType]


@function("COLUMN", kind="lazy")
def column(engine: "Engine", reference: Node | None = None) -> Value:
    ref = _origin(engine, reference, rows=False)
    assert ref.left is not None and ref.right is not None
    cols = [float(c) for c in range(ref.left, ref.right + 1)]
    return cols[0] if len(cols) == 1 else Grid([cols])
