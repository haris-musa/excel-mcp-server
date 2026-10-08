"""Tables, PivotTables and drawings when rows or columns are inserted or deleted."""

from dataclasses import replace
from typing import cast

from openpyxl.cell.cell import Cell
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor
from openpyxl.workbook import Workbook
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.filters import FilterColumn as StoredColumn
from openpyxl.worksheet.table import Table, TableColumn
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.formulas import storable_formula
from excel_mcp.operations.comparison import FormulaResults
from excel_mcp.operations.filter_stored import stored_test
from excel_mcp.operations.filters import FormulaValues, hide_failing_rows
from excel_mcp.operations.pivot_index import pivot_area, sheet_pivots, workbook_pivots
from excel_mcp.package.lines import LineEdit
from excel_mcp.refs import CellRange, cell_name, parse_cell, parse_range


def _bounds(area: CellRange, edit: LineEdit) -> tuple[int, int]:
    return (area.min_row, area.max_row) if edit.axis == "rows" else (area.min_col, area.max_col)


def table_area_after(table: Table, edit: LineEdit) -> CellRange | None:
    """Where a table of the edited sheet ends up, or None when it is deleted."""
    area = parse_range(table.ref)
    if edit.axis == "rows" and edit.delete and edit.at <= area.min_row <= edit.end:
        if edit.end >= area.max_row:
            return None
        raise InvalidArgumentError(
            f"Rows {edit.at} to {edit.end} include the header row of table "
            f"{table.displayName!r} ({table.ref}) but not all of its rows. Delete the whole "
            "table's rows, or only rows below its header."
        )
    return edit.range(area)


def keep_table_row(sheet: Worksheet, edit: LineEdit) -> LineEdit:
    """Excel keeps one empty data row in a table whose data rows are all deleted."""
    if edit.axis != "rows" or not edit.delete:
        return edit
    for table in sheet.tables.values():
        area = parse_range(table.ref)
        last_data = area.max_row - (1 if table.totalsRowCount else 0)
        if edit.at == area.min_row + 1 and edit.end >= last_data:
            for column in range(area.min_col, area.max_col + 1):
                cell = sheet._cells.get((edit.at, column))
                if isinstance(cell, Cell):
                    cell.value = None
            return replace(edit, at=edit.at + 1, count=edit.count - 1)
    return edit


def check_table_cuts(sheet: Worksheet, edit: LineEdit) -> None:
    """Excel cannot change rows or columns that run through two tables at once."""
    cut = [
        table.displayName
        for table in sheet.tables.values()
        if edit.cuts(*_bounds(parse_range(table.ref), edit))
    ]
    if len(cut) > 1:
        raise InvalidArgumentError(
            f"The edit would cut through tables {', '.join(cut)} at once, which Excel does not "
            "allow. Edit lines that run through only one of them."
        )


def dead_tables(sheet: Worksheet, edit: LineEdit) -> dict[str, frozenset[str] | None]:
    """The tables, and table columns, that the edit deletes. Raises if it cuts a table header."""
    dead: dict[str, frozenset[str] | None] = {}
    for table in sheet.tables.values():
        area = table_area_after(table, edit)
        key = table.displayName.casefold()
        if area is None:
            dead[key] = None
        elif edit.axis == "columns" and edit.delete:
            first = parse_range(table.ref).min_col
            gone = {
                column.name.casefold()
                for index, column in enumerate(table.tableColumns, start=first)
                if edit.at <= index <= edit.end
            }
            if gone:
                dead[key] = frozenset(gone)
    return dead


def update_tables(
    sheet: Worksheet, edit: LineEdit, sheet_names: list[str], results: dict[int, FormulaResults]
) -> None:
    """Call after the cells have moved."""
    for table in list(sheet.tables.values()):
        old = parse_range(table.ref)
        area = table_area_after(table, edit)
        if area is None:
            del sheet.tables[table.displayName]
            continue
        if edit.axis == "columns":
            _update_columns(sheet, table, old, edit, area)
        if edit.axis == "rows" and not edit.delete and old.min_row < edit.at <= old.max_row:
            _fill_calculated_columns(sheet, table, old, edit, sheet_names)
        if edit.axis == "rows" and edit.delete and table.totalsRowCount:
            _drop_deleted_totals(table, old, edit)
        table.ref = str(area)
        if table.autoFilter is not None:
            update_filter(sheet, table.autoFilter, old, area, edit, results)


def _filters(sheet: Worksheet) -> list[tuple[AutoFilter, CellRange]]:
    """The filters of the sheet and its tables, with the cells they cover."""
    found = []
    if sheet.auto_filter.ref:
        found.append((sheet.auto_filter, parse_range(sheet.auto_filter.ref)))
    for table in sheet.tables.values():
        if table.autoFilter is not None:
            area = parse_range(table.ref)
            last = area.max_row - 1 if table.totalsRowCount else area.max_row
            found.append(
                (table.autoFilter, CellRange(area.min_row, area.min_col, last, area.max_col))
            )
    return found


def filter_results(
    sheet: Worksheet, edit: LineEdit, results: FormulaValues
) -> dict[int, FormulaResults]:
    """Formula results of the data of each filter that loses a criteria column but keeps others.

    Excel applies the remaining criteria again; this needs the results as they are now."""
    found: dict[int, FormulaResults] = {}
    if edit.axis != "columns" or not edit.delete:
        return found
    for auto_filter, area in _filters(sheet):
        lost = [c for c in auto_filter.filterColumn if edit.index(area.min_col + c.colId) is None]
        if lost and len(lost) < len(auto_filter.filterColumn):
            found[id(auto_filter)] = results(
                CellRange(area.min_row + 1, area.min_col, area.max_row, area.max_col)
            )
    return found


def _moved_results(results: FormulaResults, edit: LineEdit) -> FormulaResults:
    moved: FormulaResults = {}
    for name, value in results.items():
        row, column = parse_cell(name)
        line = edit.index(column)
        if line is not None:
            moved[cell_name(row, line)] = value
    return moved


def update_filter(
    sheet: Worksheet,
    auto_filter: AutoFilter,
    old: CellRange,
    new: CellRange | None,
    edit: LineEdit,
    results: dict[int, FormulaResults],
) -> None:
    """Move a filter's range, the columns that have criteria, and its sort."""
    if new is None:
        auto_filter.ref = None
        auto_filter.filterColumn = []
        auto_filter.sortState = None
        return
    if edit.axis == "columns":
        kept = []
        for column in auto_filter.filterColumn:
            index = edit.index(old.min_col + column.colId)
            if index is not None:
                column.colId = index - new.min_col
                kept.append(column)
        if len(kept) < len(auto_filter.filterColumn):
            _reapply(sheet, new, kept, _moved_results(results.get(id(auto_filter), {}), edit))
        auto_filter.filterColumn = kept
    sort = auto_filter.sortState
    if sort is not None:
        sort_area = edit.range(parse_range(sort.ref)) if sort.ref else None
        if sort_area is None:
            auto_filter.sortState = None
        else:
            sort.ref = str(sort_area)
            for condition in sort.sortCondition:
                moved = edit.range(parse_range(condition.ref))
                condition.ref = str(moved) if moved else str(sort_area)
    auto_filter.ref = str(new)


def _reapply(
    sheet: Worksheet, area: CellRange, columns: list[StoredColumn], values: FormulaResults
) -> None:
    """Show every row, then hide those the remaining criteria hide."""
    for row in range(area.min_row + 1, area.max_row + 1):
        if row in sheet.row_dimensions:
            sheet.row_dimensions[row].hidden = False
    tests = [(c.colId, stored_test(sheet, area, c, values)) for c in columns]
    hide_failing_rows(sheet, area, tests, values)


def _drop_deleted_totals(table: Table, old: CellRange, edit: LineEdit) -> None:
    """A table whose totals row is deleted no longer has one."""
    if edit.at <= old.max_row <= edit.end:
        table.totalsRowCount = None
        for column in table.tableColumns:
            column.totalsRowFunction = column.totalsRowLabel = column.totalsRowFormula = None


def _fill_calculated_columns(
    sheet: Worksheet, table: Table, old: CellRange, edit: LineEdit, sheet_names: list[str]
) -> None:
    """Rows inserted into a table get the formulas of its calculated columns."""
    for index, column in enumerate(table.tableColumns, start=old.min_col):
        formula = column.calculatedColumnFormula
        if formula is not None and formula.attr_text:
            text = storable_formula(f"={formula.attr_text}", sheet_names)
            for row in range(edit.at, edit.at + edit.count):
                cast(Cell, sheet.cell(row, index)).value = text


def _update_columns(
    sheet: Worksheet, table: Table, old: CellRange, edit: LineEdit, area: CellRange
) -> None:
    columns = table.tableColumns
    if edit.delete:
        table.tableColumns = [
            column
            for index, column in enumerate(columns, start=old.min_col)
            if not edit.at <= index <= edit.end
        ]
    elif old.min_col < edit.at <= old.max_col:
        names = {column.name.casefold() for column in columns}
        number = 0
        for offset in range(edit.count):
            while f"column{number + 1}" in names:
                number += 1
            number += 1
            name = f"Column{number}"
            names.add(name.casefold())
            column = TableColumn(id=max(c.id for c in columns) + 1 + offset, name=name)
            columns.insert(edit.at - old.min_col + offset, column)
            cast(Cell, sheet.cell(area.min_row, edit.at + offset)).value = name
        table.tableColumns = columns


def update_pivots(workbook: Workbook, sheet: Worksheet, edit: LineEdit) -> None:
    """Move the PivotTables of the sheet and re-point the sources on it."""
    pivots = sheet_pivots(sheet)
    for pivot in list(pivots):
        area = pivot_area(pivot)
        if edit.cuts(*_bounds(area, edit)):
            raise InvalidArgumentError(
                f"PivotTable {pivot.name!r} occupies {area}, and the edit would cut through it. "
                "Excel does not allow that either: delete the PivotTable (delete_pivot_table) "
                "or edit rows or columns outside it."
            )
        body = edit.range(parse_range(pivot.location.ref))
        if body is None:
            pivots.remove(pivot)
        else:
            pivot.location.ref = str(body)
    caches = {id(pivot.cache): pivot.cache for pivot in workbook_pivots(workbook)}
    for cache in caches.values():
        source = cache.cacheSource.worksheetSource
        if (
            source is None
            or not source.ref
            or (source.sheet or "").casefold() != edit.sheet.casefold()
        ):
            continue
        moved = edit.range(parse_range(source.ref))
        if moved is None:
            raise InvalidArgumentError(
                f"The edit deletes all of {source.sheet}!{source.ref}, the data of a PivotTable. "
                "Delete the PivotTable first."
            )
        source.ref = str(moved)


def update_anchors(sheet: Worksheet, edit: LineEdit) -> None:
    """Pictures and charts move with their cells; two-cell ones also grow and shrink."""
    for drawing in (*sheet._images, *sheet._charts):  # pyright: ignore[reportAttributeAccessIssue]
        anchor = drawing.anchor
        if isinstance(anchor, str):
            row, column = parse_cell(anchor)
            row, column = _moved_cell(row, column, edit)
            drawing.anchor = cell_name(row, column)
        elif isinstance(anchor, OneCellAnchor | TwoCellAnchor):
            _move_marker(anchor._from, edit, end=False)
            if isinstance(anchor, TwoCellAnchor) and anchor.editAs != "oneCell":
                _move_marker(anchor.to, edit, end=True)


def _moved_cell(row: int, column: int, edit: LineEdit) -> tuple[int, int]:
    if edit.axis == "rows":
        return edit.start(row), column
    return row, edit.start(column)


def _move_marker(marker, edit: LineEdit, *, end: bool) -> None:
    move = edit.stop if end else edit.start
    if edit.axis == "rows":
        marker.row = move(marker.row + 1) - 1
    else:
        marker.col = move(marker.col + 1) - 1
