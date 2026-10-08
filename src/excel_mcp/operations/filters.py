"""AutoFilter on a sheet or a table, with criteria applied the way Excel applies them.

Excel stores the criteria and also which rows are hidden, so after the criteria are set the
rows that fail them are hidden here too and the file opens filtered.
"""

from collections.abc import Callable

from openpyxl.worksheet import filters as xl
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import Field

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.inputs import InputModel
from excel_mcp.operations.comparison import FormulaResults, cell_value
from excel_mcp.operations.filter_criteria import FilterColumn, Test, column_filter
from excel_mcp.refs import CellRange, parse_range

FormulaValues = Callable[[CellRange], FormulaResults]


class AutoFilter(InputModel):
    range: str = Field(
        description="Header row and data, e.g. 'A1:F100', or the name of a table on the sheet."
    )
    filters: list[FilterColumn] = Field(
        default=[], description="Criteria; they replace earlier ones. A row must meet all."
    )
    remove: bool = Field(default=False, description="Remove the filter and show its rows.")


def apply_auto_filter(
    sheet: Worksheet, spec: AutoFilter, formula_values: FormulaValues, max_cells: int
) -> None:
    table = next(
        (t for t in sheet.tables.values() if t.displayName.casefold() == spec.range.casefold()),
        None,
    )
    area = _table_area(table) if table else parse_range(spec.range)
    if table is None:
        _check_not_in_table(sheet, area)
    _release_previous(sheet, table, area)
    if spec.remove:
        return
    if area.rows < 2:
        raise InvalidArgumentError("A filter needs a header row and at least one data row.")
    if spec.filters:
        area.within(max_cells)
    values = formula_values(area) if spec.filters else {}
    columns = [column_filter(sheet, area, item, values) for item in spec.filters]
    if table is None:
        auto_filter = sheet.auto_filter
        auto_filter.ref = str(area)
    else:
        auto_filter = table.autoFilter = xl.AutoFilter(ref=str(area))
    auto_filter.filterColumn = [column for column, _, _ in columns]
    _hide_failing_rows(sheet, area, [(index, test) for _, index, test in columns], values)


def _table_area(table: Table) -> CellRange:
    area = parse_range(table.ref)
    last = area.max_row - 1 if table.totalsRowCount else area.max_row
    return CellRange(area.min_row, area.min_col, last, area.max_col)


def _check_not_in_table(sheet: Worksheet, area: CellRange) -> None:
    for table in sheet.tables.values():
        if area.overlaps(parse_range(table.ref)):
            raise InvalidArgumentError(
                f"{area} overlaps table {table.displayName!r}; pass the table's name as the "
                "range to filter it."
            )


def _release_previous(sheet: Worksheet, table: Table | None, area: CellRange) -> None:
    """Show the rows an earlier filter hid. A sheet has one filter, so a new range replaces it."""
    if table is None and sheet.auto_filter.ref:
        area = parse_range(sheet.auto_filter.ref)
    for row in range(area.min_row + 1, area.max_row + 1):
        if (dimension := sheet.row_dimensions.get(row)) is not None:
            dimension.hidden = False
    if table is not None:
        table.autoFilter = None
    else:
        sheet.auto_filter = xl.AutoFilter()


def _hide_failing_rows(
    sheet: Worksheet,
    area: CellRange,
    tests: list[tuple[int, Test]],
    values: FormulaResults,
) -> None:
    for row in range(area.min_row + 1, area.max_row + 1):
        cells = [(sheet.cell(row, area.min_col + index), test) for index, test in tests]
        if any(not test(cell, cell_value(cell, values)) for cell, test in cells):
            sheet.row_dimensions[row].hidden = True
