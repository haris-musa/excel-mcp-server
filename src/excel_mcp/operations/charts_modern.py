"""Series and option checks for the Excel 2016 chart types (waterfall, histogram and so on)."""

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.operations.charts_check import MAX_SERIES, check_labels, require
from excel_mcp.operations.charts_data import (
    Plot,
    one_line,
    qualified,
    series_name,
    split_sheet,
)
from excel_mcp.operations.charts_options import (
    Bins,
    ChartOptions,
    ChartType,
    DataLabels,
    SeriesSpec,
)
from excel_mcp.refs import CellRange, parse_range

_MANY = ("box_whisker",)
_HIERARCHIES = ("treemap", "sunburst")
_LABELLED = ("waterfall", "histogram", "pareto")
_VALUE_AXIS = ("waterfall", "histogram", "pareto", "box_whisker")
_ONLY = {
    "totals": ("waterfall",),
    "connector_lines": ("waterfall",),
    "bins": ("histogram", "pareto"),
    "box": ("box_whisker",),
    "parent_labels": ("treemap",),
}
_CLASSIC_ONLY = (
    "title_size", "grouping", "colors", "markers", "smooth", "scatter_style", "style",
    "plot_color",
)  # fmt: skip
_SERIES_CLASSIC = (
    "sizes", "type", "secondary_axis", "color", "line_width_pt", "marker", "marker_size",
    "trendline", "error_bars",
)  # fmt: skip
_Y_AXIS = {"title", "min", "max", "major_unit", "number_format", "major_gridlines"}


def resolve_modern(
    workbook: Workbook,
    sheet: str,
    chart_type: ChartType,
    source: str | None,
    series: list[SeriesSpec],
    categories: str | None,
) -> list[Plot]:
    if (source is None) == (not series):
        raise InvalidArgumentError(
            "Pass either source (a block of data) or series (explicit ranges), not both "
            "and not neither."
        )
    if source is not None:
        series, categories = _block(workbook, sheet, chart_type, source), None
    plots = [_plot(workbook, sheet, chart_type, spec, categories) for spec in series]
    if chart_type not in _MANY and len(plots) != 1:
        raise InvalidArgumentError(f"A {chart_type} chart plots one series, got {len(plots)}.")
    return plots


def _plot(
    workbook: Workbook, sheet: str, chart_type: str, spec: SeriesSpec, categories: str | None
) -> Plot:
    values, points = one_line(workbook, sheet, spec.values, "values")
    labels = spec.categories or categories
    if labels is None:
        found = None
    elif chart_type in _HIERARCHIES:
        title, cells = split_sheet(workbook, sheet, labels)
        found = qualified(title, parse_range(cells))
    else:
        found = one_line(workbook, sheet, labels, "categories")[0]
    return Plot(spec, values, found, None, series_name(workbook, spec.name), points)


def _block(workbook: Workbook, default: str, chart_type: str, source: str) -> list[SeriesSpec]:
    """Series for a block with a header row, read as Excel reads it: leading text columns are
    labels (one, or the levels of a treemap or sunburst), the number columns are the series."""
    sheet, cells = split_sheet(workbook, default, source)
    area = parse_range(cells)
    if area.rows < 2:
        raise InvalidArgumentError("source needs a header row and at least one row of data.")
    labels = _text_columns(workbook[sheet], area)
    if labels == area.cols:
        raise InvalidArgumentError(f"source {source!r} holds no column of numbers.")
    if chart_type not in _HIERARCHIES:
        labels = min(labels, 1)

    def column(offset: int, header: bool) -> str:
        top = area.min_row + (0 if header else 1)
        col = area.min_col + offset
        return qualified(sheet, CellRange(top, col, top if header else area.max_row, col))

    levels = None
    if labels:
        body = CellRange(area.min_row + 1, area.min_col, area.max_row, area.min_col + labels - 1)
        levels = qualified(sheet, body)
    return [
        SeriesSpec(values=column(offset, False), name=column(offset, True), categories=levels)
        for offset in range(labels, area.cols)
    ]


def _text_columns(sheet: Worksheet, area: CellRange) -> int:
    """How many columns from the left of the block hold text below the header row."""
    count = 0
    for col in range(area.min_col, area.max_col + 1):
        values = (sheet.cell(row, col).value for row in range(area.min_row + 1, area.max_row + 1))
        if not any(isinstance(v, str) and not v.startswith("=") for v in values):
            break
        count += 1
    return count


def check_modern(plots: list[Plot], chart_type: ChartType, options: ChartOptions) -> None:
    if len(plots) > MAX_SERIES:
        raise InvalidArgumentError(f"A chart has at most {MAX_SERIES} series, got {len(plots)}.")
    _check_options(options, chart_type)
    for number, plot in enumerate(plots, start=1):
        try:
            _check_series(plot, chart_type)
        except InvalidArgumentError as error:
            raise InvalidArgumentError(f"Series {number}: {error}") from None
    plot = plots[0]
    if chart_type == "waterfall" and any(not 1 <= n <= plot.points for n in options.totals):
        raise InvalidArgumentError(
            f"totals are positions from 1 to {plot.points} (the points of the series)."
        )
    if plots[0].categories and chart_type == "pareto" and options.bins != Bins():
        raise InvalidArgumentError(
            "bins apply to a pareto chart of numbers; with categories it adds them up."
        )
    if options.bins.width is not None and options.bins.count is not None:
        raise InvalidArgumentError("bins: give a width or a count, not both.")
    if (
        options.bins.underflow is not None
        and options.bins.overflow is not None
        and options.bins.underflow >= options.bins.overflow
    ):
        raise InvalidArgumentError("bins.underflow must be below bins.overflow.")


def check_only_for(options: ChartOptions, chart_type: ChartType) -> None:
    """Reject the options of the Excel 2016 chart types on any other chart."""
    for name, kinds in _ONLY.items():
        require(_changed(options, name), name, chart_type, kinds)


def _check_options(options: ChartOptions, chart_type: ChartType) -> None:
    for name in _CLASSIC_ONLY:
        if _changed(options, name):
            raise InvalidArgumentError(f"{name} does not apply to {chart_type} charts.")
    x = {"title"} if chart_type in _VALUE_AXIS else set()
    y = _Y_AXIS if chart_type in _VALUE_AXIS else set()
    percent = {"title", "min", "max"} if chart_type == "pareto" else set()
    for name, axis, allowed in (
        ("x_axis", options.x_axis, x),
        ("y_axis", options.y_axis, y),
        ("secondary_y_axis", options.secondary_y_axis, percent),
    ):
        extra = {f for f in type(axis).model_fields if _changed(axis, f)} - allowed
        if extra:
            raise InvalidArgumentError(
                f"{name}.{sorted(extra)[0]} does not apply to {chart_type} charts. "
                f"Settable: {', '.join(sorted(allowed)) or 'nothing'}."
            )
    if options.data_labels:
        _check_labels(options.data_labels, chart_type)


def _check_series(plot: Plot, chart_type: ChartType) -> None:
    spec = plot.spec
    for name in _SERIES_CLASSIC:
        if _changed(spec, name):
            raise InvalidArgumentError(f"{name} does not apply to {chart_type} charts.")
    if spec.data_labels:
        _check_labels(spec.data_labels, chart_type)


def _check_labels(labels: DataLabels, chart_type: str) -> None:
    if labels.position is not None and chart_type not in _LABELLED:
        raise InvalidArgumentError(
            f"{chart_type} charts place their labels themselves; leave out position."
        )
    check_labels(labels, "column", False)


def _changed(model: BaseModel, field: str) -> bool:
    default = type(model).model_fields[field].get_default(call_default_factory=True)
    return getattr(model, field) != default
