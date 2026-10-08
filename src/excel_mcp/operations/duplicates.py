"""Remove Duplicates: keep the first of each set of rows that agree on the chosen columns."""

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.operations.columns import find_column
from excel_mcp.operations.comparison import FormulaResults, as_number, cell_value
from excel_mcp.operations.sorting import save_cells, write_rows
from excel_mcp.refs import CellRange


def remove_duplicates(
    sheet: Worksheet,
    area: CellRange,
    columns: list[str] | None,
    has_header: bool,
    results: FormulaResults,
) -> int:
    """Move the rows to keep up, empty the rows left at the bottom, and return how many went.

    Rows are equal when the chosen columns hold equal values: text ignores case, and text
    and numbers differ. Other columns do not matter. Like Excel, only the range's cells move.
    """
    if any(
        area.overlaps(CellRange(m.min_row, m.min_col, m.max_row, m.max_col))
        for m in sheet.merged_cells.ranges
    ):
        raise InvalidArgumentError(
            f"{area} contains merged cells, which cannot be processed. Unmerge them first."
        )
    first_row = area.min_row + has_header
    chosen = (
        [find_column(sheet, area, name, has_header) for name in columns]
        if columns
        else list(range(area.min_col, area.max_col + 1))
    )
    seen: set[tuple[object, ...]] = set()
    kept = []
    for cells in sheet.iter_rows(
        min_row=first_row, max_row=area.max_row, min_col=area.min_col, max_col=area.max_col
    ):
        key = tuple(_comparable(cells[col - area.min_col], results) for col in chosen)
        if key not in seen:
            seen.add(key)
            kept.append(save_cells(cells))
    removed = area.max_row - first_row + 1 - len(kept)
    write_rows(sheet, area.min_col, first_row, kept)
    for cells in sheet.iter_rows(
        min_row=first_row + len(kept),
        max_row=area.max_row,
        min_col=area.min_col,
        max_col=area.max_col,
    ):
        for cell in cells:
            cell.value = None
            cell.style = "Normal"
            cell.comment = None
            cell.hyperlink = None
    return removed


def _comparable(cell: Cell | MergedCell, results: FormulaResults) -> tuple[str, object]:
    value = cell_value(cell, results)
    if isinstance(value, bool):
        return "boolean", value
    if (number := as_number(value)) is not None:
        return "number", number
    if isinstance(value, str):
        return "text", value.casefold()
    return "blank", None
