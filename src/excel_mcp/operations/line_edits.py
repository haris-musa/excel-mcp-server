"""Inserting and deleting rows and columns, with every reference updated as Excel does."""

from copy import copy
from typing import cast

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.styles import Border, Side
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.merge import MergedCellRange
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError, UnsafeFormulaError
from excel_mcp.operations.cells import stored_cells, used_range
from excel_mcp.operations.comparison import FormulaResults
from excel_mcp.operations.filters import FormulaValues
from excel_mcp.operations.line_edit_objects import (
    check_table_cuts,
    dead_tables,
    filter_results,
    keep_table_row,
    update_anchors,
    update_filter,
    update_pivots,
    update_tables,
)
from excel_mcp.operations.line_edit_rules import update_rules
from excel_mcp.operations.shifting import Shifter
from excel_mcp.operations.spans import format_span, parse_span
from excel_mcp.operations.workbook_rewrite import (
    gated_operand,
    rewrite_charts,
    rewrite_formulas,
    rewrite_names,
)
from excel_mcp.package.lines import Axis, LineEdit
from excel_mcp.package.references import rewrite_lines
from excel_mcp.refs import CellRange, cell_name, parse_range
from excel_mcp.workspace import worksheets


def edit_lines(
    workbook: Workbook,
    sheet: Worksheet,
    axis: Axis,
    at: int,
    count: int,
    *,
    delete: bool,
    formula_values: FormulaValues,
) -> None:
    """Insert ``count`` rows or columns before ``at``, or delete them from ``at``."""
    edit = keep_table_row(sheet, LineEdit(sheet.title, axis, at, count, delete))
    if not delete:
        _check_room(sheet, edit)
    _check_arrays(sheet, edit)
    check_table_cuts(sheet, edit)
    shifter = Shifter(edit, dead_tables(sheet, edit))
    names = workbook.sheetnames
    results = filter_results(sheet, edit, formula_values)
    try:
        rewrite_formulas(workbook, shifter, names)
        rewrite_names(workbook, shifter, names)
        for other in worksheets(workbook):
            update_rules(other, shifter, names)
        rewrite_charts(workbook, shifter, names)
    except TokenizerError as error:
        raise UnsafeFormulaError(f"A formula could not be parsed: {error}.") from None
    update_pivots(workbook, sheet, edit)
    rewrite_lines(workbook, edit, lambda text, host: gated_operand(text, host, shifter, names))
    _move_cells(sheet, edit)
    _move_dimensions(sheet, edit)
    _move_merged_ranges(sheet, edit)
    _inherit_cell_styles(sheet, edit)
    update_tables(sheet, edit, names, results)
    update_anchors(sheet, edit)
    _move_sheet_settings(sheet, edit, results)


def _check_room(sheet: Worksheet, edit: LineEdit) -> None:
    area = used_range(sheet)
    needed = (area.max_row if edit.axis == "rows" else area.max_col) + edit.count
    if needed > edit.limit:
        raise InvalidArgumentError(
            f"Inserting would push data past the last of Excel's {edit.limit:,} {edit.axis}."
        )


def _check_arrays(sheet: Worksheet, edit: LineEdit) -> None:
    for cell in stored_cells(sheet):
        if isinstance(cell.value, ArrayFormula):
            area = parse_range(cell.value.ref)
            first, last = (
                (area.min_row, area.max_row)
                if edit.axis == "rows"
                else (area.min_col, area.max_col)
            )
            if edit.cuts(first, last):
                raise InvalidArgumentError(
                    f"The edit would change part of the array formula in {cell.value.ref}, "
                    "which Excel does not allow. Rewrite it as a normal formula first."
                )


def _move_cells(sheet: Worksheet, edit: LineEdit) -> None:
    rows = edit.axis == "rows"
    moved: dict[tuple[int, int], Cell | MergedCell] = {}
    for (row, column), cell in sheet._cells.items():
        index = edit.index(row if rows else column)
        if index is None:
            continue
        cell.row, cell.column = (index, column) if rows else (row, index)
        moved[cell.row, cell.column] = cell
        if isinstance(cell, Cell):
            if cell.hyperlink is not None:
                cell.hyperlink.ref = cell.coordinate
            if isinstance(cell.value, ArrayFormula):
                cell.value.ref = str(edit.range(parse_range(cell.value.ref)))
    sheet._cells = moved
    sheet._current_row = sheet.max_row  # pyright: ignore[reportAttributeAccessIssue]


def _move_dimensions(sheet: Worksheet, edit: LineEdit) -> None:
    if edit.axis == "rows":
        rows = dict(sheet.row_dimensions)
        sheet.row_dimensions.clear()
        for index, dimension in rows.items():
            moved = edit.index(index)
            if moved is not None:
                dimension.index = moved
                sheet.row_dimensions[moved] = dimension
        _move_breaks(sheet.row_breaks.brk, edit)
        _inherit_row_dimension(sheet, edit)
        return
    columns = dict(sheet.column_dimensions)
    sheet.column_dimensions.clear()
    for letter, dimension in columns.items():
        first = dimension.min or column_index_from_string(letter)
        span = edit.span(first, dimension.max or first)
        if span is not None:
            dimension.index = get_column_letter(span[0])
            dimension.min, dimension.max = span
            sheet.column_dimensions[dimension.index] = dimension
    _move_breaks(sheet.col_breaks.brk, edit)
    _inherit_column_dimensions(sheet, edit)


def _inherit_cell_styles(sheet: Worksheet, edit: LineEdit) -> None:
    """Inserted cells are formatted like the cells above or to the left, as in Excel's default.

    Excel keeps a side border only where the cell on the far side of the insertion has the
    same one, and the borders along the insertion belong to the neighbours.
    """
    if edit.delete:
        return
    rows = edit.axis == "rows"
    for (row, column), cell in list(sheet._cells.items()):
        if (row if rows else column) != edit.at - 1:
            continue
        beyond = edit.at + edit.count
        after = sheet._cells.get((beyond, column) if rows else (row, beyond))
        for line in range(edit.at, edit.at + edit.count):
            clone = sheet.cell(line, column) if rows else sheet.cell(row, line)
            clone._style = copy(cell._style)
            clone.border = _inherited_border(
                cast(Border, copy(cell.border)),
                cast(Border, copy(after.border)) if after else None,
                rows,
            )


def _inherited_border(above: Border, after: Border | None, rows: bool) -> Border:
    sides = {name: Side() for name in ("left", "right", "top", "bottom")}
    across = ("left", "right") if rows else ("top", "bottom")
    trailing, leading = ("bottom", "top") if rows else ("right", "left")
    if after is not None:
        for name in across:
            if getattr(above, name).style == getattr(after, name).style:
                sides[name] = getattr(above, name)
        if getattr(above, trailing).style == getattr(after, leading).style:
            sides[trailing] = sides[leading] = getattr(above, trailing)
    return Border(
        left=sides["left"], right=sides["right"], top=sides["top"], bottom=sides["bottom"]
    )


def _inherit_row_dimension(sheet: Worksheet, edit: LineEdit) -> None:
    """Inserted rows take the height and outline level of their neighbour, but are not hidden."""
    neighbour = sheet.row_dimensions.get(edit.at - 1)
    if edit.delete or neighbour is None:
        return
    for index in range(edit.at, edit.at + edit.count):
        clone = copy(neighbour)
        clone.index, clone.hidden = index, False
        sheet.row_dimensions[index] = clone


def _inherit_column_dimensions(sheet: Worksheet, edit: LineEdit) -> None:
    """Inserted columns take the width of their neighbour, but are not hidden."""
    if edit.delete:
        return
    new = range(edit.at, edit.at + edit.count)
    source = edit.at - 1
    for dimension in list(sheet.column_dimensions.values()):
        first, last = dimension.min or 0, dimension.max or 0
        if not first <= source <= last:
            continue
        spans_insertion = first < edit.at <= new[-1] < last
        if spans_insertion and not dimension.hidden:
            continue  # An unhidden definition that spans the insertion already covers it.
        tail = copy(dimension)
        dimension.max = edit.at - 1
        inserted = copy(dimension)
        inserted.min, inserted.max, inserted.hidden = new[0], new[-1], False
        inserted.index = get_column_letter(new[0])
        sheet.column_dimensions[inserted.index] = inserted
        if spans_insertion:
            tail.min = new[-1] + 1
            tail.index = get_column_letter(tail.min)
            sheet.column_dimensions[tail.index] = tail


def _move_breaks(breaks: list[Break], edit: LineEdit) -> None:
    for page_break in list(breaks):
        moved = edit.index(page_break.id or 0)
        if moved is None:
            breaks.remove(page_break)
        else:
            page_break.id = moved


def _move_merged_ranges(sheet: Worksheet, edit: LineEdit) -> None:
    areas = [edit.range(CellRange.of(merged)) for merged in sheet.merged_cells.ranges]
    sheet.merged_cells = MultiCellRange()
    for area in areas:
        if area is None or area.size == 1:
            continue
        sheet.merged_cells.add(MergedCellRange(sheet, str(area)))
        top_left = (area.min_row, area.min_col)
        for row in range(area.min_row, area.max_row + 1):
            for column in range(area.min_col, area.max_col + 1):
                if (row, column) != top_left and (row, column) not in sheet._cells:
                    sheet._cells[row, column] = MergedCell(sheet, row, column)


def _move_sheet_settings(
    sheet: Worksheet, edit: LineEdit, results: dict[int, FormulaResults]
) -> None:
    _move_freeze_panes(sheet, edit)
    _move_print_settings(sheet, edit)
    if sheet.auto_filter.ref:
        old = parse_range(sheet.auto_filter.ref)
        update_filter(sheet, sheet.auto_filter, old, edit.range(old), edit, results)


def _move_freeze_panes(sheet: Worksheet, edit: LineEdit) -> None:
    pane = sheet.sheet_view.pane
    if pane is None or pane.state not in ("frozen", "frozenSplit"):
        return
    rows = edit.axis == "rows"
    frozen = int((pane.ySplit if rows else pane.xSplit) or 0)
    if not frozen:
        return
    span = edit.span(1, frozen)
    now = 1 if span is None else span[1]  # Excel keeps one frozen line when all are deleted.
    if rows:
        pane.ySplit = now
    else:
        pane.xSplit = now
    pane.topLeftCell = cell_name(int(pane.ySplit or 0) + 1, int(pane.xSplit or 0) + 1)


def _move_print_settings(sheet: Worksheet, edit: LineEdit) -> None:
    areas = [edit.range(CellRange.of(area)) for area in sheet._print_area.ranges]  # pyright: ignore[reportAttributeAccessIssue]
    sheet.print_area = [str(area) for area in areas if area]
    rows = edit.axis == "rows"
    current = sheet.print_title_rows if rows else sheet.print_title_cols
    if not current:
        return
    span = edit.span(*parse_span(current.replace("$", ""), edit.axis))
    # openpyxl's setters ignore None, so clearing titles needs the attributes themselves.
    if rows:
        sheet._print_rows = None  # pyright: ignore[reportAttributeAccessIssue]
    else:
        sheet._print_cols = None  # pyright: ignore[reportAttributeAccessIssue]
    if span:
        if rows:
            sheet.print_title_rows = format_span(*span, "rows")
        else:
            sheet.print_title_cols = format_span(*span, "columns")
