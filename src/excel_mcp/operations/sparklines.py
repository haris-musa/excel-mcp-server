"""Sparklines: adding, replacing, listing and deleting groups, as Insert > Sparklines does.

Excel keeps them in the sheet's Excel 2010 extension, which openpyxl drops; the package
layer carries that text through edits (see `excel_mcp.package`).
"""

import re
from collections.abc import Callable
from typing import Literal, cast

from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.formulas import split_top_level, unquote
from excel_mcp.operations.sparkline_style import SparklineInfo, SparklineStyle, check_style
from excel_mcp.operations.sparkline_xml import build_group, read_groups
from excel_mcp.operations.x14_xml import X14, XM, sheet_prefix
from excel_mcp.package import state_of
from excel_mcp.package.extensions import SPARKLINES, edit_sparklines
from excel_mcp.refs import CellRange, cell_name, parse_range
from excel_mcp.workspace import get_sheet

_LOCATION = re.compile(r"<xm:sqref>([^<]*)</xm:sqref>")
_GROUPS = re.compile(r"<x14:sparklineGroups\b[^>]*>")


def add_sparklines(
    sheet: Worksheet, cells: str, source: str, style: SparklineStyle, max_cells: int
) -> tuple[CellRange, int]:
    """Add a group of sparklines, replacing those already in the cells.

    Returns the cells and how many sparklines they replaced."""
    check_style(style)
    workbook = cast(Workbook, sheet.parent)
    area = parse_range(cells).within(max_cells)
    if min(area.rows, area.cols) != 1:
        raise InvalidArgumentError("The range must be one row or one column of cells.")
    title, data = _source(workbook, sheet.title, source)
    lines = _lines(area, data)
    if style.dates:
        style = style.model_copy(
            update={"dates": _dates(workbook, sheet.title, style.dates, lines)}
        )
    prefix = sheet_prefix(title)
    group = build_group(style, [(cell, f"{prefix}!{line}") for cell, line in lines.items()])
    extensions = state_of(workbook).sheet(sheet).extensions
    replaced = _delete(extensions, lambda cell: cell in lines)
    _insert(extensions, group)
    return area, replaced


def delete_sparklines(sheet: Worksheet, area: CellRange) -> int:
    """Delete the sparklines in the cells of ``area``; the others in their groups stay."""
    extensions = state_of(cast(Workbook, sheet.parent)).sheet(sheet).extensions
    removed = _delete(extensions, lambda cell: area.overlaps(parse_range(cell)))
    if not removed:
        raise InvalidArgumentError(f"There are no sparklines in {sheet.title}!{area}.")
    return removed


def list_sparklines(sheet: Worksheet) -> list[SparklineInfo]:
    package = state_of(cast(Workbook, sheet.parent)).sheets.get(sheet)
    xml = package.extensions.get(SPARKLINES) if package else None
    return read_groups(xml) if xml else []


def _source(workbook: Workbook, default: str, text: str) -> tuple[str, CellRange]:
    """The sheet title and cells of ``Data!B2:F9`` or ``B2:F9`` (on the sheet itself)."""
    parts = split_top_level(text.strip(), "!")
    if len(parts) > 2:
        raise InvalidArgumentError(f"Invalid source {text!r}. Use 'B2:F9' or 'Data!B2:F9'.")
    title = get_sheet(workbook, unquote(parts[0])).title if len(parts) == 2 else default
    return title, parse_range(parts[-1])


def _lines(area: CellRange, source: CellRange) -> dict[str, CellRange]:
    """Each location cell with the row or column of data it plots, as Excel pairs them: by
    row if the data has as many rows as there are cells, otherwise by column."""
    cells = [
        cell_name(row, col)
        for row in range(area.min_row, area.max_row + 1)
        for col in range(area.min_col, area.max_col + 1)
    ]
    if area.size == source.rows:
        rows = range(source.min_row, source.max_row + 1)
        return dict(
            zip(cells, (CellRange(r, source.min_col, r, source.max_col) for r in rows), strict=True)
        )
    if area.size == source.cols:
        cols = range(source.min_col, source.max_col + 1)
        return dict(
            zip(cells, (CellRange(source.min_row, c, source.max_row, c) for c in cols), strict=True)
        )
    raise InvalidArgumentError(
        f"{area.size} cells need {area.size} rows or columns of data, but {source} has "
        f"{source.rows} rows and {source.cols} columns."
    )


def _span(source: CellRange, axis: Literal["rows", "cols"]) -> range:
    if axis == "rows":
        return range(source.min_row, source.max_row + 1)
    return range(source.min_col, source.max_col + 1)


def _dates(workbook: Workbook, default: str, text: str, lines: dict[str, CellRange]) -> str:
    title, dates = _source(workbook, default, text)
    points = next(iter(lines.values())).size
    if min(dates.rows, dates.cols) != 1 or dates.size != points:
        raise InvalidArgumentError(f"dates must be one row or column of {points} cells.")
    return f"{sheet_prefix(title)}!{dates}"


def _delete(extensions: dict[str, str], doomed: Callable[[str], bool]) -> int:
    removed = 0

    def handle(item: str) -> list[str]:
        nonlocal removed
        location = _LOCATION.search(item)
        assert location is not None
        if doomed(location[1]):
            removed += 1
            return []
        return [item]

    edit_sparklines(extensions, handle)
    return removed


def _insert(extensions: dict[str, str], group: str) -> None:
    """Add a group as Excel does: ahead of the others."""
    xml = extensions.get(SPARKLINES)
    if xml is None:
        extensions[SPARKLINES] = (
            f'<ext uri="{SPARKLINES}" xmlns:x14="{X14}">'
            f'<x14:sparklineGroups xmlns:xm="{XM}">{group}</x14:sparklineGroups></ext>'
        )
    else:
        extensions[SPARKLINES] = _GROUPS.sub(lambda m: m[0] + group, xml, count=1)
