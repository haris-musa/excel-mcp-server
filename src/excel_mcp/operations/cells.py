"""Reading, writing, clearing, copying and searching cell contents."""

from collections.abc import Iterator

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from excel_mcp.errors import InvalidArgumentError, LimitExceededError
from excel_mcp.operations.hyperlinks import Link, set_link
from excel_mcp.operations.spill import spill_formulas
from excel_mcp.refs import (
    MAX_COLUMN,
    MAX_ROW,
    CellRange,
    cell_name,
    parse_cell,
    parse_clamped_range,
    parse_range,
)
from excel_mcp.spill import show_spills
from excel_mcp.values import CellValue, date_number_format, to_cell, to_json, typed_value
from excel_mcp.workspace import sheet_names


class RangeData(BaseModel):
    range: str
    values: list[list[CellValue]]
    next_range: str | None = None
    uncalculated: dict[str, str] | None = None


class WriteResult(BaseModel):
    sheet: str
    range: str
    cells_written: int
    blocked: list[str] | None = None


class FindResult(BaseModel):
    matches: dict[str, dict[str, CellValue]]
    truncated: bool = False


def stored_cells(sheet: Worksheet) -> Iterator[Cell]:
    """Cells that hold a value, row by row.

    Iterating the stored cells directly avoids creating empty cells for every
    position up to max_row x max_column, as iter_rows() would.
    """
    for _, cell in sorted(sheet._cells.items()):
        if cell.value is not None:
            yield cell


def used_range(sheet: Worksheet) -> CellRange:
    """The smallest range containing every cell that has a value.

    Worksheet.max_row/max_column also count cells that only carry formatting.
    """
    cells = list(stored_cells(sheet))
    if not cells:
        return CellRange(1, 1, 1, 1)
    rows = [cell.row for cell in cells]
    cols = [cell.column for cell in cells]
    return CellRange(min(rows), min(cols), max(rows), max(cols))


def streamed_cells(sheet: ReadOnlyWorksheet) -> Iterator[tuple[int, int, object]]:
    """Cells that hold a value as (row, column, value), row by row, in constant memory."""
    # A wrong <dimension> in the file would otherwise cut the pass short.
    sheet.reset_dimensions()
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        for col_number, value in enumerate(row, start=1):
            if value is not None:
                yield row_number, col_number, value


def streamed_used_range(sheet: ReadOnlyWorksheet) -> CellRange:
    """Like used_range, for a streamed sheet: one full pass over its cells."""
    min_row = min_col = MAX_ROW + 1
    max_row = max_col = 0
    for row, col, _ in streamed_cells(sheet):
        min_row, max_row = min(min_row, row), max(max_row, row)
        min_col, max_col = min(min_col, col), max(max_col, col)
    if max_row == 0:
        return CellRange(1, 1, 1, 1)
    return CellRange(min_row, min_col, max_row, max_col)


def writable_cell(sheet: Worksheet, row: int, col: int) -> Cell:
    if row > MAX_ROW or col > MAX_COLUMN:
        raise InvalidArgumentError(f"{cell_name(row, col)} is outside the worksheet limits.")
    cell = sheet.cell(row, col)
    if isinstance(cell, MergedCell):
        raise InvalidArgumentError(
            f"{cell.coordinate} is inside a merged range; write to its top-left cell."
        )
    return cell


def store_value(cell: Cell, value: CellValue) -> None:
    """Store a computed value, keeping text as text even if it starts with '='."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def read_window(sheet: ReadOnlyWorksheet, ref: str | None) -> CellRange:
    """The range to read: ``ref`` (whole columns and rows allowed), else the used range."""
    if ref is None:
        return streamed_used_range(sheet)
    return parse_clamped_range(ref, lambda: streamed_used_range(sheet))


def _displayed(value: CellValue) -> CellValue:
    return show_spills(value) if isinstance(value, str) and value.startswith("=") else value


def store_typed(cell: Cell, text: str, names: list[str]) -> None:
    """Store ``text`` as if it were typed into the cell: a number, date, formula or text.

    Cells formatted as Text keep it as text; a General cell takes the number format the
    typed value implies (a percent sign gives a percentage).
    """
    if cell.number_format == "@":
        store_value(cell, text)
        return
    value, number_format = typed_value(text, names)
    if text.startswith("="):
        cell.value = value  # pyright: ignore[reportArgumentType]
    else:
        store_value(cell, value)  # pyright: ignore[reportArgumentType]
    if number_format and cell.number_format == "General":
        cell.number_format = number_format


def read_range(sheet: ReadOnlyWorksheet, ref: str | None, max_cells: int) -> RangeData:
    target = read_window(sheet, ref)
    if target.cols > max_cells:
        raise LimitExceededError(
            f"Range {target} has {target.cols} columns; at most {max_cells} cells can be "
            "read per call. Read fewer columns."
        )
    last_row = min(target.max_row, target.min_row + max_cells // target.cols - 1)
    rows = sheet.iter_rows(
        min_row=target.min_row,
        max_row=last_row,
        min_col=target.min_col,
        max_col=target.max_col,
        values_only=True,
    )
    values = [[_displayed(to_json(value)) for value in row] for row in rows]
    for row in values:
        while row and row[-1] is None:
            row.pop()
    while values and not values[-1]:
        values.pop()
    read = CellRange(target.min_row, target.min_col, last_row, target.max_col)
    remaining = CellRange(last_row + 1, target.min_col, target.max_row, target.max_col)
    return RangeData(
        range=str(read),
        values=values,
        next_range=str(remaining) if last_row < target.max_row else None,
    )


def write_range(
    sheet: Worksheet,
    at: str,
    rows: list[list[CellValue]],
    links: list[Link],
    max_cells: int,
) -> WriteResult:
    start_row, start_col = parse_cell(at)
    if not any(rows):
        raise InvalidArgumentError("rows must contain at least one value.")
    count = sum(len(row) for row in rows)
    if count > max_cells:
        raise LimitExceededError(f"{count:,} cells exceeds the limit of {max_cells:,} per call.")
    width = max(len(row) for row in rows)
    written = CellRange(start_row, start_col, start_row + len(rows) - 1, start_col + width - 1)
    if written.max_row > MAX_ROW or written.max_col > MAX_COLUMN:
        raise InvalidArgumentError(f"Writing {written} would go past the worksheet limits.")

    for link in links:
        row, col = parse_cell(link.cell)
        if not (
            written.min_row <= row <= written.max_row and written.min_col <= col <= written.max_col
        ):
            raise InvalidArgumentError(f"Link cell {link.cell} is outside the written {written}.")
    names = sheet_names(sheet)
    converted = [[to_cell(value, names) for value in row] for row in rows]
    formulas = []
    for row_offset, row in enumerate(converted):
        for col_offset, value in enumerate(row):
            cell = writable_cell(sheet, start_row + row_offset, start_col + col_offset)
            cell.value = value
            if number_format := date_number_format(value):
                cell.number_format = number_format
            if cell.data_type == "f":
                formulas.append(cell)
    for link in links:
        set_link(writable_cell(sheet, *parse_cell(link.cell)), link)
    blocked = spill_formulas(sheet, formulas, max_cells)
    return WriteResult(
        sheet=sheet.title, range=str(written), cells_written=count, blocked=blocked or None
    )


def clear_range(sheet: Worksheet, ref: str, contents: bool, formats: bool, max_cells: int) -> str:
    target = parse_range(ref).within(max_cells)
    for row in sheet.iter_rows(
        min_row=target.min_row,
        max_row=target.max_row,
        min_col=target.min_col,
        max_col=target.max_col,
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if contents:
                cell.value = None
            if formats:
                cell.style = "Normal"
            if contents and formats:
                cell.hyperlink = None
    return str(target)


def find_cells(
    sheets: list[ReadOnlyWorksheet],
    query: str,
    exact: bool,
    case_sensitive: bool,
    max_results: int,
) -> FindResult:
    def normalize(text: str) -> str:
        return text if case_sensitive else text.casefold()

    needle = normalize(query)
    matches: dict[str, dict[str, CellValue]] = {}
    count = 0
    for sheet in sheets:
        for row, col, value in streamed_cells(sheet):
            text = normalize(str(value))
            if not (text == needle if exact else needle in text):
                continue
            if count == max_results:
                return FindResult(matches=matches, truncated=True)
            matches.setdefault(sheet.title, {})[cell_name(row, col)] = to_json(value)
            count += 1
    return FindResult(matches=matches)
