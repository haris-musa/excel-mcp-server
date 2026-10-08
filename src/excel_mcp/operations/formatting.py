"""Cell formatting and merging."""

import re
from copy import copy
from typing import Literal

from openpyxl.styles import Alignment, Border, PatternFill, Protection, Side
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import Field

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.inputs import InputModel
from excel_mcp.refs import parse_range

HorizontalAlignment = Literal["general", "left", "center", "right", "fill", "justify"]
VerticalAlignment = Literal["top", "center", "bottom", "justify"]
BorderStyle = Literal["none", "thin", "medium", "thick", "double", "dashed", "dotted"]

_HEX_COLOR = re.compile(r"#?([0-9A-Fa-f]{6})")


class CellFormat(InputModel):
    """Fields left out keep the cell's current setting."""

    bold: bool | None = None
    italic: bool | None = None
    underline: bool | None = None
    strikethrough: bool | None = None
    font_name: str | None = Field(default=None, description="e.g. 'Calibri'.")
    font_size: float | None = Field(default=None, gt=0, le=409, description="Points.")
    font_color: str | None = Field(default=None, description="Hex, e.g. '#1F4E78'.")
    fill_color: str | None = Field(default=None, description="Hex.")
    number_format: str | None = Field(
        default=None, description="Excel code: '#,##0.00', '0%', 'yyyy-mm-dd' or '@' (text)."
    )
    horizontal_alignment: HorizontalAlignment | None = None
    vertical_alignment: VerticalAlignment | None = None
    wrap_text: bool | None = None
    border_style: BorderStyle | None = Field(
        default=None, description="On all four sides of every cell; 'none' removes it."
    )
    border_color: str | None = Field(default=None, description="Hex. Default: black.")
    locked: bool | None = Field(
        default=None,
        description="Locked cells (the default) cannot be edited on a protected sheet.",
    )
    formula_hidden: bool | None = Field(
        default=None, description="Hide the formula in the formula bar on a protected sheet."
    )


def parse_color(value: str) -> str:
    """Accept ``RRGGBB`` or ``#RRGGBB`` and return openpyxl's ``FFRRGGBB``."""
    match = _HEX_COLOR.fullmatch(value.strip())
    if not match:
        raise InvalidArgumentError(f"Invalid color {value!r}. Use a hex color like '#1F4E78'.")
    return f"FF{match.group(1).upper()}"


def format_range(sheet: Worksheet, ref: str, style: CellFormat, max_cells: int) -> str:
    target = parse_range(ref).within(max_cells)
    fill = _fill(style)
    border = _border(style)
    for row in sheet.iter_rows(
        min_row=target.min_row,
        max_row=target.max_row,
        min_col=target.min_col,
        max_col=target.max_col,
    ):
        for cell in row:
            cell.font = _font(cell.font, style)
            cell.alignment = _alignment(cell.alignment, style)
            cell.protection = _protection(cell.protection, style)
            if fill:
                cell.fill = fill
            if border:
                cell.border = border
            if style.number_format is not None:
                cell.number_format = style.number_format
    return str(target)


def _font(current, style: CellFormat):
    font = copy(current)
    if style.bold is not None:
        font.bold = style.bold
    if style.italic is not None:
        font.italic = style.italic
    if style.underline is not None:
        font.underline = "single" if style.underline else None
    if style.strikethrough is not None:
        font.strike = style.strikethrough
    if style.font_name is not None:
        font.name = style.font_name
        font.scheme = None  # a scheme makes Excel use the theme's font instead of the name
    if style.font_size is not None:
        font.size = style.font_size
    if style.font_color is not None:
        font.color = parse_color(style.font_color)
    return font


def _alignment(current, style: CellFormat) -> Alignment:
    alignment = copy(current)
    if style.horizontal_alignment is not None:
        alignment.horizontal = style.horizontal_alignment
    if style.vertical_alignment is not None:
        alignment.vertical = style.vertical_alignment
    if style.wrap_text is not None:
        alignment.wrap_text = style.wrap_text
    return alignment


def _protection(current, style: CellFormat) -> Protection:
    protection = copy(current)
    if style.locked is not None:
        protection.locked = style.locked
    if style.formula_hidden is not None:
        protection.hidden = style.formula_hidden
    return protection


def _fill(style: CellFormat) -> PatternFill | None:
    if style.fill_color is None:
        return None
    color = parse_color(style.fill_color)
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def _border(style: CellFormat) -> Border | None:
    if style.border_style is None:
        return None
    if style.border_style == "none":
        return Border()
    color = parse_color(style.border_color or "000000")
    side = Side(style=style.border_style, color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def merge_cells(sheet: Worksheet, ref: str, max_cells: int) -> str:
    target = parse_range(ref).within(max_cells)
    if target.size == 1:
        raise InvalidArgumentError("Merging needs a range of at least two cells.")
    for merged in sheet.merged_cells.ranges:
        if target.overlaps(parse_range(merged.coord)):
            raise InvalidArgumentError(f"{target} overlaps the merged range {merged.coord}.")
    sheet.merge_cells(str(target))
    return str(target)


def unmerge_cells(sheet: Worksheet, ref: str, max_cells: int) -> str:
    target = parse_range(ref).within(max_cells)
    if str(target) not in {str(merged) for merged in sheet.merged_cells.ranges}:
        raise InvalidArgumentError(f"{target} is not a merged range.")
    sheet.unmerge_cells(str(target))
    return str(target)
