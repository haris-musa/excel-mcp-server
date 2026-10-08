"""Fill Down, Fill Right and Fill Series."""

import calendar
import datetime as dt
from typing import Literal

from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.operations.cells import writable_cell
from excel_mcp.operations.paste import copy_range
from excel_mcp.refs import MAX_COLUMN, MAX_ROW, CellRange, cell_name
from excel_mcp.values import parse_iso_date

Direction = Literal["down", "right"]
Series = Literal["copy", "linear", "growth", "date"]
DateUnit = Literal["day", "weekday", "month", "year"]


def fill_range(
    sheet: Worksheet,
    area: CellRange,
    direction: Direction,
    series: Series,
    step: float,
    stop: str | float | None,
    unit: DateUnit,
    max_cells: int,
) -> CellRange:
    """Fill from the first row (down) or first column (right) of ``area`` across the rest.

    ``copy`` repeats the first line as Fill Down and Fill Right do, shifting relative
    references. Series continue each seed cell by ``step``. A single seed cell with a
    ``stop`` fills as far as the stop value allows. Returns the range filled.
    """
    down = direction == "down"
    if series == "copy":
        return _copy_fill(sheet, area, down, max_cells)
    if area.size == 1:
        if stop is None:
            raise InvalidArgumentError("A series from one cell needs a range of cells or a stop.")
        area = _extend(sheet, area, down, series, step, stop, unit, max_cells)
    count = area.rows if down else area.cols
    if count < 2:
        raise InvalidArgumentError(f"Select at least two {'rows' if down else 'columns'} to fill.")
    limit = _limit(series, stop, unit)
    for first_row, first_col in _seed_positions(area, down):
        seed = sheet.cell(first_row, first_col)
        for index, value in enumerate(_values(seed.value, series, step, unit, count, limit)):
            if index:
                row, col = (
                    (first_row + index, first_col) if down else (first_row, first_col + index)
                )
                cell = writable_cell(sheet, row, col)
                cell.value = value
                cell._style = seed._style
    return area


def _copy_fill(sheet: Worksheet, area: CellRange, down: bool, max_cells: int) -> CellRange:
    count = area.rows if down else area.cols
    if count < 2:
        raise InvalidArgumentError(f"Select at least two {'rows' if down else 'columns'} to fill.")
    area.within(max_cells)
    first = (
        CellRange(area.min_row, area.min_col, area.min_row, area.max_col)
        if down
        else CellRange(area.min_row, area.min_col, area.max_row, area.min_col)
    )
    for index in range(1, count):
        target = (
            (area.min_row + index, area.min_col) if down else (area.min_row, area.min_col + index)
        )
        copy_range(
            sheet,
            str(first),
            sheet,
            cell_name(*target),
            paste="all",
            transpose=False,
            skip_blanks=False,
            results={},
            max_cells=max_cells,
        )
    return area


def _seed_positions(area: CellRange, down: bool) -> list[tuple[int, int]]:
    if down:
        return [(area.min_row, col) for col in range(area.min_col, area.max_col + 1)]
    return [(row, area.min_col) for row in range(area.min_row, area.max_row + 1)]


def _extend(
    sheet: Worksheet,
    area: CellRange,
    down: bool,
    series: Series,
    step: float,
    stop: str | float,
    unit: DateUnit,
    max_cells: int,
) -> CellRange:
    seed = sheet.cell(area.min_row, area.min_col)
    limit = _limit(series, stop, unit)
    count = len(_values(seed.value, series, step, unit, max_cells + 1, limit))
    if count > max_cells:
        raise InvalidArgumentError(f"The series would fill more than {max_cells:,} cells.")
    last_row = min(area.min_row + count - 1, MAX_ROW) if down else area.min_row
    last_col = area.min_col if down else min(area.min_col + count - 1, MAX_COLUMN)
    return CellRange(area.min_row, area.min_col, last_row, last_col)


def _limit(series: Series, stop: str | float | None, unit: DateUnit) -> float | None:
    if stop is None:
        return None
    if series == "date":
        date = parse_iso_date(stop) if isinstance(stop, str) else None
        if date is None:
            raise InvalidArgumentError("For a date series, stop is a date such as '2026-12-31'.")
        return float(to_excel(date))
    if isinstance(stop, str):
        raise InvalidArgumentError("For a number series, stop is a number.")
    return float(stop)


def _values(
    seed: object, series: Series, step: float, unit: DateUnit, count: int, limit: float | None
) -> list[int | float | dt.date]:
    """The seed and its successors, up to ``count`` values and not past ``limit``."""
    if series == "date":
        if not isinstance(seed, dt.date):
            raise InvalidArgumentError("A date series starts from a date cell.")
        if step != int(step):
            raise InvalidArgumentError("A date series needs a whole-number step.")
        values: list[int | float | dt.date] = [
            _add_dates(seed, unit, int(step) * index) for index in range(count)
        ]
        numbers = [float(to_excel(value)) for value in values]
    else:
        if isinstance(seed, bool) or not isinstance(seed, int | float):
            raise InvalidArgumentError("A number series starts from a number cell.")
        if series == "growth" and step == 0:
            raise InvalidArgumentError("A growth series needs a step other than 0.")
        try:
            numbers = [
                float(f"{seed + step * index if series == 'linear' else seed * step**index:.15g}")
                for index in range(count)
            ]
        except OverflowError:
            raise InvalidArgumentError("The series grows too large for a cell.") from None
        values = [int(n) if n == int(n) and abs(n) < 1e15 else n for n in numbers]
    if limit is None:
        return values
    rising = len(numbers) > 1 and numbers[1] >= numbers[0]
    for index, number in enumerate(numbers):
        if number > limit if rising else number < limit:
            return values[:index]
    return values


def _add_dates(seed: dt.date, unit: DateUnit, amount: int) -> dt.date:
    if unit == "day":
        return seed + dt.timedelta(days=amount)
    if unit == "weekday":
        return _add_weekdays(seed, amount)
    months = amount * 12 if unit == "year" else amount
    year, month = divmod(seed.year * 12 + seed.month - 1 + months, 12)
    day = min(seed.day, calendar.monthrange(year, month + 1)[1])
    return seed.replace(year=year, month=month + 1, day=day)


def _add_weekdays(seed: dt.date, amount: int) -> dt.date:
    direction = 1 if amount >= 0 else -1
    result = seed
    for _ in range(abs(amount)):
        result += dt.timedelta(days=direction)
        while result.weekday() >= 5:
            result += dt.timedelta(days=direction)
    return result
