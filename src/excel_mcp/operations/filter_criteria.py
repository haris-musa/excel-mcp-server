"""Filter criteria: what each kind of Excel filter stores and which cells it shows."""

import datetime as dt
import operator
from collections.abc import Callable
from typing import Literal

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet import filters as xl
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import Field

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.inputs import InputModel
from excel_mcp.operations.columns import find_column
from excel_mcp.operations.comparison import FormulaResults, as_number, cell_value
from excel_mcp.operations.formatting import parse_color
from excel_mcp.refs import CellRange
from excel_mcp.values import parse_iso_date, typed_value

Operator = Literal[
    "equals",
    "not_equals",
    "greater_than",
    "greater_or_equal",
    "less_than",
    "less_or_equal",
    "begins_with",
    "not_begins_with",
    "ends_with",
    "not_ends_with",
    "contains",
    "not_contains",
]
Test = Callable[[Cell | MergedCell, object], bool]

_COMPARISONS = {
    "equals": ("equal", operator.eq),
    "not_equals": ("notEqual", operator.ne),
    "greater_than": ("greaterThan", operator.gt),
    "greater_or_equal": ("greaterThanOrEqual", operator.ge),
    "less_than": ("lessThan", operator.lt),
    "less_or_equal": ("lessThanOrEqual", operator.le),
}
# Excel stores text operators as wildcard patterns: (stored operator, pattern, test).
_TEXT_OPERATORS = {
    "begins_with": ("equal", "{}*", str.startswith),
    "not_begins_with": ("notEqual", "{}*", str.startswith),
    "ends_with": ("equal", "*{}", str.endswith),
    "not_ends_with": ("notEqual", "*{}", str.endswith),
    "contains": ("equal", "*{}*", str.__contains__),
    "not_contains": ("notEqual", "*{}*", str.__contains__),
}
_USED = {
    "values": {"values"},
    "compare": {"conditions", "combine"},
    "top": {"count", "percent"},
    "bottom": {"count", "percent"},
    "above_average": set(),
    "below_average": set(),
    "color": {"color"},
}


class FilterCondition(InputModel):
    operator: Operator
    value: str | float = Field(description="Text, number or date '2026-01-31'.")


class FilterColumn(InputModel):
    column: str = Field(description="Header text or column letter.")
    type: Literal["values", "compare", "top", "bottom", "above_average", "below_average", "color"]
    values: list[str] | None = Field(
        default=None,
        description="values: to show; '' is blank, dates '2026-01-31'. Formatted numbers "
        "need compare.",
    )
    conditions: list[FilterCondition] | None = Field(
        default=None, min_length=1, max_length=2, description="compare."
    )
    combine: Literal["and", "or"] = Field(default="and", description="compare: 2 conditions.")
    count: int | None = Field(
        default=None, ge=1, le=1000, description="top, bottom: items, or percent (1-100)."
    )
    percent: bool = Field(default=False, description="top, bottom.")
    color: str | None = Field(default=None, description="color: fill color to show.")


def column_filter(
    sheet: Worksheet, area: CellRange, item: FilterColumn, values: FormulaResults
) -> tuple[xl.FilterColumn, int, Test]:
    item.reject_unused(_USED[item.type] | {"column", "type"}, f"A {item.type} filter")
    index = find_column(sheet, area, item.column, has_header=True) - area.min_col
    xml = xl.FilterColumn(colId=index)
    numbers = [
        n
        for row in range(area.min_row + 1, area.max_row + 1)
        if (n := as_number(cell_value(sheet.cell(row, area.min_col + index), values))) is not None
    ]
    match item.type:
        case "values":
            xml.filters, test = values_filter(item.values)
        case "compare":
            xml.customFilters, test = compare_filter(item.conditions, item.combine)
        case "top" | "bottom":
            xml.top10, test = top_filter(item, numbers)
        case "color":
            xml.colorFilter, test = _color_filter(sheet, item.color)
        case _:
            xml.dynamicFilter, test = average_filter(item.type, numbers)
    return xml, index, test


def _text(value: object) -> str:
    """The text Excel shows for a cell in the General format."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value == int(value) and abs(value) < 1e15:
        return str(int(value))
    return str(value)


def _stored(number: float) -> str:
    return str(int(number)) if number == int(number) else repr(number)


def values_filter(wanted: list[str] | None) -> tuple[xl.Filters, Test]:
    if not wanted:
        raise InvalidArgumentError("A values filter needs values.")
    dates = {date for text in wanted if (date := parse_iso_date(text))}
    texts = {text.casefold() for text in wanted if parse_iso_date(text) is None}
    xml = xl.Filters(
        blank=True if "" in wanted else None,
        filter=[text for text in wanted if text and parse_iso_date(text) is None],
        dateGroupItem=[
            xl.DateGroupItem(year=d.year, month=d.month, day=d.day, dateTimeGrouping="day")
            for d in sorted(dates)
        ],
    )

    def test(cell: Cell | MergedCell, value: object) -> bool:
        if isinstance(value, dt.date):
            return (value.date() if isinstance(value, dt.datetime) else value) in dates
        if as_number(value) is not None and cell.number_format != "General":
            raise InvalidArgumentError(
                f"{cell.coordinate} is a number with a number format; Excel lists numbers as "
                "formatted text. Filter this column with a compare filter instead."
            )
        return _text(value).casefold() in texts

    return xml, test


def compare_filter(
    conditions: list[FilterCondition] | None, combine: str
) -> tuple[xl.CustomFilters, Test]:
    if not conditions:
        raise InvalidArgumentError("A compare filter needs conditions.")
    built = [_condition(condition) for condition in conditions]
    combine_all = combine == "and"
    xml = xl.CustomFilters(_and=True if combine_all and len(built) > 1 else None)
    xml.customFilter = [filter_ for filter_, _ in built]

    def test(cell: Cell | MergedCell, value: object) -> bool:
        results = (check(value) for _, check in built)
        return all(results) if combine_all else any(results)

    return xml, test


def _condition(condition: FilterCondition) -> tuple[xl.CustomFilter, Callable[[object], bool]]:
    name = condition.operator
    if name in _TEXT_OPERATORS:
        stored, pattern, match = _TEXT_OPERATORS[name]
        text = _text(condition.value).casefold()
        escaped = _text(condition.value).replace("~", "~~").replace("*", "~*").replace("?", "~?")
        negate = stored == "notEqual"
        return xl.CustomFilter(
            operator=stored, val=pattern.format(escaped)
        ), lambda value: match(_text(value).casefold(), text) != negate
    stored, compare = _COMPARISONS[name]
    wanted = _wanted(condition.value)
    if wanted == "":
        if name != "not_equals":
            raise InvalidArgumentError("To match blanks, use a values filter with ''.")
        return xl.CustomFilter(operator=stored, val=" "), lambda value: value not in (None, "")
    if isinstance(wanted, str):
        val = wanted.replace("~", "~~").replace("*", "~*").replace("?", "~?")
        check = _text_check(compare, wanted.casefold(), name)
    else:
        val, check = _stored(wanted), _numeric_check(compare, wanted, name)
    return xl.CustomFilter(operator=stored, val=val), check


def _wanted(value: str | float) -> float | str:
    """A number (dates as serial numbers) or the text to compare with."""
    if isinstance(value, str) and value.startswith("="):
        return value
    parsed = typed_value(value, [])[0] if isinstance(value, str) else value
    if isinstance(parsed, bool) or parsed is None:
        return str(value)
    number = as_number(parsed)
    return str(value) if number is None else number


def _numeric_check(compare: Callable[[float, float], bool], wanted: float, name: str):
    def check(value: object) -> bool:
        number = as_number(value)
        return name == "not_equals" if number is None else compare(number, wanted)

    return check


def _text_check(compare: Callable[[str, str], bool], wanted: str, name: str):
    def check(value: object) -> bool:
        if not isinstance(value, str):
            return name == "not_equals"
        return compare(value.casefold(), wanted)

    return check


def top_filter(item: FilterColumn, numbers: list[float]) -> tuple[xl.Top10, Test]:
    if item.count is None or (item.percent and item.count > 100):
        raise InvalidArgumentError("top and bottom need a count (1-100 when percent).")
    if not numbers:
        raise InvalidArgumentError("The column has no numbers to rank.")
    top = item.type == "top"
    wanted = max(1, len(numbers) * item.count // 100) if item.percent else item.count
    ordered = sorted(numbers, reverse=top)
    edge = ordered[min(wanted, len(ordered)) - 1]
    xml = xl.Top10(
        top=None if top else False, percent=item.percent or None, val=item.count, filterVal=edge
    )

    def test(cell: Cell | MergedCell, value: object) -> bool:
        number = as_number(value)
        return number is not None and (number >= edge if top else number <= edge)

    return xml, test


def average_filter(kind: str, numbers: list[float]) -> tuple[xl.DynamicFilter, Test]:
    if not numbers:
        raise InvalidArgumentError("The column has no numbers to average.")
    average = float(f"{sum(numbers) / len(numbers):.15g}")
    above = kind == "above_average"
    xml = xl.DynamicFilter(type="aboveAverage" if above else "belowAverage", val=average)

    def test(cell: Cell | MergedCell, value: object) -> bool:
        number = as_number(value)
        return number is not None and (number > average if above else number < average)

    return xml, test


def _color_filter(sheet: Worksheet, color: str | None) -> tuple[xl.ColorFilter, Test]:
    if color is None:
        raise InvalidArgumentError("A color filter needs a color.")
    rgb = parse_color(color)
    fill = PatternFill(fill_type="solid", fgColor=rgb, bgColor="FFFFFFFF")
    dxf_id = sheet.parent._differential_styles.add(DifferentialStyle(fill=fill))  # pyright: ignore[reportOptionalMemberAccess, reportAttributeAccessIssue]

    def test(cell: Cell | MergedCell, value: object) -> bool:
        return cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb == rgb

    return xl.ColorFilter(dxfId=dxf_id), test
