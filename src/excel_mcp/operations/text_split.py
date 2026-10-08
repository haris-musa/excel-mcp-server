"""Text to Columns: split the text of one column into the columns to its right."""

from itertools import accumulate

from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError, LimitExceededError
from excel_mcp.operations.cells import store_typed, writable_cell
from excel_mcp.refs import MAX_COLUMN, CellRange
from excel_mcp.workspace import sheet_names

_NAMED_DELIMITERS = {"tab": "\t", "semicolon": ";", "comma": ",", "space": " "}


def text_to_columns(
    sheet: Worksheet,
    area: CellRange,
    delimiters: list[str] | None,
    fixed_widths: list[int] | None,
    qualifier: str,
    merge_delimiters: bool,
    max_cells: int,
) -> int:
    """Split each text cell; pieces are typed as in Excel (numbers, dates, formulas, text).

    Returns the number of cells split. Cells that are not text are left alone, and the cells
    the pieces go into must be empty.
    """
    if area.cols != 1:
        raise InvalidArgumentError(f"Text to columns needs a single column, not {area}.")
    if bool(delimiters) == bool(fixed_widths):
        raise InvalidArgumentError("Give either delimiters or fixed_widths.")
    separators = {_NAMED_DELIMITERS.get(name, name) for name in delimiters or []}
    if any(len(separator) != 1 for separator in separators):
        raise InvalidArgumentError(
            "Each delimiter is 'tab', 'semicolon', 'comma', 'space' or a single character."
        )
    split = [
        (row, _split(text, separators, fixed_widths, qualifier, merge_delimiters))
        for row in range(area.min_row, area.max_row + 1)
        if isinstance(text := sheet.cell(row, area.min_col).value, str)
        and sheet.cell(row, area.min_col).data_type == "s"
    ]
    if sum(len(pieces) for _, pieces in split) > max_cells:
        raise LimitExceededError(f"Splitting would fill more than {max_cells:,} cells.")
    names = sheet_names(sheet)
    for row, pieces in split:
        if area.min_col + len(pieces) - 1 > MAX_COLUMN:
            raise InvalidArgumentError(f"Row {row} would split past the last column.")
        for offset in range(1, len(pieces)):
            neighbour = writable_cell(sheet, row, area.min_col + offset)
            if neighbour.value is not None:
                raise InvalidArgumentError(
                    f"{neighbour.coordinate} is not empty; Excel would ask before replacing it. "
                    "Clear or move it first."
                )
    for row, pieces in split:
        for offset, piece in enumerate(pieces):
            store_typed(writable_cell(sheet, row, area.min_col + offset), piece, names)
    return len(split)


def _split(
    text: str,
    separators: set[str],
    fixed_widths: list[int] | None,
    qualifier: str,
    merge_delimiters: bool,
) -> list[str]:
    if fixed_widths:
        edges = [0, *accumulate(fixed_widths)]
        pieces = [
            text[start:end] for start, end in zip(edges, [*edges[1:], len(text)], strict=True)
        ]
        return [piece for piece in pieces if piece] or [text]
    fields: list[str] = []
    field = ""
    quoted = False
    at_start = True
    index = 0
    while index < len(text):
        char = text[index]
        if quoted:
            if char == qualifier and text[index + 1 : index + 2] == qualifier:
                field += char
                index += 1
            elif char == qualifier:
                quoted = False
            else:
                field += char
        elif at_start and qualifier and char == qualifier:
            quoted = True
            at_start = False
        elif char in separators:
            fields.append(field)
            field, at_start = "", True
        else:
            field += char
            at_start = False
        index += 1
    fields.append(field)
    return [f for position, f in enumerate(fields) if f or position == 0 or not merge_delimiters]
