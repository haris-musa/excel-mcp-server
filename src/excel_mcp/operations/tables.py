"""Native Excel tables."""

import re

from openpyxl.workbook import Workbook
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.operations.table_options import TableOptions, apply_options
from excel_mcp.refs import parse_range

_TABLE_NAME = re.compile(r"[A-Za-z_][A-Za-z0-9_.]{0,254}")
_CELL_LIKE = re.compile(r"[A-Za-z]{1,3}[0-9]+|[RrCc]([0-9]*|[Rr]?[0-9]*[Cc][0-9]*)")
DEFAULT_STYLE = "TableStyleMedium2"
_BUILTIN_STYLE = re.compile(
    r"TableStyle(Light([1-9]|1[0-9]|2[01])|Medium([1-9]|1[0-9]|2[0-8])|Dark([1-9]|1[01]))"
)


def create_table(
    workbook: Workbook,
    sheet: Worksheet,
    ref: str,
    name: str | None,
    options: TableOptions,
    max_cells: int,
) -> tuple[str, str]:
    area = parse_range(ref)
    if area.rows < 2:
        raise InvalidArgumentError("A table needs a header row and at least one data row.")
    table_name = name or _next_table_name(workbook)
    _validate_table_name(workbook, table_name)
    if options.style is not None and not _BUILTIN_STYLE.fullmatch(options.style):
        raise InvalidArgumentError(
            f"Unknown table style {options.style!r}. Use a built-in style such as "
            "'TableStyleMedium2' (Light1-21, Medium1-28, Dark1-11)."
        )
    for existing in sheet.tables.values():
        if area.overlaps(parse_range(existing.ref)):
            raise InvalidArgumentError(f"{area} overlaps table {existing.displayName!r}.")
    if sheet.auto_filter.ref and area.overlaps(parse_range(sheet.auto_filter.ref)):
        raise InvalidArgumentError(
            f"{area} overlaps the sheet's auto filter; tables have their own filter."
        )
    headers = _validate_headers(sheet, area.min_row, area.min_col, area.max_col)

    table = Table(displayName=table_name, ref=str(area))
    table.tableColumns = [TableColumn(id=i, name=h) for i, h in enumerate(headers, start=1)]
    table.autoFilter = AutoFilter(ref=str(area))
    table.tableStyleInfo = TableStyleInfo(name=DEFAULT_STYLE, showRowStripes=True)
    sheet.add_table(table)
    apply_options(sheet, table, options, None, max_cells)
    return table_name, table.ref


def edit_table(
    sheet: Worksheet, name: str, ref: str | None, options: TableOptions, max_cells: int
) -> str:
    found = {table.displayName.casefold(): table for table in sheet.tables.values()}
    table = found.get(name.casefold())
    if table is None:
        listed = ", ".join(table.displayName for table in sheet.tables.values()) or "none"
        raise InvalidArgumentError(
            f"Sheet {sheet.title!r} has no table {name!r}. Tables: {listed}."
        )
    apply_options(sheet, table, options, parse_range(ref) if ref else None, max_cells)
    return table.ref


def is_valid_name(name: str) -> bool:
    """Whether ``name`` can name a table or a defined name."""
    return bool(_TABLE_NAME.fullmatch(name)) and not _CELL_LIKE.fullmatch(name)


def table_names(workbook: Workbook) -> set[str]:
    return {name.casefold() for sheet in workbook.worksheets for name in sheet.tables}


def _next_table_name(workbook: Workbook) -> str:
    existing = table_names(workbook)
    number = 1
    while f"table{number}" in existing:
        number += 1
    return f"Table{number}"


def _validate_table_name(workbook: Workbook, name: str) -> None:
    if not is_valid_name(name):
        raise InvalidArgumentError(
            f"Invalid table name {name!r}. Start with a letter or underscore and use only "
            "letters, digits, underscores and periods."
        )
    if name.casefold() in table_names(workbook) or name in workbook.defined_names:
        raise InvalidArgumentError(f"The name {name!r} is already used in this workbook.")


def _validate_headers(sheet: Worksheet, row: int, first_col: int, last_col: int) -> list[str]:
    headers = [sheet.cell(row, col).value for col in range(first_col, last_col + 1)]
    labels = [header for header in headers if isinstance(header, str) and header.strip()]
    if len(labels) != len(headers):
        raise InvalidArgumentError("Every header cell in the table's first row must be text.")
    if len({label.casefold() for label in labels}) != len(labels):
        raise InvalidArgumentError("Table headers must be unique.")
    return labels
