"""Results for formula cells that Excel has not calculated yet."""

import datetime as dt
from typing import cast

from openpyxl.cell.cell import Cell
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.calc.engine import Engine
from excel_mcp.calc.values import ExcelError, Scalar, UncalculableError
from excel_mcp.errors import InvalidArgumentError
from excel_mcp.operations import cells
from excel_mcp.operations.cells import RangeData
from excel_mcp.operations.comparison import FormulaResults
from excel_mcp.operations.spill import is_free
from excel_mcp.refs import CellRange, cell_name, parse_range
from excel_mcp.values import CellValue, to_json
from excel_mcp.workspace import Workspace, get_sheet, get_streamed_sheet

MAX_LISTED = 20


def read_calculated(
    workspace: Workspace, path: str, sheet: str, ref: str | None, max_cells: int
) -> RangeData:
    """Read values as Excel last stored them; calculate formulas that have no stored result.

    The window is streamed. Only when it holds such formulas is the workbook loaded in full,
    because evaluating a formula needs the cells it refers to, wherever they are.
    """
    with workspace.stream(path) as formulas:
        formula_sheet = get_streamed_sheet(formulas, sheet)
        # A used range found from stored results alone would miss formulas never calculated.
        window = str(cells.read_window(formula_sheet, ref))
    with workspace.stream(path, data_only=True) as stored:
        data = cells.read_range(get_streamed_sheet(stored, sheet), window, max_cells)
    area = parse_range(data.range)
    with workspace.stream(path) as formulas:
        formula_cells = _formula_cells(get_streamed_sheet(formulas, sheet), area)
    missing = [(row, col) for row, col in formula_cells if _stored(data, area, row, col) is None]
    if missing:
        with workspace.stream(path, data_only=True) as stored:
            blanks = _stored_empty_text(get_streamed_sheet(stored, sheet), area, missing)
        for row, col in blanks:
            _put(data, area, row, col, "")
        pending = [position for position in missing if position not in blanks]
        if pending:
            _fill(workspace, path, sheet, data, pending)
    return data


def formula_values(
    workspace: Workspace, path: str, sheet: Worksheet, area: CellRange
) -> FormulaResults:
    """The results of the formula cells in ``area``, with dates as datetimes.

    Reads the saved file, so call it before changes made in the same edit are saved.
    """
    formula_cells = [
        (cell, row_number, col_number)
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=area.min_row,
                max_row=area.max_row,
                min_col=area.min_col,
                max_col=area.max_col,
            ),
            start=area.min_row,
        )
        for col_number, cell in enumerate(row, start=area.min_col)
        if cell.data_type == "f"
    ]
    if not formula_cells:
        return {}
    data = read_calculated(workspace, path, sheet.title, str(area), area.size)
    if data.uncalculated:
        cells_list = ", ".join(data.uncalculated)
        raise InvalidArgumentError(
            f"The result of formulas in {cells_list} is unknown until Excel recalculates, so "
            "this cannot use them."
        )
    results: FormulaResults = {}
    for cell, row, col in formula_cells:
        value = _stored(data, parse_range(data.range), row, col)
        if isinstance(value, str) and is_date_format(cell.number_format):
            value = dt.datetime.fromisoformat(value)
        results[cell.coordinate] = value
    return results


def _formula_cells(sheet: ReadOnlyWorksheet, area: CellRange) -> list[tuple[int, int]]:
    found = []
    for row in sheet.iter_rows(
        min_row=area.min_row, max_row=area.max_row, min_col=area.min_col, max_col=area.max_col
    ):
        found.extend((cell.row, cell.column) for cell in row if cell.data_type == "f")
    return found


def _stored_empty_text(
    sheet: ReadOnlyWorksheet, area: CellRange, positions: list[tuple[int, int]]
) -> set[tuple[int, int]]:
    """The cells whose stored result is empty text, which openpyxl reads as no value at all.

    Excel stores a formula that returns "" as a text result without content; a formula that was
    never calculated has no result element, and only that one needs calculating.
    """
    wanted = set(positions)
    found = set()
    for row in sheet.iter_rows(
        min_row=area.min_row, max_row=area.max_row, min_col=area.min_col, max_col=area.max_col
    ):
        found.update(
            (cell.row, cell.column)
            for cell in row
            if cell.data_type == "str" and (cell.row, cell.column) in wanted
        )
    return found


def _stored(data: RangeData, area: CellRange, row: int, col: int) -> CellValue:
    values = data.values
    line = values[row - area.min_row] if row - area.min_row < len(values) else []
    return line[col - area.min_col] if col - area.min_col < len(line) else None


def _fill(
    workspace: Workspace, path: str, sheet: str, data: RangeData, pending: list[tuple[int, int]]
) -> None:
    area = parse_range(data.range)
    unresolved: dict[str, str] = {}
    with workspace.read(path, data_only=True) as stored, workspace.read(path) as formulas:
        target = get_sheet(formulas, sheet)
        engine = Engine(stored, formulas)
        for row, col in pending:
            cell = target._cells[(row, col)]
            try:
                if isinstance(cell.value, ArrayFormula):
                    _fill_spill(engine, cast(Cell, cell), data, area)
                else:
                    result = engine.calculate(target, row, col)
                    _put(data, area, row, col, _json(result, cell.number_format))
            except UncalculableError as reason:
                unresolved[cell_name(row, col)] = str(reason)
    if unresolved:
        listed = dict(list(unresolved.items())[:MAX_LISTED])
        if len(unresolved) > MAX_LISTED:
            listed["..."] = f"{len(unresolved) - MAX_LISTED} more"
        data.uncalculated = listed


def _fill_spill(engine: Engine, anchor: Cell, data: RangeData, area: CellRange) -> None:
    """Show an array formula and what it spills into, as Excel would once it calculated it."""
    sheet = cast(Worksheet, anchor.parent)
    formula = cast(ArrayFormula, anchor.value)
    grid = engine.spill(sheet, anchor.row, anchor.column, str(formula.text))
    if parse_range(formula.ref).size > 1:
        # A legacy array formula: the other cells of its range hold their own results.
        _put(data, area, anchor.row, anchor.column, _json(grid.rows[0][0], anchor.number_format))
        return
    spilled = CellRange(
        anchor.row, anchor.column, anchor.row + grid.height - 1, anchor.column + grid.width - 1
    )
    if not is_free(sheet, spilled, anchor):
        _put(data, area, anchor.row, anchor.column, ExcelError("#SPILL!").code)
        return
    for row_offset, line in enumerate(grid.rows):
        for col_offset, item in enumerate(line):
            row, col = anchor.row + row_offset, anchor.column + col_offset
            inside = area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col
            if inside and (row_offset == col_offset == 0 or _stored(data, area, row, col) is None):
                _put(data, area, row, col, _json(item, anchor.number_format))


def _put(data: RangeData, area: CellRange, row: int, col: int, value: CellValue) -> None:
    values = data.values
    while len(values) <= row - area.min_row:
        values.append([])
    line = values[row - area.min_row]
    while len(line) <= col - area.min_col:
        line.append(None)
    line[col - area.min_col] = value


def _json(result: Scalar, number_format: str) -> CellValue:
    if isinstance(result, ExcelError):
        return result.code
    if isinstance(result, bool | str):
        return result
    if is_date_format(number_format):
        return to_json(from_excel(result))
    number = float(f"{result:.15g}")
    return int(number) if number == int(number) and abs(number) < 1e15 else number
