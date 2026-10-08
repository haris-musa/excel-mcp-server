"""Finding a column of a range by its header text or letter."""

from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.refs import CellRange


def find_column(sheet: Worksheet, area: CellRange, key: str, has_header: bool) -> int:
    headers = [sheet.cell(area.min_row, col).value for col in range(area.min_col, area.max_col + 1)]
    if has_header:
        wanted = key.strip().casefold()
        matches = [
            area.min_col + index
            for index, header in enumerate(headers)
            if str(header).strip().casefold() == wanted
        ]
        if len(matches) > 1:
            raise InvalidArgumentError(
                f"Several columns are headed {key!r}; use the column letter instead."
            )
        if matches:
            return matches[0]
    if key.isalpha() and len(key) <= 3:
        column = column_index_from_string(key.upper())
        if area.min_col <= column <= area.max_col:
            return column
    where = f"a header ({', '.join(repr(header) for header in headers)}) or " if has_header else ""
    raise InvalidArgumentError(
        f"Column {key!r} is not in {area}. Use {where}a column letter from "
        f"{get_column_letter(area.min_col)} to {get_column_letter(area.max_col)}."
    )
