"""Spans of whole rows ('3:5') or whole columns ('B:D')."""

import re
from itertools import groupby
from typing import Literal

from openpyxl.utils.cell import column_index_from_string, get_column_letter

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.refs import MAX_COLUMN, MAX_ROW

Axis = Literal["rows", "columns"]

MAX_ROW_SPAN = 10_000

_ROWS = re.compile(r"(\d+)(?::(\d+))?")
_COLUMNS = re.compile(r"([A-Za-z]{1,3})(?::([A-Za-z]{1,3}))?")


def parse_span(text: str, axis: Axis) -> tuple[int, int]:
    """Parse '3', '3:5', 'B' or 'B:D' into 1-based first and last indices."""
    example = "'3' or '3:5'" if axis == "rows" else "'B' or 'B:D'"
    match = (_ROWS if axis == "rows" else _COLUMNS).fullmatch(text.strip())
    if match is None:
        raise InvalidArgumentError(f"Invalid {axis} span {text!r}. Use {example}.")
    ends = [match[1], match[2] or match[1]]
    if axis == "rows":
        first, last = (int(end) for end in ends)
        limit = MAX_ROW
    else:
        first, last = (column_index_from_string(end.upper()) for end in ends)
        limit = MAX_COLUMN
    first, last = min(first, last), max(first, last)
    if first < 1 or last > limit:
        raise InvalidArgumentError(f"The {axis} span {text!r} is outside the worksheet.")
    if axis == "rows" and last - first >= MAX_ROW_SPAN:
        raise InvalidArgumentError(f"A rows span can cover at most {MAX_ROW_SPAN:,} rows.")
    return first, last


def format_span(first: int, last: int, axis: Axis) -> str:
    if axis == "rows":
        return f"{first}:{last}"
    return f"{get_column_letter(first)}:{get_column_letter(last)}"


def to_spans(indices: set[int], axis: Axis) -> list[str]:
    """Condense indices into spans such as ['2:4', '9:9']."""
    spans = []
    for _, run in groupby(enumerate(sorted(indices)), key=lambda pair: pair[1] - pair[0]):
        members = [index for _, index in run]
        spans.append(format_span(members[0], members[-1], axis))
    return spans
