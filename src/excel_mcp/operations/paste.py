"""Copying a range the way Excel's Copy and Paste Special do."""

import datetime as dt
import re
from copy import copy
from typing import Any, Literal

from openpyxl.cell.cell import Cell
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.formulas import storable_formula
from excel_mcp.operations.cells import store_value, writable_cell
from excel_mcp.operations.comparison import FormulaResults
from excel_mcp.refs import MAX_COLUMN, MAX_ROW, CellRange, cell_name, parse_cell, parse_range
from excel_mcp.workspace import sheet_names

PasteMode = Literal["all", "values", "formulas", "formats"]

_CELL = re.compile(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)")
_WHOLE_LINE = re.compile(r"\$?[A-Za-z]{1,3}|\$?\d+")


def copy_range(
    source: Worksheet,
    ref: str,
    target: Worksheet,
    target_cell: str,
    *,
    paste: PasteMode,
    transpose: bool,
    skip_blanks: bool,
    results: FormulaResults,
    max_cells: int,
) -> str:
    """Copy cells; relative references in formulas shift the way Excel shifts them.

    ``results`` holds the calculated results of the source's formulas, for pasting values.
    """
    area = parse_range(ref).within(max_cells)
    first_row, first_col = parse_cell(target_cell)
    rows, cols = (area.cols, area.rows) if transpose else (area.rows, area.cols)
    destination = CellRange(first_row, first_col, first_row + rows - 1, first_col + cols - 1)
    snapshot = [
        (cell.row, cell.column, cell.value, cell.data_type, copy(cell._style))
        for row in source.iter_rows(
            min_row=area.min_row, max_row=area.max_row, min_col=area.min_col, max_col=area.max_col
        )
        for cell in row
    ]
    names = sheet_names(target)
    for row, col, value, data_type, style in snapshot:
        if skip_blanks and value is None:
            continue
        down, right = row - area.min_row, col - area.min_col
        if transpose:
            down, right = right, down
        cell = writable_cell(target, first_row + down, first_col + right)
        if paste != "formats":
            origin = (row, col)
            _paste_content(cell, origin, value, data_type, paste, transpose, results, names)
        if paste in ("all", "formats"):
            cell._style = style
    return str(destination)


def _paste_content(
    cell: Cell,
    origin: tuple[int, int],
    value: Any,
    data_type: str,
    paste: PasteMode,
    transpose: bool,
    results: FormulaResults,
    names: list[str],
) -> None:
    if isinstance(value, ArrayFormula):
        raise InvalidArgumentError(
            f"{get_column_letter(origin[1])}{origin[0]} holds an array formula, "
            "which cannot be copied."
        )
    if data_type == "f" and paste == "values":
        value = results[cell_name(*origin)]
        cell.value = _serial(value) if isinstance(value, dt.date) else value
        if isinstance(value, str):
            cell.data_type = "s"
    elif data_type == "f":
        formula = _moved_formula(str(value), origin, (cell.row, cell.column), transpose)
        cell.value = storable_formula(formula, names)
    elif paste == "all":
        cell.value = value
        cell.data_type = data_type
    elif isinstance(value, dt.date | dt.time | dt.timedelta):
        cell.value = _serial(value)
    else:
        store_value(cell, value)


def _serial(value: dt.date | dt.time | dt.timedelta) -> float:
    return float(to_excel(value))


def _moved_formula(
    formula: str, origin: tuple[int, int], destination: tuple[int, int], transpose: bool
) -> str:
    start = f"{get_column_letter(origin[1])}{origin[0]}"
    end = f"{get_column_letter(destination[1])}{destination[0]}"
    if not transpose:
        return Translator(formula, origin=start).translate_formula(end)
    tokens = Tokenizer(formula).items
    moved = (
        _transposed_reference(token.value, origin, destination)
        if token.type == Token.OPERAND and token.subtype == Token.RANGE
        else token.value
        for token in tokens
    )
    return "=" + "".join(moved)


def _transposed_reference(
    reference: str, origin: tuple[int, int], destination: tuple[int, int]
) -> str:
    """Excel turns a relative offset of (rows, columns) into (columns, rows) when transposing.

    References with a "$" stay as they are.
    """
    sheet, bang, cells = reference.rpartition("!")
    corners = cells.split(":")
    if len(corners) == 2 and all(_WHOLE_LINE.fullmatch(corner) for corner in corners):
        raise InvalidArgumentError(
            "Formulas with whole row or column references cannot be transposed."
        )
    moved = []
    for corner in corners:
        match = _CELL.fullmatch(corner)
        if match is None:
            return reference
        column_fixed, letters, row_fixed, digits = match.groups()
        if column_fixed or row_fixed:
            moved.append(corner)
            continue
        column, row = column_index_from_string(letters.upper()), int(digits)
        new_row = destination[0] + column - origin[1]
        new_col = destination[1] + row - origin[0]
        if not (1 <= new_row <= MAX_ROW and 1 <= new_col <= MAX_COLUMN):
            raise InvalidArgumentError(f"Transposing {reference!r} would move it off the sheet.")
        moved.append(f"{get_column_letter(new_col)}{new_row}")
    return sheet + bang + ":".join(moved)
