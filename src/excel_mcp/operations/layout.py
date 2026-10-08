"""Sheet layout: sizes, hidden and grouped lines, panes, filters, tab color, visibility."""

from copy import copy
from typing import Literal, cast

from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, Field

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.operations.cells import stored_cells
from excel_mcp.operations.formatting import parse_color
from excel_mcp.operations.print_setup import PrintSetup, apply_print_setup
from excel_mcp.operations.protection import Protection, apply_protection
from excel_mcp.operations.spans import Axis, parse_span, to_spans
from excel_mcp.refs import MAX_ROW, parse_cell, parse_range

MIN_AUTOFIT_WIDTH = 8
MAX_AUTOFIT_WIDTH = 80
MAX_OUTLINE_LEVEL = 7


class ColumnWidth(BaseModel):
    """The width of one column."""

    column: str = Field(description="Column letter, e.g. 'B'.")
    width: float = Field(gt=0, le=255, description="Width in characters.")


class RowHeight(BaseModel):
    """The height of one row."""

    row: int = Field(ge=1, le=MAX_ROW, description="1-based row number.")
    height: float = Field(gt=0, le=409, description="Height in points.")


class LineAction(BaseModel):
    """A change to a span of rows or columns."""

    span: str = Field(description="Rows like '3' or '3:5'; columns like 'B' or 'B:D'.")
    action: Literal["hide", "show", "group", "ungroup"] = Field(
        description="Group adds one outline level (up to 7), ungroup removes one."
    )


class SheetLayout(BaseModel):
    """Layout changes. Fields left as null are not changed."""

    column_widths: list[ColumnWidth] | None = Field(default=None, description="Widths to set.")
    row_heights: list[RowHeight] | None = Field(default=None, description="Heights to set.")
    autofit_columns: list[str] | None = Field(
        default=None, description="Column letters to size to their content, e.g. ['A', 'C']."
    )
    freeze_panes: str | None = Field(
        default=None, description="First unfrozen cell: 'A2' freezes row 1, 'A1' unfreezes."
    )
    auto_filter: str | None = Field(
        default=None, description="Range with filter buttons, e.g. 'A1:F100'."
    )
    tab_color: str | None = Field(default=None, description="Sheet tab hex color.")
    rows: list[LineAction] | None = Field(default=None, description="Hide, show or group rows.")
    columns: list[LineAction] | None = Field(
        default=None, description="Hide, show or group columns."
    )
    visibility: Literal["visible", "hidden"] | None = Field(
        default=None, description="Show or hide the whole sheet; one sheet must stay visible."
    )
    print_setup: PrintSetup | None = Field(default=None, description="Page setup for printing.")
    protection: Protection | None = Field(
        default=None, description="Protect or unprotect the sheet."
    )


def apply_layout(sheet: Worksheet, layout: SheetLayout) -> None:
    if layout.column_widths or layout.autofit_columns or layout.columns:
        split_column_dimensions(sheet)
    for column_width in layout.column_widths or []:
        sheet.column_dimensions[column_letter(column_width.column)].width = column_width.width
    for row_height in layout.row_heights or []:
        sheet.row_dimensions[row_height.row].height = row_height.height
    for column in layout.autofit_columns or []:
        letter = column_letter(column)
        sheet.column_dimensions[letter].width = estimate_width(sheet, letter)
    if layout.freeze_panes is not None:
        parse_cell(layout.freeze_panes)
        sheet.freeze_panes = layout.freeze_panes.upper()
    if layout.auto_filter is not None:
        area = parse_range(layout.auto_filter)
        for table in sheet.tables.values():
            if area.overlaps(parse_range(table.ref)):
                raise InvalidArgumentError(
                    f"{area} overlaps table {table.displayName!r}, which has its own filter."
                )
        sheet.auto_filter.ref = str(area)
    if layout.tab_color is not None:
        sheet.sheet_properties.tabColor = parse_color(layout.tab_color)
    for action in layout.rows or []:
        change_lines(sheet, "rows", action)
    for action in layout.columns or []:
        change_lines(sheet, "columns", action)
    if layout.print_setup is not None:
        apply_print_setup(sheet, layout.print_setup)
    if layout.visibility is not None:
        set_visibility(sheet, layout.visibility)
    if layout.protection is not None:
        apply_protection(sheet, layout.protection)


def split_column_dimensions(sheet: Worksheet) -> None:
    """Give every column of a multi-column definition its own, as openpyxl expects.

    Excel stores equal neighbouring columns as one definition; changing one of them
    separately would otherwise leave overlapping definitions that Excel rejects.
    """
    dimensions = sheet.column_dimensions
    for dimension in list(dimensions.values()):
        if dimension.min and dimension.max and dimension.max > dimension.min:
            for index in range(dimension.min, dimension.max + 1):
                piece = copy(dimension)
                piece.index = get_column_letter(index)
                piece.min = piece.max = index
                dimensions[piece.index] = piece


def change_lines(sheet: Worksheet, axis: Axis, change: LineAction) -> None:
    first, last = parse_span(change.span, axis)
    for index in range(first, last + 1):
        dimension = (
            sheet.row_dimensions[index]
            if axis == "rows"
            else sheet.column_dimensions[get_column_letter(index)]
        )
        if change.action in ("hide", "show"):
            dimension.hidden = change.action == "hide"
        else:
            step = 1 if change.action == "group" else -1
            level = (dimension.outlineLevel or 0) + step
            dimension.outlineLevel = min(max(level, 0), MAX_OUTLINE_LEVEL)
    if axis == "rows":
        levels = (row.outlineLevel or 0 for row in sheet.row_dimensions.values())
        sheet.sheet_format.outlineLevelRow = max(levels, default=0)


def set_visibility(sheet: Worksheet, visibility: Literal["visible", "hidden"]) -> None:
    if visibility == "visible":
        sheet.sheet_state = "visible"
        return
    if sheet.sheet_state != "visible":
        return
    workbook = cast(Workbook, sheet.parent)
    others = [other for other in workbook.worksheets if other is not sheet]
    shown = [other for other in others if other.sheet_state == "visible"]
    if not shown:
        raise InvalidArgumentError("A workbook must keep at least one visible worksheet.")
    sheet.sheet_state = "hidden"
    if workbook.active is sheet:
        workbook.active = shown[0]
    sheet.sheet_view.tabSelected = False
    shown[0].sheet_view.tabSelected = True


def hidden_lines(sheet: Worksheet, axis: Axis) -> list[str]:
    if axis == "rows":
        indices = {index for index, row in sheet.row_dimensions.items() if row.hidden}
    else:
        indices = set()
        for column in sheet.column_dimensions.values():
            if column.hidden:
                first = column.min or column_index_from_string(column.index)
                indices.update(range(first, (column.max or first) + 1))
    return to_spans(indices, axis)


def column_letter(column: str) -> str:
    try:
        return get_column_letter(column_index_from_string(column.strip().upper()))
    except ValueError:
        raise InvalidArgumentError(
            f"Invalid column {column!r}. Use a letter such as 'B'."
        ) from None


def estimate_width(sheet: Worksheet, letter: str) -> float:
    """Approximate a fitting width from the longest text line in the column."""
    index = column_index_from_string(letter)
    lines = [
        line
        for cell in stored_cells(sheet)
        if cell.column == index
        for line in str(cell.value).splitlines()
    ]
    longest = max((len(line) for line in lines), default=0)
    return min(max(longest + 2, MIN_AUTOFIT_WIDTH), MAX_AUTOFIT_WIDTH)
