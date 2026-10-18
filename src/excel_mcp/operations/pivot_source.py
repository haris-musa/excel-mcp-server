"""Reading the data a PivotTable summarizes."""

import datetime as dt
from collections.abc import Callable
from dataclasses import dataclass
from functools import cached_property
from typing import Literal, TypeGuard, cast

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.calc.values import ERRORS, ExcelError
from excel_mcp.errors import InvalidArgumentError, LimitExceededError
from excel_mcp.lazy_workbook import LazyCells
from excel_mcp.operations.comparison import FormulaResults
from excel_mcp.refs import CellRange, cell_name, parse_range
from excel_mcp.text import quoted

Value = str | int | float | bool | dt.datetime | ExcelError | None
"""A cell as a PivotTable sees it. Dates, times and durations are moments counted from `EPOCH`."""
Kind = Literal["number", "date", "other"]
Results = Callable[[Worksheet, CellRange], FormulaResults]

EPOCH = dt.datetime(1899, 12, 30)
"""Where the cache counts date serial numbers from; times and durations are dates of 1899-1900."""


def is_number(value: Value) -> TypeGuard[int | float]:
    return isinstance(value, int | float) and not isinstance(value, bool)


def kind_of(value: Value) -> Kind | None:
    if value is None:
        return None
    if is_number(value):
        return "number"
    return "date" if isinstance(value, dt.datetime) else "other"


def serial(moment: dt.datetime) -> float:
    return (moment - EPOCH) / dt.timedelta(days=1)


def moment_of(number: float) -> dt.datetime:
    return EPOCH + dt.timedelta(milliseconds=round(number * 86_400_000))


@dataclass(frozen=True)
class Column:
    name: str
    values: list[Value]
    number_format: str
    number_format_id: int
    """The format of every cell of the column, or General when they differ, as Excel takes it."""

    @cached_property
    def kinds(self) -> set[Kind]:
        return {kind for value in self.values if (kind := kind_of(value))}

    @cached_property
    def has_errors(self) -> bool:
        return any(isinstance(value, ExcelError) for value in self.values)


@dataclass(frozen=True)
class Source:
    sheet: str
    ref: str
    columns: list[Column]

    @property
    def record_count(self) -> int:
        return len(self.columns[0].values)

    def field_index(self, name: str) -> int:
        folded = [column.name.strip().casefold() for column in self.columns]
        if name.strip().casefold() not in folded:
            raise InvalidArgumentError(
                f"Field {name!r} not found. "
                f"Available fields: {quoted(column.name for column in self.columns)}."
            )
        return folded.index(name.strip().casefold())


def read_source(sheet: Worksheet, source_range: str, max_cells: int, results: Results) -> Source:
    """Read the block; ``results`` gives the results of its formulas (when it has any)."""
    area = parse_range(source_range)
    if area.size > max_cells:
        raise LimitExceededError(
            f"Source {area} has {area.size:,} cells; PivotTables can read at most "
            f"{max_cells:,}. Use a smaller source, or raise --max-pivot-cells "
            "(environment variable EXCEL_MCP_MAX_PIVOT_CELLS)."
        )
    if area.rows < 2:
        raise InvalidArgumentError(f"source {area} needs a header row and at least one data row.")
    if isinstance(sheet._cells, LazyCells) and (
        stored := sheet._cells.stored_rows(area.min_row, area.max_row)
    ):
        return _read_stored(sheet, area, stored)
    rows = [
        list(row)
        for row in sheet.iter_rows(
            min_row=area.min_row, max_row=area.max_row, min_col=area.min_col, max_col=area.max_col
        )
    ]
    names = [_header(cell) for cell in rows[0]]
    folded = [name.strip().casefold() for name in names]
    if len(set(folded)) < len(folded):
        raise InvalidArgumentError(f"Header names must be unique, got {names}.")
    calculated: FormulaResults = {}
    if any(cell.data_type == "f" for row in rows[1:] for cell in row):
        calculated = results(sheet, area)
    columns = [
        _column(name, [row[index] for row in rows[1:]], calculated)
        for index, name in enumerate(names)
    ]
    return Source(sheet.title, str(area), columns)


def _read_stored(
    sheet: Worksheet, area: CellRange, rows: list[dict[int, tuple[int, object, str]]]
) -> Source:
    """`read_source` of a large sheet read as the file stores it (`LazyCells.stored_rows`):
    the same columns, without making a cell of each of the sheet's millions."""
    book = cast(Workbook, sheet.parent)
    formats: dict[int, tuple[str, int]] = {}

    def format_of(style: int) -> tuple[str, int]:
        if style not in formats:
            probe = Cell(sheet, row=1, column=1, style_array=book._cell_styles[style])  # pyright: ignore[reportAttributeAccessIssue]
            formats[style] = (probe.number_format, probe._style.numFmtId)
        return formats[style]

    columns = []
    for column in range(area.min_col, area.max_col + 1):
        cells = [row.get(column) for row in rows]
        title = cells[0]
        if title is None or title[2] != "s" or not str(title[1]).strip():
            raise InvalidArgumentError(
                f"Header cell {cell_name(area.min_row, column)} must contain text: the first row "
                "of source is the header row, with one label per column."
            )
        data = cells[1:]
        # A position the file has no cell for is a blank cell of the default format.
        used = {format_of(cell[0] if cell else 0) for cell in data}
        shared = used.pop() if len(used) == 1 else ("General", 0)
        columns.append(Column(str(title[1]), [_stored_value(cell) for cell in data], *shared))
    names = [column.name for column in columns]
    folded = [name.strip().casefold() for name in names]
    if len(set(folded)) < len(folded):
        raise InvalidArgumentError(f"Header names must be unique, got {names}.")
    return Source(sheet.title, str(area), columns)


def _stored_value(cell: tuple[int, object, str] | None) -> Value:
    if cell is None:
        return None
    _, value, data_type = cell
    return ExcelError(str(value)) if data_type == "e" else _plain(value)


def _header(cell: Cell | MergedCell) -> str:
    if cell.data_type != "s" or not str(cell.value).strip():
        raise InvalidArgumentError(
            f"Header cell {cell.coordinate} must contain text: the first row of "
            "source is the header row, with one label per column."
        )
    return str(cell.value)


def _column(name: str, cells: list[Cell | MergedCell], calculated: FormulaResults) -> Column:
    formats = {(cell.number_format, cell._style.numFmtId) for cell in cells}
    shared = formats.pop() if len(formats) == 1 else ("General", 0)
    return Column(name, [_value(cell, calculated) for cell in cells], *shared)


def _value(cell: Cell | MergedCell, calculated: FormulaResults) -> Value:
    if cell.data_type == "f":
        return _result(calculated.get(cell.coordinate))
    if cell.data_type == "e":
        return ExcelError(str(cell.value))
    return _plain(cell.value)


def _plain(value: object) -> Value:
    if isinstance(value, dt.datetime | dt.date | dt.time | dt.timedelta):
        return moment_of(to_excel(value))
    return cast(Value, value)


def _result(result: object) -> Value:
    """A calculated result: errors come as their text, dates and times as moments."""
    if isinstance(result, str) and result in ERRORS:
        return ERRORS[result]
    return _plain(result)
