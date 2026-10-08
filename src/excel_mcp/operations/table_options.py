"""Table options as Excel's Table Design tab offers them: header and totals rows, calculated
columns, banding, emphasis, filter buttons and resizing."""

from typing import Annotated, Literal

from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import Field, StringConstraints

from excel_mcp.errors import InvalidArgumentError, LimitExceededError
from excel_mcp.inputs import InputModel
from excel_mcp.operations.cells import store_value, writable_cell
from excel_mcp.operations.shifting import Mover
from excel_mcp.refs import MAX_ROW, CellRange, cell_name, parse_range
from excel_mcp.structured import StructuredRef, qualify_formula
from excel_mcp.values import to_cell
from excel_mcp.workspace import sheet_names

TotalFunction = Literal[
    "none", "sum", "average", "count", "count_numbers", "max", "min", "std_dev", "var"
]
# What Excel stores: the totalsRowFunction and the SUBTOTAL function number of the cell.
_FUNCTIONS = {
    "sum": ("sum", 109),
    "average": ("average", 101),
    "count": ("count", 103),
    "count_numbers": ("countNums", 102),
    "max": ("max", 104),
    "min": ("min", 105),
    "std_dev": ("stdDev", 107),
    "var": ("var", 110),
}


class ColumnOptions(InputModel):
    name: str = Field(description="The column's header text.")
    formula: str | None = Field(
        default=None,
        description="Makes it a calculated column: the formula fills every row, and rows added "
        "later. Use structured references, e.g. '=[@Price]*[@Qty]' (the same row) or "
        "'=Sales[Price]'; cell references must be absolute.",
    )
    total: TotalFunction | Annotated[str, StringConstraints(pattern=r"^=")] | None = Field(
        default=None,
        description="Totals row cell: a function (shown as SUBTOTAL, so it ignores filtered-out "
        "rows), 'none', or a formula like '=SUM([Price])/2'.",
    )
    total_label: str | None = Field(
        default=None, description="Totals row text instead of a total, e.g. 'Total'."
    )


class TableOptions(InputModel):
    """Fields left out keep the table's setting; a new table gets Excel's defaults."""

    style: str | None = Field(
        default=None,
        description="Built-in style, e.g. 'TableStyleMedium9' (Light1-21, Medium1-28, Dark1-11).",
    )
    header_row: bool | None = Field(
        default=None,
        description="Turning it off deletes the header cells and the table shrinks to its data; "
        "turning it on needs empty cells above the table.",
    )
    totals_row: bool | None = Field(
        default=None,
        description="A row below the table, which must be empty; its first cell says 'Total'.",
    )
    striped_rows: bool | None = None
    striped_columns: bool | None = None
    first_column: bool | None = Field(default=None, description="Emphasize the first column.")
    last_column: bool | None = Field(default=None, description="Emphasize the last column.")
    filter_button: bool | None = Field(
        default=None, description="Filter buttons in the header row."
    )
    columns: list[ColumnOptions] = Field(
        default=[], description="Calculated columns and totals, by header text."
    )


def _data_rows(table: Table) -> tuple[int, int]:
    area = parse_range(table.ref)
    return area.min_row + (table.headerRowCount != 0), area.max_row - bool(table.totalsRowCount)


def apply_options(
    sheet: Worksheet, table: Table, options: TableOptions, resize: CellRange | None, max_cells: int
) -> None:
    """Make the table as the options and the new range say, in the order Excel's steps go."""
    if resize is not None:
        _resize(sheet, table, resize, max_cells)
    if options.header_row is not None and options.header_row != (table.headerRowCount != 0):
        _set_header(sheet, table, options.header_row)
    if options.totals_row is not None and options.totals_row != bool(table.totalsRowCount):
        _set_totals_row(sheet, table, options.totals_row)
    for column in options.columns:
        _set_column(sheet, table, column, max_cells)
    style = table.tableStyleInfo or TableStyleInfo()
    table.tableStyleInfo = style
    for field, attribute in (
        ("striped_rows", "showRowStripes"),
        ("striped_columns", "showColumnStripes"),
        ("first_column", "showFirstColumn"),
        ("last_column", "showLastColumn"),
    ):
        if (value := getattr(options, field)) is not None:
            setattr(style, attribute, value)
    if options.style is not None:
        style.name = options.style
    _set_filter(table, options.filter_button)


def _set_filter(table: Table, wanted: bool | None) -> None:
    """Filter buttons cover the header and data rows, not the totals row."""
    if table.headerRowCount == 0:
        if wanted:
            raise InvalidArgumentError("filter_button needs a header row.")
        table.autoFilter = None
        return
    if wanted is False:
        table.autoFilter = None
    elif wanted or table.autoFilter is not None:
        area = parse_range(table.ref)
        last = area.max_row - bool(table.totalsRowCount)
        buttons = table.autoFilter or AutoFilter()
        buttons.ref = str(CellRange(area.min_row, area.min_col, last, area.max_col))
        buttons.filterColumn = [c for c in buttons.filterColumn if c.colId < area.cols]
        table.autoFilter = buttons


def _set_header(sheet: Worksheet, table: Table, shown: bool) -> None:
    area = parse_range(table.ref)
    if not shown:
        for column in range(area.min_col, area.max_col + 1):
            writable_cell(sheet, area.min_row, column).value = None
        table.ref = str(CellRange(area.min_row + 1, area.min_col, area.max_row, area.max_col))
        table.headerRowCount = 0
        return
    if area.min_row == 1:
        raise InvalidArgumentError("There is no row above the table for the header row.")
    row = area.min_row - 1
    _require_empty(sheet, row, area.min_col, area.max_col, "the header row")
    for index, column in enumerate(table.tableColumns, start=area.min_col):
        store_value(writable_cell(sheet, row, index), column.name)
    table.ref = str(CellRange(row, area.min_col, area.max_row, area.max_col))
    table.headerRowCount = 1
    table.autoFilter = AutoFilter()


def _set_totals_row(sheet: Worksheet, table: Table, shown: bool) -> None:
    area = parse_range(table.ref)
    if not shown:
        for index, column in enumerate(table.tableColumns, start=area.min_col):
            writable_cell(sheet, area.max_row, index).value = None
            _clear_total(column)
        table.ref = str(CellRange(area.min_row, area.min_col, area.max_row - 1, area.max_col))
        table.totalsRowCount = None
        return
    row = area.max_row + 1
    if row > MAX_ROW:
        raise InvalidArgumentError("There is no row below the table for the totals row.")
    _require_empty(sheet, row, area.min_col, area.max_col, "the totals row")
    table.ref = str(CellRange(area.min_row, area.min_col, row, area.max_col))
    table.totalsRowCount = 1
    first = table.tableColumns[0]
    first.totalsRowLabel = "Total"
    store_value(writable_cell(sheet, row, area.min_col), "Total")


def _require_empty(sheet: Worksheet, row: int, first: int, last: int, what: str) -> None:
    for column in range(first, last + 1):
        if sheet.cell(row, column).value is not None:
            raise InvalidArgumentError(
                f"{cell_name(row, column)} is not empty; {what} needs row {row} to be."
            )


def _clear_total(column: TableColumn) -> None:
    column.totalsRowFunction = column.totalsRowLabel = column.totalsRowFormula = None


def _set_column(sheet: Worksheet, table: Table, options: ColumnOptions, max_cells: int) -> None:
    names = [column.name.casefold() for column in table.tableColumns]
    if options.name.casefold() not in names:
        listed = ", ".join(column.name for column in table.tableColumns)
        raise InvalidArgumentError(f"The table has no column {options.name!r}. Columns: {listed}.")
    index = names.index(options.name.casefold())
    column = table.tableColumns[index]
    position = parse_range(table.ref).min_col + index
    if options.formula is not None:
        _set_formula(sheet, table, column, position, options.formula, max_cells)
    if options.total is not None or options.total_label is not None:
        _set_total(sheet, table, column, position, options)


def _set_formula(
    sheet: Worksheet, table: Table, column: TableColumn, position: int, formula: str, max_cells: int
) -> None:
    text = qualify_formula(f"={formula.removeprefix('=')}", table.displayName)
    if Mover(1, 0).operand(text, sheet.title) != text:
        raise InvalidArgumentError(
            "A calculated column formula cannot use relative cell references like A2: write "
            "'[@Price]' for the same row's Price, or '$A$2' for one fixed cell."
        )
    first, last = _data_rows(table)
    if last - first + 1 > max_cells:
        raise LimitExceededError(f"The column has more than {max_cells:,} rows.")
    stored = to_cell(text, sheet_names(sheet))
    column.calculatedColumnFormula = TableFormula(attr_text=text.removeprefix("="))
    for row in range(first, last + 1):
        writable_cell(sheet, row, position).value = stored


def _set_total(
    sheet: Worksheet, table: Table, column: TableColumn, position: int, options: ColumnOptions
) -> None:
    if not table.totalsRowCount:
        raise InvalidArgumentError("Totals need a totals row: set totals_row to true.")
    if options.total is not None and options.total_label is not None:
        raise InvalidArgumentError(f"Column {column.name!r}: give a total or a total_label.")
    cell = writable_cell(sheet, parse_range(table.ref).max_row, position)
    _clear_total(column)
    cell.value = None
    total = options.total
    if options.total_label is not None:
        column.totalsRowLabel = options.total_label
        store_value(cell, options.total_label)
    elif total is not None and total.startswith("="):
        text = qualify_formula(total, table.displayName)
        column.totalsRowFunction = "custom"
        column.totalsRowFormula = TableFormula(attr_text=text.removeprefix("="))
        cell.value = to_cell(text, sheet_names(sheet))
    elif total in _FUNCTIONS:
        function, number = _FUNCTIONS[total]
        column.totalsRowFunction = function  # pyright: ignore[reportAttributeAccessIssue]
        found = StructuredRef(None, (), column.name, None).long_form(table.displayName)
        cell.value = f"=SUBTOTAL({number},{found})"


def _resize(sheet: Worksheet, table: Table, new: CellRange, max_cells: int) -> None:
    """Move the table's last row and column, as dragging its resize handle does."""
    old = parse_range(table.ref)
    if (new.min_row, new.min_col) != (old.min_row, old.min_col):
        raise InvalidArgumentError(f"The table's top-left cell stays {old.top_left}.")
    if table.totalsRowCount:
        raise InvalidArgumentError(
            "Resizing a table with a totals row is not supported: set totals_row to false, "
            "resize, and turn it on again."
        )
    for other in sheet.tables.values():
        if other is not table and new.overlaps(parse_range(other.ref)):
            raise InvalidArgumentError(f"{new} overlaps table {other.displayName!r}.")
    if new.rows < 1 + (table.headerRowCount != 0):
        raise InvalidArgumentError("A table needs at least one data row.")
    columns = table.tableColumns
    del columns[new.cols :]
    for index in range(len(columns), new.cols):
        columns.append(_new_column(sheet, table, old, new.min_col + index, columns))
    table.tableColumns = columns
    first = old.max_row + 1
    table.ref = str(new)
    for index, column in enumerate(columns, start=new.min_col):
        formula = column.calculatedColumnFormula
        if formula is not None and formula.attr_text and new.max_row >= first:
            stored = to_cell(f"={formula.attr_text}", sheet_names(sheet))
            for row in range(first, new.max_row + 1):
                writable_cell(sheet, row, index).value = stored


def _new_column(
    sheet: Worksheet, table: Table, old: CellRange, position: int, columns: list[TableColumn]
) -> TableColumn:
    """A column named by the header text in its cell, or Excel's 'ColumnN' if there is none."""
    taken = {column.name.casefold() for column in columns}
    header = table.headerRowCount != 0
    cell = sheet.cell(old.min_row, position)
    name = cell.value if header and cell.value != "" else None
    if name is None:
        number = 1
        while f"column{number}" in taken:
            number += 1
        name = f"Column{number}"
    elif not isinstance(name, str):
        raise InvalidArgumentError(f"{cell.coordinate} must be text to be a column header.")
    elif name.casefold() in taken:
        raise InvalidArgumentError(f"The header {name!r} in {cell.coordinate} is not unique.")
    if header:
        store_value(writable_cell(sheet, old.min_row, position), name)
    return TableColumn(id=max((c.id for c in columns), default=0) + 1, name=name)
