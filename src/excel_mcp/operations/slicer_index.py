"""Finding the slicers, timelines and caches a workbook has."""

from dataclasses import dataclass

from openpyxl.pivot.table import TableDefinition
from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.errors import InvalidArgumentError
from excel_mcp.operations import slicer_xml as xml
from excel_mcp.operations.pivot_index import sheet_pivots
from excel_mcp.package import Link, PackageState, Part, state_of
from excel_mcp.package.consistency import slicer_caches
from excel_mcp.text import quoted
from excel_mcp.workspace import get_sheet, worksheets


@dataclass(frozen=True)
class Entry:
    """A slicer or timeline on a sheet."""

    sheet: Worksheet
    kind: str
    values: dict[str, str]
    part: Part
    link: Link

    @property
    def name(self) -> str:
        return self.values["name"]


def entries(workbook: Workbook, sheet: Worksheet | None = None) -> list[Entry]:
    state = state_of(workbook)
    found = []
    for candidate in [sheet] if sheet else worksheets(workbook):
        package = state.sheets.get(candidate)
        for link in package.links if package else []:
            part = link.target
            if isinstance(part, Part) and part.content_type in (xml.SLICER_PART, xml.TIMELINE_PART):
                kind = "timeline" if part.content_type == xml.TIMELINE_PART else "slicer"
                found += [
                    Entry(candidate, kind, values, part, link)
                    for values in xml.read_entries(part.data.decode("utf-8"))
                ]
    return found


def caches(workbook: Workbook) -> dict[str, tuple[Link, xml.CacheInfo]]:
    """The slicer and timeline caches by name."""
    return {
        name: (link, xml.read_cache(link.target.data.decode("utf-8")))  # pyright: ignore[reportAttributeAccessIssue]
        for name, link in slicer_caches(state_of(workbook)).items()
    }


def sheet_number(state: PackageState, sheet: Worksheet) -> int:
    """The number the sheet has in the file; a sheet made since gets a free one."""
    if sheet not in state.sheet_ids:
        state.sheet_ids[sheet] = max(state.sheet_ids.values(), default=0) + 1
    return state.sheet_ids[sheet]


def table_number(state: PackageState, table: Table) -> int:
    if table not in state.table_ids:
        state.table_ids[table] = max(state.table_ids.values(), default=0) + 1
    return state.table_ids[table]


def sheet_of(state: PackageState, number: int) -> Worksheet | None:
    return next((s for s, n in state.sheet_ids.items() if n == number), None)  # pyright: ignore[reportReturnType]


def pivot_of(workbook: Workbook, number: int, name: str) -> tuple[Worksheet, TableDefinition]:
    sheet = sheet_of(state_of(workbook), number)
    if isinstance(sheet, Worksheet):
        for pivot in sheet_pivots(sheet):
            if pivot.name == name:
                return sheet, pivot
    raise InvalidArgumentError(f"PivotTable {name!r} of a slicer was not found.")


def table_of(workbook: Workbook, number: int) -> tuple[Worksheet, Table]:
    state = state_of(workbook)
    for sheet in worksheets(workbook):
        for table in sheet.tables.values():
            if state.table_ids.get(table) == number:
                return sheet, table
    raise InvalidArgumentError("The table of a slicer was not found.")


def find_source(
    workbook: Workbook, sheet_name: str, name: str
) -> tuple[Worksheet, Table | TableDefinition]:
    """The table or PivotTable called ``name`` on a sheet."""
    sheet = get_sheet(workbook, sheet_name)
    wanted = name.strip().casefold()
    tables = [t for t in sheet.tables.values() if t.displayName.casefold() == wanted]
    pivots = [p for p in sheet_pivots(sheet) if p.name.casefold() == wanted]
    if len(tables) + len(pivots) == 1:
        return sheet, (tables or pivots)[0]
    if tables and pivots:
        raise InvalidArgumentError(f"{name!r} is both a table and a PivotTable on {sheet_name!r}.")
    available = [
        *(t.displayName for t in sheet.tables.values()),
        *(p.name for p in sheet_pivots(sheet)),
    ]
    raise InvalidArgumentError(
        f"Sheet {sheet_name!r} has no table or PivotTable {name!r}. "
        f"Available: {quoted(available) if available else 'none'}."
    )
