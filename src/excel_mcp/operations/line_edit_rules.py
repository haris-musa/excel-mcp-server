"""Conditional formats and data validation when rows or columns are inserted or deleted.

Their formulas are written for the top-left cell of the range and apply to every cell
with relative references. Excel moves every cell's references on its own, so a range
that the edit cuts between a cell and the cells it refers to ends up with different
formulas in different parts: such a rule is split into one rule per part.
"""

from copy import copy
from dataclasses import dataclass, replace

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.cell_range import CellRange as SheetRange
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.formulas import storable_operand
from excel_mcp.operations.shifting import Mover, Shifter
from excel_mcp.operations.workbook_rewrite import rewritten_operand, rule_points
from excel_mcp.refs import CellRange
from excel_mcp.rewrite import ReferenceRewriter


@dataclass
class _Group:
    areas: list[CellRange]
    formulas: list[str]


def update_rules(worksheet: Worksheet, shifter: Shifter, sheet_names: list[str]) -> None:
    """Rewrite the conditional formats and data validations of one sheet."""
    host = worksheet.title

    def moved(areas: list[SheetRange], formulas: list[str]) -> list[_Group]:
        groups = _relocate([CellRange.of(a) for a in areas], formulas, host, shifter)
        for group in groups:
            group.formulas = [
                f if f in formulas else storable_operand(f, sheet_names) for f in group.formulas
            ]
        return groups

    rules = ConditionalFormattingList()
    for entry in worksheet.conditional_formatting:
        for rule in entry.rules:
            for point in rule_points(rule):
                point.val = rewritten_operand(str(point.val), host, shifter, sheet_names)
            for group in moved(list(entry.sqref.ranges), list(rule.formula or [])):
                clone = copy(rule)
                clone.formula = group.formulas
                rules.add(_sqref(group.areas), clone)
    worksheet.conditional_formatting = rules

    validations = []
    for validation in worksheet.data_validations.dataValidation:
        operands = [f for f in (validation.formula1, validation.formula2) if f]
        for group in moved(list(validation.sqref.ranges), operands):
            clone = copy(validation)
            clone.sqref = _sqref(group.areas)
            if validation.formula1:
                clone.formula1 = group.formulas[0]
            if validation.formula2:
                clone.formula2 = group.formulas[-1]
            validations.append(clone)
    worksheet.data_validations.dataValidation = validations


def _relocate(
    areas: list[CellRange], formulas: list[str], host: str, shifter: Shifter
) -> list[_Group]:
    edit = shifter.edit
    rows = edit.axis == "rows"
    moves = host.casefold() == edit.sheet.casefold()
    corner = (min(a.min_row for a in areas), min(a.min_col for a in areas))
    offsets = {0} if moves else set()
    for formula in formulas:
        finder = _OffsetFinder(shifter, corner[0] if rows else corner[1])
        finder.operand(formula, host)
        offsets |= finder.offsets
    cuts = {edit.at - d for d in offsets}
    if edit.delete:
        cuts |= {edit.end + 1 - d for d in offsets}

    pieces: list[tuple[CellRange, list[str]]] = []
    for area in areas:
        low, high = (area.min_row, area.max_row) if rows else (area.min_col, area.max_col)
        starts = [low, *sorted(c for c in cuts if low < c <= high)]
        for start, last in zip(starts, [s - 1 for s in starts[1:]] + [high], strict=True):
            span = edit.span(start, last) if moves else (start, last)
            if span is None:
                continue
            if moves and not edit.delete and last == edit.at - 1:
                span = (span[0], span[1] + edit.count)  # Excel extends a range over inserted lines.
            cell = (start, area.min_col) if rows else (area.min_row, start)
            texts = [shifter.operand(_retarget(f, corner, cell, host), host) for f in formulas]
            new = (
                replace(area, min_row=span[0], max_row=span[1])
                if rows
                else replace(area, min_col=span[0], max_col=span[1])
            )
            pieces.append((new, texts))
    return _group(pieces, host)


def _retarget(formula: str, origin: tuple[int, int], cell: tuple[int, int], host: str) -> str:
    """The formula as written for ``cell`` instead of ``origin``, as when copying."""
    return Mover(cell[0] - origin[0], cell[1] - origin[1]).operand(formula, host)


def _group(pieces: list[tuple[CellRange, list[str]]], host: str) -> list[_Group]:
    """Parts whose formulas are equal for their own first cells share one rule."""
    groups: list[tuple[tuple[int, int], _Group]] = []
    for area, formulas in pieces:
        corner = (area.min_row, area.min_col)
        for group_corner, group in groups:
            if [_retarget(f, corner, group_corner, host) for f in formulas] == group.formulas:
                group.areas.append(area)
                break
        else:
            groups.append((corner, _Group([area], formulas)))
    return [_Group(_merge(group.areas), group.formulas) for _, group in groups]


def _merge(areas: list[CellRange]) -> list[CellRange]:
    """Join areas that touch along a full side."""
    merged: list[CellRange] = []
    for area in sorted(areas, key=lambda a: (a.min_row, a.min_col)):
        for index, other in enumerate(merged):
            if _touch(other, area):
                merged[index] = CellRange(
                    min(other.min_row, area.min_row),
                    min(other.min_col, area.min_col),
                    max(other.max_row, area.max_row),
                    max(other.max_col, area.max_col),
                )
                break
        else:
            merged.append(area)
    return merged


def _touch(a: CellRange, b: CellRange) -> bool:
    same_columns = (a.min_col, a.max_col) == (b.min_col, b.max_col)
    same_rows = (a.min_row, a.max_row) == (b.min_row, b.max_row)
    return (same_columns and a.max_row + 1 >= b.min_row and b.max_row + 1 >= a.min_row) or (
        same_rows and a.max_col + 1 >= b.min_col and b.max_col + 1 >= a.min_col
    )


def _sqref(areas: list[CellRange]) -> str:
    return " ".join(str(area) for area in areas)


class _OffsetFinder(ReferenceRewriter):
    """Collects how far the relative references of a formula lie from the rule's first cell."""

    def __init__(self, shifter: Shifter, origin: int) -> None:
        self.shifter = shifter
        self.origin = origin
        self.offsets: set[int] = set()

    def reference(self, text: str, host: str) -> str:
        target = self.shifter.target(text, host)
        for part in target[1] if target else []:
            position, fixed = self.shifter.position(part)
            if position is not None and not fixed:
                self.offsets.add(position - self.origin)
        return text
