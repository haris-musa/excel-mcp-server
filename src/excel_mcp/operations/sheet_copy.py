"""Copying a worksheet the way Excel's Move or Copy > Create a copy does."""

from copy import copy, deepcopy
from io import BytesIO

from openpyxl.drawing.image import Image
from openpyxl.workbook import Workbook
from openpyxl.worksheet.copier import WorksheetCopy
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.formulas import storable_formula, storable_operand
from excel_mcp.operations import chartex
from excel_mcp.operations.cells import stored_cells
from excel_mcp.operations.sheet_refs import SheetCopyRefs
from excel_mcp.operations.sheet_shapes import copy_shapes, copy_sheet_elements
from excel_mcp.operations.sheets import validate_sheet_name
from excel_mcp.operations.slicer_manage import copy_slicers
from excel_mcp.operations.tables import table_names
from excel_mcp.package import arrays, extensions, state_of
from excel_mcp.rewrite import chart_references
from excel_mcp.workspace import get_sheet


def copy_sheet(workbook: Workbook, name: str, new_name: str) -> list[str]:
    """Add a copy of a sheet at the end of the workbook. Returns what could not be copied."""
    source = get_sheet(workbook, name)
    validate_sheet_name(new_name, workbook.sheetnames)
    target = workbook.create_sheet(new_name)
    WorksheetCopy(source, target).copy_worksheet()
    tables = _copy_tables(workbook, source, target)
    refs = SheetCopyRefs(source.title, target.title, tables)
    _copy_formulas(workbook, target, refs)
    _copy_rules(workbook, source, target, refs)
    _copy_extensions(workbook, source, target, refs)
    _copy_names(workbook, source, target, refs)
    _copy_sheet_settings(source, target)
    _copy_drawings(workbook, source, target, refs)
    _copy_pivots(source, target)
    arrays.copy_marks(source, target)
    copy_slicers(workbook, source, target, tables)
    return copy_sheet_elements(workbook, source, target)


def _copy_tables(workbook: Workbook, source: Worksheet, target: Worksheet) -> dict[str, str]:
    """Copies get new names, Excel-style: the table's name followed by the next free number."""
    taken = table_names(workbook) | {name.casefold() for name in workbook.defined_names}
    renamed: dict[str, str] = {}
    for table in source.tables.values():
        number = 2
        while f"{table.displayName}{number}".casefold() in taken:
            number += 1
        new_name = f"{table.displayName}{number}"
        taken.add(new_name.casefold())
        renamed[table.displayName.casefold()] = new_name
        clone = deepcopy(table)
        clone.name = clone.displayName = new_name
        target.add_table(clone)
    return renamed


def _copy_formulas(workbook: Workbook, target: Worksheet, refs: SheetCopyRefs) -> None:
    names = workbook.sheetnames
    for cell in stored_cells(target):
        value = cell.value
        if isinstance(value, ArrayFormula):
            cell.value = ArrayFormula(
                value.ref, storable_formula(refs.formula(str(value.text), target.title), names)
            )
        elif cell.data_type == "f" and isinstance(value, str):
            cell.value = storable_formula(refs.formula(value, target.title), names)
    for table in target.tables.values():
        for column in table.tableColumns:
            for formula in (column.calculatedColumnFormula, column.totalsRowFormula):
                if formula is not None and formula.attr_text:
                    formula.attr_text = refs.operand(formula.attr_text, target.title)


def _copy_rules(
    workbook: Workbook, source: Worksheet, target: Worksheet, refs: SheetCopyRefs
) -> None:
    names = workbook.sheetnames
    for validation in source.data_validations.dataValidation:
        clone = deepcopy(validation)
        for field in ("formula1", "formula2"):
            operand = getattr(clone, field)
            if operand:
                setattr(clone, field, storable_operand(refs.operand(operand, target.title), names))
        target.add_data_validation(clone)
    for entry in source.conditional_formatting:
        for rule in entry.rules:
            clone = deepcopy(rule)
            clone.formula = [
                storable_operand(refs.operand(f, target.title), names) for f in rule.formula
            ]
            target.conditional_formatting.add(str(entry.sqref), clone)


def _copy_extensions(
    workbook: Workbook, source: Worksheet, target: Worksheet, refs: SheetCopyRefs
) -> None:
    """Sparklines and the Excel 2010 half of conditional formats. The copy declares the prefixes
    of the original's XML, which the shapes and parts copied later rely on."""
    state = state_of(workbook)
    package = state.sheets.get(source)
    if package is None:
        return
    found, rules = extensions.copied(
        package.extensions,
        package.rule_extensions,
        lambda text: refs.operand(text, target.title),
    )
    copy = state.sheet(target)
    copy.namespaces = dict(package.namespaces)
    copy.drawing_namespaces = dict(package.drawing_namespaces)
    copy.vml_namespaces = dict(package.vml_namespaces)
    copy.extensions.update(found)
    copy.rule_extensions.update(rules)


def _copy_names(
    workbook: Workbook, source: Worksheet, target: Worksheet, refs: SheetCopyRefs
) -> None:
    for name, defined in source.defined_names.items():
        clone = copy(defined)
        clone.attr_text = storable_operand(
            refs.operand(defined.attr_text or "", target.title), workbook.sheetnames
        )
        target.defined_names[name] = clone


def _copy_sheet_settings(source: Worksheet, target: Worksheet) -> None:
    """What openpyxl's own worksheet copy leaves out."""
    target.views = deepcopy(source.views)
    for view in target.views.sheetView:
        view.tabSelected = False
    target.sheet_properties.codeName = None
    target.auto_filter = copy(source.auto_filter)
    target.protection = copy(source.protection)
    target.HeaderFooter = deepcopy(source.HeaderFooter)
    target.row_breaks = deepcopy(source.row_breaks)
    target.col_breaks = deepcopy(source.col_breaks)
    target.print_area = source.print_area
    target.print_title_rows = source.print_title_rows
    target.print_title_cols = source.print_title_cols


def _copy_drawings(
    workbook: Workbook, source: Worksheet, target: Worksheet, refs: SheetCopyRefs
) -> None:
    for image in source._images:  # pyright: ignore[reportAttributeAccessIssue]
        clone = Image(BytesIO(image.ref.getvalue()))
        clone.width, clone.height = image.width, image.height
        clone.anchor = deepcopy(image.anchor)
        target.add_image(clone)
    for chart in source._charts:  # pyright: ignore[reportAttributeAccessIssue]
        clone = deepcopy(chart)
        for plot in clone._charts:  # pyright: ignore[reportAttributeAccessIssue]
            for reference in chart_references(plot):
                reference.f = refs.operand(str(reference.f), target.title)
        target.add_chart(clone)
    chartex.copy(source, target, refs)
    copy_shapes(workbook, source, target, refs)


def _copy_pivots(source: Worksheet, target: Worksheet) -> None:
    """Copies share the source's data cache, as in Excel."""
    for pivot in source._pivots:  # pyright: ignore[reportAttributeAccessIssue]
        target.add_pivot(deepcopy(pivot, {id(pivot.cache): pivot.cache}))
