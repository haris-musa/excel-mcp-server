"""Dynamic array formulas: storing a formula that spills the way Excel does."""

from decimal import Decimal
from typing import cast

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from excel_mcp.calc.dynamic import is_dynamic_array
from excel_mcp.calc.engine import Engine
from excel_mcp.calc.parser import Name
from excel_mcp.calc.values import ExcelError, Grid, Scalar, UncalculableError
from excel_mcp.errors import LimitExceededError
from excel_mcp.package import CellMark, arrays, state_of
from excel_mcp.refs import MAX_COLUMN, MAX_ROW, CellRange
from excel_mcp.values import store_value

_STORED_ERRORS = frozenset(["#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"])


def spill_formulas(sheet: Worksheet, cells: list[Cell], max_cells: int) -> list[str]:
    """Store each of the formulas that Excel would spill as a dynamic array formula.

    The formula's cell gets the array formula with the range it fills, and the cells in
    that range get the results, as Excel saves them. Where the calculator cannot tell what
    a formula returns, only its own cell is declared; Excel works out the rest when it
    opens the file. A range that holds data is not overwritten: Excel shows #SPILL! there,
    and the cell is returned among those whose formulas are blocked.
    """
    workbook = cast(Workbook, sheet.parent)
    resolve = _name_resolver(workbook, sheet)
    engine: Engine | None = None
    blocked = []
    for cell in cells:
        formula = cell.value
        if not (isinstance(formula, str) and is_dynamic_array(formula, resolve)):
            continue
        engine = engine or Engine(Workbook(), workbook)
        try:
            grid: Grid | None = engine.spill(sheet, cell.row, cell.column, formula)
        except UncalculableError:
            grid = None
        if grid is not None and _has_new_errors(grid):
            grid = None
        if not _store(sheet, cell, formula, grid, max_cells):
            blocked.append(cell.coordinate)
    return blocked


def _has_new_errors(grid: Grid) -> bool:
    """Whether a result holds an error such as #CALC! that a file cannot store as a value.

    Excel saves those with a rich value this server does not write; it calculates them
    again when it opens the file, so nothing is stored for the formula.
    """
    return any(
        isinstance(item, ExcelError) and item.code not in _STORED_ERRORS
        for line in grid.rows
        for item in line
    )


def _store(sheet: Worksheet, cell: Cell, formula: str, grid: Grid | None, max_cells: int) -> bool:
    """Store a formula; False if cells in its way hold data, so that it could not spill."""
    package = state_of(cast(Workbook, sheet.parent)).sheet(sheet)
    for earlier in (m for m in package.marks if m.cell is cell):
        arrays.release(earlier)
    height, width = (grid.height, grid.width) if grid else (1, 1)
    if height * width > max_cells:
        raise LimitExceededError(
            f"The formula in {cell.coordinate} returns {height * width:,} values; "
            f"at most {max_cells:,} cells can be written per call."
        )
    area = CellRange(cell.row, cell.column, cell.row + height - 1, cell.column + width - 1)
    fits = grid is not None and is_free(sheet, area, cell)
    if not fits:
        area = CellRange(cell.row, cell.column, cell.row, cell.column)
    cell.value = ArrayFormula(str(area), formula)
    mark = CellMark(cell, dynamic=True, value=(cell.value, cell.data_type))
    if grid is not None and fits:
        mark.cached = _cached(grid.rows[0][0])
        mark.spill = _fill(sheet, cell, grid)
    package.mark(mark)
    return grid is None or fits


def is_free(sheet: Worksheet, area: CellRange, anchor: Cell) -> bool:
    """Whether the cells a formula would spill into are empty, so that it can."""
    if area.max_row > MAX_ROW or area.max_col > MAX_COLUMN:
        return False
    for row in range(area.min_row, area.max_row + 1):
        for col in range(area.min_col, area.max_col + 1):
            other = sheet._cells.get((row, col))
            if other is None or other is anchor:
                continue
            if isinstance(other, MergedCell) or other.value is not None:
                return False
    return True


def _fill(sheet: Worksheet, anchor: Cell, grid: Grid) -> list[tuple[Cell, object]]:
    spilled = []
    for row_offset, line in enumerate(grid.rows):
        for col_offset, item in enumerate(line):
            if row_offset == 0 and col_offset == 0:
                continue
            target = cast(Cell, sheet.cell(anchor.row + row_offset, anchor.column + col_offset))
            if isinstance(item, ExcelError):
                target.value = item.code
            else:
                store_value(target, _stored(item))
                if target.data_type == "s":
                    target.data_type = "str"  # how Excel stores the text a formula spilled
            spilled.append((target, (target.value, target.data_type)))
    return spilled


def _stored(item: Scalar) -> str | int | float | bool:
    if isinstance(item, ExcelError):
        return item.code
    if item is None:
        return 0
    if isinstance(item, Decimal):
        return float(item)
    return item if isinstance(item, bool | str) else _number(item)


def _number(value: float) -> int | float:
    number = float(f"{value:.15g}")
    return int(number) if number == int(number) and abs(number) < 1e15 else number


def _cached(item: Scalar) -> tuple[str, str]:
    """A result as the ``t`` and ``<v>`` Excel stores for a formula cell."""
    stored = _stored(item)
    if isinstance(item, ExcelError):
        return "e", item.code
    if isinstance(stored, bool):
        return "b", "1" if stored else "0"
    if isinstance(stored, str):
        return "str", stored
    return "n", repr(stored)


def _name_resolver(workbook: Workbook, sheet: Worksheet):
    def resolve(name: Name) -> str | None:
        scope = workbook[name.sheet] if name.sheet in workbook.sheetnames else sheet
        for names in (scope.defined_names, workbook.defined_names):
            for key, defined in names.items():
                if key.casefold() == name.name.casefold() and defined.attr_text:
                    return defined.attr_text
        return None

    return resolve
