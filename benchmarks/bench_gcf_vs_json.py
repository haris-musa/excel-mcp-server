#!/usr/bin/env python3
"""Benchmark: GCF vs JSON token counts for typical Excel MCP responses.

Generates realistic spreadsheet-shaped payloads at various sizes and compares
the token cost of returning them as GCF (generic profile) vs indented JSON.

Token estimation: max(1, len(text) // 4)  (standard BPE approximation)
"""

import json
import random
import string
import sys
import os
import datetime

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def estimate_tokens(text: str) -> int:
    return max(1, len(text) // 4)


def random_string(length: int = 8) -> str:
    return "".join(random.choices(string.ascii_letters, k=length))


def generate_cells(rows: int, cols: int) -> dict:
    """Build a read_data_from_excel-style response dict."""
    headers = [f"Col_{i}" for i in range(1, cols + 1)]
    cells = []
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            from openpyxl.utils import get_column_letter
            addr = f"{get_column_letter(c)}{r}"
            if r == 1:
                val = headers[c - 1]
            else:
                # Mix of types: strings, ints, floats, formulas
                kind = random.choice(["str", "int", "float", "formula"])
                if kind == "str":
                    val = random_string(random.randint(4, 16))
                elif kind == "int":
                    val = random.randint(0, 100_000)
                elif kind == "float":
                    val = round(random.uniform(0, 100_000), 2)
                else:
                    val = f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{r})"
            cells.append({
                "address": addr,
                "value": val,
                "row": r,
                "column": c,
                "validation": {"has_validation": False},
            })
    return {
        "range": f"A1:{get_column_letter(cols)}{rows}",
        "sheet_name": "Sheet1",
        "cells": cells,
    }


# ---------------------------------------------------------------------------
# Benchmark runner
# ---------------------------------------------------------------------------

def run_benchmark():
    from gcf import encode_generic

    scenarios = [
        ("Small (10 rows x 5 cols)", 10, 5),
        ("Medium (50 rows x 10 cols)", 50, 10),
        ("Large (100 rows x 15 cols)", 100, 15),
        ("XL (250 rows x 20 cols)", 250, 20),
        ("XXL (500 rows x 20 cols)", 500, 20),
    ]

    lines = []
    lines.append("=" * 72)
    lines.append("GCF vs JSON Token Benchmark for excel-mcp-server")
    lines.append(f"Date: {datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    lines.append("Token estimator: max(1, len(text) // 4)")
    lines.append("=" * 72)
    lines.append("")
    lines.append(f"{'Scenario':<30} {'JSON tok':>10} {'GCF tok':>10} {'Saved':>8} {'Reduction':>10}")
    lines.append("-" * 72)

    total_json = 0
    total_gcf = 0

    for label, rows, cols in scenarios:
        data = generate_cells(rows, cols)

        json_text = json.dumps(data, indent=2, default=str)
        gcf_text = encode_generic(data)

        json_tok = estimate_tokens(json_text)
        gcf_tok = estimate_tokens(gcf_text)
        saved = json_tok - gcf_tok
        pct = (saved / json_tok) * 100 if json_tok else 0

        total_json += json_tok
        total_gcf += gcf_tok

        lines.append(f"{label:<30} {json_tok:>10,} {gcf_tok:>10,} {saved:>+8,} {pct:>9.1f}%")

    lines.append("-" * 72)
    total_saved = total_json - total_gcf
    total_pct = (total_saved / total_json) * 100 if total_json else 0
    lines.append(f"{'TOTAL':<30} {total_json:>10,} {total_gcf:>10,} {total_saved:>+8,} {total_pct:>9.1f}%")
    lines.append("")
    lines.append("GCF = Graph Compact Format (https://github.com/blackwell-systems/gcf)")
    lines.append("gcf-python v2.1.0 | encode_generic(data)")

    report = "\n".join(lines) + "\n"
    print(report)
    return report


if __name__ == "__main__":
    report = run_benchmark()

    out_path = os.path.join(os.path.dirname(__file__), "results-2026-06-17.txt")
    with open(out_path, "w") as f:
        f.write(report)
    print(f"Results written to {out_path}")
