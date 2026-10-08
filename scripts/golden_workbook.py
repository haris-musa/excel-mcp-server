"""Builds the golden workbook and runs the calculator on it (shared by the generator and tests)."""

import math
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from excel_mcp.calc.engine import Engine
from excel_mcp.calc.values import ExcelError, UncalculableError
from excel_mcp.lazy_workbook import Parts, load_lazy
from excel_mcp.operations.cells import write_range
from excel_mcp.operations.table_options import TableOptions
from excel_mcp.operations.tables import create_table
from excel_mcp.workspace import save_atomically

MAX_CELLS = 1_000_000
DEFAULT_TOLERANCE = 1e-12
CASE_STRIDE = 25  # rows between cases, so that what a case spills never reaches the next one


def case_row(index: int) -> int:
    """The row of 'Cases' column A that holds the formula with this 0-based index."""
    return 1 + index * CASE_STRIDE


def build_workbook(
    path: Path,
    inputs: dict[str, list[list[Any]]],
    names: dict[str, str],
    formulas: list[str],
    hidden_rows: dict[str, list[int]],
    tables: list[dict[str, Any]],
) -> None:
    """Write the inputs and the formulas into 'Cases' column A, as the server would."""
    workbook = Workbook()
    workbook.remove(workbook.active)  # pyright: ignore[reportArgumentType]
    for title in inputs:
        workbook.create_sheet(title)
    for title, rows in inputs.items():
        if rows:
            write_range(workbook[title], "A1", rows, [], MAX_CELLS)  # pyright: ignore[reportArgumentType]
    for table in tables:
        sheet = workbook[table["sheet"]]
        options = TableOptions(**table["options"])
        create_table(workbook, sheet, table["range"], table["name"], options, MAX_CELLS)
    cases = workbook["Cases"]
    for index, formula in enumerate(formulas):
        write_range(cases, f"A{case_row(index)}", [[formula]], [], MAX_CELLS)  # pyright: ignore[reportArgumentType]
    for title, rows_hidden in hidden_rows.items():
        for row in rows_hidden:
            workbook[title].row_dimensions[row].hidden = True
    for name, target in names.items():
        workbook.defined_names[name] = DefinedName(name, attr_text=target)
    save_atomically(workbook, path)


def calculate_cases(path: Path, count: int, *, lazy: bool = False) -> list[object]:
    """The calculator's value for each of the first ``count`` cases, or None if uncalculated.

    ``lazy`` loads the workbook's cells row by row as the calculation asks for them.
    """
    if lazy:
        parts: Parts = {}
        formulas = load_lazy(path, parts, data_only=False)
        cached = load_lazy(path, parts, data_only=True)
    else:
        formulas = load_workbook(path)
        cached = load_workbook(path, data_only=True)
    engine = Engine(cached, formulas, str(path))
    sheet = formulas["Cases"]
    results: list[object] = []
    for index in range(count):
        try:
            value = engine.calculate(sheet, case_row(index), 1)
        except UncalculableError:
            results.append(None)
            continue
        results.append(value.code if isinstance(value, ExcelError) else value)
    return results


def same(expected: object, actual: object, tolerance: float = DEFAULT_TOLERANCE) -> bool:
    if isinstance(expected, bool) or isinstance(actual, bool):
        return type(expected) is type(actual) and expected == actual
    if isinstance(expected, int | float) and isinstance(actual, int | float):
        return math.isclose(expected, actual, rel_tol=tolerance, abs_tol=tolerance)
    return expected == actual
