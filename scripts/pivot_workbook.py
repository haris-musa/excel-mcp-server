"""Builds the golden PivotTable workbook and reads back its cells (shared by generator and test)."""

import datetime as dt
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.worksheet import Worksheet
from pivot_cases import SOURCES

from excel_mcp.calc.engine import Engine
from excel_mcp.calc.values import ExcelError
from excel_mcp.config import Limits
from excel_mcp.operations import pivot
from excel_mcp.operations.pivot import PivotRequest
from excel_mcp.operations.pivot_index import pivot_area, sheet_pivots
from excel_mcp.operations.pivot_options import CalculatedField, PivotField, PivotValue
from excel_mcp.refs import CellRange

LIMITS = Limits()


def source_range(name: str) -> str:
    rows = SOURCES[name]
    return f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"


def calculated(workbook: Workbook, sheet: Worksheet, area: CellRange) -> dict[str, Any]:
    """The results of the formulas of a sheet, as the server's calculator gives them."""
    engine = Engine(Workbook(), workbook)
    found: dict[str, Any] = {}
    for row in sheet.iter_rows(
        min_row=area.min_row, max_row=area.max_row, min_col=area.min_col, max_col=area.max_col
    ):
        for cell in row:
            if cell.data_type == "f":
                result = engine.calculate(sheet, cell.row, cell.column)
                found[cell.coordinate] = result.code if isinstance(result, ExcelError) else result
    return found


def add_sources(workbook: Workbook) -> None:
    """One sheet for each source, the first named Data."""
    for position, (name, rows) in enumerate(SOURCES.items()):
        sheet = workbook.worksheets[0] if position == 0 else workbook.create_sheet()
        sheet.title = name
        for row in rows:
            sheet.append(row)


def build_workbook(cases: list[dict[str, Any]]) -> Workbook:
    """The sources, then one sheet per case with the PivotTable the server creates for it."""
    workbook = Workbook()
    add_sources(workbook)
    for case in cases:
        request = PivotRequest(
            rows=case.get("rows", []),
            columns=case.get("columns", []),
            values=[PivotValue(**value) for value in case["values"]],
            filters=case.get("filters", []),
            fields=[PivotField(**field) for field in case.get("fields", [])],
            calculated=[CalculatedField(**field) for field in case.get("calculated_fields", [])],
            layout=case.get("layout", "tabular"),
            subtotals=case.get("subtotals", True),
            values_in=case.get("values_in", "columns"),
            name=case["name"],
        )
        name = case.get("source", "Data")
        target = workbook.create_sheet(case["name"])
        pivot.create_pivot(
            workbook,
            workbook[name],
            source_range(name),
            target,
            "A1",
            request,
            LIMITS,
            lambda sheet, area: calculated(workbook, sheet, area),
        )
    return workbook


def stored_grids(workbook: Workbook) -> dict[str, list[list[Any]]]:
    """The cells the server wrote for each PivotTable, as Excel's Value2 would show them."""
    grids = {}
    for sheet in workbook.worksheets[len(SOURCES) :]:
        area = pivot_area(sheet_pivots(sheet)[0])
        grids[sheet.title] = [
            [_value(cell.value, cell.data_type) for cell in row]
            for row in sheet.iter_rows(
                min_row=area.min_row,
                max_row=area.max_row,
                min_col=area.min_col,
                max_col=area.max_col,
            )
        ]
    return grids


def _value(value: Any, data_type: str) -> Any:
    if data_type == "e":
        return value
    if isinstance(value, dt.datetime):
        return to_excel(value)
    return value


def same(expected: Any, actual: Any) -> bool:
    if expected in (None, "") and actual in (None, ""):
        return True
    if isinstance(expected, int | float) and isinstance(actual, int | float):
        return abs(expected - actual) <= 1e-9 * max(1.0, abs(expected), abs(actual))
    return expected == actual
