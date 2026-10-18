"""The workbook and cases behind tests/fixtures/audit_golden.json (shared by recorder and test).

Formula tracing and Evaluate Formula run on the 'Sales' sheet; each Goal Seek case has a sheet
of its own, so that the cases cannot disturb one another.
"""

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from excel_mcp.config import Limits
from excel_mcp.operations.cells import write_range
from excel_mcp.operations.table_options import TableOptions
from excel_mcp.operations.tables import create_table
from excel_mcp.workspace import save_atomically

MAX_CELLS = 1_000_000
FILE = "audit.xlsx"

ASSUMPTIONS = [["Tax rate", 0.2], ["Fee", 5]]
SALES = [
    ["Region", "Units", "Price", "Revenue"],
    *[
        [region, units, price, "=tblSales[[#This Row],[Units]]*tblSales[[#This Row],[Price]]"]
        for region, units, price in [
            ("North", 10, 1.5),
            ("South", 5, 2.0),
            ("East", 7, 1.5),
            ("West", 3, 2.5),
            ("North", 12, 3.25),
            ("South", 8, 1.5),
            ("East", 9, 2.0),
            ("West", 4, 2.25),
        ]
    ],
]
FORMULAS = {
    "F1": "=SUM(tblSales[Revenue])",
    "F2": "=F1*TaxRate",
    "F3": '=SUMIFS(tblSales[Revenue],tblSales[Region],"North")',
    "F4": "=SUM(SalesRange)",
    "F5": "=SUM(H1#)",
    "F6": "=INDEX(B2:B9,3)+MIN(C2:C9)",
    "F7": "=IF(F1>100,F2,F3)",
    "F8": "=F7+Assumptions!B2",
    "F9": "=SUM(B:B)",
    "F10": "=AVERAGE(Sales!C2:C9)*Assumptions!B1",
    "F11": '=MAX(IF(A2:A9="North",D2:D9))',
    "G1": "=F1+F4",
    "G2": "=G1*2",
    "G3": '=IFERROR(VLOOKUP("East",A2:D9,4,FALSE)/(B2-10),"none")',
    "G4": '=ROUND(SUMPRODUCT(B2:B9,C2:C9)/COUNT(B2:B9),2)&" avg"',
    "G5": "=-(B2+3)^2%",
    "H1": "=SORT(B2:B9)",
}
NAMES = {"TaxRate": "Assumptions!$B$1", "SalesRange": "Sales!$B$2:$B$9"}

PRECEDENTS = [
    ("F1", False), ("F2", False), ("F3", False), ("F4", False), ("F5", False),
    ("F6", False), ("F7", False), ("F7", True), ("F8", False), ("F9", False),
    ("F10", False), ("F11", False), ("G1", False), ("G2", False), ("G2", True),
    ("H1", False), ("D2", False), ("D2:D4", False), ("F1:F4", False), ("G3", False),
]  # fmt: skip
DEPENDENTS = [
    ("B2", False), ("B2", True), ("C2", False), ("C2", True), ("D2", False), ("D2:D9", True),
    ("F1", False), ("F1", True), ("F4", False), ("F4", True), ("A2", False), ("B2:B3", False),
    ("H1", False), ("H1", True), ("G1", True),
]  # fmt: skip
EVALUATE = [
    "F1", "F2", "F3", "F5", "F6", "F7", "F8", "F9", "F10", "F11", "G1", "G2", "G3", "G4", "G5",
]  # fmt: skip

# name, sheet cells (rows from A1), goal cell, target, changing cell
GOAL_NAMES = {"Rate_In": "G11!$A$1"}
GOALS: list[dict[str, Any]] = [
    {"name": "loan rate", "rows": [[0.05, "=PMT(A1/12,360,-200000)"]], "target": 1200},
    {"name": "loan rate from 1%", "rows": [[0.01, "=PMT(A1/12,360,-200000)"]], "target": 1200},
    {"name": "loan amount", "rows": [[100000, "=PMT(0.005,360,-A1)"]], "target": 1500},
    {
        "name": "break-even units",
        "rows": [
            [0, "=A1*C1", 10, 6, 5000],
            [None, "=A1*D1", None, None, None],
            [None, "=B1-B2-E1", None, None, None],
        ],
        "goal": "B3",
        "target": 0,
    },
    {
        "name": "npv of zero",
        "rows": [[0.1, "=-1000+NPV(A1,300,400,500,200)"]],
        "target": 0,
    },
    {"name": "x squared from 1", "rows": [[1, "=A1^2-4"]], "target": 0},
    {"name": "x squared from -1", "rows": [[-1, "=A1^2-4"]], "target": 0},
    {"name": "x squared from 0", "rows": [[0, "=A1^2-4"]], "target": 0},
    {"name": "no solution", "rows": [[1, "=A1^2"]], "target": -1},
    {"name": "cubic stuck", "rows": [[0, "=A1^3-2*A1-5"]], "target": 0},
    {"name": "named input", "rows": [[0.04, "=PMT(Rate_In/12,180,-150000)"]], "target": 1000},
    {"name": "sine", "rows": [[0, "=SIN(A1)"]], "target": 0.5},
    {"name": "exponential", "rows": [[0, "=EXP(A1)"]], "target": 1000},
    {"name": "square root", "rows": [[1, "=SQRT(A1)"]], "target": 3},
    {"name": "absolute value", "rows": [[3, "=ABS(A1-2)"]], "target": 0},
    {"name": "independent", "rows": [[1, "=5"]], "target": 10},
    {
        "name": "across sheets",
        "rows": [[2, "=A1*Assumptions!B2+Sales!B2"]],
        "target": 100,
    },
    {"name": "blank input", "rows": [[None, "=A1*3+1"]], "target": 10},
    {
        "name": "interest on balance",
        "rows": [
            [1000, 0.05, "=A1*(1+B1)^10"],
            [None, None, "=C1-A1"],
        ],
        "goal": "C2",
        "target": 1000,
    },
]


def goal_sheet(index: int) -> str:
    return f"G{index + 1:02d}"


def goal_cell(case: dict[str, Any]) -> str:
    return case.get("goal", "B1")


def goal_changing(case: dict[str, Any]) -> str:
    return case.get("changing", "A1")


def build_workbook(path: Path) -> None:
    workbook = Workbook()
    sales = workbook.active
    assert sales is not None
    sales.title = "Sales"
    assumptions = workbook.create_sheet("Assumptions")
    sheets = [workbook.create_sheet(goal_sheet(i)) for i in range(len(GOALS))]
    write_range(assumptions, "A1", ASSUMPTIONS, [], MAX_CELLS)  # pyright: ignore[reportArgumentType]
    write_range(sales, "A1", SALES, [], MAX_CELLS)  # pyright: ignore[reportArgumentType]
    create_table(workbook, sales, "A1:D9", "tblSales", TableOptions(), MAX_CELLS)
    for name, target in NAMES.items():
        workbook.defined_names[name] = DefinedName(name, attr_text=target)
    for cell, formula in FORMULAS.items():
        write_range(sales, cell, [[formula]], [], MAX_CELLS)  # pyright: ignore[reportArgumentType]
    for name, target in GOAL_NAMES.items():
        workbook.defined_names[name] = DefinedName(name, attr_text=target)
    for sheet, case in zip(sheets, GOALS, strict=True):
        write_range(sheet, "A1", case["rows"], [], MAX_CELLS)  # pyright: ignore[reportArgumentType]
    save_atomically(workbook, path, Limits().max_file_bytes)
